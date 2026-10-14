from fastapi import APIRouter, Request, HTTPException
from pathlib import Path
import sqlite3, datetime

from .qtgh_v02 import (
    ui04_employee_suggestions,
    ui04_lot_suggestions,
    delivery_context,
    qtgh_v082_finalize_selected,
    receiving_banks,
    add_receiving_bank,
    qtgh_v085_receive_cash,
    return_notes,
    receive_return,
    customer_debts,
    collect_customer_debt,
    employee_debts,
    plan_employee_debt_collection,
    adjustments,
    create_adjustment,
    approve_adjustment,
    cash_inventory,
    lock_cash_inventory,
)

router = APIRouter(prefix="/api/qtgh-stable", tags=["QTGH Stable"])
ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "miniapp.db"

def _conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    c.execute("PRAGMA busy_timeout=8000")
    return c

def _num(v):
    try:
        if v is None or v == "":
            return 0.0
        return float(v)
    except Exception:
        return 0.0

def _money(v):
    return int(round(_num(v)))

def _now():
    return datetime.datetime.now().isoformat(timespec="seconds")

def _get_draft(c, did):
    h = c.execute("SELECT * FROM qtgh_draft_order WHERE id=?", (did,)).fetchone()
    if not h:
        raise ValueError("Không tìm thấy quyết toán nháp")
    lines = [dict(r) for r in c.execute(
        "SELECT * FROM qtgh_draft_line WHERE draft_order_id=? ORDER BY id", (did,)
    ).fetchall()]
    out = dict(h)
    out["lines"] = lines
    return out

def _trace(c, source, order):
    c.execute("""CREATE TABLE IF NOT EXISTS qtgh_ui04_draft_trace(
      source_system TEXT NOT NULL,
      ma_don_nguon TEXT NOT NULL,
      status TEXT NOT NULL DEFAULT 'QT_NHAP',
      saved_at TEXT DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY(source_system,ma_don_nguon)
    )""")
    c.execute("""INSERT INTO qtgh_ui04_draft_trace(source_system,ma_don_nguon,status,saved_at)
      VALUES(?,?,'QT_NHAP',CURRENT_TIMESTAMP)
      ON CONFLICT(source_system,ma_don_nguon)
      DO UPDATE SET status='QT_NHAP',saved_at=CURRENT_TIMESTAMP""", (source, order))

def _ensure_draft(source, order):
    import datetime as _dt
    import sqlite3 as _sqlite3
    from pathlib import Path as _Path
    _db_path=_Path(__file__).resolve().parent.parent/"data"/"miniapp.db"
    c=_sqlite3.connect(_db_path)
    c.row_factory=_sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    c.execute("PRAGMA busy_timeout=5000")
    source=str(source or "DMS_MBW").strip() or "DMS_MBW"
    order=str(order or "").strip()
    if not order:
        c.close()
        raise ValueError("Thiếu mã đơn")

    ex=c.execute(
        "SELECT id FROM qtgh_draft_order WHERE source_system=? AND ma_don_nguon=?",
        (source,order)
    ).fetchone()
    if ex:
        did=ex["id"]
        out=_get_draft(c,did)
        c.close()
        return out

    src=c.execute("""SELECT *
                     FROM qtgh_source_projection
                     WHERE source_system='DMS_MBW' AND ma_don_nguon=?
                     ORDER BY source_batch_id DESC,source_row_no,id""",
                  (order,)).fetchall()
    if not src:
        c.close()
        raise ValueError("Không tìm thấy dữ liệu nguồn QTGH")

    # If same order exists in more than one canonical batch, use newest batch only.
    newest_batch=max([int(r["source_batch_id"] or 0) for r in src] or [0])
    src=[r for r in src if int(r["source_batch_id"] or 0)==newest_batch]

    tl=max([float(r["thuong_tich_luy_goc"] or 0) for r in src] or [0])
    tb=max([float(r["thuong_trung_bay_goc"] or 0) for r in src] or [0])
    now=_dt.datetime.now().isoformat(timespec="seconds")

    cur=c.execute("""INSERT INTO qtgh_draft_order(
      source_system,ma_don_nguon,status,vo_type,
      accumulation_after,display_after,other_order_discount_after,
      created_at,updated_at
    ) VALUES(?,?,'NHAP','GIAO_DU',?,?,0,?,?)""",
    ("DMS_MBW",order,round(tl),round(tb),now,now))
    did=cur.lastrowid

    from decimal import Decimal, ROUND_HALF_UP
    subtotal=0
    for r in src:
        gia=float(r["don_gia_co_vat"] or 0)
        ck=float(r["ck_km_dong_goc"] or 0)
        hb=float(r["sl_hang_ban_dat"] or 0)
        km=float(r["sl_hang_km_dat"] or 0)

        # Preserve exact decimals from projection.
        # Round only the final line amount to VND.
        thanh=int((Decimal(str(r["sl_hang_ban_dat"] or 0))*Decimal(str(r["don_gia_co_vat"] or 0))-Decimal(str(r["ck_km_dong_goc"] or 0))).quantize(Decimal("1"),rounding=ROUND_HALF_UP))
        if thanh<0:
            thanh=0

        c.execute("""INSERT INTO qtgh_draft_line(
          draft_order_id,ma_sku,ten_sku,
          sl_hang_ban_dat,sl_hang_ban_tra,sl_hang_ban_thuc_giao,
          sl_hang_km_dat,sl_hang_km_tra,sl_hang_km_thuc_giao,
          don_gia_co_vat,ck_km_dong_goc,ck_km_dong_sau,thanh_tien_sau
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (did,r["ma_sku"],r["ten_sku"],
         str(r["sl_hang_ban_dat"] or "0"),"0",str(r["sl_hang_ban_dat"] or "0"),
         str(r["sl_hang_km_dat"] or "0"),"0",str(r["sl_hang_km_dat"] or "0"),
         gia,ck,ck,int(thanh)))

        subtotal+=int(thanh)

    due=max(0,int(subtotal)-round(tl)-round(tb))
    c.execute("""UPDATE qtgh_draft_order
                 SET so_phai_thu_goc=?,
                     so_phai_thu_sau_qt=?,
                     tien_mat=?,
                     tien_ck_xac_nhan=0,
                     cong_no_kh=0,
                     other_order_discount_after=0
                 WHERE id=?""",
              (due,due,due,did))
    c.commit()
    out=_get_draft(c,did)
    c.close()
    return out


def _assert_draft_editable(c, did):
    row = c.execute("SELECT status FROM qtgh_draft_order WHERE id=?", (did,)).fetchone()
    if not row:
        raise ValueError("Không tìm thấy nháp")
    if str(row["status"] or "").upper() == "DA_KHOA":
        raise ValueError("QTGH_LOCK_001: Quyết toán đã khóa, không được sửa")


def _apply_draft_payload(c, did, payload):
    _assert_draft_editable(c, did)
    from decimal import Decimal, ROUND_HALF_UP
    import datetime as _dt

    def D(v):
        try:
            if v in (None, ""):
                return Decimal("0")
            value = Decimal(str(v).replace(",", ""))
            if not value.is_finite():
                raise ValueError("Giá trị số không hữu hạn")
            return value
        except Exception as exc:
            raise ValueError("Giá trị số nhập vào không hợp lệ") from exc

    def q0(v):
        return D(v).quantize(Decimal("1"), rounding=ROUND_HALF_UP)

    h = c.execute("SELECT * FROM qtgh_draft_order WHERE id=?", (did,)).fetchone()
    if not h:
        raise ValueError("Không tìm thấy draft")

    for key in ("tien_mat", "tien_ck_xac_nhan", "so_du_co_su_dung", "can_tru"):
        if key in payload and D(payload[key]) < 0:
            raise ValueError("QTGH_PAY_002: Các khoản thanh toán không được âm")
    line_payloads = {}
    raw_lines = payload.get("lines")
    if raw_lines is None:
        raw_lines = []
    if not isinstance(raw_lines, list):
        raise ValueError("Danh sách dòng sửa phải là một danh sách")
    for x in raw_lines:
        if not isinstance(x, dict) or x.get("id") is None or isinstance(x.get("id"), bool):
            raise ValueError("Dòng sửa thiếu ID hoặc ID không hợp lệ")
        value = x["id"]
        lid = int(value)
        if lid <= 0 or (isinstance(value, float) and value != lid):
            raise ValueError("ID dòng phải là số nguyên dương")
        if lid in line_payloads:
            raise ValueError("ID dòng sửa bị trùng")
        line_payloads[lid] = x

    rows = c.execute(
        "SELECT * FROM qtgh_draft_line WHERE draft_order_id=? ORDER BY id",
        (did,)
    ).fetchall()

    owned_ids = {int(row["id"]) for row in rows}
    if set(line_payloads) - owned_ids:
        raise ValueError("Dòng hàng không thuộc quyết toán đang sửa")

    subtotal = Decimal("0")

    for row in rows:
        x = line_payloads.get(int(row["id"]), {})

        hb_dat = D(row["sl_hang_ban_dat"])
        hb_tra = D(x.get("sl_hang_ban_tra", row["sl_hang_ban_tra"]))
        hb_tra = max(Decimal("0"), min(hb_dat, hb_tra))
        hb_thuc = hb_dat - hb_tra

        km_dat = D(row["sl_hang_km_dat"])
        km_tra = D(x.get("sl_hang_km_tra", row["sl_hang_km_tra"]))
        km_tra = max(Decimal("0"), min(km_dat, km_tra))
        km_thuc = km_dat - km_tra

        gia = D(row["don_gia_co_vat"])
        ck_goc = D(row["ck_km_dong_goc"])

        sent_ck = x.get("ck_km_dong_sau", row["ck_km_dong_sau"])
        ck_sent = D(sent_ck)

        # UI hiển thị CK theo đồng tròn. Nếu người dùng không thực sự đổi CK,
        # dùng lại CK gốc chính xác có phần thập phân.
        if q0(ck_sent) == q0(ck_goc):
            ck_sau = ck_goc
        else:
            ck_sau = ck_sent

        if hb_thuc <= 0:
            ck_sau = Decimal("0")

        thanh = (hb_thuc * gia - ck_sau).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
        if thanh < 0:
            thanh = Decimal("0")

        c.execute(
            """UPDATE qtgh_draft_line
               SET sl_hang_ban_tra=?,
                   sl_hang_ban_thuc_giao=?,
                   sl_hang_km_tra=?,
                   sl_hang_km_thuc_giao=?,
                   ck_km_dong_sau=?,
                   thanh_tien_sau=?
               WHERE id=?""",
            (
                str(hb_tra), str(hb_thuc),
                str(km_tra), str(km_thuc),
                float(ck_sau), int(thanh),
                row["id"]
            )
        )
        subtotal += thanh

    accumulation = int(q0(payload.get(
        "accumulation_after", h["accumulation_after"]
    )))
    display = int(q0(payload.get(
        "display_after", h["display_after"]
    )))

    # BA V1.1: không còn "Giảm trừ khác".
    other = 0

    due = int(q0(subtotal)) - accumulation - display
    if due < 0:
        due = 0

    cash = int(q0(payload.get("tien_mat", h["tien_mat"])))
    bank = int(q0(payload.get(
        "tien_ck_xac_nhan", h["tien_ck_xac_nhan"]
    )))
    used = int(q0(payload.get(
        "so_du_co_su_dung", h["so_du_co_su_dung"]
    )))
    cantru = int(q0(payload.get("can_tru", h["can_tru"])))

    old_due = int(h["so_phai_thu_sau_qt"] or 0)
    old_cash = int(h["tien_mat"] or 0)
    old_bank = int(h["tien_ck_xac_nhan"] or 0)

    # Giữ cơ chế 100% tiền mặt tự động khi người dùng chưa đổi thanh toán.
    if old_bank == 0 and old_cash == old_due and cash == old_cash and bank == 0:
        cash = due

    mode = str(payload.get("payment_mode") or "").upper()
    if mode == "AUTO_CASH":
        cash, bank = due - used - cantru, 0
    elif mode == "BANK_100":
        cash, bank = 0, due - used - cantru
    if min(cash, bank, used, cantru) < 0:
        raise ValueError("QTGH_PAY_002: Các khoản thanh toán không được âm")
    if cash + bank + used + cantru > due:
        raise ValueError("QTGH_PAY_001: Tổng thanh toán vượt số phải thu sau QT")
    debt = max(0, due - cash - bank - used - cantru)

    sets = [
        "vo_type=?",
        "accumulation_after=?",
        "display_after=?",
        "other_order_discount_after=0",
        "so_phai_thu_sau_qt=?",
        "tien_mat=?",
        "tien_ck_xac_nhan=?",
        "so_du_co_su_dung=?",
        "can_tru=?",
        "cong_no_kh=?",
    ]
    vals = [
        str(payload.get("vo_type", h["vo_type"] or "GIAO_DU")),
        accumulation,
        display,
        due,
        cash,
        bank,
        used,
        cantru,
        debt,
    ]

    if "bank_txn_id" in payload:
        sets.append("bank_txn_id=?")
        vals.append(payload.get("bank_txn_id") or None)

    if "bank_account" in payload:
        sets.append("bank_account=?")
        vals.append(payload.get("bank_account") or None)

    sets.append("updated_at=?")
    vals.append(_dt.datetime.now().isoformat(timespec="seconds"))
    vals.append(did)

    c.execute(
        f"UPDATE qtgh_draft_order SET {','.join(sets)} WHERE id=?",
        vals
    )

def _recalc(c, did, payment_mode=None):
    _assert_draft_editable(c, did)
    h = c.execute("SELECT * FROM qtgh_draft_order WHERE id=?", (did,)).fetchone()
    if not h:
        raise ValueError("Không tìm thấy nháp")
    from decimal import Decimal, ROUND_HALF_UP
    total = 0
    lines = c.execute(
        "SELECT * FROM qtgh_draft_line WHERE draft_order_id=? ORDER BY id",(did,)
    ).fetchall()
    for r in lines:
        hb = max(0, _num(r["sl_hang_ban_dat"]) - _num(r["sl_hang_ban_tra"]))
        km = max(0, _num(r["sl_hang_km_dat"]) - _num(r["sl_hang_km_tra"]))
        ck = _num(r["ck_km_dong_sau"])
        price = _num(r["don_gia_co_vat"])
        # Promotional quantities are delivered free; only sale quantity is payable.
        sale_qty = max(Decimal("0"), Decimal(str(r["sl_hang_ban_dat"] or 0)) - Decimal(str(r["sl_hang_ban_tra"] or 0)))
        amount = int(max(Decimal("0"), sale_qty * Decimal(str(r["don_gia_co_vat"] or 0)) - Decimal(str(r["ck_km_dong_sau"] or 0))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        c.execute("""UPDATE qtgh_draft_line
          SET sl_hang_ban_thuc_giao=?,sl_hang_km_thuc_giao=?,thanh_tien_sau=?
          WHERE id=?""", (str(hb),str(km),_money(amount),r["id"]))
        total += amount

    due = total - _num(h["accumulation_after"]) - _num(h["display_after"]) - _num(h["other_order_discount_after"])
    if -0.5 < due < 0:
        due = 0
    due_i = max(0, int(Decimal(str(due)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)))

    cash = _money(h["tien_mat"])
    bank = _money(h["tien_ck_xac_nhan"])
    used = _money(h["so_du_co_su_dung"])
    offset = _money(h["can_tru"])
    remaining = due_i - used - offset
    mode = (payment_mode or "").upper()
    if mode == "AUTO_CASH":
        cash,bank = remaining,0
    elif mode == "BANK_100":
        cash,bank = 0,remaining

    if cash > due_i:
        raise ValueError("Tiền mặt không được lớn hơn số phải thu")

    if min(cash, bank, used, offset) < 0:
        raise ValueError("QTGH_PAY_002: Các khoản thanh toán không được âm")
    if cash + bank + used + offset > due_i:
        raise ValueError("QTGH_PAY_001: Tổng thanh toán vượt số phải thu sau QT")
    debt = due_i - cash - bank - used - offset

    c.execute("""UPDATE qtgh_draft_order
      SET so_phai_thu_sau_qt=?,tien_mat=?,tien_ck_xac_nhan=?,cong_no_kh=?,updated_at=?
      WHERE id=?""", (due_i,cash,bank,debt,_now(),did))

def _pending(ma_nvgh, ngay_giao=None, ma_lo=None, ma_don_nguon=None):
    if not (ma_nvgh or "").strip() and not (ma_don_nguon or "").strip():
        raise ValueError("Nhập NVGH hoặc mã đơn nguồn chính xác")
    c = _conn()
    sql = """SELECT
      s.source_system, MAX(s.loai_don) loai_don, s.ma_don_nguon,
      MIN(s.ngay_don) ngay_don, MIN(s.ngay_giao) ngay_giao,
      MAX(s.ma_nvgh) ma_nvgh, MAX(s.ma_lo) ma_lo,
      MAX(s.ma_kh_chuan) ma_kh_chuan, MAX(s.ma_kh_mbw) ma_kh_mbw,
      MAX(s.ma_kh_dms) ma_kh_dms, COUNT(*) so_sku,
      SUM(CAST(s.sl_hang_ban_dat AS REAL)*CAST(s.don_gia_co_vat AS REAL)
          -CAST(s.ck_km_dong_goc AS REAL)) line_due,
      MAX(CAST(s.thuong_tich_luy_goc AS REAL)) tl,
      MAX(CAST(s.thuong_trung_bay_goc AS REAL)) tb
      FROM qtgh_source_projection s
      WHERE s.source_system='DMS_MBW'"""
    p=[]
    if ma_don_nguon:
        sql += " AND UPPER(s.ma_don_nguon)=UPPER(?)"
        p.append(ma_don_nguon.strip())
    if ma_nvgh:
        sql += " AND s.ma_nvgh=?"
        p.append(ma_nvgh)
    if ngay_giao:
        sql += " AND s.ngay_giao=?"
        p.append(ngay_giao)
    if ma_lo:
        sql += " AND s.ma_lo=?"
        p.append(ma_lo)
    sql += " GROUP BY s.ma_don_nguon ORDER BY ngay_giao,ma_lo,ma_don_nguon"
    src=[dict(r) for r in c.execute(sql,p).fetchall()]

    settled=set()
    try:
        for r in c.execute("""SELECT d.ma_don_nguon
          FROM qtgh_session_order so
          JOIN qtgh_session se ON se.id=so.session_id
          JOIN qtgh_draft_order d ON d.id=so.draft_order_id
          WHERE se.status='DA_KHOA'"""):
            settled.add(r["ma_don_nguon"])
    except Exception:
        pass

    drafts={}
    for r in c.execute("""SELECT id,source_system,ma_don_nguon,status,vo_type,
      so_phai_thu_sau_qt,tien_mat,tien_ck_xac_nhan,cong_no_kh
      FROM qtgh_draft_order"""):
        drafts[r["ma_don_nguon"]]=dict(r)

    out=[]
    for r in src:
        order=r["ma_don_nguon"]
        if order in settled:
            continue
        d=drafts.get(order)
        if d:
            due=_money(d["so_phai_thu_sau_qt"])
            cash=_money(d["tien_mat"]); bank=_money(d["tien_ck_xac_nhan"]); debt=_money(d["cong_no_kh"])
            r.update({
                "draft_id":d["id"],"source_system":"DMS_MBW",
                "so_phai_thu":due,"tien_mat":cash,"tien_ck":bank,"cong_no_kh":debt,
                "qt_state":"QT nháp","status_qt":"QT_NHAP"
            })
        else:
            due=_money(_num(r["line_due"])-_num(r["tl"])-_num(r["tb"]))
            r.update({
                "draft_id":None,"source_system":"DMS_MBW",
                "so_phai_thu":due,"tien_mat":due,"tien_ck":0,"cong_no_kh":0,
                "qt_state":"Bình thường","status_qt":"BINH_THUONG"
            })
        out.append(r)
    c.close()
    return {"rows":out,"count":len(out)}

@router.get("/context")
def context():
    try:
        return delivery_context()
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui04/suggest-nvgh")
def suggest_nvgh(q:str="",ngay_giao:str=None):
    try:
        return ui04_employee_suggestions(q,ngay_giao)
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui04/suggest-lo")
def suggest_lo(ma_nvgh:str,q:str="",ngay_giao:str=None):
    try:
        return ui04_lot_suggestions(ma_nvgh,q,ngay_giao)
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui04/pending-orders")
def pending_orders(ma_nvgh:str="",ngay_giao:str=None,ma_lo:str=None,ma_don_nguon:str=None):
    try:
        return _pending(ma_nvgh,ngay_giao,ma_lo,ma_don_nguon)
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui04/open-draft")
async def open_draft(request:Request):
    try:
        x=await request.json()
        return _ensure_draft(x.get("source_system","DMS_MBW"),x.get("ma_don_nguon"))
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui04/draft/{did}/recalc")
async def recalc_draft(did:int,request:Request):
    c = None
    try:
        x=await request.json()
        c=_conn()
        _apply_draft_payload(c,did,x)
        _recalc(c,did,x.get("payment_mode"))
        h=c.execute("SELECT source_system,ma_don_nguon FROM qtgh_draft_order WHERE id=?",(did,)).fetchone()
        if h:_trace(c,h["source_system"],h["ma_don_nguon"])
        c.commit()
        out=_get_draft(c,did)
        c.close()
        return out
    except Exception as e:
        if c is not None:
            c.rollback()
        raise HTTPException(400,str(e))
    finally:
        if c is not None:
            c.close()

@router.post("/ui04/draft/{did}/payment")
async def save_payment(did:int,request:Request):
    c = None
    try:
        x=await request.json()
        bank=_money(x.get("tien_ck_xac_nhan",0))
        cash=_money(x.get("tien_mat",0))

        c0=_conn()
        h0=c0.execute("SELECT so_phai_thu_sau_qt FROM qtgh_draft_order WHERE id=?",(did,)).fetchone()
        c0.close()
        due0=_money(h0["so_phai_thu_sau_qt"] if h0 else 0)

        if cash > due0:
            raise ValueError("Tiền mặt không được lớn hơn số phải thu")

        if bank>0 and (not str(x.get("bank_account","")).strip() or not str(x.get("bank_txn_id","")).strip()):
            raise ValueError("Có chuyển khoản: bắt buộc chọn tài khoản nhận và nhập mã giao dịch")
        c=_conn()
        _apply_draft_payload(c,did,x)
        _recalc(c,did,None)
        h=c.execute("SELECT source_system,ma_don_nguon FROM qtgh_draft_order WHERE id=?",(did,)).fetchone()
        if h:_trace(c,h["source_system"],h["ma_don_nguon"])
        c.commit()
        out=_get_draft(c,did)
        c.close()
        return out
    except Exception as e:
        if c is not None:
            c.rollback()
        raise HTTPException(400,str(e))
    finally:
        if c is not None:
            c.close()

@router.post("/ui04/draft/{did}/full-break")
def full_break(did:int):
    c = None
    try:
        c=_conn()
        _assert_draft_editable(c, did)
        rows=c.execute("SELECT * FROM qtgh_draft_line WHERE draft_order_id=?",(did,)).fetchall()
        for r in rows:
            c.execute("""UPDATE qtgh_draft_line
              SET sl_hang_ban_tra=sl_hang_ban_dat,sl_hang_km_tra=sl_hang_km_dat,
                  ck_km_dong_sau=0
              WHERE id=?""",(r["id"],))
        _recalc(c,did,"AUTO_CASH")
        h=c.execute("SELECT source_system,ma_don_nguon FROM qtgh_draft_order WHERE id=?",(did,)).fetchone()
        if h:_trace(c,h["source_system"],h["ma_don_nguon"])
        c.commit()
        out=_get_draft(c,did)
        c.close()
        return out
    except Exception as e:
        if c is not None:
            c.rollback()
        raise HTTPException(400,str(e))
    finally:
        if c is not None:
            c.close()

@router.post("/ui04/finalize-selected")
async def finalize_selected(request:Request):
    try:
        x=await request.json()
        items=x.get("items",[])
        for item in items:
            _ensure_draft("DMS_MBW",item.get("ma_don_nguon"))
        return qtgh_v082_finalize_selected(x["ma_nvgh"],[
            {"source_system":"DMS_MBW","ma_don_nguon":i.get("ma_don_nguon")} for i in items
        ])
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/banks")
def banks():
    try:
        return receiving_banks()
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/banks")
async def add_bank(request:Request):
    try:
        x=await request.json()
        return add_receiving_bank(x.get("bank_name"),x.get("account_no"),x.get("account_name",""))
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui05/cash")
def cash_list():
    c=_conn()
    rows=[dict(r) for r in c.execute("SELECT * FROM qtgh_cash_receipt ORDER BY id DESC").fetchall()]
    c.close()
    return {"rows":rows}

@router.post("/ui05/cash/{rid}/receive")
async def cash_receive(rid:int,request:Request):
    try:
        x=await request.json()
        resolution=str(x.get("resolution_code") or "").strip().upper()
        if resolution=="XIN_Y_KIEN":
            resolution="XIN_Y_KIEN_LANH_DAO"
        return qtgh_v085_receive_cash(
            rid,x.get("actual_amount",0),resolution,x.get("resolution_note","")
        )
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui05/cash/{rid}/print-data")
def print_data(rid:int):
    # INTERNAL_RECEIPT_V1
    from .receipt_internal import print_data as internal_print_data
    c=None
    try:
        c=_conn()
        return internal_print_data(c,rid)
    except Exception as e:
        raise HTTPException(400,str(e))
    finally:
        if c is not None:c.close()

@router.get("/ui06/returns")
def ui06_returns():
    try:return return_notes()
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/ui06/returns/{nid}/receive")
async def ui06_receive(nid:int,request:Request):
    try:return receive_return(nid,(await request.json()).get("lines",[]))
    except Exception as e:raise HTTPException(400,str(e))

@router.get("/ui07/debts")
def ui07_debts():
    try:return customer_debts()
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/ui08/debt/{did}/collect")
async def ui08_collect(did:int,request:Request):
    try:
        x=await request.json()
        return collect_customer_debt(did,x.get("cash_amount",0),x.get("bank_amount",0),x.get("reason",""))
    except Exception as e:raise HTTPException(400,str(e))

@router.get("/ui09/debts")
def ui09_debts():
    try:return employee_debts()
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/ui09/debt/{did}/plan")
async def ui09_plan(did:int,request:Request):
    try:return plan_employee_debt_collection(did,(await request.json()).get("amount",0))
    except Exception as e:raise HTTPException(400,str(e))

@router.get("/ui10/adjustments")
def ui10_list():
    try:return adjustments()
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/ui10/adjustment")
async def ui10_create(request:Request):
    try:
        raise ValueError("Luồng phiếu điều chỉnh cũ đã đóng; dùng phiếu hàng hoặc phiếu tiền mới")
        x=await request.json()
        return create_adjustment(x["source_doc_type"],x["source_doc_id"],x["adjustment_type"],x["delta_amount"],x["reason"])
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/ui10/adjustment/{aid}/approve")
def ui10_approve(aid:int):
    raise HTTPException(409,"Phiếu cũ chỉ để đối chiếu; không duyệt qua luồng này")
    try:return approve_adjustment(aid)
    except Exception as e:raise HTTPException(400,str(e))

@router.get("/ui11/cash-inventory/{date}")
def ui11_inventory(date:str):
    try:return cash_inventory(date)
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/ui11/cash-inventory/{date}/lock")
async def ui11_lock(date:str,request:Request):
    try:return lock_cash_inventory(date,(await request.json()).get("actual_count",0))
    except Exception as e:raise HTTPException(400,str(e))


# ===== QTGH STABLE V1.1.4 UI06 GROUPED WAREHOUSE RETURNS =====

def _ui06_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS qtgh_wh_return_request(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      request_no TEXT UNIQUE,
      ma_nvgh TEXT NOT NULL,
      status TEXT NOT NULL DEFAULT 'CHO_NHAN',
      created_at TEXT NOT NULL,
      received_at TEXT,
      note TEXT
    );

    CREATE TABLE IF NOT EXISTS qtgh_wh_return_request_order(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      request_id INTEGER NOT NULL,
      draft_order_id INTEGER NOT NULL,
      ma_don_nguon TEXT NOT NULL,
      UNIQUE(draft_order_id),
      FOREIGN KEY(request_id) REFERENCES qtgh_wh_return_request(id)
    );

    CREATE TABLE IF NOT EXISTS qtgh_wh_return_request_line(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      request_id INTEGER NOT NULL,
      ma_sku TEXT NOT NULL,
      ten_sku TEXT,
      sl_hb_phai_tra REAL NOT NULL DEFAULT 0,
      sl_km_phai_tra REAL NOT NULL DEFAULT 0,
      sl_phai_tra REAL NOT NULL DEFAULT 0,
      sl_thuc_nhan REAL NOT NULL DEFAULT 0,
      FOREIGN KEY(request_id) REFERENCES qtgh_wh_return_request(id)
    );

    CREATE INDEX IF NOT EXISTS ix_qtgh_wh_return_request_status
      ON qtgh_wh_return_request(status,ma_nvgh);
    """)
    c.commit()
    c.close()

_ui06_init()

def _ui06_finalized_return_orders(ma_nvgh=None):
    c=_conn()

    sql = """
    SELECT
      d.id draft_order_id,
      d.ma_don_nguon,
      MAX(p.ma_nvgh) ma_nvgh,
      MAX(p.ngay_giao) ngay_giao,
      MAX(p.ma_lo) ma_lo
    FROM qtgh_draft_order d
    JOIN qtgh_draft_line l ON l.draft_order_id=d.id
    JOIN qtgh_source_projection p
      ON p.source_system='DMS_MBW'
     AND p.ma_don_nguon=d.ma_don_nguon
     AND p.ma_sku=l.ma_sku
    WHERE
      (CAST(l.sl_hang_ban_tra AS REAL)>0 OR CAST(l.sl_hang_km_tra AS REAL)>0)
      AND EXISTS(
        SELECT 1
        FROM qtgh_session_order so
        JOIN qtgh_session se ON se.id=so.session_id
        WHERE so.draft_order_id=d.id
          AND se.status='DA_KHOA'
      )
      AND NOT EXISTS(
        SELECT 1 FROM qtgh_wh_return_request_order ro
        WHERE ro.draft_order_id=d.id
      )
    """

    params=[]
    if ma_nvgh:
        sql += " AND p.ma_nvgh=?"
        params.append(ma_nvgh)

    sql += """
    GROUP BY d.id,d.ma_don_nguon
    ORDER BY ma_nvgh,ngay_giao,ma_lo,d.ma_don_nguon
    """

    heads=[dict(r) for r in c.execute(sql,params).fetchall()]
    out=[]

    for h in heads:
        lines=[dict(r) for r in c.execute("""
          SELECT
            l.id,
            l.ma_sku,
            l.ten_sku,
            CAST(l.sl_hang_ban_tra AS REAL) sl_hb_tra,
            CAST(l.sl_hang_km_tra AS REAL) sl_km_tra,
            CAST(l.sl_hang_ban_tra AS REAL)+CAST(l.sl_hang_km_tra AS REAL) sl_tra
          FROM qtgh_draft_line l
          WHERE l.draft_order_id=?
            AND (CAST(l.sl_hang_ban_tra AS REAL)>0 OR CAST(l.sl_hang_km_tra AS REAL)>0)
          ORDER BY l.ma_sku
        """,(h["draft_order_id"],)).fetchall()]
        h["lines"]=lines
        h["total_return_qty"]=sum(_num(x["sl_tra"]) for x in lines)
        out.append(h)

    c.close()
    return out

@router.get("/ui06/pending-orders")
def ui06_pending_orders(ma_nvgh:str=None):
    try:
        rows=_ui06_finalized_return_orders(ma_nvgh)
        groups={}
        for r in rows:
            nv=r.get("ma_nvgh") or ""
            groups.setdefault(nv,{"ma_nvgh":nv,"orders":[],"order_count":0,"total_return_qty":0})
            groups[nv]["orders"].append(r)
            groups[nv]["order_count"]+=1
            groups[nv]["total_return_qty"]+=r["total_return_qty"]
        return {"rows":list(groups.values()),"order_count":len(rows)}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui06/group")
async def ui06_group_orders(request:Request):
    try:
        x=await request.json()
        ma_nvgh=str(x.get("ma_nvgh") or "").strip()
        draft_ids=[int(v) for v in x.get("draft_order_ids",[])]

        if not ma_nvgh:
            raise ValueError("Thiếu NVGH")
        if not draft_ids:
            raise ValueError("Chưa chọn đơn để gom")

        eligible={r["draft_order_id"]:r for r in _ui06_finalized_return_orders(ma_nvgh)}
        bad=[v for v in draft_ids if v not in eligible]
        if bad:
            raise ValueError("Có đơn không còn đủ điều kiện gom: "+",".join(map(str,bad)))

        c=_conn()
        now=_now()
        seq=c.execute("SELECT COALESCE(MAX(id),0)+1 FROM qtgh_wh_return_request").fetchone()[0]
        req_no=f"YCTK-{datetime.datetime.now():%Y%m%d}-{seq:04d}"

        cur=c.execute("""INSERT INTO qtgh_wh_return_request(
          request_no,ma_nvgh,status,created_at
        ) VALUES(?,?,'CHO_NHAN',?)""",(req_no,ma_nvgh,now))
        rid=cur.lastrowid

        sku={}
        for did in draft_ids:
            order=eligible[did]
            c.execute("""INSERT INTO qtgh_wh_return_request_order(
              request_id,draft_order_id,ma_don_nguon
            ) VALUES(?,?,?)""",(rid,did,order["ma_don_nguon"]))

            for l in order["lines"]:
                key=l["ma_sku"]
                row=sku.setdefault(key,{
                  "ma_sku":key,
                  "ten_sku":l.get("ten_sku") or "",
                  "hb":0.0,"km":0.0
                })
                row["hb"]+=_num(l["sl_hb_tra"])
                row["km"]+=_num(l["sl_km_tra"])

        for row in sku.values():
            total=row["hb"]+row["km"]
            c.execute("""INSERT INTO qtgh_wh_return_request_line(
              request_id,ma_sku,ten_sku,
              sl_hb_phai_tra,sl_km_phai_tra,
              sl_phai_tra,sl_thuc_nhan
            ) VALUES(?,?,?,?,?,?,?)""",
            (rid,row["ma_sku"],row["ten_sku"],
             row["hb"],row["km"],total,total))

        c.commit()
        c.close()
        return {"ok":True,"request_id":rid,"request_no":req_no}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui06/requests")
def ui06_requests(status:str=None):
    try:
        c=_conn()
        sql="""SELECT r.*,
          (SELECT COUNT(*) FROM qtgh_wh_return_request_order o WHERE o.request_id=r.id) order_count,
          (SELECT COALESCE(SUM(sl_phai_tra),0) FROM qtgh_wh_return_request_line l WHERE l.request_id=r.id) total_required,
          (SELECT COALESCE(SUM(sl_thuc_nhan),0) FROM qtgh_wh_return_request_line l WHERE l.request_id=r.id) total_actual
          FROM qtgh_wh_return_request r"""
        params=[]
        if status:
            sql+=" WHERE r.status=?"
            params.append(status)
        sql+=" ORDER BY r.id DESC"
        rows=[dict(r) for r in c.execute(sql,params).fetchall()]
        for r in rows:
            r["total_actual_original"]=r["total_actual"]
            r["total_actual"]+=_num(c.execute("""SELECT COALESCE(SUM(a.delta_qty),0)
              FROM qtgh_wh_quantity_adjustment a WHERE a.request_id=? AND a.status='DA_DUYET'""",
              (r["id"],)).fetchone()[0])
        c.close()
        return {"rows":rows}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/ui06/requests/{rid}")
def ui06_request_detail(rid:int):
    import json  # VG_RETURN_PRINT_JSON_FIX
    try:
        c=_conn()
        h=c.execute("SELECT * FROM qtgh_wh_return_request WHERE id=?",(rid,)).fetchone()
        if not h:
            c.close()
            raise ValueError("Không tìm thấy phiếu yêu cầu trả kho")
        out=dict(h)
        # VG_PRINT_FORMS_20261003: display only, actual return employee code.
        out['ma_nvgh_name']=''
        staff_code=str(out.get('ma_nvgh') or '').strip().upper()
        for staff_row in c.execute("SELECT business_key,payload_json FROM master_data_generic WHERE category='dm02-staff'"):
            staff=json.loads(staff_row['payload_json'] or '{}')
            code=str(staff.get('MA_NV') or staff_row['business_key'] or '').strip().upper()
            if staff_code and code==staff_code:
                out['ma_nvgh_name']=str(staff.get('TEN_NV') or staff.get('HO_TEN') or staff.get('TEN_NHAN_VIEN') or '').strip()
                break
        out["orders"]=[dict(r) for r in c.execute("""
          SELECT draft_order_id,ma_don_nguon
          FROM qtgh_wh_return_request_order
          WHERE request_id=?
          ORDER BY ma_don_nguon
        """,(rid,)).fetchall()]
        out["lines"]=[dict(r) for r in c.execute("""
          SELECT * FROM qtgh_wh_return_request_line
          WHERE request_id=?
          ORDER BY ma_sku
        """,(rid,)).fetchall()]
        for line in out["lines"]:
            line["sl_thuc_nhan_hieu_luc"]=_ui10_effective_warehouse_qty(
                c,line["id"],line["sl_thuc_nhan"])
        c.close()
        return out
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui06/requests/{rid}/receive")
async def ui06_receive_request(rid:int,request:Request):
    try:
        x=await request.json()
        values={int(v["id"]):_num(v.get("sl_thuc_nhan",0)) for v in x.get("lines",[])}

        c=_conn()
        h=c.execute("SELECT * FROM qtgh_wh_return_request WHERE id=?",(rid,)).fetchone()
        if not h:
            c.close()
            raise ValueError("Không tìm thấy phiếu")
        if h["status"]=="DA_NHAN":
            c.close()
            raise ValueError("Phiếu đã nhận và đã khóa")

        lines=c.execute("""
          SELECT * FROM qtgh_wh_return_request_line
          WHERE request_id=?
        """,(rid,)).fetchall()

        for l in lines:
            actual=values.get(l["id"],_num(l["sl_thuc_nhan"]))
            required=_num(l["sl_phai_tra"])
            if actual<0:
                c.close()
                raise ValueError(f"SL-thực nhận không được âm: {l['ma_sku']}")
            if actual>required:
                c.close()
                raise ValueError(f"SL-thực nhận vượt SL-phải trả: {l['ma_sku']}")
            c.execute("""UPDATE qtgh_wh_return_request_line
              SET sl_thuc_nhan=?
              WHERE id=?""",(actual,l["id"]))

        c.execute("""UPDATE qtgh_wh_return_request
          SET status='DA_NHAN',received_at=?
          WHERE id=?""",(_now(),rid))

        c.commit()
        c.close()

        return ui06_request_detail(rid)
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== QTGH STABLE V1.1.6 UI07 CUSTOMER DEBT BY NVGH =====

def _ui07_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS qtgh_customer_collect_batch_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      ma_nvgh TEXT NOT NULL,
      status TEXT NOT NULL DEFAULT 'OPEN',
      cash_total REAL NOT NULL DEFAULT 0,
      bank_total REAL NOT NULL DEFAULT 0,
      cash_receipt_id INTEGER,
      created_at TEXT NOT NULL,
      closed_at TEXT
    );

    CREATE TABLE IF NOT EXISTS qtgh_customer_collect_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      batch_id INTEGER NOT NULL,
      ma_kh_key TEXT NOT NULL,
      cash_amount REAL NOT NULL DEFAULT 0,
      bank_amount REAL NOT NULL DEFAULT 0,
      bank_account TEXT,
      bank_txn_id TEXT,
      created_at TEXT NOT NULL,
      FOREIGN KEY(batch_id) REFERENCES qtgh_customer_collect_batch_stable(id)
    );

    CREATE TABLE IF NOT EXISTS qtgh_customer_collect_alloc_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      collect_id INTEGER NOT NULL,
      debt_id INTEGER NOT NULL,
      ma_don_nguon TEXT,
      allocated_amount REAL NOT NULL DEFAULT 0,
      FOREIGN KEY(collect_id) REFERENCES qtgh_customer_collect_stable(id)
    );
    """)
    c.commit()
    c.close()

_ui07_init()

def _ui07_debt_rows():
    x=customer_debts()
    return x.get("rows",[]) if isinstance(x,dict) else []

def _ui07_customer_key(r):
    return str(
      r.get("ma_kh_chuan")
      or r.get("ma_kh_nguon")
      or r.get("ma_kh_dms")
      or r.get("ma_kh_mbw")
      or ""
    ).strip()

def _ui07_balance(r):
    for k in ("balance","con_no","so_con_no","amount_due","debt_balance"):
        if k in r and r.get(k) is not None:
            return _num(r.get(k))
    return 0.0

def _ui07_date(r):
    return str(r.get("ngay_don") or r.get("ngay_giao") or r.get("created_at") or "")

def _ui07_open_batch(ma_nvgh):
    c=_conn()
    r=c.execute("""SELECT * FROM qtgh_customer_collect_batch_stable
      WHERE ma_nvgh=? AND status='OPEN'
      ORDER BY id DESC LIMIT 1""",(ma_nvgh,)).fetchone()
    if r:
        out=dict(r)
        c.close()
        return out
    cur=c.execute("""INSERT INTO qtgh_customer_collect_batch_stable(
      ma_nvgh,status,created_at
    ) VALUES(?,'OPEN',?)""",(ma_nvgh,_now()))
    bid=cur.lastrowid
    c.commit()
    out=dict(c.execute("SELECT * FROM qtgh_customer_collect_batch_stable WHERE id=?",(bid,)).fetchone())
    c.close()
    return out

@router.get("/ui07/customer-groups")
def ui07_customer_groups(ma_nvgh:str):
    try:
        rows=[r for r in _ui07_debt_rows()
              if str(r.get("ma_nvgh") or "").strip()==ma_nvgh
              and _ui07_balance(r)>0]

        groups={}
        for r in rows:
            key=_ui07_customer_key(r)
            g=groups.setdefault(key,{
              "ma_kh_key":key,
              "ma_kh_chuan":r.get("ma_kh_chuan") or "",
              "ma_kh_nguon":r.get("ma_kh_nguon") or r.get("ma_kh_dms") or r.get("ma_kh_mbw") or "",
              "ten_kh":r.get("ten_kh") or r.get("customer_name") or "",
              "total_debt":0,
              "invoices":[]
            })
            inv=dict(r)
            inv["balance"]=_ui07_balance(r)
            inv["ngay_sort"]=_ui07_date(r)
            g["invoices"].append(inv)
            g["total_debt"]+=inv["balance"]

        for g in groups.values():
            g["invoices"].sort(key=lambda x:(x.get("ngay_sort") or "",str(x.get("ma_don_nguon") or "")))

        batch=_ui07_open_batch(ma_nvgh)

        c=_conn()
        temp=[dict(r) for r in c.execute("""
          SELECT c.*, COUNT(a.id) alloc_count
          FROM qtgh_customer_collect_stable c
          LEFT JOIN qtgh_customer_collect_alloc_stable a ON a.collect_id=c.id
          WHERE c.batch_id=?
          GROUP BY c.id
          ORDER BY c.id
        """,(batch["id"],)).fetchall()]
        c.close()

        return {
          "rows":list(groups.values()),
          "batch":batch,
          "temp_collections":temp
        }
    except Exception as e:
        raise HTTPException(400,str(e))

def _ui07_update_debt_after_alloc(c,debt_id,amount):
    # Dynamically adapt to current qtgh_customer_debt schema.
    row=c.execute("SELECT * FROM qtgh_customer_debt WHERE id=?",(debt_id,)).fetchone()
    if not row:
        raise ValueError(f"Không tìm thấy công nợ #{debt_id}")
    d=dict(row)
    cols=set(d.keys())

    balance_col=next((k for k in ("balance","con_no","so_con_no","debt_balance") if k in cols),None)
    paid_col=next((k for k in ("paid_amount","collected_amount","da_thu","amount_paid") if k in cols),None)

    if balance_col:
        old=_num(d.get(balance_col))
        newbal=max(0,old-amount)
        sets=[f'"{balance_col}"=?']
        vals=[newbal]
        if paid_col:
            sets.append(f'"{paid_col}"=?')
            vals.append(_num(d.get(paid_col))+amount)
        if "status" in cols:
            sets.append('"status"=?')
            vals.append("DA_THU_DU" if newbal<=0.0001 else "CON_NO")
        vals.append(debt_id)
        c.execute(f'UPDATE qtgh_customer_debt SET {",".join(sets)} WHERE id=?',vals)
    else:
        raise ValueError("Không xác định được cột số dư công nợ")

@router.post("/ui07/customer-collect")
async def ui07_customer_collect(request:Request):
    try:
        x=await request.json()
        ma_nvgh=str(x.get("ma_nvgh") or "").strip()
        ma_kh_key=str(x.get("ma_kh_key") or "").strip()
        cash=_num(x.get("cash_amount",0))
        bank=_num(x.get("bank_amount",0))
        bank_account=str(x.get("bank_account") or "").strip()
        bank_txn_id=str(x.get("bank_txn_id") or "").strip()

        if cash<0 or bank<0:
            raise ValueError("Số tiền không được âm")
        if cash+bank<=0:
            raise ValueError("Chưa nhập số tiền thu nợ")
        if bank>0 and (not bank_account or not bank_txn_id):
            raise ValueError("Có chuyển khoản: bắt buộc chọn tài khoản nhận và mã giao dịch")

        rows=[r for r in _ui07_debt_rows()
              if str(r.get("ma_nvgh") or "").strip()==ma_nvgh
              and _ui07_customer_key(r)==ma_kh_key
              and _ui07_balance(r)>0]

        rows.sort(key=lambda r:(_ui07_date(r),str(r.get("ma_don_nguon") or "")))
        total_debt=sum(_ui07_balance(r) for r in rows)
        total_pay=cash+bank

        if total_pay>total_debt:
            raise ValueError("Số tiền thu nợ vượt tổng công nợ khách hàng")

        batch=_ui07_open_batch(ma_nvgh)

        c=_conn()
        cur=c.execute("""INSERT INTO qtgh_customer_collect_stable(
          batch_id,ma_kh_key,cash_amount,bank_amount,bank_account,bank_txn_id,created_at
        ) VALUES(?,?,?,?,?,?,?)""",
        (batch["id"],ma_kh_key,cash,bank,bank_account,bank_txn_id,_now()))
        cid=cur.lastrowid

        remain=total_pay
        allocs=[]
        for r in rows:
            if remain<=0:
                break
            bal=_ui07_balance(r)
            a=min(remain,bal)
            debt_id=int(r["id"])
            c.execute("""INSERT INTO qtgh_customer_collect_alloc_stable(
              collect_id,debt_id,ma_don_nguon,allocated_amount
            ) VALUES(?,?,?,?)""",(cid,debt_id,r.get("ma_don_nguon"),a))
            _ui07_update_debt_after_alloc(c,debt_id,a)
            allocs.append({
              "debt_id":debt_id,
              "ma_don_nguon":r.get("ma_don_nguon"),
              "allocated_amount":a
            })
            remain-=a

        c.execute("""UPDATE qtgh_customer_collect_batch_stable
          SET cash_total=cash_total+?,bank_total=bank_total+?
          WHERE id=?""",(cash,bank,batch["id"]))

        c.commit()
        c.close()

        return {
          "ok":True,
          "batch_id":batch["id"],
          "collect_id":cid,
          "allocations":allocs
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui07/finalize-cash-receipt")
async def ui07_finalize_cash_receipt(request:Request):
    try:
        x=await request.json()
        ma_nvgh=str(x.get("ma_nvgh") or "").strip()
        batch=_ui07_open_batch(ma_nvgh)

        c=_conn()
        b=c.execute("SELECT * FROM qtgh_customer_collect_batch_stable WHERE id=?",(batch["id"],)).fetchone()
        if not b:
            c.close()
            raise ValueError("Không tìm thấy phiên thu nợ")

        if _num(b["cash_total"])<=0:
            c.close()
            raise ValueError("Phiên này không có tiền mặt để lập phiếu thu")

        # Dynamically insert a cash receipt using current qtgh_cash_receipt schema.
        cols=[r["name"] for r in c.execute("PRAGMA table_info(qtgh_cash_receipt)").fetchall()]
        values={}
        now=_now()

        candidates={
          "receipt_kind":"THU_NO_KH",
          "source_type":"THU_NO_KH",
          "source_id":int(batch["id"]),
          "session_key":f"THU_NO_KH:{batch['id']}",
          "ma_nvgh":ma_nvgh,
          "expected_amount":_money(b["cash_total"]),
          "actual_amount":0,
          "variance_amount":_money(b["cash_total"]),
          "status":"CHO_THU",
          "created_at":now,
          "updated_at":now
        }
        for k,v in candidates.items():
            if k in cols:
                values[k]=v

        if not values:
            c.close()
            raise ValueError("Không xác định được cấu trúc phiếu thu")

        fields=",".join(f'"{k}"' for k in values)
        qs=",".join("?" for _ in values)
        cur=c.execute(
          f'INSERT INTO qtgh_cash_receipt({fields}) VALUES({qs})',
          list(values.values())
        )
        rid=cur.lastrowid

        c.execute("""UPDATE qtgh_customer_collect_batch_stable
          SET status='CLOSED',cash_receipt_id=?,closed_at=?
          WHERE id=?""",(rid,now,batch["id"]))

        c.commit()
        c.close()

        return {
          "ok":True,
          "cash_receipt_id":rid,
          "cash_total":_money(b["cash_total"]),
          "bank_total":_money(b["bank_total"])
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== QTGH STABLE V1.1.8 UI09 EMPLOYEE DEBT COLLECTION =====

def _ui09_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS qtgh_employee_collect_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      employee_debt_id INTEGER NOT NULL,
      ma_nvgh TEXT NOT NULL,
      ten_nv TEXT,
      cash_amount REAL NOT NULL DEFAULT 0,
      bank_amount REAL NOT NULL DEFAULT 0,
      bank_account TEXT,
      bank_txn_id TEXT,
      cash_receipt_id INTEGER,
      status TEXT NOT NULL DEFAULT 'OPEN',
      created_at TEXT NOT NULL
    );
    """)
    c.commit()
    c.close()

_ui09_init()

def _ui09_employee_name(ma_nv):
    c=_conn()
    name=""

    # Direct relational tables first.
    tables=[r["name"] for r in c.execute(
      "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
    ).fetchall()]

    for t in tables:
        try:
            cols=[r["name"] for r in c.execute(f'PRAGMA table_info("{t}")').fetchall()]
        except Exception:
            continue

        up={x.upper():x for x in cols}
        ma_col=up.get("MA_NV")
        ten_col=up.get("TEN_NV") or up.get("HO_TEN") or up.get("TEN_NHAN_VIEN") or up.get("NAME")

        if ma_col and ten_col:
            try:
                r=c.execute(
                  f'SELECT "{ten_col}" ten FROM "{t}" WHERE UPPER(TRIM("{ma_col}"))=? LIMIT 1',
                  (str(ma_nv).strip().upper(),)
                ).fetchone()
                if r and r["ten"]:
                    name=str(r["ten"]).strip()
                    break
            except Exception:
                pass

    # Fallback generic JSON master.
    if not name and "master_data_generic" in tables:
        try:
            rows=c.execute("""
              SELECT payload_json
              FROM master_data_generic
              WHERE payload_json LIKE ?
              LIMIT 100
            """,(f'%{ma_nv}%',)).fetchall()

            import json as _json
            for rr in rows:
                try:
                    p=_json.loads(rr["payload_json"] or "{}")
                except Exception:
                    continue
                if str(p.get("MA_NV") or "").strip().upper()!=str(ma_nv).strip().upper():
                    continue
                name=str(
                    p.get("TEN_NV")
                    or p.get("HO_TEN")
                    or p.get("TEN_NHAN_VIEN")
                    or ""
                ).strip()
                if name:
                    break
        except Exception:
            pass

    c.close()
    return name

def _ui09_debt_row(c,debt_id):
    r=c.execute("SELECT * FROM qtgh_employee_debt WHERE id=?",(debt_id,)).fetchone()
    if not r:
        raise ValueError("Không tìm thấy khoản nợ nhân viên")
    return dict(r)

def _ui09_balance_info(d):
    keys=set(d.keys())
    bal_col=next((k for k in ("balance","con_no","so_con_no","debt_balance","amount_due") if k in keys),None)
    if not bal_col:
        raise ValueError("Không xác định được cột số dư nợ NVGH")
    return bal_col,_num(d.get(bal_col))

def _ui09_reduce_debt(c,debt_id,amount):
    if amount<=0:
        return
    d=_ui09_debt_row(c,debt_id)
    bal_col,bal=_ui09_balance_info(d)
    newbal=max(0,bal-amount)

    sets=[f'"{bal_col}"=?']
    vals=[newbal]
    paid_col=next((k for k in ("collected_amount","paid_amount","da_thu","amount_paid") if k in d),None)
    if paid_col:
        sets.append(f'"{paid_col}"=?')
        vals.append(_num(d.get(paid_col))+min(amount,max(0,bal)))

    if "status" in d:
        sets.append('"status"=?')
        vals.append("DA_THU_DU" if newbal<=0.0001 else "CON_NO")

    vals.append(debt_id)
    c.execute(f'UPDATE qtgh_employee_debt SET {",".join(sets)} WHERE id=?',vals)

def _ui09_insert_cash_receipt(c,source_id,ma_nvgh,ten_nv,amount):
    cols=[r["name"] for r in c.execute("PRAGMA table_info(qtgh_cash_receipt)").fetchall()]
    now=_now()

    desc=f"Thu nợ nhân viên {ma_nvgh}"
    if ten_nv:
        desc+=f" - {ten_nv}"

    candidates={
      "receipt_kind":"THU_NO_NVGH",
      "source_type":"THU_NO_NVGH",
      "source_id":int(source_id),
      "session_key":f"THU_NO_NVGH:{source_id}",
      "ma_nvgh":ma_nvgh,
      "payer_code":ma_nvgh,
      "payer_name":ten_nv,
      "description":desc,
      "content":desc,
      "expected_amount":_money(amount),
      "actual_amount":0,
      "variance_amount":_money(amount),
      "status":"CHO_THU",
      "created_at":now,
      "updated_at":now
    }

    values={k:v for k,v in candidates.items() if k in cols}
    if not values:
        raise ValueError("Không xác định được cấu trúc phiếu thu")

    fields=",".join(f'"{k}"' for k in values)
    qs=",".join("?" for _ in values)
    cur=c.execute(
      f'INSERT INTO qtgh_cash_receipt({fields}) VALUES({qs})',
      list(values.values())
    )
    return cur.lastrowid

@router.get("/ui09/debts-stable")
def ui09_debts_stable():
    try:
        x=employee_debts()
        rows=x.get("rows",[]) if isinstance(x,dict) else []

        out=[]
        for r in rows:
            d=dict(r)
            ma=str(d.get("ma_nvgh") or d.get("ma_nv") or "").strip()
            d["ten_nv"]=_ui09_employee_name(ma)
            out.append(d)

        return {"rows":out}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui09/collect")
async def ui09_collect(request:Request):
    try:
        x=await request.json()
        debt_id=int(x.get("employee_debt_id"))
        cash=_num(x.get("cash_amount",0))
        bank=_num(x.get("bank_amount",0))
        bank_account=str(x.get("bank_account") or "").strip()
        bank_txn_id=str(x.get("bank_txn_id") or "").strip()

        if cash<0 or bank<0:
            raise ValueError("Số tiền không được âm")
        if cash+bank<=0:
            raise ValueError("Chưa nhập số tiền thu nợ")
        if bank>0 and (not bank_account or not bank_txn_id):
            raise ValueError("Có chuyển khoản: bắt buộc chọn tài khoản nhận và mã giao dịch")

        c=_conn()
        d=_ui09_debt_row(c,debt_id)
        bal_col,balance=_ui09_balance_info(d)

        if cash+bank>balance:
            c.close()
            raise ValueError("Tổng số tiền thu vượt số nợ nhân viên")

        ma_nvgh=str(d.get("ma_nvgh") or d.get("ma_nv") or "").strip()
        ten_nv=_ui09_employee_name(ma_nvgh)

        # Bank transfer is already in company bank => reduce debt immediately.
        if bank>0:
            _ui09_reduce_debt(c,debt_id,bank)

        cur=c.execute("""INSERT INTO qtgh_employee_collect_stable(
          employee_debt_id,ma_nvgh,ten_nv,
          cash_amount,bank_amount,bank_account,bank_txn_id,
          status,created_at
        ) VALUES(?,?,?,?,?,?,?,'OPEN',?)""",
        (debt_id,ma_nvgh,ten_nv,cash,bank,bank_account,bank_txn_id,_now()))
        collect_id=cur.lastrowid

        receipt_id=None
        if cash>0:
            receipt_id=_ui09_insert_cash_receipt(c,collect_id,ma_nvgh,ten_nv,cash)
            c.execute("""UPDATE qtgh_employee_collect_stable
              SET cash_receipt_id=?,status='WAIT_CASH'
              WHERE id=?""",(receipt_id,collect_id))
        else:
            c.execute("""UPDATE qtgh_employee_collect_stable
              SET status='DONE'
              WHERE id=?""",(collect_id,))

        c.commit()
        c.close()

        return {
          "ok":True,
          "collect_id":collect_id,
          "cash_receipt_id":receipt_id,
          "cash_amount":_money(cash),
          "bank_amount":_money(bank),
          "ma_nvgh":ma_nvgh,
          "ten_nv":ten_nv
        }
    except Exception as e:
        raise HTTPException(400,str(e))

def _ui09_apply_cash_receipt_if_needed(rid,actual_amount):
    c=_conn()
    row=c.execute("""SELECT *
      FROM qtgh_employee_collect_stable
      WHERE cash_receipt_id=?
      LIMIT 1""",(rid,)).fetchone()

    if not row:
        c.close()
        return

    d=dict(row)
    if d["status"]=="DONE":
        c.close()
        return

    actual=min(_num(actual_amount),_num(d["cash_amount"]))
    if actual>0:
        _ui09_reduce_debt(c,int(d["employee_debt_id"]),actual)

    # If Treasurer receives less than planned cash, only actual cash reduces debt.
    c.execute("""UPDATE qtgh_employee_collect_stable
      SET status='DONE'
      WHERE id=?""",(d["id"],))

    c.commit()
    c.close()

# Override stable UI05 receive with employee-debt integration.
@router.post("/ui05/cash/{rid}/receive-stable")
async def cash_receive_stable(rid:int,request:Request):
    try:
        x=await request.json()
        resolution=str(x.get("resolution_code") or "").strip().upper()
        if resolution=="XIN_Y_KIEN":
            resolution="XIN_Y_KIEN_LANH_DAO"

        result=qtgh_v085_receive_cash(
            rid,x.get("actual_amount",0),resolution,x.get("resolution_note","")
        )

        _ui09_apply_cash_receipt_if_needed(rid,x.get("actual_amount",0))
        return result
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== QTGH STABLE V1.1.9 UI11 CASH COUNT BY DATE =====

def _ui11_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS qtgh_cash_count_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      count_date TEXT NOT NULL UNIQUE,
      status TEXT NOT NULL DEFAULT 'DRAFT',
      received_total REAL NOT NULL DEFAULT 0,
      counted_total REAL NOT NULL DEFAULT 0,
      variance REAL NOT NULL DEFAULT 0,
      amount_words TEXT,
      loose_amount INTEGER NOT NULL DEFAULT 0,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL,
      locked_at TEXT
    );

    CREATE TABLE IF NOT EXISTS qtgh_cash_count_line_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      count_id INTEGER NOT NULL,
      denomination INTEGER NOT NULL,
      note_count INTEGER NOT NULL DEFAULT 0,
      amount REAL NOT NULL DEFAULT 0,
      UNIQUE(count_id,denomination),
      FOREIGN KEY(count_id) REFERENCES qtgh_cash_count_stable(id)
    );
    """)
    cols={r["name"] for r in c.execute("PRAGMA table_info(qtgh_cash_count_stable)")}
    if "loose_amount" not in cols:
        c.execute("ALTER TABLE qtgh_cash_count_stable ADD COLUMN loose_amount INTEGER NOT NULL DEFAULT 0")
    c.commit()
    c.close()

_ui11_init()

def _ui11_date(v):
    s=str(v or "").strip()
    return s[:10] if len(s)>=10 else ""

def _ui11_received_status(status):
    return str(status or "").strip().upper() in (
      "DA_THU_DU","THU_THIEU","DA_THU_THIEU"
    )

@router.get("/ui11/day-summary")
def ui11_day_summary(date:str):
    try:
        day=str(date or "").strip()
        if not day:
            raise ValueError("Thiếu ngày kiểm kê")

        c=_conn()
        rows=[dict(r) for r in c.execute(
          "SELECT * FROM qtgh_cash_receipt ORDER BY id DESC"
        ).fetchall()]
        c.close()

        created=[]
        received_today=[]
        for r in rows:
            created_date=_ui11_date(r.get("created_at") or r.get("ngay_phieu"))
            received_date=_ui11_date(r.get("received_at"))
            if not received_date and _ui11_received_status(r.get("status")):
                received_date=_ui11_date(r.get("updated_at"))

            r["created_date"]=created_date
            r["received_date"]=received_date
            r["is_received"]=_ui11_received_status(r.get("status"))

            if created_date==day:
                created.append(r)
            if r["is_received"] and received_date==day:
                received_today.append(r)

        received_total=sum(_num(r.get("actual_amount",0)) for r in received_today)

        return {
          "date":day,
          "receipts_created_count":len(created),
          "receipts_created":created,
          "receipts_received_count":len(received_today),
          "receipts_received":received_today,
          "received_total":_money(received_total)
        }
    except Exception as e:
        raise HTTPException(400,str(e))

def _ui11_get_count(day):
    c=_conn()
    h=c.execute(
      "SELECT * FROM qtgh_cash_count_stable WHERE count_date=?",(day,)
    ).fetchone()
    if not h:
        c.close()
        return None

    out=dict(h)
    out["lines"]=[dict(r) for r in c.execute("""
      SELECT * FROM qtgh_cash_count_line_stable
      WHERE count_id=?
      ORDER BY denomination DESC
    """,(h["id"],)).fetchall()]
    c.close()
    return out

@router.get("/ui11/count")
def ui11_count(date:str):
    try:
        day=str(date or "").strip()
        if not day:
            raise ValueError("Thiếu ngày")
        saved=_ui11_get_count(day)
        return {"count":saved}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui11/count/save")
async def ui11_count_save(request:Request):
    try:
        x=await request.json()
        day=str(x.get("date") or "").strip()
        words=str(x.get("amount_words") or "").strip()
        lines=x.get("lines",[])
        loose_amount=_num(x.get("loose_amount",0))
        if loose_amount<0 or loose_amount>=100:
            raise ValueError("Tiền lẻ khác phải từ 0 đến dưới 100 đồng")

        if not day:
            raise ValueError("Thiếu ngày kiểm kê")

        summary=ui11_day_summary(day)
        received=_num(summary["received_total"])

        clean=[]
        total=0.0
        for l in lines:
            denom=int(_num(l.get("denomination",0)))
            count=int(_num(l.get("note_count",0)))
            if denom<=0:
                continue
            if count<0:
                raise ValueError("Số tờ không được âm")
            amount=denom*count
            clean.append((denom,count,amount))
            total+=amount

        total+=loose_amount
        variance=total-received
        now=_now()

        c=_conn()
        old=c.execute(
          "SELECT * FROM qtgh_cash_count_stable WHERE count_date=?",(day,)
        ).fetchone()

        if old and old["status"]=="LOCKED":
            c.close()
            raise ValueError("Bảng kiểm kê đã chốt và đã khóa")

        if old:
            cid=old["id"]
            c.execute("""UPDATE qtgh_cash_count_stable
              SET received_total=?,counted_total=?,variance=?,
                  amount_words=?,loose_amount=?,updated_at=?
              WHERE id=?""",
              (received,total,variance,words,loose_amount,now,cid))
            c.execute(
              "DELETE FROM qtgh_cash_count_line_stable WHERE count_id=?",(cid,)
            )
        else:
            cur=c.execute("""INSERT INTO qtgh_cash_count_stable(
              count_date,status,received_total,counted_total,variance,
              amount_words,loose_amount,created_at,updated_at
            ) VALUES(?,'DRAFT',?,?,?,?,?,?,?)""",
            (day,received,total,variance,words,loose_amount,now,now))
            cid=cur.lastrowid

        for denom,count,amount in clean:
            c.execute("""INSERT INTO qtgh_cash_count_line_stable(
              count_id,denomination,note_count,amount
            ) VALUES(?,?,?,?)""",(cid,denom,count,amount))

        c.commit()
        c.close()
        return _ui11_get_count(day)
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui11/count/lock")
async def ui11_count_lock(request:Request):
    try:
        x=await request.json()
        day=str(x.get("date") or "").strip()

        if not day:
            raise ValueError("Thiếu ngày kiểm kê")

        c=_conn()
        h=c.execute(
          "SELECT * FROM qtgh_cash_count_stable WHERE count_date=?",(day,)
        ).fetchone()

        if not h:
            c.close()
            raise ValueError("Chưa lưu bảng kiểm kê")

        if h["status"]=="LOCKED":
            c.close()
            return _ui11_get_count(day)

        c.execute("""UPDATE qtgh_cash_count_stable
          SET status='LOCKED',locked_at=?,updated_at=?
          WHERE id=?""",(_now(),_now(),h["id"]))
        c.commit()
        c.close()

        # Keep legacy aggregate inventory in sync when possible.
        try:
            lock_cash_inventory(day,_money(h["counted_total"]))
        except Exception:
            pass

        return _ui11_get_count(day)
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== QTGH STABLE V1.1.11 UI10 CONTROLLED ADJUSTMENTS =====

def _ui10_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS qtgh_adjustment_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      adjustment_no TEXT UNIQUE NOT NULL,
      source_doc_type TEXT NOT NULL,
      source_doc_id INTEGER NOT NULL,
      source_ref TEXT,
      ma_nvgh TEXT,
      adjustment_type TEXT NOT NULL,
      delta_amount REAL NOT NULL DEFAULT 0,
      reason TEXT NOT NULL,
      note TEXT,
      status TEXT NOT NULL DEFAULT 'NHAP',
      before_value REAL,
      after_value REAL,
      effect_target TEXT,
      effect_applied INTEGER NOT NULL DEFAULT 0,
      created_at TEXT NOT NULL,
      submitted_at TEXT,
      approved_at TEXT,
      approved_by TEXT
    );

    CREATE INDEX IF NOT EXISTS ix_qtgh_adjustment_stable_status
      ON qtgh_adjustment_stable(status,source_doc_type,source_doc_id);
    """)
    c.commit()
    c.close()

_ui10_init()


def _ui10_warehouse_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS qtgh_wh_quantity_adjustment(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      adjustment_no TEXT UNIQUE NOT NULL,
      request_id INTEGER NOT NULL,
      line_id INTEGER NOT NULL,
      delta_qty REAL NOT NULL,
      before_qty REAL NOT NULL,
      after_qty REAL NOT NULL,
      reason TEXT NOT NULL,
      note TEXT,
      status TEXT NOT NULL DEFAULT 'NHAP',
      created_at TEXT NOT NULL,
      submitted_at TEXT,
      approved_at TEXT,
      approved_by TEXT,
      FOREIGN KEY(request_id) REFERENCES qtgh_wh_return_request(id),
      FOREIGN KEY(line_id) REFERENCES qtgh_wh_return_request_line(id)
    );
    CREATE INDEX IF NOT EXISTS ix_qtgh_wh_quantity_adjustment_line
      ON qtgh_wh_quantity_adjustment(line_id,status);
    """)
    c.close()


_ui10_warehouse_init()


def _ui10_effective_warehouse_qty(c,line_id,base):
    delta=c.execute("""SELECT COALESCE(SUM(delta_qty),0)
      FROM qtgh_wh_quantity_adjustment WHERE line_id=? AND status='DA_DUYET'""",
      (line_id,)).fetchone()[0]
    return round(_num(base)+_num(delta),6)


@router.get("/ui10/warehouse/{rid}")
def ui10_warehouse_detail(rid:int):
    c=_conn()
    try:
        h=c.execute("SELECT * FROM qtgh_wh_return_request WHERE id=?",(rid,)).fetchone()
        if not h:
            raise ValueError("Không tìm thấy phiếu nhận kho")
        lines=[]
        for row in c.execute("SELECT * FROM qtgh_wh_return_request_line WHERE request_id=? ORDER BY ma_sku,id",(rid,)):
            d=dict(row)
            d["effective_qty"]=_ui10_effective_warehouse_qty(c,d["id"],d["sl_thuc_nhan"])
            lines.append(d)
        return {"request":dict(h),"lines":lines}
    except ValueError as e:
        raise HTTPException(400,str(e))
    finally:
        c.close()


@router.get("/ui10/warehouse-adjustments")
def ui10_warehouse_adjustments():
    c=_conn()
    try:
        rows=[dict(r) for r in c.execute("""SELECT a.*,l.ma_sku,l.ten_sku,h.request_no
          FROM qtgh_wh_quantity_adjustment a
          JOIN qtgh_wh_return_request_line l ON l.id=a.line_id
          JOIN qtgh_wh_return_request h ON h.id=a.request_id
          ORDER BY a.id DESC""").fetchall()]
        return {"rows":rows}
    finally:
        c.close()


@router.post("/ui10/warehouse-adjustment")
async def ui10_create_warehouse_adjustment(request:Request):
    x=await request.json()
    c=_conn()
    try:
        rid=int(x.get("request_id"))
        lid=int(x.get("line_id"))
        delta=float(x.get("delta_qty"))
        reason=str(x.get("reason") or "").strip()
        note=str(x.get("note") or "").strip()
        if not reason or not delta or not (-1e9<delta<1e9):
            raise ValueError("Bắt buộc nhập lý do và số lượng điều chỉnh hợp lệ, khác 0")
        c.execute("BEGIN IMMEDIATE")
        h=c.execute("SELECT status FROM qtgh_wh_return_request WHERE id=?",(rid,)).fetchone()
        l=c.execute("SELECT * FROM qtgh_wh_return_request_line WHERE id=? AND request_id=?",(lid,rid)).fetchone()
        if not h or h["status"]!="DA_NHAN" or not l:
            raise ValueError("Chỉ điều chỉnh SKU thuộc phiếu nhận kho đã xác nhận")
        before=_ui10_effective_warehouse_qty(c,lid,l["sl_thuc_nhan"])
        after=round(before+delta,6)
        if after<0 or after>_num(l["sl_phai_tra"]):
            raise ValueError("Số thực nhận sau điều chỉnh phải từ 0 đến số phải trả")
        cur=c.execute("""INSERT INTO qtgh_wh_quantity_adjustment(
          adjustment_no,request_id,line_id,delta_qty,before_qty,after_qty,reason,note,created_at
        ) VALUES('',?,?,?,?,?,?,?,?)""",(rid,lid,delta,before,after,reason,note,_now()))
        aid=cur.lastrowid
        c.execute("UPDATE qtgh_wh_quantity_adjustment SET adjustment_no=? WHERE id=?",
                  (f"DCH-{datetime.datetime.now():%Y%m%d}-{aid:06d}",aid))
        c.commit()
        return {"id":aid,"adjustment_no":f"DCH-{datetime.datetime.now():%Y%m%d}-{aid:06d}"}
    except (ValueError,TypeError,OverflowError) as e:
        c.rollback()
        raise HTTPException(400,str(e))
    finally:
        c.close()


@router.post("/ui10/warehouse-adjustment/{aid}/submit")
def ui10_submit_warehouse_adjustment(aid:int):
    c=_conn()
    try:
        c.execute("BEGIN IMMEDIATE")
        cur=c.execute("""UPDATE qtgh_wh_quantity_adjustment
          SET status='CHO_DUYET',submitted_at=? WHERE id=? AND status='NHAP'""",(_now(),aid))
        if cur.rowcount!=1:
            raise ValueError("Chỉ phiếu nháp mới được gửi duyệt")
        c.commit()
        return {"ok":True}
    except ValueError as e:
        c.rollback();raise HTTPException(400,str(e))
    finally:
        c.close()


@router.post("/ui10/warehouse-adjustment/{aid}/approve")
async def ui10_approve_warehouse_adjustment(aid:int,request:Request):
    x=await request.json()
    from .access_control_v1 import auth_user
    user=auth_user(request.headers.get("authorization"))
    if str(user.get("role") or "").upper()!="SUPER_ADMIN":
        raise HTTPException(403,"Chỉ quản trị cao nhất được duyệt điều chỉnh hàng")
    approver=user["username"]
    c=_conn()
    try:
        c.execute("BEGIN IMMEDIATE")
        a=c.execute("SELECT * FROM qtgh_wh_quantity_adjustment WHERE id=?",(aid,)).fetchone()
        if not a or a["status"]!="CHO_DUYET":
            raise ValueError("Chỉ phiếu chờ duyệt mới được duyệt")
        h=c.execute("SELECT status FROM qtgh_wh_return_request WHERE id=?",(a["request_id"],)).fetchone()
        l=c.execute("SELECT * FROM qtgh_wh_return_request_line WHERE id=? AND request_id=?",
                    (a["line_id"],a["request_id"])).fetchone()
        if not h or h["status"]!="DA_NHAN" or not l:
            raise ValueError("Phiếu nhận kho nguồn không còn hợp lệ")
        current=_ui10_effective_warehouse_qty(c,a["line_id"],l["sl_thuc_nhan"])
        after=round(current+_num(a["delta_qty"]),6)
        if abs(current-_num(a["before_qty"]))>0.000001:
            raise ValueError("SKU đã được điều chỉnh sau khi lập phiếu; cần lập phiếu mới")
        if after<0 or after>_num(l["sl_phai_tra"]):
            raise ValueError("Số lượng sau điều chỉnh ngoài phạm vi được phép")
        c.execute("""UPDATE qtgh_wh_quantity_adjustment SET
          status='DA_DUYET',approved_at=?,approved_by=? WHERE id=?""",(_now(),approver,aid))
        c.commit()
        return {"ok":True,"effective_qty":after}
    except ValueError as e:
        c.rollback();raise HTTPException(400,str(e))
    finally:
        c.close()

def _ui10_ref(prefix, idv):
    return f"{prefix}#{idv}"

def _ui10_doc_rows():
    c=_conn()
    out=[]

    # 1. Customer debt
    try:
        rows=c.execute("SELECT * FROM qtgh_customer_debt ORDER BY id DESC").fetchall()
        for rr in rows:
            d=dict(rr)
            bal=0
            for k in ("balance","con_no","so_con_no","debt_balance","amount_due"):
                if k in d and d.get(k) is not None:
                    bal=_num(d.get(k)); break
            out.append({
              "source_doc_type":"CUSTOMER_DEBT",
              "source_doc_id":d["id"],
              "source_ref":d.get("ma_don_nguon") or _ui10_ref("CNKH",d["id"]),
              "ma_nvgh":d.get("ma_nvgh") or "",
              "description":"Công nợ khách hàng",
              "current_value":bal,
              "status":d.get("status") or ""
            })
    except Exception:
        pass

    # 2. Employee debt
    try:
        rows=c.execute("SELECT * FROM qtgh_employee_debt ORDER BY id DESC").fetchall()
        for rr in rows:
            d=dict(rr)
            bal=0
            for k in ("balance","con_no","so_con_no","debt_balance","amount_due"):
                if k in d and d.get(k) is not None:
                    bal=_num(d.get(k)); break
            out.append({
              "source_doc_type":"EMPLOYEE_DEBT",
              "source_doc_id":d["id"],
              "source_ref":d.get("ma_nvgh") or d.get("ma_nv") or _ui10_ref("CNNV",d["id"]),
              "ma_nvgh":d.get("ma_nvgh") or d.get("ma_nv") or "",
              "description":"Nợ nhân viên với Công ty",
              "current_value":bal,
              "status":d.get("status") or ""
            })
    except Exception:
        pass

    # 3. Cash receipts
    try:
        rows=c.execute("SELECT * FROM qtgh_cash_receipt ORDER BY id DESC").fetchall()
        for rr in rows:
            d=dict(rr)
            val=_num(d.get("actual_amount") if d.get("actual_amount") is not None else d.get("expected_amount"))
            out.append({
              "source_doc_type":"CASH_RECEIPT",
              "source_doc_id":d["id"],
              "source_ref":_ui10_ref("PT",d["id"]),
              "ma_nvgh":d.get("ma_nvgh") or "",
              "description":d.get("receipt_kind") or d.get("source_type") or "Phiếu thu",
              "current_value":val,
              "status":d.get("status") or ""
            })
    except Exception:
        pass

    # 4. Warehouse return requests
    try:
        rows=c.execute("SELECT * FROM qtgh_wh_return_request ORDER BY id DESC").fetchall()
        for rr in rows:
            d=dict(rr)
            val=c.execute(
              "SELECT COALESCE(SUM(sl_thuc_nhan),0) FROM qtgh_wh_return_request_line WHERE request_id=?",
              (d["id"],)
            ).fetchone()[0]
            val+=_num(c.execute("""SELECT COALESCE(SUM(delta_qty),0)
              FROM qtgh_wh_quantity_adjustment WHERE request_id=? AND status='DA_DUYET'""",
              (d["id"],)).fetchone()[0])
            out.append({
              "source_doc_type":"WAREHOUSE_RETURN",
              "source_doc_id":d["id"],
              "source_ref":d.get("request_no") or _ui10_ref("YCTK",d["id"]),
              "ma_nvgh":d.get("ma_nvgh") or "",
              "description":"Phiếu yêu cầu trả kho",
              "current_value":_num(val),
              "status":d.get("status") or ""
            })
    except Exception:
        pass

    # 5. Settled sessions
    try:
        rows=c.execute("SELECT * FROM qtgh_session ORDER BY id DESC").fetchall()
        for rr in rows:
            d=dict(rr)
            if str(d.get("status") or "")!="DA_KHOA":
                continue
            out.append({
              "source_doc_type":"SETTLEMENT_SESSION",
              "source_doc_id":d["id"],
              "source_ref":d.get("session_no") or d.get("session_key") or _ui10_ref("QT",d["id"]),
              "ma_nvgh":d.get("ma_nvgh") or "",
              "description":"Phiên quyết toán đã khóa",
              "current_value":_num(d.get("so_phai_thu") or d.get("expected_cash") or 0),
              "status":d.get("status") or ""
            })
    except Exception:
        pass

    c.close()
    return out

@router.get("/ui10/documents")
def ui10_documents(doc_type:str=None,ma_nvgh:str=None,q:str=None):
    try:
        rows=_ui10_doc_rows()
        if doc_type:
            rows=[r for r in rows if r["source_doc_type"]==doc_type]
        if ma_nvgh:
            rows=[r for r in rows if str(r.get("ma_nvgh") or "")==str(ma_nvgh)]
        if q:
            qq=str(q).upper()
            rows=[r for r in rows if qq in (
              str(r.get("source_ref") or "")+" "+
              str(r.get("description") or "")+" "+
              str(r.get("ma_nvgh") or "")
            ).upper()]
        return {"rows":rows}
    except Exception as e:
        raise HTTPException(400,str(e))

def _ui10_current_value(c,doc_type,doc_id):
    table={
      "CUSTOMER_DEBT":"qtgh_customer_debt",
      "EMPLOYEE_DEBT":"qtgh_employee_debt"
    }.get(doc_type)

    if not table:
        return None,None,None

    r=c.execute(f'SELECT * FROM "{table}" WHERE id=?',(doc_id,)).fetchone()
    if not r:
        raise ValueError("Không tìm thấy chứng từ nguồn")

    d=dict(r)
    bal_col=next((k for k in ("balance","con_no","so_con_no","debt_balance","amount_due") if k in d),None)
    if not bal_col:
        raise ValueError("Không xác định được cột giá trị cần điều chỉnh")

    return table,bal_col,_num(d.get(bal_col))

@router.get("/ui10/adjustments-stable")
def ui10_adjustments_stable(status:str=None):
    try:
        c=_conn()
        sql="SELECT * FROM qtgh_adjustment_stable"
        p=[]
        if status:
            sql+=" WHERE status=?"
            p.append(status)
        sql+=" ORDER BY id DESC"
        rows=[dict(r) for r in c.execute(sql,p).fetchall()]
        c.close()
        return {"rows":rows}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui10/adjustment-stable")
async def ui10_create_stable(request:Request):
    try:
        x=await request.json()

        doc_type=str(x.get("source_doc_type") or "").strip()
        doc_id=int(x.get("source_doc_id"))
        source_ref=str(x.get("source_ref") or "").strip()
        ma_nvgh=str(x.get("ma_nvgh") or "").strip()
        adj_type=str(x.get("adjustment_type") or "").strip()
        delta=_num(x.get("delta_amount",0))
        reason=str(x.get("reason") or "").strip()
        note=str(x.get("note") or "").strip()

        raise ValueError("Luồng điều chỉnh tiền đang được xây dựng; chỉ dùng luồng điều chỉnh hàng theo SKU trong bản UAT này")
        if not doc_type or not adj_type:
            raise ValueError("Thiếu loại chứng từ hoặc loại điều chỉnh")
        if not reason:
            raise ValueError("Bắt buộc nhập lý do điều chỉnh")
        if delta==0:
            raise ValueError("Giá trị điều chỉnh phải khác 0")

        c=_conn()
        table,col,before=_ui10_current_value(c,doc_type,doc_id)

        if before is None:
            before=0
            effect_target="AUDIT_ONLY"
            after=delta
        else:
            effect_target=f"{table}.{col}"
            after=before+delta
            if after<0:
                c.close()
                raise ValueError("Giá trị sau điều chỉnh không được âm")

        seq=c.execute("SELECT COALESCE(MAX(id),0)+1 FROM qtgh_adjustment_stable").fetchone()[0]
        no=f"DC-{datetime.datetime.now():%Y%m%d}-{seq:04d}"

        cur=c.execute("""INSERT INTO qtgh_adjustment_stable(
          adjustment_no,source_doc_type,source_doc_id,source_ref,ma_nvgh,
          adjustment_type,delta_amount,reason,note,status,
          before_value,after_value,effect_target,effect_applied,created_at
        ) VALUES(?,?,?,?,?,?,?,?,?,'NHAP',?,?,?,?,?)""",
        (no,doc_type,doc_id,source_ref,ma_nvgh,
         adj_type,delta,reason,note,before,after,effect_target,0,_now()))

        aid=cur.lastrowid
        c.commit()
        out=dict(c.execute("SELECT * FROM qtgh_adjustment_stable WHERE id=?",(aid,)).fetchone())
        c.close()
        return out
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui10/adjustment/{aid}/submit")
def ui10_submit(aid:int):
    raise HTTPException(409,"Phiếu cũ chỉ để đối chiếu; không gửi duyệt qua luồng này")
    try:
        c=_conn()
        r=c.execute("SELECT * FROM qtgh_adjustment_stable WHERE id=?",(aid,)).fetchone()
        if not r:
            c.close(); raise ValueError("Không tìm thấy phiếu điều chỉnh")
        if r["status"]!="NHAP":
            c.close(); raise ValueError("Chỉ phiếu NHAP mới được gửi duyệt")

        c.execute("""UPDATE qtgh_adjustment_stable
          SET status='CHO_DUYET',submitted_at=?
          WHERE id=?""",(_now(),aid))
        c.commit()
        out=dict(c.execute("SELECT * FROM qtgh_adjustment_stable WHERE id=?",(aid,)).fetchone())
        c.close()
        return out
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/ui10/adjustment/{aid}/approve-stable")
async def ui10_approve_stable(aid:int,request:Request):
    raise HTTPException(409,"Phiếu cũ chỉ để đối chiếu; không duyệt qua luồng này")
    try:
        x=await request.json()
        approver=str(x.get("approved_by") or "ADMIN").strip() or "ADMIN"

        c=_conn()
        r=c.execute("SELECT * FROM qtgh_adjustment_stable WHERE id=?",(aid,)).fetchone()
        if not r:
            c.close(); raise ValueError("Không tìm thấy phiếu điều chỉnh")
        if r["status"]!="CHO_DUYET":
            c.close(); raise ValueError("Chỉ phiếu CHO_DUYET mới được duyệt")

        d=dict(r)
        effect_applied=0

        table,col,current=_ui10_current_value(c,d["source_doc_type"],d["source_doc_id"])

        if table and col:
            expected_before=_num(d["before_value"])
            if abs(current-expected_before)>0.0001:
                c.close()
                raise ValueError(
                  f"Giá trị chứng từ nguồn đã thay đổi từ {expected_before} thành {current}. "
                  "Cần lập lại phiếu điều chỉnh."
                )

            newval=current+_num(d["delta_amount"])
            if newval<0:
                c.close(); raise ValueError("Giá trị sau điều chỉnh không được âm")

            # Update value and status when supported.
            cols=[r["name"] for r in c.execute(f'PRAGMA table_info("{table}")').fetchall()]
            sets=[f'"{col}"=?']
            vals=[newval]

            if "status" in cols:
                sets.append('"status"=?')
                vals.append("DA_THU_DU" if newval<=0.0001 else "CON_NO")

            vals.append(d["source_doc_id"])
            c.execute(f'UPDATE "{table}" SET {",".join(sets)} WHERE id=?',vals)
            effect_applied=1

        c.execute("""UPDATE qtgh_adjustment_stable
          SET status='DA_DUYET',approved_at=?,approved_by=?,
              effect_applied=?
          WHERE id=?""",
          (_now(),approver,effect_applied,aid))

        c.commit()
        out=dict(c.execute("SELECT * FROM qtgh_adjustment_stable WHERE id=?",(aid,)).fetchone())
        c.close()
        return out
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.0 =====
import json as _json_kq
from collections import defaultdict as _dd_kq

def _kq_json(s):
    try:
        x=_json_kq.loads(s or "{}")
        return x if isinstance(x,dict) else {}
    except Exception:
        return {}

def _kq_month(v):
    s=str(v or "").strip()
    return s[:7] if len(s)>=7 else ""

def _kq_num(v):
    try:return float(v or 0)
    except:return 0.0

def _kq_master_maps():
    c=_conn()
    staff={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm02-staff'"):
        p=_kq_json(r["payload_json"]); ma=str(p.get("MA_NV") or "").strip()
        if ma: staff[ma]=p

    staff_map={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm02a-staff-map'"):
        p=_kq_json(r["payload_json"])
        nv=str(p.get("MA_NV") or "").strip()
        code=str(p.get("MA_NGHIEP_VU") or "").strip()
        if nv and code:
            staff_map[code]=nv

    products={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm01-product'"):
        p=_kq_json(r["payload_json"]); sku=str(p.get("MA_SKU") or "").strip()
        if sku: products[sku]=p

    customers=[]
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm03-customer'"):
        customers.append(_kq_json(r["payload_json"]))
    c.close()
    return staff,staff_map,products,customers

def _kq_active_batch_rows():
    c=_conn()
    b=c.execute("""
      SELECT batch_id FROM xb_dms_mbw_batch
      WHERE UPPER(COALESCE(status,''))='COMMITTED'
        AND superseded_by_batch_id IS NULL
      ORDER BY rowid DESC LIMIT 1
    """).fetchone()
    if not b:
        c.close(); return []
    rows=[_kq_json(r["payload_json"]) for r in c.execute("""
      SELECT payload_json FROM xb_dms_mbw_row
      WHERE batch_id=? ORDER BY row_no
    """,(b["batch_id"],))]
    c.close()
    return rows

def _kq_settled_actual():
    c=_conn()
    rows=c.execute("""
      SELECT d.ma_don_nguon,l.ma_sku,
             SUM(COALESCE(l.sl_hang_ban_dat,0)) sl_dat,
             SUM(COALESCE(l.sl_hang_ban_thuc_giao,0)) sl_thuc
      FROM qtgh_draft_order d
      JOIN qtgh_draft_line l ON l.draft_order_id=d.id
      WHERE EXISTS(
        SELECT 1 FROM qtgh_session_order so
        JOIN qtgh_session se ON se.id=so.session_id
        WHERE so.draft_order_id=d.id AND se.status='DA_KHOA'
      )
      GROUP BY d.ma_don_nguon,l.ma_sku
    """).fetchall()
    c.close()

    actual={}
    totals=_dd_kq(lambda:[0.0,0.0])
    for r in rows:
        order=str(r["ma_don_nguon"] or "").strip()
        sku=str(r["ma_sku"] or "").strip()
        dat=_kq_num(r["sl_dat"]); thuc=_kq_num(r["sl_thuc"])
        actual[(order,sku)]=thuc
        totals[order][0]+=dat; totals[order][1]+=thuc

    classes={}; success=set()
    for order,(dat,thuc) in totals.items():
        if thuc<=0: classes[order]="BROKEN_100"
        elif thuc<dat:
            classes[order]="PARTIAL"; success.add(order)
        else:
            classes[order]="FULL"; success.add(order)
    return actual,success,classes

def _kq_customer_key(p):
    for k in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
        v=str(p.get(k) or "").strip()
        if v:return v
    return ""

@router.get("/kq-nvkd/meta")
def kq_nvkd_meta():
    try:
        staff,staff_map,products,customers=_kq_master_maps()
        emp=[]
        for ma,p in staff.items():
            if str(p.get("NHOM_NV") or "").upper()=="KINH_DOANH":
                emp.append({
                    "ma_nv":ma,
                    "ten_nv":p.get("TEN_NV") or "",
                    "ma_vung_kd":p.get("MA_VUNG_KD") or "",
                    "bo_phan":p.get("BO_PHAN") or ""
                })
        emp.sort(key=lambda x:x["ma_nv"])
        bus=sorted({str(p.get("BU") or "").strip() for p in products.values() if str(p.get("BU") or "").strip()})
        return {"employees":emp,"bus":bus}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/kq-nvkd/report")
def kq_nvkd_report(month:str,ma_nv:str=None,ma_vung_kd:str=None,bu:str=None):
    try:
        month=_kq_month(month)
        if not month: raise ValueError("Thiếu tháng báo cáo")

        staff,staff_map,products,customers=_kq_master_maps()
        src=_kq_active_batch_rows()
        actual,success_orders,classes=_kq_settled_actual()

        assigned=_dd_kq(set)
        aliases={}
        for kh in customers:
            canonical=_kq_customer_key(kh)
            if canonical:
                for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
                    cv=str(kh.get(ck) or "").strip()
                    if cv: aliases[cv]=canonical
            for fld in ("MA_NVBH_01","MA_NVBH_02"):
                code=str(kh.get(fld) or "").strip()
                if code:
                    owner=staff_map.get(code,code if code in staff else "")
                    if owner and canonical: assigned[owner].add(canonical)

        by_emp=_dd_kq(lambda:{
            "order_sales":0.0,"delivered_sales":0.0,
            "order_set":set(),"success_set":set(),
            "full_set":set(),"partial_set":set(),"broken_set":set(),
            "eco_customers":set(),
            "by_bu":_dd_kq(lambda:{
                "order_sales":0.0,"delivered_sales":0.0,
                "order_set":set(),"success_set":set()
            })
        })

        for p in src:
            day=str(p.get("NGAY_DON") or p.get("NGAY") or p.get("NGAY_GIAO") or "")
            if _kq_month(day)!=month: continue

            src_emp=str(p.get("MA_NV_KD_NGHIEP_VU") or p.get("MA_NV_KD") or "").strip()
            emp=staff_map.get(src_emp,src_emp if src_emp in staff else "")
            if not emp: continue
            sp=staff.get(emp,{})
            if ma_nv and emp!=ma_nv: continue
            if ma_vung_kd and str(sp.get("MA_VUNG_KD") or "")!=ma_vung_kd: continue

            order=str(p.get("SO_HOA_DON") or p.get("MA_DON_DMS") or p.get("MA_DON") or "").strip()
            sku=str(p.get("MA_SKU") or "").strip()
            prod=products.get(sku,{})
            row_bu=str(prod.get("BU") or p.get("BU") or "").strip() or "KHAC"
            if bu and row_bu!=bu: continue

            pack=_kq_num(prod.get("SO_SU_THUNG") or 0)
            qty_order=_kq_num(p.get("SL_SU"))+_kq_num(p.get("SL_THUNG"))*pack
            price=_kq_num(p.get("DON_GIA_CHUA_VAT"))
            qty_actual=_kq_num(actual.get((order,sku),0))

            a=by_emp[emp]
            a["order_sales"]+=qty_order*price
            a["delivered_sales"]+=qty_actual*price
            if order: a["order_set"].add(order)

            b=a["by_bu"][row_bu]
            b["order_sales"]+=qty_order*price
            b["delivered_sales"]+=qty_actual*price
            if order: b["order_set"].add(order)

            if order in success_orders:
                a["success_set"].add(order); b["success_set"].add(order)
                if classes.get(order)=="FULL": a["full_set"].add(order)
                elif classes.get(order)=="PARTIAL": a["partial_set"].add(order)

                for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
                    cv=str(p.get(ck) or "").strip()
                    if cv:
                        a["eco_customers"].add(aliases.get(cv,cv))
                        break
            elif order and order in classes:
                a["broken_set"].add(order)

        rows=[]
        totals={"order_sales":0.0,"delivered_sales":0.0,
                "orders":set(),"success":set(),"full":set(),"partial":set(),"broken":set()}

        for emp,a in by_emp.items():
            sp=staff.get(emp,{})
            eco_den=assigned.get(emp,set())
            eco_num=a["eco_customers"] & eco_den if eco_den else set()

            bu_rows=[]
            for bname,x in sorted(a["by_bu"].items()):
                bu_rows.append({
                    "bu":bname,
                    "doanh_so_dat":round(x["order_sales"],2),
                    "doanh_so_giao":round(x["delivered_sales"],2),
                    "so_don_dat":len(x["order_set"]),
                    "so_don_giao_tc":len(x["success_set"])
                })

            rows.append({
                "ma_nv":emp,
                "ten_nv":sp.get("TEN_NV") or "",
                "ma_vung_kd":sp.get("MA_VUNG_KD") or "",
                "bo_phan":sp.get("BO_PHAN") or "",
                "doanh_so_dat":round(a["order_sales"],2),
                "doanh_so_giao":round(a["delivered_sales"],2),
                "so_don_dat":len(a["order_set"]),
                "so_don_giao_tc":len(a["success_set"]),
                "don_giao_du":len(a["full_set"]),
                "don_giao_mot_phan":len(a["partial_set"]),
                "don_vo_100":len(a["broken_set"]),
                "kh_co_don_giao_tc":len(eco_num),
                "tong_kh_phu_trach":len(eco_den),
                "eco":round((len(eco_num)/len(eco_den)*100) if eco_den else 0,2),
                "by_bu":bu_rows
            })

            totals["order_sales"]+=a["order_sales"]; totals["delivered_sales"]+=a["delivered_sales"]
            totals["orders"]|=a["order_set"]; totals["success"]|=a["success_set"]
            totals["full"]|=a["full_set"]; totals["partial"]|=a["partial_set"]; totals["broken"]|=a["broken_set"]

        rows.sort(key=lambda x:x["ma_nv"])
        return {
            "month":month,"rows":rows,
            "summary":{
                "doanh_so_dat":round(totals["order_sales"],2),
                "doanh_so_giao":round(totals["delivered_sales"],2),
                "so_don_dat":len(totals["orders"]),
                "so_don_giao_tc":len(totals["success"]),
                "don_giao_du":len(totals["full"]),
                "don_giao_mot_phan":len(totals["partial"]),
                "don_vo_100":len(totals["broken"])
            }
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.2 CLEAN UI SUPPORT =====

@router.get("/kq-nvkd/periods-v102")
def kq_nvkd_periods_v102():
    try:
        rows=_kq_active_batch_rows()
        months=set()
        source_rows=0
        orders=set()
        for p in rows:
            day=str(p.get("NGAY_DON") or p.get("NGAY") or p.get("NGAY_GIAO") or "")
            m=_kq_month(day)
            if m:
                months.add(m)
            source_rows+=1
            order=str(p.get("SO_HOA_DON") or p.get("MA_DON_DMS") or p.get("MA_DON") or "").strip()
            if order:
                orders.add(order)

        actual,success,classes=_kq_settled_actual()
        locked_orders=set(o for o,_ in actual.keys())

        return {
          "months":sorted(months,reverse=True),
          "source_rows":source_rows,
          "source_orders":len(orders),
          "locked_orders":len(locked_orders),
          "successful_orders_locked":len(success),
          "broken_100_locked":sum(1 for v in classes.values() if v=="BROKEN_100")
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.3 ECO DETAIL =====

@router.get("/kq-nvkd/eco-detail-v103")
def kq_nvkd_eco_detail_v103(month:str,ma_nv:str):
    try:
        month=_kq_month(month)
        ma_nv=str(ma_nv or "").strip()
        if not month or not ma_nv:
            raise ValueError("Thiếu tháng hoặc MA_NV")

        staff,staff_map,products,customers=_kq_master_maps()
        src=_kq_active_batch_rows()
        actual,success_orders,classes=_kq_settled_actual()

        # Customer aliases + assigned customers for employee
        assigned={}
        alias_to_canonical={}
        for kh in customers:
            canonical=_kq_customer_key(kh)
            if not canonical:
                continue

            for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
                cv=str(kh.get(ck) or "").strip()
                if cv:
                    alias_to_canonical[cv]=canonical

            owner_codes=[]
            for fld in ("MA_NVBH_01","MA_NVBH_02"):
                code=str(kh.get(fld) or "").strip()
                if code:
                    owner_codes.append(code)

            owners=set()
            for code in owner_codes:
                owner=staff_map.get(code,code if code in staff else "")
                if owner:
                    owners.add(owner)

            if ma_nv in owners:
                assigned[canonical]={
                  "ma_kh_chuan":canonical,
                  "ma_kh_dms":kh.get("MA_KH_DMS") or "",
                  "ma_kh_mbw":kh.get("MA_KH_MBW") or "",
                  "ten_kh":kh.get("TEN_KH") or "",
                  "ma_dia_ban":kh.get("MA_DIA_BAN") or "",
                  "trang_thai":kh.get("TRANG_THAI") or "",
                  "co_giao_tc":False,
                  "so_don_giao_tc":0,
                  "doanh_so_giao":0.0
                }

        # Successful delivered orders of employee in selected month.
        order_seen=set()
        for p in src:
            day=str(p.get("NGAY_DON") or p.get("NGAY") or p.get("NGAY_GIAO") or "")
            if _kq_month(day)!=month:
                continue

            src_emp=str(p.get("MA_NV_KD_NGHIEP_VU") or p.get("MA_NV_KD") or "").strip()
            emp=staff_map.get(src_emp,src_emp if src_emp in staff else "")
            if emp!=ma_nv:
                continue

            order=str(p.get("SO_HOA_DON") or p.get("MA_DON_DMS") or p.get("MA_DON") or "").strip()
            if order not in success_orders:
                continue

            customer=""
            for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
                cv=str(p.get(ck) or "").strip()
                if cv:
                    customer=alias_to_canonical.get(cv,cv)
                    break

            if customer not in assigned:
                continue

            sku=str(p.get("MA_SKU") or "").strip()
            qty_actual=_kq_num(actual.get((order,sku),0))
            price=_kq_num(p.get("DON_GIA_CHUA_VAT"))
            assigned[customer]["doanh_so_giao"] += qty_actual*price

            key=(customer,order)
            if key not in order_seen:
                order_seen.add(key)
                assigned[customer]["so_don_giao_tc"]+=1
                assigned[customer]["co_giao_tc"]=True

        rows=list(assigned.values())
        rows.sort(key=lambda x:(not x["co_giao_tc"],x["ma_kh_chuan"]))

        yes=sum(1 for x in rows if x["co_giao_tc"])
        total=len(rows)

        return {
          "month":month,
          "ma_nv":ma_nv,
          "ten_nv":staff.get(ma_nv,{}).get("TEN_NV") or "",
          "tong_kh_phu_trach":total,
          "kh_co_giao_tc":yes,
          "kh_chua_giao_tc":total-yes,
          "eco":round((yes/total*100) if total else 0,2),
          "rows":rows
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.4 ECO OWNER-MAP FIX =====

from collections import Counter as _Counter_kq104, defaultdict as _DefaultDict_kq104

def _kq104_owner_map(staff,staff_map,customers,src):
    """
    Infer DM03 MA_NVBH_01/02 -> canonical MA_NV from real transaction relations:
      source MA_KH_NGUON -> DM03 customer -> MA_NVBH_01
      source MA_NV_KD_NGHIEP_VU -> DM02A -> MA_NV

    This is intentionally separate from DM02A because MA_NVBH_01 is a different code domain.
    """
    # Alias customer code -> DM03 row
    cust_alias={}
    for kh in customers:
        for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
            v=str(kh.get(ck) or "").strip()
            if v:
                cust_alias[v]=kh

    votes=_DefaultDict_kq104(_Counter_kq104)

    for p in src:
        emp_code=str(
            p.get("MA_NV_KD_NGHIEP_VU")
            or p.get("MA_NV_KD")
            or ""
        ).strip()
        emp=staff_map.get(emp_code,emp_code if emp_code in staff else "")
        if not emp:
            continue

        customer_values=[]
        for ck in ("MA_KH_NGUON","MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
            v=str(p.get(ck) or "").strip()
            if v and v not in customer_values:
                customer_values.append(v)

        kh=None
        for v in customer_values:
            if v in cust_alias:
                kh=cust_alias[v]
                break
        if not kh:
            continue

        for fld in ("MA_NVBH_01","MA_NVBH_02"):
            owner_code=str(kh.get(fld) or "").strip()
            if owner_code:
                votes[owner_code][emp]+=1

    inferred={}
    diagnostics=[]
    for owner_code,counter in votes.items():
        ranked=counter.most_common()
        if not ranked:
            continue
        best_emp,best_n=ranked[0]
        total=sum(counter.values())
        second_n=ranked[1][1] if len(ranked)>1 else 0
        confidence=(best_n/total) if total else 0

        # Accept if clearly dominant. Most owner codes should be 100% consistent.
        if best_n>second_n and confidence>=0.60:
            inferred[owner_code]=best_emp

        diagnostics.append({
            "owner_code":owner_code,
            "ma_nv":best_emp,
            "votes":best_n,
            "total_votes":total,
            "confidence":round(confidence*100,2),
            "accepted":owner_code in inferred,
            "alternatives":[{"ma_nv":e,"votes":n} for e,n in ranked[1:4]]
        })

    return inferred,diagnostics

def _kq104_assigned_customers(staff,staff_map,customers,src):
    owner_map,diag=_kq104_owner_map(staff,staff_map,customers,src)
    assigned=_DefaultDict_kq104(set)

    for kh in customers:
        canonical=_kq_customer_key(kh)
        if not canonical:
            continue

        for fld in ("MA_NVBH_01","MA_NVBH_02"):
            owner_code=str(kh.get(fld) or "").strip()
            if not owner_code:
                continue

            # IMPORTANT:
            # 1) first use inferred customer-owner map;
            # 2) only fallback to DM02A if it really contains that code.
            owner=owner_map.get(owner_code) or staff_map.get(owner_code,"")
            if owner:
                assigned[owner].add(canonical)

    return assigned,owner_map,diag

@router.get("/kq-nvkd/eco-owner-map-v104")
def kq_nvkd_eco_owner_map_v104():
    try:
        staff,staff_map,products,customers=_kq_master_maps()
        src=_kq_active_batch_rows()
        assigned,owner_map,diag=_kq104_assigned_customers(
            staff,staff_map,customers,src
        )
        return {
            "mapped_owner_codes":len(owner_map),
            "employee_customer_counts":{
                k:len(v) for k,v in sorted(assigned.items())
            },
            "owner_map":owner_map,
            "diagnostics":diag
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/kq-nvkd/report-v104")
def kq_nvkd_report_v104(month:str,ma_nv:str=None,ma_vung_kd:str=None,bu:str=None):
    try:
        month=_kq_month(month)
        if not month:
            raise ValueError("Thiếu tháng báo cáo")

        staff,staff_map,products,customers=_kq_master_maps()
        src=_kq_active_batch_rows()
        actual,success_orders,classes=_kq_settled_actual()

        assigned,owner_map,owner_diag=_kq104_assigned_customers(
            staff,staff_map,customers,src
        )

        aliases={}
        for kh in customers:
            canonical=_kq_customer_key(kh)
            if not canonical:
                continue
            for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
                cv=str(kh.get(ck) or "").strip()
                if cv:
                    aliases[cv]=canonical

        by_emp=_dd_kq(lambda:{
            "order_sales":0.0,
            "delivered_sales":0.0,
            "order_set":set(),
            "success_set":set(),
            "full_set":set(),
            "partial_set":set(),
            "broken_set":set(),
            "eco_customers":set(),
            "by_bu":_dd_kq(lambda:{
                "order_sales":0.0,
                "delivered_sales":0.0,
                "order_set":set(),
                "success_set":set()
            })
        })

        for p in src:
            day=str(
                p.get("NGAY_DON")
                or p.get("NGAY")
                or p.get("NGAY_GIAO")
                or ""
            )
            if _kq_month(day)!=month:
                continue

            src_emp=str(
                p.get("MA_NV_KD_NGHIEP_VU")
                or p.get("MA_NV_KD")
                or ""
            ).strip()
            emp=staff_map.get(src_emp,src_emp if src_emp in staff else "")
            if not emp:
                continue

            sp=staff.get(emp,{})
            if ma_nv and emp!=ma_nv:
                continue
            if ma_vung_kd and str(sp.get("MA_VUNG_KD") or "")!=ma_vung_kd:
                continue

            order=str(
                p.get("SO_HOA_DON")
                or p.get("MA_DON_DMS")
                or p.get("MA_DON")
                or ""
            ).strip()
            sku=str(p.get("MA_SKU") or "").strip()

            prod=products.get(sku,{})
            row_bu=str(prod.get("BU") or p.get("BU") or "").strip() or "KHAC"
            if bu and row_bu!=bu:
                continue

            pack=_kq_num(prod.get("SO_SU_THUNG") or 0)
            qty_order=_kq_num(p.get("SL_SU")) + _kq_num(p.get("SL_THUNG"))*pack
            price=_kq_num(p.get("DON_GIA_CHUA_VAT"))
            qty_actual=_kq_num(actual.get((order,sku),0))

            a=by_emp[emp]
            a["order_sales"] += qty_order*price
            a["delivered_sales"] += qty_actual*price
            if order:
                a["order_set"].add(order)

            b=a["by_bu"][row_bu]
            b["order_sales"] += qty_order*price
            b["delivered_sales"] += qty_actual*price
            if order:
                b["order_set"].add(order)

            if order in success_orders:
                a["success_set"].add(order)
                b["success_set"].add(order)

                if classes.get(order)=="FULL":
                    a["full_set"].add(order)
                elif classes.get(order)=="PARTIAL":
                    a["partial_set"].add(order)

                # Source uses MA_KH_NGUON in the active batch.
                customer_raw=""
                for ck in (
                    "MA_KH_NGUON","MA_KH_CHUAN",
                    "MA_KH_DMS","MA_KH_MBW","MA_KH"
                ):
                    cv=str(p.get(ck) or "").strip()
                    if cv:
                        customer_raw=cv
                        break

                if customer_raw:
                    a["eco_customers"].add(
                        aliases.get(customer_raw,customer_raw)
                    )

            elif order and order in classes:
                a["broken_set"].add(order)

        rows=[]
        totals={
            "order_sales":0.0,"delivered_sales":0.0,
            "orders":set(),"success":set(),
            "full":set(),"partial":set(),"broken":set()
        }

        for emp,a in by_emp.items():
            sp=staff.get(emp,{})
            eco_den=assigned.get(emp,set())

            # Numerator = customers assigned to employee AND delivered successfully.
            eco_num=a["eco_customers"] & eco_den if eco_den else set()

            bu_rows=[]
            for bname,x in sorted(a["by_bu"].items()):
                bu_rows.append({
                    "bu":bname,
                    "doanh_so_dat":round(x["order_sales"],2),
                    "doanh_so_giao":round(x["delivered_sales"],2),
                    "so_don_dat":len(x["order_set"]),
                    "so_don_giao_tc":len(x["success_set"])
                })

            rows.append({
                "ma_nv":emp,
                "ten_nv":sp.get("TEN_NV") or "",
                "ma_vung_kd":sp.get("MA_VUNG_KD") or "",
                "bo_phan":sp.get("BO_PHAN") or "",
                "doanh_so_dat":round(a["order_sales"],2),
                "doanh_so_giao":round(a["delivered_sales"],2),
                "so_don_dat":len(a["order_set"]),
                "so_don_giao_tc":len(a["success_set"]),
                "don_giao_du":len(a["full_set"]),
                "don_giao_mot_phan":len(a["partial_set"]),
                "don_vo_100":len(a["broken_set"]),
                "kh_co_don_giao_tc":len(eco_num),
                "tong_kh_phu_trach":len(eco_den),
                "eco":round(
                    (len(eco_num)/len(eco_den)*100)
                    if eco_den else 0,
                    2
                ),
                "eco_success_source_total":len(a["eco_customers"]),
                "by_bu":bu_rows
            })

            totals["order_sales"]+=a["order_sales"]
            totals["delivered_sales"]+=a["delivered_sales"]
            totals["orders"]|=a["order_set"]
            totals["success"]|=a["success_set"]
            totals["full"]|=a["full_set"]
            totals["partial"]|=a["partial_set"]
            totals["broken"]|=a["broken_set"]

        rows.sort(key=lambda x:x["ma_nv"])

        return {
            "month":month,
            "rows":rows,
            "summary":{
                "doanh_so_dat":round(totals["order_sales"],2),
                "doanh_so_giao":round(totals["delivered_sales"],2),
                "so_don_dat":len(totals["orders"]),
                "so_don_giao_tc":len(totals["success"]),
                "don_giao_du":len(totals["full"]),
                "don_giao_mot_phan":len(totals["partial"]),
                "don_vo_100":len(totals["broken"])
            },
            "eco_owner_map_count":len(owner_map)
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/kq-nvkd/eco-detail-v104")
def kq_nvkd_eco_detail_v104(month:str,ma_nv:str):
    try:
        month=_kq_month(month)
        ma_nv=str(ma_nv or "").strip()
        if not month or not ma_nv:
            raise ValueError("Thiếu tháng hoặc MA_NV")

        staff,staff_map,products,customers=_kq_master_maps()
        src=_kq_active_batch_rows()
        actual,success_orders,classes=_kq_settled_actual()

        assigned_all,owner_map,diag=_kq104_assigned_customers(
            staff,staff_map,customers,src
        )
        assigned_set=assigned_all.get(ma_nv,set())

        alias_to_canonical={}
        cust_by_can={}
        for kh in customers:
            canonical=_kq_customer_key(kh)
            if not canonical:
                continue
            cust_by_can[canonical]=kh
            for ck in ("MA_KH_CHUAN","MA_KH_DMS","MA_KH_MBW","MA_KH"):
                cv=str(kh.get(ck) or "").strip()
                if cv:
                    alias_to_canonical[cv]=canonical

        result={}
        for canonical in assigned_set:
            kh=cust_by_can.get(canonical,{})
            result[canonical]={
                "ma_kh_chuan":canonical,
                "ma_kh_dms":kh.get("MA_KH_DMS") or "",
                "ma_kh_mbw":kh.get("MA_KH_MBW") or "",
                "ten_kh":kh.get("TEN_KH") or "",
                "ma_dia_ban":kh.get("MA_DIA_BAN") or "",
                "trang_thai":kh.get("TRANG_THAI") or "",
                "co_giao_tc":False,
                "so_don_giao_tc":0,
                "doanh_so_giao":0.0
            }

        order_seen=set()

        for p in src:
            day=str(
                p.get("NGAY_DON")
                or p.get("NGAY")
                or p.get("NGAY_GIAO")
                or ""
            )
            if _kq_month(day)!=month:
                continue

            src_emp=str(
                p.get("MA_NV_KD_NGHIEP_VU")
                or p.get("MA_NV_KD")
                or ""
            ).strip()
            emp=staff_map.get(src_emp,src_emp if src_emp in staff else "")
            if emp!=ma_nv:
                continue

            order=str(
                p.get("SO_HOA_DON")
                or p.get("MA_DON_DMS")
                or p.get("MA_DON")
                or ""
            ).strip()
            if order not in success_orders:
                continue

            raw=""
            for ck in (
                "MA_KH_NGUON","MA_KH_CHUAN",
                "MA_KH_DMS","MA_KH_MBW","MA_KH"
            ):
                cv=str(p.get(ck) or "").strip()
                if cv:
                    raw=cv
                    break

            canonical=alias_to_canonical.get(raw,raw)
            if canonical not in result:
                continue

            sku=str(p.get("MA_SKU") or "").strip()
            qty_actual=_kq_num(actual.get((order,sku),0))
            price=_kq_num(p.get("DON_GIA_CHUA_VAT"))

            result[canonical]["doanh_so_giao"] += qty_actual*price

            key=(canonical,order)
            if key not in order_seen:
                order_seen.add(key)
                result[canonical]["so_don_giao_tc"] += 1
                result[canonical]["co_giao_tc"]=True

        rows=list(result.values())
        rows.sort(key=lambda x:(not x["co_giao_tc"],x["ma_kh_chuan"]))

        yes=sum(1 for x in rows if x["co_giao_tc"])
        total=len(rows)

        return {
            "month":month,
            "ma_nv":ma_nv,
            "ten_nv":staff.get(ma_nv,{}).get("TEN_NV") or "",
            "tong_kh_phu_trach":total,
            "kh_co_giao_tc":yes,
            "kh_chua_giao_tc":total-yes,
            "eco":round((yes/total*100) if total else 0,2),
            "owner_map_count":len(owner_map),
            "rows":rows
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.7 TARGET IMPORT =====

import io as _io_kq107
import re as _re_kq107
from openpyxl import load_workbook as _load_wb_kq107, Workbook as _Wb_kq107
from fastapi.responses import StreamingResponse as _StreamingResponse_kq107

def _kq107_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS kq_nvkd_target_stable(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      thang TEXT NOT NULL,
      ma_nv TEXT NOT NULL,
      ten_nv TEXT,
      ct_hc REAL NOT NULL DEFAULT 0,
      ct_pc REAL NOT NULL DEFAULT 0,
      ct_be REAL NOT NULL DEFAULT 0,
      ct_f REAL NOT NULL DEFAULT 0,
      ct_dstl REAL NOT NULL DEFAULT 0,
      ct_dstt REAL NOT NULL DEFAULT 0,
      ct_eco REAL NOT NULL DEFAULT 0,
      source_filename TEXT,
      imported_at TEXT NOT NULL,
      UNIQUE(thang,ma_nv)
    );
    CREATE INDEX IF NOT EXISTS ix_kq_nvkd_target_stable_month
      ON kq_nvkd_target_stable(thang,ma_nv);
    """)
    c.commit()
    c.close()

_kq107_init()

def _kq107_target_rows(month=None):
    c=_conn()
    if month:
        rows=[dict(r) for r in c.execute("""
          SELECT * FROM kq_nvkd_target_stable
          WHERE thang=?
          ORDER BY ma_nv
        """,(month,)).fetchall()]
    else:
        rows=[dict(r) for r in c.execute("""
          SELECT * FROM kq_nvkd_target_stable
          ORDER BY thang DESC,ma_nv
        """).fetchall()]
    c.close()
    return rows

@router.get("/kq-nvkd/targets-v107")
def kq_nvkd_targets_v107(month:str=None):
    try:
        rows=_kq107_target_rows(month)
        return {
          "month":month,
          "rows":rows,
          "count":len(rows),
          "months":sorted({r["thang"] for r in _kq107_target_rows()}, reverse=True)
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/kq-nvkd/target-template-v107")
def kq_nvkd_target_template_v107():
    try:
        wb=_Wb_kq107()
        ws=wb.active
        ws.title="CHI_TIEU_NVKD"
        headers=["THANG","MA_NV","TEN_NV","CT_HC","CT_PC","CT_BE","CT_F","CT_DSTL","CT_DSTT","CT_ECO"]
        ws.append(headers)
        ws.append(["2026-09","VG0060","Nguyễn Công Hiền",100000000,80000000,60000000,40000000,
                   "=SUM(D2:G2)",320000000,0.90])
        ws["J2"].number_format="0.00%"
        for c in range(1,11):
            ws.cell(1,c).font=ws.cell(1,c).font.copy(bold=True)
        widths=[12,12,24,15,15,15,15,16,16,12]
        for i,w in enumerate(widths,1):
            ws.column_dimensions[chr(64+i)].width=w
        ws.freeze_panes="A2"

        hd=wb.create_sheet("HUONG_DAN")
        hd.append(["Trường","Quy tắc"])
        hd.append(["THANG","YYYY-MM"])
        hd.append(["MA_NV","Mã nhân viên chuẩn DM02"])
        hd.append(["CT_HC/PC/BE/F","Chỉ tiêu theo BU"])
        hd.append(["CT_DSTL","Tự tính = HC+PC+BE+F"])
        hd.append(["CT_DSTT","Nhập độc lập"])
        hd.append(["CT_ECO","Nhập dạng %, ví dụ 90%"])
        hd.append(["% đạt ECO","ECO thực tế / CT_ECO"])

        bio=_io_kq107.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return _StreamingResponse_kq107(
            bio,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":'attachment; filename="Bieu_mau_nap_chi_tieu_NVKD.xlsx"'}
        )
    except Exception as e:
        raise HTTPException(400,str(e))

def _kq107_parse_xlsx(raw,filename=""):
    wb=_load_wb_kq107(_io_kq107.BytesIO(raw),data_only=False)
    if "CHI_TIEU_NVKD" not in wb.sheetnames:
        raise ValueError("Không tìm thấy sheet CHI_TIEU_NVKD")

    ws=wb["CHI_TIEU_NVKD"]

    header_row=None
    headers={}
    for r in range(1,min(ws.max_row,20)+1):
        vals=[str(ws.cell(r,c).value or "").strip().upper() for c in range(1,min(ws.max_column,20)+1)]
        if "THANG" in vals and "MA_NV" in vals:
            header_row=r
            headers={v:i+1 for i,v in enumerate(vals) if v}
            break

    if not header_row:
        raise ValueError("Không tìm thấy dòng tiêu đề THANG / MA_NV")

    required=["THANG","MA_NV","CT_HC","CT_PC","CT_BE","CT_F","CT_DSTT","CT_ECO"]
    missing=[x for x in required if x not in headers]
    if missing:
        raise ValueError("Thiếu cột: "+", ".join(missing))

    staff,_,_,_=_kq_master_maps()

    rows=[]
    errors=[]
    seen=set()

    for r in range(header_row+1,ws.max_row+1):
        def val(name):
            col=headers.get(name)
            return ws.cell(r,col).value if col else None

        thang=str(val("THANG") or "").strip()
        ma_nv=str(val("MA_NV") or "").strip()
        ten_nv=str(val("TEN_NV") or "").strip()

        if not thang and not ma_nv:
            continue

        row_errors=[]
        if not _re_kq107.fullmatch(r"\d{4}-\d{2}",thang):
            row_errors.append("THANG phải dạng YYYY-MM")

        if not ma_nv:
            row_errors.append("Thiếu MA_NV")
        elif ma_nv not in staff:
            row_errors.append("MA_NV không tồn tại trong DM02")
        elif str(staff.get(ma_nv,{}).get("NHOM_NV") or "").upper()!="KINH_DOANH":
            row_errors.append("MA_NV không thuộc nhóm KINH_DOANH")

        key=(thang,ma_nv)
        if key in seen:
            row_errors.append("Trùng THANG + MA_NV trong file")
        seen.add(key)

        nums={}
        for name in ("CT_HC","CT_PC","CT_BE","CT_F","CT_DSTT"):
            try:
                x=float(val(name) or 0)
                if x<0:
                    row_errors.append(f"{name} không được âm")
                nums[name]=x
            except Exception:
                row_errors.append(f"{name} không hợp lệ")
                nums[name]=0

        try:
            eco=float(val("CT_ECO") or 0)
            # Excel percentage normally returns 0.9. Accept 90 as convenience.
            if eco>1 and eco<=100:
                eco=eco/100
            if eco<0 or eco>1:
                row_errors.append("CT_ECO phải từ 0% đến 100%")
        except Exception:
            eco=0
            row_errors.append("CT_ECO không hợp lệ")

        ct_dstl=nums["CT_HC"]+nums["CT_PC"]+nums["CT_BE"]+nums["CT_F"]

        rec={
          "row_no":r,
          "thang":thang,
          "ma_nv":ma_nv,
          "ten_nv":staff.get(ma_nv,{}).get("TEN_NV") or ten_nv,
          "ct_hc":nums["CT_HC"],
          "ct_pc":nums["CT_PC"],
          "ct_be":nums["CT_BE"],
          "ct_f":nums["CT_F"],
          "ct_dstl":ct_dstl,
          "ct_dstt":nums["CT_DSTT"],
          "ct_eco":eco,
          "errors":row_errors
        }
        rows.append(rec)
        if row_errors:
            errors.append({"row_no":r,"errors":row_errors})

    if not rows:
        raise ValueError("File không có dòng chỉ tiêu")

    return {
      "filename":filename,
      "rows":rows,
      "errors":errors,
      "valid_count":sum(1 for x in rows if not x["errors"]),
      "error_count":sum(1 for x in rows if x["errors"])
    }

@router.post("/kq-nvkd/target-preview-v107")
async def kq_nvkd_target_preview_v107(request:Request,filename:str=""):
    try:
        raw=await request.body()
        if not raw:
            raise ValueError("File rỗng")
        return _kq107_parse_xlsx(raw,filename)
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/kq-nvkd/target-commit-v107")
async def kq_nvkd_target_commit_v107(request:Request,filename:str=""):
    try:
        raw=await request.body()
        if not raw:
            raise ValueError("File rỗng")

        parsed=_kq107_parse_xlsx(raw,filename)
        if parsed["error_count"]:
            raise ValueError("File còn lỗi. Hãy sửa trước khi ghi nhận.")

        c=_conn()
        now=_now()
        for x in parsed["rows"]:
            c.execute("""
            INSERT INTO kq_nvkd_target_stable(
              thang,ma_nv,ten_nv,ct_hc,ct_pc,ct_be,ct_f,
              ct_dstl,ct_dstt,ct_eco,source_filename,imported_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(thang,ma_nv) DO UPDATE SET
              ten_nv=excluded.ten_nv,
              ct_hc=excluded.ct_hc,
              ct_pc=excluded.ct_pc,
              ct_be=excluded.ct_be,
              ct_f=excluded.ct_f,
              ct_dstl=excluded.ct_dstl,
              ct_dstt=excluded.ct_dstt,
              ct_eco=excluded.ct_eco,
              source_filename=excluded.source_filename,
              imported_at=excluded.imported_at
            """,(
              x["thang"],x["ma_nv"],x["ten_nv"],
              x["ct_hc"],x["ct_pc"],x["ct_be"],x["ct_f"],
              x["ct_dstl"],x["ct_dstt"],x["ct_eco"],
              filename,now
            ))
        c.commit()
        c.close()

        months=sorted({x["thang"] for x in parsed["rows"]})
        return {
          "ok":True,
          "committed":len(parsed["rows"]),
          "months":months
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/kq-nvkd/report-v107")
def kq_nvkd_report_v107(month:str,ma_nv:str=None,ma_vung_kd:str=None,bu:str=None):
    try:
        base=kq_nvkd_report_v104(month,ma_nv,ma_vung_kd,bu)
        targets={r["ma_nv"]:r for r in _kq107_target_rows(_kq_month(month))}

        for r in base.get("rows",[]):
            t=targets.get(r["ma_nv"])
            if not t:
                r["target"]=None
                continue

            ds=_kq_num(r.get("doanh_so_giao"))
            eco_actual=_kq_num(r.get("eco"))/100.0
            ct_dstl=_kq_num(t.get("ct_dstl"))
            ct_dstt=_kq_num(t.get("ct_dstt"))
            ct_eco=_kq_num(t.get("ct_eco"))

            r["target"]={
              "ct_hc":t["ct_hc"],
              "ct_pc":t["ct_pc"],
              "ct_be":t["ct_be"],
              "ct_f":t["ct_f"],
              "ct_dstl":ct_dstl,
              "ct_dstt":ct_dstt,
              "ct_eco":ct_eco,
              "pct_dstl":round(ds/ct_dstl*100,2) if ct_dstl else 0,
              "pct_dstt":round(ds/ct_dstt*100,2) if ct_dstt else 0,
              "pct_eco":round(eco_actual/ct_eco*100,2) if ct_eco else 0
            }

        base["target_count"]=len(targets)
        return base
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.9 BU TARGET MAPPING =====

def _kq109_target_for_bu(target, bu):
    if not target:
        return 0.0
    key={
      "HC":"ct_hc",
      "PC":"ct_pc",
      "BE":"ct_be",
      "F":"ct_f"
    }.get(str(bu or "").upper())
    return _kq_num(target.get(key)) if key else 0.0

@router.get("/kq-nvkd/report-v109")
def kq_nvkd_report_v109(month:str,ma_nv:str=None,ma_vung_kd:str=None,bu:str=None):
    try:
        base=kq_nvkd_report_v107(month,ma_nv,ma_vung_kd,bu)

        for r in base.get("rows",[]):
            t=r.get("target") or None

            # Enrich each BU result with its matching BU target.
            for b in r.get("by_bu",[]):
                ct_bu=_kq109_target_for_bu(t,b.get("bu"))
                ds_bu=_kq_num(b.get("doanh_so_giao"))
                b["ct_bu"]=ct_bu
                b["pct_bu"]=round(ds_bu/ct_bu*100,2) if ct_bu else 0

            # If report is filtered to one BU, expose selected BU target separately.
            if bu:
                selected_ct=_kq109_target_for_bu(t,bu)
                selected_ds=sum(
                    _kq_num(b.get("doanh_so_giao"))
                    for b in r.get("by_bu",[])
                    if str(b.get("bu") or "").upper()==str(bu).upper()
                )
                r["selected_bu_target"]={
                  "bu":str(bu).upper(),
                  "doanh_so_giao":round(selected_ds,2),
                  "ct_bu":selected_ct,
                  "pct_bu":round(selected_ds/selected_ct*100,2) if selected_ct else 0
                }
            else:
                r["selected_bu_target"]=None

        return base
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== KQ NVKD STABLE V1.0.10 UNIFIED TARGET REPORT =====

@router.get("/kq-nvkd/report-v110")
def kq_nvkd_report_v110(month:str,ma_nv:str=None,ma_vung_kd:str=None,bu:str=None):
    try:
        base=kq_nvkd_report_v104(month,ma_nv,ma_vung_kd,bu)
        targets={r["ma_nv"]:r for r in _kq107_target_rows(_kq_month(month))}

        for r in base.get("rows",[]):
            t=targets.get(r["ma_nv"])
            if not t:
                r["target"]=None
                for b in r.get("by_bu",[]):
                    b["ct_bu"]=0
                    b["pct_bu"]=0
                continue

            ds=_kq_num(r.get("doanh_so_giao"))
            eco_actual_pct=_kq_num(r.get("eco"))   # already 0..100
            ct_dstl=_kq_num(t.get("ct_dstl"))
            ct_dstt=_kq_num(t.get("ct_dstt"))
            ct_eco_ratio=_kq_num(t.get("ct_eco")) # 0..1
            ct_eco_pct=ct_eco_ratio*100

            r["target"]={
              "ct_hc":_kq_num(t.get("ct_hc")),
              "ct_pc":_kq_num(t.get("ct_pc")),
              "ct_be":_kq_num(t.get("ct_be")),
              "ct_f":_kq_num(t.get("ct_f")),
              "ct_dstl":ct_dstl,
              "ct_dstt":ct_dstt,
              "ct_eco_ratio":ct_eco_ratio,
              "ct_eco_pct":round(ct_eco_pct,2),
              "pct_dstl":round(ds/ct_dstl*100,2) if ct_dstl else 0,
              "pct_dstt":round(ds/ct_dstt*100,2) if ct_dstt else 0,
              "pct_eco":round(eco_actual_pct/ct_eco_pct*100,2) if ct_eco_pct else 0
            }

            target_map={
              "HC":_kq_num(t.get("ct_hc")),
              "PC":_kq_num(t.get("ct_pc")),
              "BE":_kq_num(t.get("ct_be")),
              "F":_kq_num(t.get("ct_f"))
            }

            for b in r.get("by_bu",[]):
                bname=str(b.get("bu") or "").upper()
                ct=target_map.get(bname,0)
                th=_kq_num(b.get("doanh_so_giao"))
                b["ct_bu"]=ct
                b["pct_bu"]=round(th/ct*100,2) if ct else 0

        base["target_count"]=len(targets)
        return base
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== DM02A STABLE V1.0.12 TEMPORAL IMPORT =====
import io as _io112, json as _json112, datetime as _dt112
from openpyxl import load_workbook as _load112, Workbook as _WB112
from fastapi.responses import StreamingResponse as _Stream112

def _j112(s):
    try:
        x=_json112.loads(s or "{}")
        return x if isinstance(x,dict) else {}
    except Exception:return {}

def _d112(v):
    if v in (None,""):return ""
    if isinstance(v,_dt112.datetime):return v.date().isoformat()
    if isinstance(v,_dt112.date):return v.isoformat()
    s=str(v).strip()
    for f in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y"):
        try:return _dt112.datetime.strptime(s[:10],f).date().isoformat()
        except Exception:pass
    try:return _dt112.datetime.fromisoformat(s.replace("Z","")).date().isoformat()
    except Exception:return ""

def _staff112():
    c=_conn(); out={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm02-staff'"):
        p=_j112(r["payload_json"]); ma=str(p.get("MA_NV") or "").strip()
        if ma:out[ma]=p
    c.close(); return out

def _rows112():
    c=_conn(); out=[]
    for r in c.execute("SELECT business_key,payload_json FROM master_data_generic WHERE category='dm02a-staff-map'"):
        p=_j112(r["payload_json"])
        out.append({
          "business_key":r["business_key"],
          "ma_nv":str(p.get("MA_NV") or "").strip(),
          "loai_ma":str(p.get("LOAI_MA") or "").strip(),
          "ma_nghiep_vu":str(p.get("MA_NGHIEP_VU") or "").strip(),
          "he_thong":str(p.get("HE_THONG") or "").strip().upper(),
          "tu_ngay":_d112(p.get("TU_NGAY")) or "0001-01-01",
          "den_ngay":_d112(p.get("DEN_NGAY")) or "9999-12-31",
          "trang_thai":str(p.get("TRANG_THAI") or "DANG_HIEU_LUC").strip().upper()
        })
    c.close(); return out

def _ov112(a,b):
    if (a["ma_nghiep_vu"],a["he_thong"],a["loai_ma"]) != (b["ma_nghiep_vu"],b["he_thong"],b["loai_ma"]):return False
    if a["trang_thai"] in ("HUY","XOA","CANCELLED") or b["trang_thai"] in ("HUY","XOA","CANCELLED"):return False
    return a["tu_ngay"]<=b["den_ngay"] and b["tu_ngay"]<=a["den_ngay"]

def _bk112(x):
    return "||".join([x["ma_nv"],x["loai_ma"],x["ma_nghiep_vu"],x["he_thong"],x["tu_ngay"]])

def _parse112(raw,filename=""):
    wb=_load112(_io112.BytesIO(raw),data_only=True)
    sn="DM02A_MA_NGHIEP_VU"
    if sn not in wb.sheetnames:raise ValueError("Không tìm thấy sheet "+sn)
    ws=wb[sn]
    hr=None; hm={}
    for r in range(1,min(ws.max_row,20)+1):
        vals=[str(ws.cell(r,c).value or "").strip().upper() for c in range(1,min(ws.max_column,30)+1)]
        if "MA_NV" in vals and "MA_NGHIEP_VU" in vals:
            hr=r; hm={v:i+1 for i,v in enumerate(vals) if v}; break
    if not hr:raise ValueError("Không tìm thấy dòng tiêu đề")
    req=["MA_NV","LOAI_MA","MA_NGHIEP_VU","HE_THONG","TU_NGAY","DEN_NGAY","TRANG_THAI"]
    miss=[x for x in req if x not in hm]
    if miss:raise ValueError("Thiếu cột: "+", ".join(miss))
    staff=_staff112(); rows=[]
    for r in range(hr+1,ws.max_row+1):
        def V(n): return ws.cell(r,hm[n]).value if n in hm else None
        ma=str(V("MA_NV") or "").strip(); code=str(V("MA_NGHIEP_VU") or "").strip()
        if not ma and not code:continue
        x={
          "row_no":r,"ma_nv":ma,"loai_ma":str(V("LOAI_MA") or "").strip(),
          "ma_nghiep_vu":code,"he_thong":str(V("HE_THONG") or "").strip().upper(),
          "tu_ngay":_d112(V("TU_NGAY")) or "",
          "den_ngay":_d112(V("DEN_NGAY")) or "9999-12-31",
          "trang_thai":str(V("TRANG_THAI") or "DANG_HIEU_LUC").strip().upper(),
          "ly_do_thay_doi":str(V("LY_DO_THAY_DOI") or "").strip(),
          "nguoi_cap_nhat":str(V("NGUOI_CAP_NHAT") or "").strip(),
          "ghi_chu":str(V("GHI_CHU") or "").strip(),"errors":[]
        }
        if not ma:x["errors"].append("Thiếu MA_NV")
        elif ma not in staff:x["errors"].append("MA_NV không tồn tại DM02")
        if not x["loai_ma"]:x["errors"].append("Thiếu LOAI_MA")
        if not code:x["errors"].append("Thiếu MA_NGHIEP_VU")
        if not x["he_thong"]:x["errors"].append("Thiếu HE_THONG")
        if not x["tu_ngay"]:x["errors"].append("TU_NGAY không hợp lệ")
        if x["tu_ngay"] and x["den_ngay"]!="9999-12-31" and x["den_ngay"]<x["tu_ngay"]:
            x["errors"].append("DEN_NGAY nhỏ hơn TU_NGAY")
        x["business_key"]=_bk112(x) if x["tu_ngay"] else ""
        rows.append(x)

    seen={}
    for x in rows:
        if x["business_key"] in seen:x["errors"].append("Trùng khóa trong file")
        else:seen[x["business_key"]]=x["row_no"]
    for i,a in enumerate(rows):
        for b in rows[i+1:]:
            if a["business_key"]!=b["business_key"] and _ov112(a,b):
                a["errors"].append("Chồng hiệu lực với dòng "+str(b["row_no"]))
                b["errors"].append("Chồng hiệu lực với dòng "+str(a["row_no"]))
    existing=_rows112()

    # V1.0.12A:
    # Khi file đồng thời đóng dòng cũ và mở dòng mới, phải kiểm tra overlap
    # theo trạng thái đề xuất sau import, không phải trạng thái DB trước import.
    proposed_keys={x["business_key"] for x in rows if x.get("business_key")}
    existing_for_check=[
        old for old in existing
        if old["business_key"] not in proposed_keys
    ]

    for x in rows:
        for old in existing_for_check:
            if _ov112(x,old):
                x["errors"].append(
                    f"Chồng DB: {old['ma_nv']} {old['tu_ngay']}.."
                    f"{'' if old['den_ngay']=='9999-12-31' else old['den_ngay']}"
                )
                break
    return {"filename":filename,"rows":rows,
            "valid_count":sum(not x["errors"] for x in rows),
            "error_count":sum(bool(x["errors"]) for x in rows)}

# Historical rows marked NGUNG_HIEU_LUC are still valid within TU_NGAY..DEN_NGAY.
def _resolve112(code,ngay,he_thong="DMS",loai_ma=None,strict=True):
    code=str(code or "").strip(); day=_d112(ngay); sys=str(he_thong or "").upper()
    m=[]
    for x in _rows112():
        if x["ma_nghiep_vu"]!=code:continue
        if sys and x["he_thong"] and x["he_thong"]!=sys:continue
        if loai_ma and x["loai_ma"] and x["loai_ma"]!=loai_ma:continue
        if x["trang_thai"] in ("HUY","XOA","CANCELLED"):continue
        if x["tu_ngay"]<=day<=x["den_ngay"]:m.append(x)
    if len(m)==1:return m[0]["ma_nv"]
    if not m:
        if strict:raise ValueError(f"Không có mapping hiệu lực {code} tại {day}")
        return ""
    raise ValueError(f"Mapping chồng hiệu lực {code} tại {day}")

try:_tm111_resolve=_resolve112
except Exception:pass

@router.get("/dm02a/template-v112")
def template112():
    wb=_WB112(); ws=wb.active; ws.title="DM02A_MA_NGHIEP_VU"
    ws.append(["MA_NV","LOAI_MA","MA_NGHIEP_VU","HE_THONG","TU_NGAY","DEN_NGAY","TRANG_THAI","LY_DO_THAY_DOI","NGUOI_CAP_NHAT","GHI_CHU"])
    ws.append(["VG0060","MA_MAY_BAN_HANG","24151","DMS","2026-01-01","2026-08-31","NGUNG_HIEU_LUC","Đổi tuyến","",""])
    ws.append(["VG0051","MA_MAY_BAN_HANG","24151","DMS","2026-09-01","","DANG_HIEU_LUC","Nhận tuyến","",""])
    bio=_io112.BytesIO(); wb.save(bio); bio.seek(0)
    return _Stream112(bio,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
      headers={"Content-Disposition":'attachment; filename="Template_DM02A_Ma_nghiep_vu_co_thoi_han.xlsx"'})

@router.get("/dm02a/list-v112")
def list112(q:str=None):
    rows=_rows112()
    if q:
        qq=q.upper()
        rows=[x for x in rows if qq in (x["ma_nv"]+" "+x["ma_nghiep_vu"]+" "+x["he_thong"]+" "+x["loai_ma"]).upper()]
    rows.sort(key=lambda x:(x["ma_nghiep_vu"],x["tu_ngay"],x["ma_nv"]))
    for x in rows:
        if x["den_ngay"]=="9999-12-31":x["den_ngay"]=""
    return {"rows":rows,"count":len(rows)}

@router.post("/dm02a/preview-v112")
async def preview112(request:Request,filename:str=""):
    try:
        raw=await request.body()
        if not raw:raise ValueError("File rỗng")
        return _parse112(raw,filename)
    except Exception as e:raise HTTPException(400,str(e))

@router.post("/dm02a/commit-v112")
async def commit112(request:Request,filename:str=""):
    try:
        raw=await request.body()
        p=_parse112(raw,filename)
        if p["error_count"]:raise ValueError("File còn lỗi/chồng hiệu lực")
        c=_conn(); now=_now()
        for x in p["rows"]:
            payload={"MA_NV":x["ma_nv"],"LOAI_MA":x["loai_ma"],"MA_NGHIEP_VU":x["ma_nghiep_vu"],
              "HE_THONG":x["he_thong"],"TU_NGAY":x["tu_ngay"],
              "DEN_NGAY":"" if x["den_ngay"]=="9999-12-31" else x["den_ngay"],
              "TRANG_THAI":x["trang_thai"],"LY_DO_THAY_DOI":x["ly_do_thay_doi"],
              "NGUOI_CAP_NHAT":x["nguoi_cap_nhat"],"GHI_CHU":x["ghi_chu"]}
            js=_json112.dumps(payload,ensure_ascii=False)
            old=c.execute("SELECT 1 FROM master_data_generic WHERE category='dm02a-staff-map' AND business_key=?",(x["business_key"],)).fetchone()
            if old:c.execute("UPDATE master_data_generic SET payload_json=?,updated_at=? WHERE category='dm02a-staff-map' AND business_key=?",(js,now,x["business_key"]))
            else:c.execute("INSERT INTO master_data_generic(category,business_key,payload_json,updated_at) VALUES('dm02a-staff-map',?,?,?)",(x["business_key"],js,now))
        c.commit();c.close()
        return {"ok":True,"committed":len(p["rows"])}
    except Exception as e:raise HTTPException(400,str(e))


# ===== SOURCE SCOPE CLEANUP V1 =====

def _ssc1_table_exists(c,name):
    return c.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
        (name,)
    ).fetchone() is not None

def _ssc1_cols(c,name):
    return [r["name"] for r in c.execute(f'PRAGMA table_info("{name}")').fetchall()]

def _ssc1_active_batch(c,batch_table):
    if not _ssc1_table_exists(c,batch_table):
        return None
    cols=_ssc1_cols(c,batch_table)
    where=["UPPER(COALESCE(status,''))='COMMITTED'"]
    if "superseded_by_batch_id" in cols:
        where.append("superseded_by_batch_id IS NULL")
    q='SELECT * FROM "' + batch_table + '" WHERE ' + ' AND '.join(where) + ' ORDER BY rowid DESC LIMIT 1'
    return c.execute(q).fetchone()

def _ssc1_batch_count(c,row_table,batch_id):
    if not _ssc1_table_exists(c,row_table):
        return 0
    cols=_ssc1_cols(c,row_table)
    if "batch_id" not in cols:
        return c.execute(f'SELECT COUNT(*) FROM "{row_table}"').fetchone()[0]
    return c.execute(
        f'SELECT COUNT(*) FROM "{row_table}" WHERE batch_id=?',
        (batch_id,)
    ).fetchone()[0]

@router.get("/xb/source-scope-v1")
def xb_source_scope_v1():
    try:
        c=_conn()
        specs=[
            ("DMS_MBW","xb_dms_mbw_batch","xb_dms_mbw_row"),
            ("DMS_XA_LEGACY","xb_src_batch_v2","xb_src_row_v2"),
            ("MISA_BU","xb_misa_bu_batch","xb_misa_bu_row"),
        ]
        rows=[]
        for scope,bt,rt in specs:
            b=_ssc1_active_batch(c,bt)
            if not b:
                rows.append({
                    "scope":scope,"batch_table":bt,"row_table":rt,
                    "active_batch_id":None,"status":"NONE","rows":0
                })
                continue
            d=dict(b)
            bid=d.get("batch_id")
            rows.append({
                "scope":scope,
                "batch_table":bt,
                "row_table":rt,
                "active_batch_id":bid,
                "status":d.get("status"),
                "rows":_ssc1_batch_count(c,rt,bid),
                "total_rows":d.get("total_rows"),
                "valid_rows":d.get("valid_rows"),
                "error_rows":d.get("error_rows"),
                "source_type":d.get("source_type"),
                "file_name":d.get("file_name") or d.get("filename"),
                "superseded_by_batch_id":d.get("superseded_by_batch_id"),
            })

        dms=next((x for x in rows if x["scope"]=="DMS_MBW"),None)
        qtgh_ready=bool(dms and dms["status"]=="COMMITTED" and dms["rows"]>0)

        c.close()
        return {
            "sources":rows,
            "qtgh_source_scope":"DMS_MBW_ONLY",
            "qtgh_ready":qtgh_ready,
            "qtgh_expected_rows":dms["rows"] if dms else 0,
            "note":"QTGH chỉ dùng DMS/MBW COMMITTED hiện hành; DMS-Xả/MISA giữ cho Xả–Bù."
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/xb/source-scope-v1/{scope}")
def xb_source_scope_one_v1(scope:str):
    try:
        scope=str(scope or "").upper()
        allx=xb_source_scope_v1()
        if scope=="ALL":
            return allx
        x=next((r for r in allx["sources"] if r["scope"]==scope),None)
        if not x:
            raise ValueError("Scope không hợp lệ")
        return {
            "source":x,
            "qtgh_source_scope":allx["qtgh_source_scope"],
            "qtgh_expected_rows":allx["qtgh_expected_rows"]
        }
    except Exception as e:
        raise HTTPException(400,str(e))


# ===== QTGH PRINT META + UI11 DATE FIX V2 =====

@router.get("/print/meta-v2")
def qtgh_print_meta_v2():
    try:
        c=_conn()

        staff={}
        try:
            rows=c.execute(
                "SELECT payload_json FROM master_data_generic WHERE category='dm02-staff'"
            ).fetchall()
            for r in rows:
                try:
                    p=json.loads(r["payload_json"] or "{}")
                except Exception:
                    p={}
                ma=str(p.get("MA_NV") or "").strip()
                if ma:
                    staff[ma]=str(
                        p.get("TEN_NV")
                        or p.get("HO_TEN")
                        or p.get("TEN_NHAN_VIEN")
                        or ""
                    ).strip()
        except Exception:
            pass

        receipt_numbers={str(r["id"]):r["receipt_no"] for r in c.execute(
            "SELECT id,receipt_no FROM qtgh_cash_receipt WHERE receipt_no IS NOT NULL")}

        c.close()
        return {
            "staff":staff,
            "receipt_numbers":receipt_numbers
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== GH KQ STABLE V1.0.2 DELIVERY PROGRESS =====
def _gh_days_between(d1,d2):
    import datetime as _dt
    try:
        a=_dt.datetime.strptime(str(d1 or "")[:10],"%Y-%m-%d").date()
        b=_dt.datetime.strptime(str(d2 or "")[:10],"%Y-%m-%d").date()
        return (b-a).days
    except Exception:
        return None

# ===== GH KQ STABLE V1.0.0 =====

from collections import defaultdict as _ghdd

def _gh_num(v):
    try:return float(v or 0)
    except:return 0.0

def _gh_json(v):
    try:
        import json
        x=json.loads(v or "{}")
        return x if isinstance(x,dict) else {}
    except:return {}

def _gh_date(v):
    return str(v or "").strip()[:10]

def _gh_staff_map():
    c=_conn()
    staff={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm02-staff'"):
        p=_gh_json(r["payload_json"])
        ma=str(p.get("MA_NV") or "").strip()
        if ma: staff[ma]=p
    c.close()
    return staff

def _gh_products():
    c=_conn(); out={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm01-product'"):
        p=_gh_json(r["payload_json"])
        sku=str(p.get("MA_SKU") or "").strip()
        if sku: out[sku]=p
    c.close()
    return out

def _gh_weight(prod):
    # grams/kg compatibility; project can keep one of these fields.
    for k in ("TRONG_LUONG_KG","TRONG_LUONG","KHOI_LUONG_KG","KHOI_LUONG","WEIGHT_KG"):
        v=prod.get(k)
        if v not in (None,""):
            x=_gh_num(v)
            # If field name is generic and value looks like grams, normalize.
            if k in ("TRONG_LUONG","KHOI_LUONG") and x>50:
                return x/1000.0
            return x
    return 0.0

def _gh_active_raw_by_order():
    c=_conn()
    b=c.execute("""
      SELECT batch_id FROM xb_dms_mbw_batch
      WHERE UPPER(COALESCE(status,''))='COMMITTED'
        AND superseded_by_batch_id IS NULL
      ORDER BY rowid DESC LIMIT 1
    """).fetchone()
    if not b:
        c.close(); return {}

    out={}
    for r in c.execute("""
      SELECT payload_json FROM xb_dms_mbw_row
      WHERE batch_id=? ORDER BY rowid
    """,(b["batch_id"],)):
        p=_gh_json(r["payload_json"])
        order=str(p.get("SO_HOA_DON") or p.get("MA_DON_DMS") or p.get("MA_DON") or "").strip()
        # Reward-level rows can share order; prefer SKU row because it has transport dimensions.
        if not order: continue
        if order not in out or p.get("MA_SKU"):
            out[order]=p
    c.close()
    return out

def _gh_resolve_emp(code,day,staff):
    code=str(code or "").strip()
    if not code:return ""
    if code in staff:return code
    try:
        # Temporal resolver installed by DM02A V1.0.12.
        return _dm112_resolve(code,day,"DMS",None,strict=False)
    except Exception:
        try:return _tm111_resolve(code,day,"DMS",None,strict=False)
        except Exception:return ""

def _gh_locked_draft_ids():
    c=_conn()
    ids=set()
    try:
        for r in c.execute("""
          SELECT DISTINCT so.draft_order_id
          FROM qtgh_session_order so
          JOIN qtgh_session s ON s.id=so.session_id
          WHERE UPPER(COALESCE(s.status,''))='DA_KHOA'
        """):
            ids.add(r["draft_order_id"])
    except Exception:
        pass
    c.close()
    return ids

def _gh_line_columns():
    c=_conn()
    cols=[r["name"] for r in c.execute("PRAGMA table_info(qtgh_draft_line)")]
    c.close(); return cols

def _gh_line_value(d,*names):
    for n in names:
        if n in d and d[n] not in (None,""):
            return d[n]
    return 0

@router.get("/gh-kq/meta-v100")
def gh_kq_meta_v100():
    try:
        staff=_gh_staff_map()
        raw=_gh_active_raw_by_order()

        nvgh=set(); drivers=set(); vehicles=set()
        for order,p in raw.items():
            day=p.get("NGAY_GIAO") or p.get("NGAY_DON")
            gh=_gh_resolve_emp(
                p.get("MA_NV_GIAO_HANG_NGHIEP_VU") or p.get("MA_NV_GIAO_HANG") or p.get("MA_NVGH"),
                day,staff
            )
            lx=_gh_resolve_emp(
                p.get("MA_NV_LAI_XE_NGHIEP_VU") or p.get("MA_NV_LAI_XE") or p.get("MA_LAI_XE"),
                day,staff
            )
            pt=str(p.get("PHUONG_TIEN") or p.get("MA_PHUONG_TIEN") or "").strip()
            if gh:nvgh.add(gh)
            if lx:drivers.add(lx)
            if pt:vehicles.add(pt)

        def emp_rows(codes):
            return [{
                "ma_nv":x,
                "ten_nv":staff.get(x,{}).get("TEN_NV") or ""
            } for x in sorted(codes)]

        return {
            "nvgh":emp_rows(nvgh),
            "drivers":emp_rows(drivers),
            "vehicles":sorted(vehicles)
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/gh-kq/report-v100")
def gh_kq_report_v100(
    from_date:str=None,to_date:str=None,
    ma_nvgh:str=None,ma_lai_xe:str=None,phuong_tien:str=None
):
    try:
        staff=_gh_staff_map()
        products=_gh_products()
        raw_by_order=_gh_active_raw_by_order()
        locked=_gh_locked_draft_ids()
        lcols=_gh_line_columns()

        c=_conn()
        drows=[]
        if locked:
            ph=",".join("?" for _ in locked)
            drows=c.execute(
                f"SELECT * FROM qtgh_draft_order WHERE id IN ({ph}) ORDER BY id",
                tuple(locked)
            ).fetchall()
        c.close()

        groups=_ghdd(lambda:{
            "orders":set(),"success":set(),"full":set(),"partial":set(),"broken":set(),
            "d0":0,"d1":0,"d2":0,"d3p":0,"d_unknown":0,
            "customers":set(),"weight_kg":0.0,"amount":0.0,
            "by_day":_ghdd(lambda:{"orders":set(),"success":set(),"weight_kg":0.0,"amount":0.0})
        })

        details=[]
        for dr in drows:
            d=dict(dr)
            did=d.get("id")
            order=str(d.get("ma_don_nguon") or "").strip()
            raw=raw_by_order.get(order,{})
            day=_gh_date(raw.get("NGAY_GIAO") or raw.get("NGAY_DON") or d.get("ngay_giao") or d.get("ngay_don"))
            if from_date and day<from_date: continue
            if to_date and day>to_date: continue

            nvgh=_gh_resolve_emp(
                raw.get("MA_NV_GIAO_HANG_NGHIEP_VU")
                or raw.get("MA_NV_GIAO_HANG")
                or raw.get("MA_NVGH")
                or d.get("ma_nvgh"),
                day,staff
            )
            driver=_gh_resolve_emp(
                raw.get("MA_NV_LAI_XE_NGHIEP_VU")
                or raw.get("MA_NV_LAI_XE")
                or raw.get("MA_LAI_XE"),
                day,staff
            )
            vehicle=str(raw.get("PHUONG_TIEN") or raw.get("MA_PHUONG_TIEN") or d.get("ma_phuong_tien") or "").strip()

            if ma_nvgh and nvgh!=ma_nvgh: continue
            if ma_lai_xe and driver!=ma_lai_xe: continue
            if phuong_tien and vehicle!=phuong_tien: continue

            c=_conn()
            lines=[dict(r) for r in c.execute(
                "SELECT * FROM qtgh_draft_line WHERE draft_order_id=? ORDER BY id",(did,)
            ).fetchall()]
            c.close()

            ordered=0.0; actual=0.0; weight=0.0; amount=0.0
            for ln in lines:
                sku=str(ln.get("ma_sku") or "").strip()
                sl_dat=_gh_num(_gh_line_value(ln,"sl_hang_ban_dat","sl_dat"))
                sl_thuc=_gh_num(_gh_line_value(ln,"sl_hang_ban_thuc_giao","sl_thuc_giao"))
                price=_gh_num(_gh_line_value(ln,"don_gia_co_vat","gia_co_vat"))

                # Use CK/KM after settlement first; fallback to original.
                ck=_gh_num(_gh_line_value(
                    ln,
                    "ck_km_dong_sau","ck_km_sau_qt","ck_km_sau",
                    "ck_km_dong_goc","ck_km_goc"
                ))

                ordered+=sl_dat
                actual+=sl_thuc
                weight += sl_thuc * _gh_weight(products.get(sku,{}))

                # Explicitly excludes order-level display/accumulation rewards.
                amount += sl_thuc*price - ck

            if actual<=0:
                cls="BROKEN_100"
            elif actual<ordered:
                cls="PARTIAL"
            else:
                cls="FULL"

            kh=str(
                d.get("ma_kh_chuan")
                or raw.get("MA_KH_CHUAN")
                or raw.get("MA_KH_NGUON")
                or ""
            ).strip()

            order_day=_gh_date(raw.get("NGAY_DON") or d.get("ngay_don"))
            detail={
                "order":order,"day":day,"order_day":order_day,
                "ma_nvgh":nvgh,"ten_nvgh":staff.get(nvgh,{}).get("TEN_NV") or "",
                "ma_lai_xe":driver,"ten_lai_xe":staff.get(driver,{}).get("TEN_NV") or "",
                "phuong_tien":vehicle,
                "status":cls,
                "weight_kg":round(weight,3),
                "amount":round(amount,2),
                "customer":kh
            }
            details.append(detail)

            # three independent dimensions
            dimensions=[
                ("NVGH",nvgh,staff.get(nvgh,{}).get("TEN_NV") or ""),
                ("LAI_XE",driver,staff.get(driver,{}).get("TEN_NV") or ""),
                ("PHUONG_TIEN",vehicle,vehicle)
            ]
            for typ,key,name in dimensions:
                if not key: continue
                g=groups[(typ,key,name)]
                g["orders"].add(order)
                if cls!="BROKEN_100":
                    g["success"].add(order)
                if cls=="FULL":g["full"].add(order)
                elif cls=="PARTIAL":g["partial"].add(order)
                else:g["broken"].add(order)

                if cls!="BROKEN_100":
                    delta=_gh_days_between(order_day,day)
                    if delta==0:g["d0"]+=1
                    elif delta==1:g["d1"]+=1
                    elif delta==2:g["d2"]+=1
                    elif delta is not None and delta>=3:g["d3p"]+=1
                    else:g["d_unknown"]+=1
                if kh:g["customers"].add(kh)
                g["weight_kg"]+=weight
                g["amount"]+=amount

                bd=g["by_day"][day]
                bd["orders"].add(order)
                if cls!="BROKEN_100":bd["success"].add(order)
                bd["weight_kg"]+=weight
                bd["amount"]+=amount

        rows=[]
        for (typ,key,name),g in groups.items():
            rows.append({
                "type":typ,"code":key,"name":name,
                "so_don_nhan":len(g["orders"]),
                "so_don_giao_tc":len(g["success"]),
                "giao_du":len(g["full"]),
                "giao_mot_phan":len(g["partial"]),
                "vo_100":len(g["broken"]),
                "so_kh_giao_tc":len(g["customers"]),
                "khoi_luong_kg":round(g["weight_kg"],3),
                "so_tien_thuc_tinh":round(g["amount"],2),
                "d0":g["d0"],"d1":g["d1"],"d2":g["d2"],"d3p":g["d3p"],
                "d_unknown":g["d_unknown"],
                "pct_le_d2":round(
                    (g["d0"]+g["d1"]+g["d2"])/len(g["success"])*100,2
                ) if g["success"] else 0
            })

        rows.sort(key=lambda x:(x["type"],x["code"]))
        details.sort(key=lambda x:(x["day"],x["order"]))

        summary={
            "so_don":len({x["order"] for x in details}),
            "so_don_giao_tc":len({x["order"] for x in details if x["status"]!="BROKEN_100"}),
            "giao_du":len({x["order"] for x in details if x["status"]=="FULL"}),
            "giao_mot_phan":len({x["order"] for x in details if x["status"]=="PARTIAL"}),
            "vo_100":len({x["order"] for x in details if x["status"]=="BROKEN_100"}),
            "khoi_luong_kg":round(sum(x["weight_kg"] for x in details),3),
            "so_tien_thuc_tinh":round(sum(x["amount"] for x in details),2)
        }

        success_details=[x for x in details if x["status"]!="BROKEN_100"]
        summary["d0"]=sum(1 for x in success_details if _gh_days_between(x.get("order_day"),x.get("day"))==0)
        summary["d1"]=sum(1 for x in success_details if _gh_days_between(x.get("order_day"),x.get("day"))==1)
        summary["d2"]=sum(1 for x in success_details if _gh_days_between(x.get("order_day"),x.get("day"))==2)
        summary["d3p"]=sum(
            1 for x in success_details
            if (_gh_days_between(x.get("order_day"),x.get("day")) is not None
                and _gh_days_between(x.get("order_day"),x.get("day"))>=3)
        )
        summary["pct_le_d2"]=round(
            (summary["d0"]+summary["d1"]+summary["d2"])/len(success_details)*100,2
        ) if success_details else 0

        return {"summary":summary,"rows":rows,"details":details}
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== GH KQ STABLE V1.0.3 DATE + WEIGHT FIX =====

def _gh_weight_v103(prod):
    return _gh_num(prod.get("TRONG_LUONG_SU"))

@router.get("/gh-kq/report-v103")
def gh_kq_report_v103(
    from_date:str=None,to_date:str=None,
    ma_nvgh:str=None,ma_lai_xe:str=None,phuong_tien:str=None
):
    try:
        # Start from V1.0.0 business/result logic.
        base=gh_kq_report_v100(from_date,to_date,ma_nvgh,ma_lai_xe,phuong_tien)

        products=_gh_products()
        c=_conn()

        # order -> locked date + draft id
        lock_map={}
        try:
            rows=c.execute("""
              SELECT d.id AS draft_id,
                     d.ma_don_nguon AS ma_don,
                     COALESCE(d.locked_at,s.locked_at) AS locked_at
              FROM qtgh_draft_order d
              LEFT JOIN qtgh_session_order so ON so.draft_order_id=d.id
              LEFT JOIN qtgh_session s ON s.id=so.session_id
              WHERE UPPER(COALESCE(d.status,''))='DA_KHOA'
            """).fetchall()
            for r in rows:
                order=str(r["ma_don"] or "").strip()
                if order:
                    lock_map[order]={
                      "draft_id":r["draft_id"],
                      "locked_day":_gh_date(r["locked_at"])
                    }
        except Exception:
            pass

        # Recompute detail weight and delivery progress date using locked_at.
        details=[]
        for x in base.get("details",[]):
            d=dict(x)
            order=d.get("order")
            lm=lock_map.get(order,{})
            draft_id=lm.get("draft_id")
            locked_day=lm.get("locked_day") or d.get("day")

            # Order day comes from source order date already added in v1.0.2 if available.
            order_day=d.get("order_day")
            if not order_day:
                # fallback from active raw source
                raw=_gh_active_raw_by_order().get(order,{})
                order_day=_gh_date(raw.get("NGAY_DON"))

            # Weight = actual delivered SU x TRONG_LUONG_SU
            weight=0.0
            if draft_id:
                try:
                    lrows=c.execute(
                        "SELECT * FROM qtgh_draft_line WHERE draft_order_id=?",
                        (draft_id,)
                    ).fetchall()
                    for lr in lrows:
                        ln=dict(lr)
                        sku=str(ln.get("ma_sku") or "").strip()
                        sl_thuc=_gh_num(_gh_line_value(
                            ln,"sl_hang_ban_thuc_giao","sl_thuc_giao"
                        ))
                        weight += sl_thuc * _gh_weight_v103(products.get(sku,{}))
                except Exception:
                    pass

            d["order_day"]=order_day
            d["day"]=locked_day
            d["weight_kg"]=round(weight,3)

            delta=_gh_days_between(order_day,locked_day)
            d["delivery_day_index"]=delta
            details.append(d)

        c.close()

        # Reaggregate all three dimensions from corrected details.
        staff=_gh_staff_map()
        groups=_ghdd(lambda:{
            "orders":set(),"success":set(),"full":set(),"partial":set(),"broken":set(),
            "customers":set(),"weight_kg":0.0,"amount":0.0,
            "d0":0,"d1":0,"d2":0,"d3p":0,"d_unknown":0
        })

        for d in details:
            order=d.get("order")
            status=d.get("status")
            kh=d.get("customer")
            delta=d.get("delivery_day_index")

            dims=[
                ("NVGH",d.get("ma_nvgh"),d.get("ten_nvgh") or ""),
                ("LAI_XE",d.get("ma_lai_xe"),d.get("ten_lai_xe") or ""),
                ("PHUONG_TIEN",d.get("phuong_tien"),d.get("phuong_tien") or "")
            ]

            for typ,key,name in dims:
                if not key: continue
                g=groups[(typ,key,name)]
                g["orders"].add(order)

                if status!="BROKEN_100":
                    g["success"].add(order)
                    if delta==0:g["d0"]+=1
                    elif delta==1:g["d1"]+=1
                    elif delta==2:g["d2"]+=1
                    elif delta is not None and delta>=3:g["d3p"]+=1
                    else:g["d_unknown"]+=1

                if status=="FULL":g["full"].add(order)
                elif status=="PARTIAL":g["partial"].add(order)
                else:g["broken"].add(order)

                if kh:g["customers"].add(kh)
                g["weight_kg"] += _gh_num(d.get("weight_kg"))
                g["amount"] += _gh_num(d.get("amount"))

        rows=[]
        for (typ,key,name),g in groups.items():
            success_count=len(g["success"])
            rows.append({
                "type":typ,"code":key,"name":name,
                "so_don_nhan":len(g["orders"]),
                "so_don_giao_tc":success_count,
                "giao_du":len(g["full"]),
                "giao_mot_phan":len(g["partial"]),
                "vo_100":len(g["broken"]),
                "so_kh_giao_tc":len(g["customers"]),
                "khoi_luong_kg":round(g["weight_kg"],3),
                "so_tien_thuc_tinh":round(g["amount"],2),
                "d0":g["d0"],"d1":g["d1"],"d2":g["d2"],"d3p":g["d3p"],
                "d_unknown":g["d_unknown"],
                "pct_le_d2":round(
                    (g["d0"]+g["d1"]+g["d2"])/success_count*100,2
                ) if success_count else 0
            })

        rows.sort(key=lambda x:(x["type"],x["code"]))

        success=[x for x in details if x.get("status")!="BROKEN_100"]
        summary={
            "so_don":len({x["order"] for x in details}),
            "so_don_giao_tc":len({x["order"] for x in success}),
            "giao_du":len({x["order"] for x in details if x["status"]=="FULL"}),
            "giao_mot_phan":len({x["order"] for x in details if x["status"]=="PARTIAL"}),
            "vo_100":len({x["order"] for x in details if x["status"]=="BROKEN_100"}),
            "khoi_luong_kg":round(sum(_gh_num(x.get("weight_kg")) for x in details),3),
            "so_tien_thuc_tinh":round(sum(_gh_num(x.get("amount")) for x in details),2),
            "d0":sum(1 for x in success if x.get("delivery_day_index")==0),
            "d1":sum(1 for x in success if x.get("delivery_day_index")==1),
            "d2":sum(1 for x in success if x.get("delivery_day_index")==2),
            "d3p":sum(1 for x in success if x.get("delivery_day_index") is not None and x.get("delivery_day_index")>=3),
        }
        summary["pct_le_d2"]=round(
            (summary["d0"]+summary["d1"]+summary["d2"])/len(success)*100,2
        ) if success else 0

        return {"summary":summary,"rows":rows,"details":details}
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== BX MISA V2 STABLE =====
import io as _bxio2, json as _bxjson2, datetime as _bxdt2
from openpyxl import load_workbook as _bxload2
from fastapi.responses import FileResponse as _BxFileResponse2

def _bx2_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS bx_misa_v2_batch(
      batch_id INTEGER PRIMARY KEY AUTOINCREMENT,
      source_type TEXT NOT NULL,
      file_name TEXT,
      status TEXT NOT NULL,
      total_rows INTEGER DEFAULT 0,
      valid_rows INTEGER DEFAULT 0,
      error_rows INTEGER DEFAULT 0,
      created_at TEXT,
      committed_at TEXT,
      superseded_by_batch_id INTEGER
    );
    CREATE TABLE IF NOT EXISTS bx_misa_v2_row(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      batch_id INTEGER NOT NULL,
      row_no INTEGER,
      ngay_chung_tu TEXT,
      so_chung_tu TEXT,
      ma_sku TEXT,
      ten_sku TEXT,
      dvt TEXT,
      so_luong REAL DEFAULT 0,
      trong_luong_dv REAL DEFAULT 0,
      khoi_luong_kg REAL DEFAULT 0,
      ma_kh TEXT,
      ten_kh TEXT,
      payload_json TEXT,
      error_text TEXT
    );
    """)
    c.commit(); c.close()

_bx2_init()

def _bx2_date(v):
    if v in (None,""): return ""
    if isinstance(v,_bxdt2.datetime): return v.date().isoformat()
    if isinstance(v,_bxdt2.date): return v.isoformat()
    if isinstance(v,(int,float)):
        try:
            base=_bxdt2.datetime(1899,12,30)
            return (base+_bxdt2.timedelta(days=float(v))).date().isoformat()
        except Exception: pass
    s=str(v).strip()
    for f in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y"):
        try:return _bxdt2.datetime.strptime(s[:10],f).date().isoformat()
        except Exception:pass
    return ""

def _bx2_num(v):
    try:return float(v or 0)
    except:return 0.0

def _bx2_products():
    c=_conn(); out={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm01-product'"):
        try:p=_bxjson2.loads(r["payload_json"] or "{}")
        except Exception:p={}
        sku=str(p.get("MA_SKU") or "").strip()
        if sku:out[sku]=p
    c.close();return out

def _bx2_weight(prod,dvt):
    d=str(dvt or "").strip().lower()
    if "thùng" in d or "thung" in d:
        return _bx2_num(prod.get("TRONG_LUONG_THUNG"))
    return _bx2_num(prod.get("TRONG_LUONG_SU"))

def _bx2_parse(raw,source_type,file_name=""):
    wb=_bxload2(_bxio2.BytesIO(raw),data_only=True)

    if source_type=="BOC_VAO":
        ws=wb["SỔ CHI TIẾT MUA HÀNG"] if "SỔ CHI TIẾT MUA HÀNG" in wb.sheetnames else wb[wb.sheetnames[0]]
        header_row=None
        for r in range(1,min(ws.max_row,15)+1):
            vals=[str(ws.cell(r,c).value or "").strip() for c in range(1,ws.max_column+1)]
            if "Số chứng từ" in vals and "Số lượng mua" in vals:
                header_row=r;break
        qty_col_name="Số lượng mua"
        kh_code_name=None;kh_name_name=None
    else:
        ws=wb["SỔ CHI TIẾT BÁN HÀNG"] if "SỔ CHI TIẾT BÁN HÀNG" in wb.sheetnames else wb[wb.sheetnames[0]]
        header_row=None
        for r in range(1,min(ws.max_row,10)+1):
            vals=[str(ws.cell(r,c).value or "").strip() for c in range(1,ws.max_column+1)]
            if "Số chứng từ" in vals and "Tổng số lượng bán" in vals:
                header_row=r;break
        qty_col_name="Tổng số lượng bán"
        kh_code_name="Mã khách hàng";kh_name_name="Tên khách hàng"

    if not header_row:
        raise ValueError("Không nhận diện được header file MISA theo mẫu V2.")

    headers={}
    for c in range(1,ws.max_column+1):
        h=str(ws.cell(header_row,c).value or "").strip()
        if h:headers[h]=c

    required=["Ngày chứng từ","Số chứng từ","Mã hàng","Tên hàng","ĐVT",qty_col_name]
    miss=[x for x in required if x not in headers]
    if miss:raise ValueError("Thiếu cột: "+", ".join(miss))

    products=_bx2_products()
    rows=[]; errors=[]
    for r in range(header_row+1,ws.max_row+1):
        def V(name):
            col=headers.get(name)
            return ws.cell(r,col).value if col else None

        doc=str(V("Số chứng từ") or "").strip()
        sku=str(V("Mã hàng") or "").strip()
        # Skip blank and total rows.
        if not doc and not sku:continue
        if str(V("Ngày chứng từ") or "").strip().lower().startswith("tổng"):
            continue
        if str(V("Ngày hạch toán") or "").strip().lower().startswith("tổng"):
            continue

        qty=_bx2_num(V(qty_col_name))
        dvt=str(V("ĐVT") or "").strip()
        prod=products.get(sku,{})
        w=_bx2_weight(prod,dvt)
        err=[]
        if not doc:err.append("Thiếu Số chứng từ")
        if not sku:err.append("Thiếu Mã hàng")
        elif not prod:err.append("SKU chưa có trong DM01")
        if qty<0:err.append("Số lượng âm")
        if prod and w<=0:err.append("SKU thiếu trọng lượng DM01 theo ĐVT")

        rec={
          "row_no":r,
          "ngay_chung_tu":_bx2_date(V("Ngày chứng từ")),
          "so_chung_tu":doc,
          "ma_sku":sku,
          "ten_sku":str(V("Tên hàng") or "").strip(),
          "dvt":dvt,
          "so_luong":qty,
          "trong_luong_dv":w,
          "khoi_luong_kg":round(qty*w,6),
          "ma_kh":str(V(kh_code_name) or "").strip() if kh_code_name else "",
          "ten_kh":str(V(kh_name_name) or "").strip() if kh_name_name else "",
          "errors":err
        }
        rows.append(rec)
        if err:errors.append(rec)

    # Group trips = one voucher/document
    trips={}
    for x in rows:
        doc=x["so_chung_tu"]
        if not doc:continue
        g=trips.setdefault(doc,{
          "so_chung_tu":doc,"ngay_chung_tu":x["ngay_chung_tu"],
          "source_type":source_type,"so_sku":0,"tong_sl":0.0,"tong_kg":0.0,
          "mapping_ok":True
        })
        g["so_sku"]+=1
        g["tong_sl"]+=x["so_luong"]
        g["tong_kg"]+=x["khoi_luong_kg"]
        if x["errors"]:g["mapping_ok"]=False

    return {
      "file_name":file_name,
      "source_type":source_type,
      "rows":rows,
      "trips":list(trips.values()),
      "total_rows":len(rows),
      "valid_rows":sum(1 for x in rows if not x["errors"]),
      "error_rows":sum(1 for x in rows if x["errors"])
    }

@router.get("/bx-v2/template/{source}")
def bx_v2_template(source:str):
    source=source.lower()
    base=Path(__file__).resolve().parent.parent/"data"/"templates"
    if source=="boc-vao":
        p=base/"Template_MISA_Boc_vao_V2.xlsx"
    elif source=="boc-ra":
        p=base/"Template_MISA_Boc_ra_Ban_buon_V2.xlsx"
    else:
        raise HTTPException(404,"Nguồn không hợp lệ")
    if not p.exists():raise HTTPException(404,"Chưa cài template BX V2")
    return _BxFileResponse2(str(p),filename=p.name)

@router.post("/bx-v2/preview/{source}")
async def bx_v2_preview(source:str,request:Request,filename:str=""):
    try:
        st="BOC_VAO" if source.lower()=="boc-vao" else "BOC_RA_BAN_BUON" if source.lower()=="boc-ra" else None
        if not st:raise ValueError("Nguồn không hợp lệ")
        raw=await request.body()
        if not raw:raise ValueError("File rỗng")
        parsed=_bx2_parse(raw,st,filename)

        c=_conn()
        cur=c.execute("""
          INSERT INTO bx_misa_v2_batch(source_type,file_name,status,total_rows,valid_rows,error_rows,created_at)
          VALUES(?,?,?,?,?,?,?)
        """,(st,filename,"PREVIEW",parsed["total_rows"],parsed["valid_rows"],parsed["error_rows"],_now()))
        bid=cur.lastrowid

        for x in parsed["rows"]:
            payload=dict(x); payload.pop("errors",None)
            c.execute("""
              INSERT INTO bx_misa_v2_row(
                batch_id,row_no,ngay_chung_tu,so_chung_tu,ma_sku,ten_sku,dvt,
                so_luong,trong_luong_dv,khoi_luong_kg,ma_kh,ten_kh,payload_json,error_text
              ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,(
              bid,x["row_no"],x["ngay_chung_tu"],x["so_chung_tu"],x["ma_sku"],x["ten_sku"],x["dvt"],
              x["so_luong"],x["trong_luong_dv"],x["khoi_luong_kg"],x["ma_kh"],x["ten_kh"],
              _bxjson2.dumps(payload,ensure_ascii=False),"; ".join(x["errors"])
            ))
        c.commit();c.close()

        return {
          "batch_id":bid,
          "source_type":st,
          "total_rows":parsed["total_rows"],
          "valid_rows":parsed["valid_rows"],
          "error_rows":parsed["error_rows"],
          "trips":parsed["trips"],
          "errors":[{"row_no":x["row_no"],"error":"; ".join(x["errors"])} for x in parsed["rows"] if x["errors"]][:100]
        }
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post("/bx-v2/batch/{batch_id}/commit")
def bx_v2_commit(batch_id:int):
    try:
        c=_conn()
        b=c.execute("SELECT * FROM bx_misa_v2_batch WHERE batch_id=?",(batch_id,)).fetchone()
        if not b:raise ValueError("Không tìm thấy batch")
        if b["error_rows"]>0:raise ValueError("Batch còn lỗi mapping, chưa thể Commit")
        if b["status"]=="COMMITTED":
            c.close();return {"batch_id":batch_id,"status":"COMMITTED","already":True}

        prev=c.execute("""
          SELECT batch_id FROM bx_misa_v2_batch
          WHERE source_type=? AND status='COMMITTED' AND superseded_by_batch_id IS NULL
          ORDER BY batch_id DESC LIMIT 1
        """,(b["source_type"],)).fetchone()

        if prev:
            c.execute(
              "UPDATE bx_misa_v2_batch SET superseded_by_batch_id=? WHERE batch_id=?",
              (batch_id,prev["batch_id"])
            )

        c.execute(
          "UPDATE bx_misa_v2_batch SET status='COMMITTED',committed_at=? WHERE batch_id=?",
          (_now(),batch_id)
        )
        c.commit();c.close()
        return {"batch_id":batch_id,"status":"COMMITTED"}
    except Exception as e:
        raise HTTPException(400,str(e))

@router.get("/bx-v2/trips")
def bx_v2_trips(from_date:str=None,to_date:str=None,source_type:str=None):
    try:
        c=_conn()
        batches=c.execute("""
          SELECT batch_id,source_type FROM bx_misa_v2_batch
          WHERE status='COMMITTED' AND superseded_by_batch_id IS NULL
        """).fetchall()
        trips=[]
        for b in batches:
            if source_type and b["source_type"]!=source_type:continue
            rows=c.execute("""
              SELECT * FROM bx_misa_v2_row WHERE batch_id=? ORDER BY row_no
            """,(b["batch_id"],)).fetchall()
            groups={}
            for rr in rows:
                r=dict(rr)
                day=r["ngay_chung_tu"] or ""
                if from_date and day<from_date:continue
                if to_date and day>to_date:continue
                doc=r["so_chung_tu"]
                g=groups.setdefault(doc,{
                  "so_chung_tu":doc,"ngay_chung_tu":day,"source_type":b["source_type"],
                  "so_sku":0,"tong_sl":0.0,"tong_kg":0.0,"mapping_ok":True,"status":"CHO_XAC_NHAN"
                })
                g["so_sku"]+=1
                g["tong_sl"]+=_bx2_num(r["so_luong"])
                g["tong_kg"]+=_bx2_num(r["khoi_luong_kg"])
                if r["error_text"]:g["mapping_ok"]=False
            trips.extend(groups.values())
        c.close()
        trips.sort(key=lambda x:(x["ngay_chung_tu"],x["so_chung_tu"]))
        return {
          "trips":trips,
          "summary":{
            "tong_chuyen":len(trips),
            "boc_vao":sum(1 for x in trips if x["source_type"]=="BOC_VAO"),
            "boc_ra":sum(1 for x in trips if x["source_type"]=="BOC_RA_BAN_BUON"),
            "tong_kg":round(sum(x["tong_kg"] for x in trips),3)
          }
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== BX PACK V3 STABLE =====
import json as _bx3json, datetime as _bx3dt
from collections import defaultdict as _bx3dd

def _bx3_init():
    c=_conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS bx_trip_class_v3(
      source_type TEXT NOT NULL,
      so_chung_tu TEXT NOT NULL,
      ngay_chung_tu TEXT,
      loai_boc TEXT NOT NULL,
      confirmed_at TEXT,
      confirmed_by TEXT,
      PRIMARY KEY(source_type,so_chung_tu)
    );
    CREATE TABLE IF NOT EXISTS bx_pack_v3(
      pack_id INTEGER PRIMARY KEY AUTOINCREMENT,
      pack_code TEXT UNIQUE NOT NULL,
      pack_date TEXT NOT NULL,
      mode TEXT NOT NULL DEFAULT 'INTERNAL',
      status TEXT NOT NULL DEFAULT 'DRAFT',
      price_in_chan REAL NOT NULL DEFAULT 20000,
      price_in_le REAL NOT NULL DEFAULT 40000,
      price_out REAL NOT NULL DEFAULT 20000,
      price_change_reason TEXT,
      external_name TEXT,
      created_at TEXT,
      created_by TEXT,
      confirmed_at TEXT,
      confirmed_by TEXT
    );
    CREATE TABLE IF NOT EXISTS bx_pack_trip_v3(
      pack_id INTEGER NOT NULL,
      source_type TEXT NOT NULL,
      so_chung_tu TEXT NOT NULL,
      ngay_chung_tu TEXT,
      loai_boc TEXT NOT NULL,
      kg REAL NOT NULL DEFAULT 0,
      PRIMARY KEY(pack_id,source_type,so_chung_tu)
    );
    CREATE TABLE IF NOT EXISTS bx_pack_member_v3(
      pack_id INTEGER NOT NULL,
      ma_nv TEXT NOT NULL,
      PRIMARY KEY(pack_id,ma_nv)
    );
    CREATE TABLE IF NOT EXISTS bx_pack_alloc_v3(
      pack_id INTEGER NOT NULL,
      participant_key TEXT NOT NULL,
      participant_type TEXT NOT NULL,
      ma_nv TEXT,
      loai_boc TEXT NOT NULL,
      kg REAL NOT NULL DEFAULT 0,
      don_gia REAL NOT NULL DEFAULT 0,
      thanh_tien REAL NOT NULL DEFAULT 0,
      PRIMARY KEY(pack_id,participant_key,loai_boc)
    );
    """)
    c.commit(); c.close()

_bx3_init()

def _bx3_num(v):
    try:return float(v or 0)
    except:return 0.0

def _bx3_staff():
    c=_conn(); out={}
    for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm02-staff'"):
        try:p=_bx3json.loads(r["payload_json"] or "{}")
        except Exception:p={}
        ma=str(p.get("MA_NV") or "").strip()
        if ma:out[ma]=p
    c.close(); return out

def _bx3_current_trips():
    data=bx_v2_trips()
    trips=data.get("trips",[])
    c=_conn()
    classes={(r["source_type"],r["so_chung_tu"]):dict(r) for r in c.execute("SELECT * FROM bx_trip_class_v3")}
    packed={(r["source_type"],r["so_chung_tu"]) for r in c.execute(
      "SELECT pt.source_type,pt.so_chung_tu FROM bx_pack_trip_v3 pt JOIN bx_pack_v3 p ON p.pack_id=pt.pack_id WHERE UPPER(COALESCE(p.status,''))<>'CANCELLED'"
    )}
    c.close()
    out=[]
    for t in trips:
        x=dict(t); key=(x["source_type"],x["so_chung_tu"]); cl=classes.get(key)
        if x["source_type"]=="BOC_RA_BAN_BUON": x["loai_boc"]="BOC_RA_BAN_BUON"
        else: x["loai_boc"]=(cl or {}).get("loai_boc") or ""
        x["confirmed"]=bool(cl and cl.get("confirmed_at"))
        x["packed"]=key in packed
        out.append(x)
    return out

def _bx3_pack_code(day):
    c=_conn(); prefix="NO"+str(day).replace("-","")
    n=c.execute("SELECT COUNT(*) FROM bx_pack_v3 WHERE pack_code LIKE ?",(prefix+".%",)).fetchone()[0]+1
    c.close(); return f"{prefix}.{n:03d}"

def _bx3_totals(pid):
    c=_conn()
    out={"BOC_VAO_CHAN":0.0,"BOC_VAO_LE":0.0,"BOC_RA_BAN_BUON":0.0}
    for r in c.execute("SELECT loai_boc,SUM(kg) kg FROM bx_pack_trip_v3 WHERE pack_id=? GROUP BY loai_boc",(pid,)):
        out[r["loai_boc"]]=_bx3_num(r["kg"])
    c.close(); return out

def _bx3_allocate(pid):
    c=_conn(); p=c.execute("SELECT * FROM bx_pack_v3 WHERE pack_id=?",(pid,)).fetchone()
    if not p: c.close(); raise ValueError("Không tìm thấy PACK")
    totals=_bx3_totals(pid)
    prices={"BOC_VAO_CHAN":_bx3_num(p["price_in_chan"]),"BOC_VAO_LE":_bx3_num(p["price_in_le"]),"BOC_RA_BAN_BUON":_bx3_num(p["price_out"])}
    c.execute("DELETE FROM bx_pack_alloc_v3 WHERE pack_id=?",(pid,))
    if p["mode"]=="EXTERNAL":
        participants=[("EXTERNAL","EXTERNAL",None)]
    else:
        members=[r["ma_nv"] for r in c.execute("SELECT ma_nv FROM bx_pack_member_v3 WHERE pack_id=? ORDER BY ma_nv",(pid,))]
        participants=[(m,"INTERNAL",m) for m in members]
    if not participants: c.commit(); c.close(); return
    for loai,total in totals.items():
        if total<=0: continue
        n=len(participants); allocated=0.0
        for i,(key,ptype,ma) in enumerate(participants):
            kg=round(total-allocated,3) if i==n-1 else round(total/n,3)
            allocated=round(allocated+kg,3)
            price=prices[loai]; amount=round(kg/1000*price,0)
            c.execute("INSERT INTO bx_pack_alloc_v3(pack_id,participant_key,participant_type,ma_nv,loai_boc,kg,don_gia,thanh_tien) VALUES(?,?,?,?,?,?,?,?)",
                      (pid,key,ptype,ma,loai,kg,price,amount))
    c.commit(); c.close()

@router.get("/bx-pack-v3/staff")
def bx3_staff():
    s=_bx3_staff()
    return {"rows":[{"ma_nv":k,"ten_nv":v.get("TEN_NV") or ""} for k,v in sorted(s.items())]}

@router.get("/bx-pack-v3/trips")
def bx3_trips():
    return {"rows":_bx3_current_trips()}

@router.post("/bx-pack-v3/trip/confirm")
async def bx3_confirm_trip(request:Request):
    try:
        d=await request.json(); st=str(d.get("source_type") or ""); doc=str(d.get("so_chung_tu") or "")
        day=str(d.get("ngay_chung_tu") or "")[:10]; loai=str(d.get("loai_boc") or "")
        if st=="BOC_RA_BAN_BUON": loai="BOC_RA_BAN_BUON"
        elif loai not in ("BOC_VAO_CHAN","BOC_VAO_LE"): raise ValueError("Bốc vào phải chọn Chẵn hoặc Lẻ")
        c=_conn()
        c.execute("""
        INSERT INTO bx_trip_class_v3(source_type,so_chung_tu,ngay_chung_tu,loai_boc,confirmed_at,confirmed_by)
        VALUES(?,?,?,?,?,?)
        ON CONFLICT(source_type,so_chung_tu) DO UPDATE SET
          ngay_chung_tu=excluded.ngay_chung_tu,loai_boc=excluded.loai_boc,
          confirmed_at=excluded.confirmed_at,confirmed_by=excluded.confirmed_by
        """,(st,doc,day,loai,_now(),"pilot"))
        c.commit(); c.close(); return {"ok":True}
    except Exception as e: raise HTTPException(400,str(e))

@router.post("/bx-pack-v3/create")
async def bx3_create(request:Request):
    try:
        d=await request.json(); picks=d.get("trips") or []
        if not picks: raise ValueError("Chưa chọn chuyến")
        cur={(x["source_type"],x["so_chung_tu"]):x for x in _bx3_current_trips()}
        selected=[]; dates=set()
        for t in picks:
            x=cur.get((t.get("source_type"),t.get("so_chung_tu")))
            if not x: raise ValueError("Không tìm thấy chuyến")
            if not x["confirmed"]: raise ValueError(f"{x['so_chung_tu']} chưa xác nhận")
            if x["packed"]: raise ValueError(f"{x['so_chung_tu']} đã thuộc PACK")
            selected.append(x); dates.add(x["ngay_chung_tu"])
        if len(dates)!=1: raise ValueError("Một PACK chỉ gom chuyến cùng ngày")
        day=next(iter(dates)); code=_bx3_pack_code(day)
        c=_conn()
        pid=c.execute("INSERT INTO bx_pack_v3(pack_code,pack_date,mode,status,created_at,created_by) VALUES(?,?,'INTERNAL','DRAFT',?,?)",
                      (code,day,_now(),"pilot")).lastrowid
        for x in selected:
            c.execute("INSERT INTO bx_pack_trip_v3(pack_id,source_type,so_chung_tu,ngay_chung_tu,loai_boc,kg) VALUES(?,?,?,?,?,?)",
                      (pid,x["source_type"],x["so_chung_tu"],x["ngay_chung_tu"],x["loai_boc"],_bx3_num(x["tong_kg"])))
        c.commit(); c.close(); return {"pack_id":pid,"pack_code":code}
    except Exception as e: raise HTTPException(400,str(e))

@router.get("/bx-pack-v3/packs")
def bx3_packs():
    c=_conn(); staff=_bx3_staff(); rows=[]
    for p in c.execute("SELECT * FROM bx_pack_v3 ORDER BY pack_date DESC,pack_id DESC"):
        d=dict(p); pid=d["pack_id"]
        d["trips"]=[dict(r) for r in c.execute("SELECT * FROM bx_pack_trip_v3 WHERE pack_id=? ORDER BY so_chung_tu",(pid,))]
        d["members"]=[r["ma_nv"] for r in c.execute("SELECT ma_nv FROM bx_pack_member_v3 WHERE pack_id=? ORDER BY ma_nv",(pid,))]
        d["alloc"]=[dict(r) for r in c.execute("SELECT * FROM bx_pack_alloc_v3 WHERE pack_id=? ORDER BY participant_key,loai_boc",(pid,))]
        d["totals"]=_bx3_totals(pid); rows.append(d)
    c.close(); return {"rows":rows}

@router.post("/bx-pack-v3/{pid}/setup")
async def bx3_setup(pid:int,request:Request):
    try:
        d=await request.json(); mode=str(d.get("mode") or "INTERNAL")
        pc=_bx3_num(d.get("price_in_chan") or 20000); pl=_bx3_num(d.get("price_in_le") or 40000); po=_bx3_num(d.get("price_out") or 20000)
        reason=str(d.get("price_change_reason") or "").strip()
        if (pc!=20000 or pl!=40000 or po!=20000) and not reason: raise ValueError("Đơn giá khác mặc định phải nhập lý do")
        members=d.get("members") or []
        if mode=="INTERNAL" and not members: raise ValueError("PACK nội bộ phải chọn nhân sự")
        c=_conn()
        c.execute("UPDATE bx_pack_v3 SET mode=?,price_in_chan=?,price_in_le=?,price_out=?,price_change_reason=?,external_name=? WHERE pack_id=?",
                  (mode,pc,pl,po,reason,str(d.get("external_name") or "").strip(),pid))
        c.execute("DELETE FROM bx_pack_member_v3 WHERE pack_id=?",(pid,))
        if mode=="INTERNAL":
            for m in members: c.execute("INSERT INTO bx_pack_member_v3(pack_id,ma_nv) VALUES(?,?)",(pid,str(m)))
        c.commit(); c.close(); _bx3_allocate(pid); return {"ok":True}
    except Exception as e: raise HTTPException(400,str(e))

@router.post("/bx-pack-v3/{pid}/confirm")
def bx3_confirm(pid:int):
    try:
        c=_conn()
        n=c.execute("SELECT COUNT(*) FROM bx_pack_alloc_v3 WHERE pack_id=?",(pid,)).fetchone()[0]
        if n<=0: raise ValueError("PACK chưa phân bổ")
        c.execute("UPDATE bx_pack_v3 SET status='CONFIRMED',confirmed_at=?,confirmed_by=? WHERE pack_id=?",(_now(),"pilot",pid))
        c.commit(); c.close(); return {"ok":True}
    except Exception as e: raise HTTPException(400,str(e))



# ===== BX PACK V3.7 REPORT =====
@router.get("/bx-pack-v3/report-v37")
def bx_pack_report_v37(from_date:str=None,to_date:str=None,month:str=None):
    try:
        import calendar
        if month:
            y,m=map(int,month.split("-"))
            from_date=f"{month}-01"
            to_date=f"{month}-{calendar.monthrange(y,m)[1]:02d}"

        c=_conn()

        # Staff master
        staff={}
        for r in c.execute("SELECT payload_json FROM master_data_generic WHERE category='dm02-staff'"):
            try:p=_bx3json.loads(r["payload_json"] or "{}")
            except Exception:p={}
            ma=str(p.get("MA_NV") or "").strip()
            if ma: staff[ma]=str(p.get("TEN_NV") or p.get("HO_TEN") or "").strip()

        q="SELECT * FROM bx_pack_v3 WHERE status='CONFIRMED'"
        params=[]
        if from_date:
            q+=" AND pack_date>=?"
            params.append(from_date)
        if to_date:
            q+=" AND pack_date<=?"
            params.append(to_date)
        q+=" ORDER BY pack_date,pack_id"

        packs=[dict(r) for r in c.execute(q,params).fetchall()]

        people={}
        external_rows=[]
        pack_rows=[]

        for p in packs:
            pid=p["pack_id"]
            alloc=[dict(r) for r in c.execute(
                "SELECT * FROM bx_pack_alloc_v3 WHERE pack_id=? ORDER BY participant_key,loai_boc",
                (pid,)
            ).fetchall()]
            trips=[dict(r) for r in c.execute(
                "SELECT * FROM bx_pack_trip_v3 WHERE pack_id=? ORDER BY so_chung_tu",
                (pid,)
            ).fetchall()]

            sums={
              "kg_in_chan":0.0,"kg_in_le":0.0,"kg_out":0.0,
              "money_in_chan":0.0,"money_in_le":0.0,"money_out":0.0
            }

            for a in alloc:
                kg=_bx3_num(a.get("kg"))
                amt=_bx3_num(a.get("thanh_tien"))
                loai=a.get("loai_boc")

                if loai=="BOC_VAO_CHAN":
                    sums["kg_in_chan"]+=kg; sums["money_in_chan"]+=amt
                elif loai=="BOC_VAO_LE":
                    sums["kg_in_le"]+=kg; sums["money_in_le"]+=amt
                elif loai=="BOC_RA_BAN_BUON":
                    sums["kg_out"]+=kg; sums["money_out"]+=amt

                if a.get("participant_type")=="INTERNAL":
                    ma=str(a.get("ma_nv") or "").strip()
                    if not ma: continue
                    g=people.setdefault(ma,{
                      "ma_nv":ma,"ten_nv":staff.get(ma,""),
                      "packs":set(),
                      "kg_in_chan":0.0,"kg_in_le":0.0,"kg_out":0.0,
                      "money_in_chan":0.0,"money_in_le":0.0,"money_out":0.0
                    })
                    g["packs"].add(pid)
                    if loai=="BOC_VAO_CHAN":
                        g["kg_in_chan"]+=kg; g["money_in_chan"]+=amt
                    elif loai=="BOC_VAO_LE":
                        g["kg_in_le"]+=kg; g["money_in_le"]+=amt
                    elif loai=="BOC_RA_BAN_BUON":
                        g["kg_out"]+=kg; g["money_out"]+=amt

            if p.get("mode")=="EXTERNAL":
                external_rows.append({
                  "pack_id":pid,
                  "pack_code":p.get("pack_code"),
                  "pack_date":p.get("pack_date"),
                  "external_name":p.get("external_name") or "Thuê ngoài",
                  "trip_count":len(trips),
                  "kg_in_chan":round(sums["kg_in_chan"],3),
                  "kg_in_le":round(sums["kg_in_le"],3),
                  "kg_out":round(sums["kg_out"],3),
                  "kg_total":round(sums["kg_in_chan"]+sums["kg_in_le"]+sums["kg_out"],3),
                  "money_in_chan":round(sums["money_in_chan"],0),
                  "money_in_le":round(sums["money_in_le"],0),
                  "money_out":round(sums["money_out"],0),
                  "amount_total":round(sums["money_in_chan"]+sums["money_in_le"]+sums["money_out"],0)
                })

            pack_rows.append({
              "pack_id":pid,
              "pack_code":p.get("pack_code"),
              "pack_date":p.get("pack_date"),
              "mode":p.get("mode"),
              "external_name":p.get("external_name"),
              "trip_count":len(trips),
              "trips":", ".join(str(x.get("so_chung_tu") or "") for x in trips),
              "kg_in_chan":round(sums["kg_in_chan"],3),
              "kg_in_le":round(sums["kg_in_le"],3),
              "kg_out":round(sums["kg_out"],3),
              "kg_total":round(sums["kg_in_chan"]+sums["kg_in_le"]+sums["kg_out"],3),
              "money_in_chan":round(sums["money_in_chan"],0),
              "money_in_le":round(sums["money_in_le"],0),
              "money_out":round(sums["money_out"],0),
              "amount_total":round(sums["money_in_chan"]+sums["money_in_le"]+sums["money_out"],0)
            })

        people_rows=[]
        for ma,g in people.items():
            people_rows.append({
              "ma_nv":ma,"ten_nv":g["ten_nv"],"pack_count":len(g["packs"]),
              "kg_in_chan":round(g["kg_in_chan"],3),
              "kg_in_le":round(g["kg_in_le"],3),
              "kg_out":round(g["kg_out"],3),
              "kg_total":round(g["kg_in_chan"]+g["kg_in_le"]+g["kg_out"],3),
              "money_in_chan":round(g["money_in_chan"],0),
              "money_in_le":round(g["money_in_le"],0),
              "money_out":round(g["money_out"],0),
              "income_total":round(g["money_in_chan"]+g["money_in_le"]+g["money_out"],0)
            })
        people_rows.sort(key=lambda x:x["ma_nv"])

        summary={
          "pack_count":len(pack_rows),
          "internal_pack_count":sum(1 for x in pack_rows if x["mode"]=="INTERNAL"),
          "external_pack_count":sum(1 for x in pack_rows if x["mode"]=="EXTERNAL"),
          "kg_total":round(sum(x["kg_total"] for x in pack_rows),3),
          "amount_total":round(sum(x["amount_total"] for x in pack_rows),0),
          "internal_income_total":round(sum(x["income_total"] for x in people_rows),0),
          "external_amount_total":round(sum(x["amount_total"] for x in external_rows),0)
        }

        c.close()
        return {
          "summary":summary,
          "people":people_rows,
          "external":external_rows,
          "packs":pack_rows
        }
    except Exception as e:
        raise HTTPException(400,str(e))



# ===== QT UI01 STABLE V1 =====
@router.get("/admin/dashboard-v1")
def admin_dashboard_v1(from_date:str=None,to_date:str=None,month:str=None):
    try:
        import calendar
        if month:
            y,m=map(int,month.split("-"))
            from_date=f"{month}-01"
            to_date=f"{month}-{calendar.monthrange(y,m)[1]:02d}"

        c=_conn()

        def period(col):
            clauses=[]; params=[]
            if from_date:
                clauses.append(f"substr({col},1,10)>=?"); params.append(from_date)
            if to_date:
                clauses.append(f"substr({col},1,10)<=?"); params.append(to_date)
            return (" AND " + " AND ".join(clauses) if clauses else ""), tuple(params)

        def scalar(sql, params=()):
            r=c.execute(sql,params).fetchone()
            if not r:return 0
            v=list(r)[0]
            try:return float(v or 0)
            except:return v or 0

        def count_status(table,status_col,statuses):
            if not statuses:return 0
            q=",".join("?" for _ in statuses)
            return int(scalar(f'SELECT COUNT(*) FROM "{table}" WHERE UPPER(COALESCE("{status_col}",\'\')) IN ({q})',
                              tuple(str(x).upper() for x in statuses)))

        # -----------------------------
        # QTGH / CASH / DEBT
        # -----------------------------
        order_where,order_params=period('locked_at')
        debt_where,debt_params=period('created_at')
        receipt_where,receipt_params=period('created_at')
        locked_orders=int(scalar("SELECT COUNT(*) FROM qtgh_draft_order WHERE status='DA_KHOA'"+order_where,order_params))
        draft_total=float(scalar("SELECT COALESCE(SUM(so_phai_thu_sau_qt),0) FROM qtgh_draft_order WHERE status='DA_KHOA'"+order_where,order_params))
        cash_total=float(scalar("SELECT COALESCE(SUM(tien_mat),0) FROM qtgh_draft_order WHERE status='DA_KHOA'"+order_where,order_params))
        bank_total=float(scalar("SELECT COALESCE(SUM(tien_ck_xac_nhan),0) FROM qtgh_draft_order WHERE status='DA_KHOA'"+order_where,order_params))

        customer_debt=float(scalar("SELECT COALESCE(SUM(balance),0) FROM qtgh_customer_debt WHERE UPPER(COALESCE(status,''))<>'DA_THU_DU'"+debt_where,debt_params))
        employee_debt=float(scalar("SELECT COALESCE(SUM(balance),0) FROM qtgh_employee_debt WHERE UPPER(COALESCE(status,''))='CON_NO'"+debt_where,debt_params))
        pending_cash=int(scalar("SELECT COUNT(*) FROM qtgh_cash_receipt WHERE status='CHO_THU'"+receipt_where,receipt_params))
        received_cash=int(scalar("SELECT COUNT(*) FROM qtgh_cash_receipt WHERE status IN ('DA_THU_DU','THU_THIEU','DA_THU_THIEU')"+receipt_where,receipt_params))

        # -----------------------------
        # XB SOURCE / IMPORT
        # -----------------------------
        xb_dms_rows=int(scalar("""
          SELECT COUNT(*) FROM xb_dms_mbw_row r
          JOIN xb_dms_mbw_batch b ON b.batch_id=r.batch_id
          WHERE b.status='COMMITTED' AND b.superseded_by_batch_id IS NULL
        """))
        xb_misa_rows=int(scalar("""
          SELECT COUNT(*) FROM xb_misa_bu_row r
          JOIN xb_misa_bu_batch b ON b.batch_id=r.batch_id
          WHERE b.status='COMMITTED' AND b.superseded_by_batch_id IS NULL
        """))
        xb_preview=int(scalar("""
          SELECT
            (SELECT COUNT(*) FROM xb_dms_mbw_batch WHERE status='PREVIEW') +
            (SELECT COUNT(*) FROM xb_misa_bu_batch WHERE status='PREVIEW') +
            (SELECT COUNT(*) FROM xb_src_batch_v2 WHERE status='PREVIEW')
        """))
        xb_error_rows=int(scalar("""
          SELECT
            (SELECT COUNT(*) FROM xb_dms_mbw_row WHERE status='ERROR') +
            (SELECT COUNT(*) FROM xb_misa_bu_row WHERE status='ERROR') +
            (SELECT COUNT(*) FROM xb_src_row_v2 WHERE status='ERROR')
        """))

        # -----------------------------
        # EXCEPTIONS
        # -----------------------------
        qt_nhap=int(scalar("SELECT COUNT(*) FROM qtgh_ui04_draft_trace WHERE status='QT_NHAP'"))
        cho_kho=int(scalar("SELECT COUNT(*) FROM qtgh_return_note WHERE status='CHO_KHO'"))
        emp_debt_count=int(scalar("SELECT COUNT(*) FROM qtgh_employee_debt WHERE status='CON_NO'"))
        cust_debt_count=int(scalar("SELECT COUNT(*) FROM qtgh_customer_debt WHERE UPPER(COALESCE(status,''))<>'DA_THU_DU'"))
        adjustment_pending=int(scalar("""
          SELECT COUNT(*) FROM qtgh_adjustment_stable
          WHERE UPPER(COALESCE(status,'')) NOT IN ('APPROVED','DA_DUYET','DONE','COMPLETED')
        """))
        bx_draft=int(scalar("SELECT COUNT(*) FROM bx_pack_v3 WHERE status='DRAFT'"))

        # -----------------------------
        # BX SUMMARY from confirmed packs
        # -----------------------------
        bx_pack_count=int(scalar("SELECT COUNT(*) FROM bx_pack_v3 WHERE status='CONFIRMED'"))
        bx_internal=int(scalar("SELECT COUNT(*) FROM bx_pack_v3 WHERE status='CONFIRMED' AND mode='INTERNAL'"))
        bx_external=int(scalar("SELECT COUNT(*) FROM bx_pack_v3 WHERE status='CONFIRMED' AND mode='EXTERNAL'"))
        bx_kg=float(scalar("""
          SELECT COALESCE(SUM(pt.kg),0)
          FROM bx_pack_trip_v3 pt
          JOIN bx_pack_v3 p ON p.pack_id=pt.pack_id
          WHERE p.status='CONFIRMED'
        """))
        bx_money=float(scalar("""
          SELECT COALESCE(SUM(a.thanh_tien),0)
          FROM bx_pack_alloc_v3 a
          JOIN bx_pack_v3 p ON p.pack_id=a.pack_id
          WHERE p.status='CONFIRMED'
        """))

        c.close()

        return {
          "period":{"from_date":from_date,"to_date":to_date,"month":month},
          "qtgh":{
            "locked_orders":locked_orders,
            "amount_due":round(draft_total,0),
            "cash":round(cash_total,0),
            "bank":round(bank_total,0),
            "customer_debt":round(customer_debt,0),
            "employee_debt":round(employee_debt,0),
            "pending_cash_receipts":pending_cash,
            "received_cash_receipts":received_cash
          },
          "xb":{
            "dms_mbw_rows":xb_dms_rows,
            "misa_bu_rows":xb_misa_rows,
            "preview_batches":xb_preview,
            "error_rows":xb_error_rows
          },
          "bx":{
            "pack_count":bx_pack_count,
            "internal_pack_count":bx_internal,
            "external_pack_count":bx_external,
            "kg_total":round(bx_kg,3),
            "amount_total":round(bx_money,0)
          },
          "exceptions":{
            "qt_draft":qt_nhap,
            "return_waiting":cho_kho,
            "employee_debt_count":emp_debt_count,
            "customer_debt_count":cust_debt_count,
            "adjustment_pending":adjustment_pending,
            "bx_draft_pack":bx_draft,
            "cash_waiting":pending_cash,
            "xb_error_rows":xb_error_rows
          }
        }
    except Exception as e:
        raise HTTPException(400,str(e))
