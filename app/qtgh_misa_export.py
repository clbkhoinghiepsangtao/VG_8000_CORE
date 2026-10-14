"""Read-only one-lot export using the user's MISA template and effective master price."""
import datetime, io, json, re, sqlite3
from copy import copy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response

router=APIRouter(prefix='/api/qtgh-v11',tags=['QTGH MISA'])
ROOT=Path(__file__).resolve().parent.parent
TEMPLATE=Path(__file__).with_name('qtgh_misa_template.xlsx')
EXTRA=['QC','Mã nhân viên giao hàng','Tên nhân viên giao hàng']
REQUIRED={13:'Ngày hóa đơn',27:'Mã hàng (*)',28:'Tên hàng',33:'ĐVT',34:'Số lượng',35:'Đơn giá',36:'Thành tiền'}

def date_value(value):
    if isinstance(value,datetime.datetime):return value.date()
    if isinstance(value,datetime.date):return value
    s=str(value or '').strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y'):
        try:return datetime.datetime.strptime(s[:10],fmt).date()
        except ValueError:pass
    raise ValueError('Ngày không hợp lệ: '+s)

def masters(c,category):
    result=[]
    for r in c.execute('SELECT business_key,payload_json FROM master_data_generic WHERE category=?',(category,)):
        p=json.loads(r['payload_json'] or '{}');p['_key']=r['business_key'];result.append(p)
    return result

def code(v):return str(v or '').strip().upper()

def effective(p,day):
    start=date_value(p['TU_NGAY']) if p.get('TU_NGAY') else datetime.date.min
    end=date_value(p['DEN_NGAY']) if p.get('DEN_NGAY') else datetime.date.max
    return start<=day<=end

def price_map(c,day):
    # History status may be closed now; its explicit validity dates govern old invoices.
    groups={}
    for p in masters(c,'price'):
        if not effective(p,day):continue
        sku=code(p.get('MA_SKU') or str(p['_key']).split('|')[0])
        groups.setdefault(sku,[]).append(p)
    out={}
    for sku,rows in groups.items():
        if len(rows)>1:
            out[sku]=ValueError('SKU '+sku+' có nhiều bảng giá cùng hiệu lực tại '+day.isoformat());continue
        try:
            v=Decimal(str(rows[0]['DON_GIA_CO_VAT']))
            if not v.is_finite() or v<0:raise ValueError()
            out[sku]=v
        except (KeyError,ValueError,InvalidOperation):out[sku]=ValueError('SKU '+sku+' thiếu giá SU có thuế hợp lệ')
    return out

def staff(c,assigned,day):
    assigned=code(assigned)
    people={code(p.get('MA_NV') or p['_key']):p for p in masters(c,'dm02-staff')}
    ma=assigned
    if ma not in people:
        choices=set()
        for p in masters(c,'dm02a-staff-map'):
            if not effective(p,day):continue
            if any(code(p.get(k))==assigned for k in ['MA_NGHIEP_VU','MA_NV_NGHIEP_VU','MA_MAY','MA_NV_MAY']) and assigned:
                choices.add(code(p.get('MA_NV')))
        if len(choices)!=1:raise ValueError('Chưa xác định duy nhất NVGH từ mã '+assigned)
        ma=choices.pop()
    p=people.get(ma,{})
    name=str(p.get('TEN_NV') or p.get('HO_TEN') or p.get('TEN_NHAN_VIEN') or '').strip()
    if not name:raise ValueError('NVGH '+assigned+' chưa có tên trong danh mục nhân sự')
    return ma,name

def prepare_rows(c,lot_id,invoice_date,pick_function):
    day=date_value(invoice_date)
    lot=c.execute('''SELECT l.*,f.actual_delivery_staff_code,f.actual_delivery_date
        FROM qtgh_print_lot l LEFT JOIN qtgh_delivery_finalization f
        ON f.print_lot_id=l.print_lot_id WHERE l.print_lot_id=?''',(lot_id,)).fetchone()
    if not lot:raise ValueError('Không tìm thấy lô')
    assigned=lot['actual_delivery_staff_code'] or lot['planned_delivery_staff_code']
    delivery_day=date_value(lot['actual_delivery_date'] or lot['planned_delivery_date'])
    employee,name=staff(c,assigned,delivery_day)
    pick=pick_function(c,[lot_id])
    products={code(p.get('MA_SKU') or p['_key']):p for p in masters(c,'dm01-product')}
    prices=price_map(c,day)
    result=[];errors=[]
    for warehouse in ('HC','PC','KHÁC'):
        for item in pick['sheets'][warehouse]:
            sku=code(item['sku']);p=products.get(sku,{})
            unit=str(p.get('DVT_CO_SO') or '').strip()
            product_name=str(p.get('TEN_SKU') or '').strip()
            price=prices.get(sku)
            if price is None:errors.append('SKU '+sku+': không có bảng giá hiệu lực');continue
            if isinstance(price,ValueError):errors.append(str(price));continue
            if not unit or not product_name:errors.append('SKU '+sku+': thiếu tên hàng hoặc ĐVT cơ sở');continue
            qty=int(item['total_su']);qc=int(item['pack_su'])
            if qty!=int(item['cases'])*qc+int(item['loose_su']):raise ValueError('Số lượng không khớp quy cách: '+sku)
            amount=(Decimal(qty)*price).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
            result.append({'date':day,'sku':sku,'name':product_name,'unit':unit,'qty':qty,
                'price':price,'amount':amount,'qc':qc,'employee':employee,'employee_name':name})
    if errors:raise ValueError('Chưa xuất file MISA: '+ '; '.join(errors))
    if not result:raise ValueError('Lô không có hàng để xuất')
    if sum(r['qty'] for r in result)!=sum(x['total_su'] for items in pick['sheets'].values() for x in items):
        raise ValueError('Tổng SU xuất MISA không khớp 3 phiếu kho')
    return result

def workbook_bytes(rows):
    from openpyxl import load_workbook
    w=load_workbook(TEMPLATE)
    if len(w.worksheets)!=1:raise ValueError('Template MISA thay đổi số sheet')
    s=w.worksheets[0]
    if s.max_column!=62 or any(s.cell(1,k).value!=v for k,v in REQUIRED.items()):
        raise ValueError('Template MISA không khớp cấu trúc đã chốt')
    s.delete_rows(2,s.max_row-1)
    # No sample data, accounts, document codes, formulas or financial defaults are copied.
    s.data_validations.dataValidation=[]
    for k,label in enumerate(EXTRA,63):
        dest=s.cell(1,k,label);dest._style=copy(s.cell(1,28)._style)
    for i,r in enumerate(rows,2):
        values={13:datetime.datetime.combine(r['date'],datetime.time()),
            14:r['employee'],15:r['employee_name'],27:r['sku'],28:r['name'],
            33:r['unit'],34:r['qty'],35:float(r['price']),36:float(r['amount']),
            63:r['qc'],64:r['employee'],65:r['employee_name']}
        for k,v in values.items():
            cell=s.cell(i,k,v)
            if isinstance(v,str):cell.data_type='s'  # preserve leading zeros and prevent formulas
            elif k==13:cell.number_format='dd/mm/yyyy'
            elif k in (35,36):cell.number_format='#,##0.##'
            else:cell.number_format='0'
    s.column_dimensions['BK'].width=8;s.column_dimensions['BL'].width=24;s.column_dimensions['BM'].width=30
    w.calculation.fullCalcOnLoad=True
    output=io.BytesIO();w.save(output);w.close();return output.getvalue()

@router.get('/print-lot/{lot_id}/misa-export.xlsx')
def export_misa(lot_id:str,invoice_date:str):
    c=sqlite3.connect((ROOT/'data/miniapp.db').as_uri()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row
    try:
        from .qtgh_print_v11 import vehicle_pick_list
        result=prepare_rows(c,lot_id,invoice_date,vehicle_pick_list)
        raw=workbook_bytes(result)
    except (ValueError,KeyError,InvalidOperation) as e:raise HTTPException(400,str(e))
    finally:c.close()
    filename='MISA_'+re.sub(r'[^A-Za-z0-9_-]','_',lot_id)+'_'+date_value(invoice_date).strftime('%Y%m%d')+'.xlsx'
    return Response(raw,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition':'attachment; filename="'+filename+'"','Cache-Control':'no-store'})
