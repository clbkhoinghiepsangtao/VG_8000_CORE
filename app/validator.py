from io import BytesIO
from datetime import datetime, date
import hashlib, json, uuid
from openpyxl import load_workbook
from .template_defs import TEMPLATES
from .db import connect

def _blank(v): return v is None or (isinstance(v,str) and not v.strip())
def _is_num(v):
    if _blank(v): return False
    try: float(v); return True
    except: return False

def _valid_date(v):
    if isinstance(v,(datetime,date)): return True
    if isinstance(v,str):
        for f in ('%d/%m/%Y','%Y-%m-%d'):
            try: datetime.strptime(v.strip(),f); return True
            except: pass
    return False

def _mapping_errors(payload, template_key, con):
    errs=[]
    sku=payload.get('MA_SKU')
    if sku and not con.execute('SELECT 1 FROM dm_product WHERE ma_sku=?',(str(sku).strip(),)).fetchone(): errs.append(f'Thiếu mapping DM01 cho SKU {sku}')
    personnel_fields=['MA_NV_NGHIEP_VU','MA_NV_KD_NGHIEP_VU','MA_NV_GIAO_HANG_NGHIEP_VU','MA_NV_LAI_XE_NGHIEP_VU']
    for f in personnel_fields:
        v=payload.get(f)
        if v and not con.execute('SELECT 1 FROM dm_personnel_code WHERE ma_nghiep_vu=?',(str(v).strip(),)).fetchone(): errs.append(f'Thiếu mapping DM02A: {f}={v}')
    if template_key=='dms-mbw':
        v=payload.get('MA_KH_NGUON')
        if v and not con.execute('SELECT 1 FROM dm_customer_map WHERE ma_kh_nguon=?',(str(v).strip(),)).fetchone(): errs.append(f'Thiếu mapping DM03 cho khách {v}')
    return errs

def validate_workbook(template_key:str, filename:str, data:bytes):
    spec=TEMPLATES[template_key]; fh=hashlib.sha256(data).hexdigest(); batch_id='BT-'+uuid.uuid4().hex[:12].upper()
    with connect() as con:
        dup=con.execute('SELECT batch_id,status FROM import_batch WHERE file_hash=? AND template_key=?',(fh,template_key)).fetchone()
        if dup: return {'duplicate':True,'batch_id':dup['batch_id'],'status':dup['status'],'message':'File đã được upload trước đó.'}
    try: wb=load_workbook(BytesIO(data),data_only=True)
    except Exception as e: return {'rejected':True,'message':f'Không đọc được XLSX: {e}'}
    version=None
    if '_META' in wb.sheetnames:
        meta=wb['_META']; version=meta['B2'].value
    structural=[]
    if version!=spec['version']: structural.append(f'Sai/mất version template. Yêu cầu {spec["version"]}, nhận {version or "trống"}.')
    for sname, fields in spec['sheets'].items():
        if sname not in wb.sheetnames: structural.append(f'Thiếu sheet {sname}'); continue
        ws=wb[sname]; actual=[c.value for c in ws[1]][:len(fields)]; expected=[x[0] for x in fields]
        if actual!=expected: structural.append(f'Sai cấu trúc cột sheet {sname}.')
    if structural:
        return {'rejected':True,'message':'Template không hợp lệ.','errors':structural}
    rows=[]; valid=0; errors_count=0
    with connect() as con:
        for sname, fields in spec['sheets'].items():
            ws=wb[sname]; headers=[x[0] for x in fields]
            for rno, vals in enumerate(ws.iter_rows(min_row=2,max_col=len(headers),values_only=True),2):
                if all(_blank(v) for v in vals): continue
                payload={h:v for h,v in zip(headers,vals)}; errs=[]
                for (name,req,typ),v in zip(fields,vals):
                    if req and _blank(v): errs.append(f'{name}: bắt buộc')
                    if _blank(v): continue
                    if typ=='number' and not _is_num(v): errs.append(f'{name}: phải là số')
                    elif typ=='date' and not _valid_date(v): errs.append(f'{name}: sai ngày')
                    elif typ.startswith('enum:') and str(v).strip().upper() not in typ.split(':',1)[1].split('|'): errs.append(f'{name}: giá trị không hợp lệ')
                errs.extend(_mapping_errors(payload,template_key,con))
                ok=not errs; valid+=1 if ok else 0; errors_count+=0 if ok else 1
                rows.append((sname,rno,payload,ok,errs))
        con.execute('INSERT INTO import_batch(batch_id,template_key,template_version,filename,file_hash,status,total_rows,valid_rows,error_rows) VALUES(?,?,?,?,?,?,?,?,?)',
                    (batch_id,template_key,version,filename,fh,'PREVIEW_READY',len(rows),valid,errors_count))
        for sname,rno,payload,ok,errs in rows:
            con.execute('INSERT INTO import_row(batch_id,sheet_name,row_no,payload_json,is_valid,errors_json) VALUES(?,?,?,?,?,?)',
                        (batch_id,sname,rno,json.dumps(payload,ensure_ascii=False,default=str),1 if ok else 0,json.dumps(errs,ensure_ascii=False)))
    return {'duplicate':False,'batch_id':batch_id,'status':'PREVIEW_READY','template_version':version,'total_rows':len(rows),'valid_rows':valid,'error_rows':errors_count,
            'preview':[{'sheet':s,'row_no':r,'is_valid':ok,'errors':e,'payload':p} for s,r,p,ok,e in rows[:50]]}

def commit_batch(batch_id:str):
    with connect() as con:
        b=con.execute('SELECT * FROM import_batch WHERE batch_id=?',(batch_id,)).fetchone()
        if not b: return None
        if b['status']=='COMMITTED': return {'batch_id':batch_id,'status':'COMMITTED','committed_rows':b['valid_rows']}
        rows=con.execute('SELECT * FROM import_row WHERE batch_id=? AND is_valid=1',(batch_id,)).fetchall()
        for r in rows:
            con.execute('INSERT OR IGNORE INTO source_row(template_key,batch_id,sheet_name,row_no,payload_json) VALUES(?,?,?,?,?)',
                        (b['template_key'],batch_id,r['sheet_name'],r['row_no'],r['payload_json']))
        con.execute("UPDATE import_batch SET status='COMMITTED', committed_at=CURRENT_TIMESTAMP WHERE batch_id=?",(batch_id,))
        return {'batch_id':batch_id,'status':'COMMITTED','committed_rows':len(rows),'skipped_error_rows':b['error_rows']}
