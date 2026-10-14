"""UI02 bridge for raw DMS/MBW workbooks. No write occurs in prepare()."""
from datetime import timedelta
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json, unicodedata
from openpyxl import load_workbook
from .lc_core.excel import read_sheet
from .lc_core.sources import adapt, norm, DMS
from .lc_core.rewards import attach_rewards
from .lc_core.mapping import audit_mapping
from .lc_core.intake import parse_rows, day, code
from .lc_core.orders import group_orders


def dms_rows(blob):
    # Source file has an identical, redundant MA_NV column at position 14.
    # Accept only this known export layout and compare both columns before ignoring it.
    wb = load_workbook(BytesIO(blob), read_only=True, data_only=False)
    try:
        if 'Sheet1' not in wb: raise ValueError('DMS thiếu Sheet1')
        ws=wb['Sheet1'];head=[norm(c.value) for c in next(ws.iter_rows())]
        if len(head)>=14 and head[3]==head[13] and head[3]==norm('Mã nhân viên'):
            for row_no,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
                if any(x not in (None,'') for x in row[:13]) and str(row[3] or '').strip()!=str(row[13] or '').strip():
                    raise ValueError(f'DMS dòng {row_no}: hai cột Mã nhân viên không khớp')
            source=read_sheet(blob,'Sheet1',ignore_columns=(13,))
        else:
            source=read_sheet(blob,'Sheet1')
    finally:wb.close()
    mapped=adapt(source,'DMS')
    if mapped['exceptions']:raise ValueError(f'DMS có {len(mapped["exceptions"])} dòng ngoại lệ; không nạp một phần')
    rewards=read_sheet(blob,'THUONG_CAP_DON',allow_empty=True,prefix_headers=(
        'SO_HOA_DON','NGAY_DON','MA_NV_KD_NGHIEP_VU','THUONG_TRUNG_BAY','THUONG_TICH_LUY'))
    return _source_assignments(attach_rewards(mapped['rows'],rewards)[0])


def mbw_rows(blob):
    source=read_sheet(blob,'Phiếu bán hàng')
    mapped=adapt(source,'MBW')
    if mapped['exceptions']:raise ValueError(f'MBW có {len(mapped["exceptions"])} dòng ngoại lệ; không nạp một phần')
    return _source_assignments(mapped['rows'])


ASSIGNMENT_FIELDS=('MA_KH','MA_NV_GIAO_HANG_NGHIEP_VU','MA_NV_LAI_XE_NGHIEP_VU','PHUONG_TIEN','LO_GIAO_HANG','MA_TKNH')

def _source_assignments(rows):
    seen={}
    for row in rows:
        raw={norm(k):v for k,v in row['_RAW_SOURCE'].items()}
        for field in ASSIGNMENT_FIELDS:
            value=code(raw.get(field))
            if not value:raise ValueError(f"Dòng {row['_SOURCE_ROW']}: thiếu {field}")
            row[field]=value
        key=(row['LOAI_DON'],code(row['SO_HOA_DON']),day(row['NGAY_DON']))
        value=tuple(row[field] for field in ASSIGNMENT_FIELDS)
        if key in seen and seen[key]!=value:raise ValueError(f"Đơn {key[1]}: phân công hoặc khách hàng khác nhau giữa các dòng")
        seen[key]=value
    return rows


def prepare(dms,mbw,customers,products,meta):
    if not dms and not mbw:raise ValueError('Chọn file DMS hoặc MBW')
    rows=(dms_rows(dms) if dms else [])+(mbw_rows(mbw) if mbw else [])
    audit=audit_mapping(rows,customers,products)
    if audit['issue_count']:
        return {'ready':False,'mapping':audit,'orders':audit['order_count'],'lines':len(rows)}
    orders,raw=parse_rows(rows,customers,products)
    if len({o.code for o in orders})!=len(orders):
        raise ValueError('Mã đơn trùng giữa DMS và MBW; UI02 cần khóa riêng theo nguồn trước khi nhận')
    aliases={}
    customer_info={}
    for cust in customers:
        canonical=code(cust.get('MA_KH_CHUAN'))
        if canonical:
            customer_info[canonical]=cust
            for field in ('MA_KH_CHUAN','MA_KH_DMS','MA_KH_MBW'):
                if cust.get(field):aliases[code(cust[field])]=canonical
    groups={}
    for o in orders:
        item=next(r for r in rows if code(r['SO_HOA_DON'])==o.code)
        key=o.customer+'|'+o.order_date+'|'+item['LO_GIAO_HANG']
        groups.setdefault(key,[]).append(o)
    points=[{'key':key,'customer':key.split('|')[0],'order_date':key.split('|')[1],
             'area':customer_info.get(key.split('|')[0],{}).get('MA_DIA_BAN'),
             'orders':len(group),'money':str(sum((o.amount_due for o in group),Decimal(0))),
             'kg':str(sum((o.kg for o in group),Decimal(0))),
             'mixed':len({o.source for o in group})>1} for key,group in sorted(groups.items())]
    missing_staff=sorted({code(r[field]) for r in rows for field in ('MA_NV_KD_NGHIEP_VU','MA_NV_GIAO_HANG_NGHIEP_VU','MA_NV_LAI_XE_NGHIEP_VU') if code(r[field]) not in meta['staff_codes']})
    common={'orders':len(orders),'points':len(points),'point_rows':points,'lines':len(rows),
            'mixed_points':sum(p['mixed'] for p in points),
            'money':str(sum((o.amount_due for o in orders),Decimal(0)))}
    if missing_staff:return {**common,'ready':False,'missing_staff':missing_staff}
    lots={}
    for item in rows:
        canonical_customer=aliases[code(item['MA_KH_NGUON'])]
        if item['MA_KH']!=canonical_customer:
            raise ValueError(f"Đơn {item['SO_HOA_DON']}: MA_KH không khớp danh mục và mã khách nguồn")
        for field in ('MA_NV_GIAO_HANG_NGHIEP_VU','MA_NV_LAI_XE_NGHIEP_VU'):
            if item[field] not in meta['staff_codes']:raise ValueError(f"Đơn {item['SO_HOA_DON']}: {field} {item[field]} chưa mapping")
        if item['MA_TKNH'] not in meta['banks']:raise ValueError(f"Đơn {item['SO_HOA_DON']}: tài khoản nhận tiền chưa hoạt động")
        lot_key=(day(item['NGAY_DON']),item['LO_GIAO_HANG'])
        assignment=tuple(item[k] for k in ('PHUONG_TIEN','MA_NV_GIAO_HANG_NGHIEP_VU','MA_NV_LAI_XE_NGHIEP_VU'))
        if lot_key in lots and lots[lot_key]!=assignment:raise ValueError(f"Lô {lot_key[1]}: xe hoặc nhân viên khác nhau trong cùng ngày")
        lots[lot_key]=assignment
    for point in points:
        item=next(r for r in rows if aliases[code(r['MA_KH_NGUON'])]+'|'+day(r['NGAY_DON'])+'|'+r['LO_GIAO_HANG']==point['key'])
        point.update({k:item[k] for k in ASSIGNMENT_FIELDS})
    canonical=[]
    for item in rows:
        od=day(item['NGAY_DON']);pd=str(__import__('datetime').date.fromisoformat(od)+timedelta(days=1))
        source=code(item['LOAI_DON'])
        raw_price=item['DON_GIA_NGUON'] if source=='L' else item['DON_GIA_CO_VAT']
        d={k:item.get(k) for k in ('SO_HOA_DON','MA_SKU','TEN_SKU','MA_NV_KD_NGHIEP_VU','DOANH_SO_PHAI_THU','MA_KH_NGUON','SL_SU','SL_THUNG','SL_KM_SU','SL_KM_THUNG','LOAI_DON','THUONG_TRUNG_BAY','THUONG_TICH_LUY')}
        d.update({'NGAY_DON':od,'NGAY_GIAO':pd,**{k:item[k] for k in ASSIGNMENT_FIELDS if k!='MA_KH'}})
        if source=='L':d.update({'DON_GIA_CHUA_VAT':raw_price,'VAT_RATE':0.08})
        else:d['DON_GIA_CO_VAT']=raw_price
        canonical.append({'row_no':len(canonical)+2,'valid':True,'errors':[],'data':d})
    checksum=sha256((sha256(dms or b'').hexdigest()+sha256(mbw or b'').hexdigest()+json.dumps({k:v for k,v in meta.items() if k not in ('staff_codes','banks')},sort_keys=True)).encode()).hexdigest()
    return {**common,'ready':True,'checksum':checksum,'rows':canonical,'lot_count':len(lots)}

# Route mounted under /api/qtgh-v11 and protected by the app's existing UI02 role.
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from .qtgh_phase2b_v11 import master_rows as _masters, conn as _conn, commit_prepared
router=APIRouter(prefix='/api/qtgh-v11',tags=['QTGH raw DMS/MBW'])

def _metadata(text):
    try:meta=json.loads(text)
    except (ValueError,TypeError):raise ValueError('Thông tin lô không hợp lệ')
    if not isinstance(meta,dict):raise ValueError('Thông tin lô không hợp lệ')
    staff=_masters('dm02-staff');aliases=_masters('dm02a-staff-map')
    meta['staff_codes']={code(x.get('MA_NV')) for x in staff if x.get('MA_NV')}|{
        code(x.get('MA_NGHIEP_VU')) for x in aliases if x.get('MA_NGHIEP_VU')}
    with _conn() as db:
        meta['banks']={str(r[0]) for r in db.execute("SELECT ma_tknh FROM qtgh_receiving_bank_account WHERE status='ACTIVE'")}
    return meta

async def _input(dms,mbw,details):
    dms_bytes=await dms.read() if dms else None
    mbw_bytes=await mbw.read() if mbw else None
    if dms_bytes is not None and len(dms_bytes)>10*1024*1024:raise ValueError('DMS vượt 10 MB')
    if mbw_bytes is not None and len(mbw_bytes)>10*1024*1024:raise ValueError('MBW vượt 10 MB')
    meta=_metadata(details)
    products={code(p.get('MA_SKU')):p for p in _masters('dm01-product')}
    return prepare(dms_bytes,mbw_bytes,_masters('dm03-customer'),products,meta)

@router.post('/raw/preview')
async def raw_preview(details:str=Form(...),dms:UploadFile|None=File(default=None),mbw:UploadFile|None=File(default=None)):
    try:
        x=await _input(dms,mbw,details)
        if x.get('ready'):
            return {k:v for k,v in x.items() if k!='rows'}
        return x
    except Exception as e:
        raise HTTPException(400,str(e))

@router.post('/raw/commit')
async def raw_commit(token:str=Form(...),details:str=Form(...),dms:UploadFile|None=File(default=None),mbw:UploadFile|None=File(default=None)):
    try:
        x=await _input(dms,mbw,details)
        if not x.get('ready'):raise ValueError('Danh mục hoặc nhân viên chưa khớp; không được nạp một phần')
        if token!=x['checksum']:raise ValueError('File hoặc thông tin lô đã đổi sau Preview')
        orders={(r['data']['LOAI_DON'],str(r['data']['SO_HOA_DON'])) for r in x['rows']}
        with _conn() as db:
            existing={tuple(r) for r in db.execute('SELECT DISTINCT loai_don,ma_don_nguon FROM qtgh_prep_line')}
        if orders&existing:
            raise ValueError(f'{len(orders&existing)} đơn đã được nạp vào UI02; không nạp trùng')
        return commit_prepared({'filename':'+'.join(f.filename for f in (dms,mbw) if f),
                                'checksum':token,'rows':x['rows'],'expected_total_rows':x['lines'],'uploaded_by':'ui02-raw'})
    except HTTPException:raise
    except Exception as e:raise HTTPException(400,str(e))
