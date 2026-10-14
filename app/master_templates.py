
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MASTER_TEMPLATES = {
 "dm01-product": {
   "filename":"DM01_San_pham_V1.0.xlsx","sheet":"DM01_SAN_PHAM",
   "headers":["MA_SKU","TEN_SKU","MA_NHAN","BU","DVT_CO_SO","SO_SU_THUNG","TRONG_LUONG_SU","TRONG_LUONG_THUNG","VAT_RATE","TRANG_THAI"],
   "example":["SKU001","OMO Matic 3kg","101","HC","SU",6,3,18,0.08,"DANG_KINH_DOANH"],
   "rules":[
     ["MA_SKU","Có","Text","Mã SKU duy nhất dùng chung DMS/MBW/MISA"],
     ["TEN_SKU","Có","Text","Tên sản phẩm"],
     ["MA_NHAN","Có","Text","Mã nhãn, ví dụ 101/103/111"],
     ["BU","Có","Text","HC/PC/BE/F/HASU/KHAC"],
     ["DVT_CO_SO","Có","Text","Đơn vị cơ sở, mặc định SU"],
     ["SO_SU_THUNG","Có","Number","Số SU trong 1 thùng"],
     ["TRONG_LUONG_SU","Có","Number","Khối lượng 1 SU (kg)"],
     ["TRONG_LUONG_THUNG","Không","Number","Khối lượng 1 thùng (kg)"],
     ["VAT_RATE","Có","Number","0.08 = 8%"],
     ["TRANG_THAI","Có","Enum","DANG_KINH_DOANH/NGUNG"]
   ]
 },
 "price": {
   "filename":"Bang_gia_V1.0.xlsx","sheet":"BANG_GIA",
   "headers":["MA_SKU","TU_NGAY","DEN_NGAY","DON_GIA_CHUA_VAT","VAT_RATE","DON_GIA_CO_VAT","NGUON_GIA","TRANG_THAI"],
   "example":["SKU001","2026-08-01","",102000,0.08,110160,"DMS","DANG_HIEU_LUC"],
   "rules":[
     ["MA_SKU","Có","Text","Mã SKU trong DM01"],
     ["TU_NGAY","Có","Date","Ngày bắt đầu hiệu lực"],
     ["DEN_NGAY","Không","Date","Để trống nếu đang hiệu lực"],
     ["DON_GIA_CHUA_VAT","Có","Number","Đơn giá chuẩn chưa VAT"],
     ["VAT_RATE","Có","Number","Ví dụ 0.08"],
     ["DON_GIA_CO_VAT","Có","Number","Giá có VAT"],
     ["NGUON_GIA","Không","Text","DMS/MBW/MISA/KHAC"],
     ["TRANG_THAI","Có","Enum","DANG_HIEU_LUC/HET_HIEU_LUC"]
   ]
 },
 "dm02-staff": {
   "filename":"DM02_Nhan_su_V1.0.xlsx","sheet":"DM02_NHAN_SU",
   "headers":["MA_NV","TEN_NV","NHOM_NV","MA_VUNG_KD","BO_PHAN","MANG_PHU_TRACH","NGAY_VAO","TRANG_THAI"],
   "example":['TEST-MA_NV','TEST-TEN_NV',"KINH_DOANH","VG_01","KD Phân phối","U","2023-03-01","DANG_LAM"],
   "rules":[
     ["MA_NV","Có","Text","Mã nhân sự chuẩn VGxxx"],
     ["TEN_NV","Có","Text","Họ tên"],
     ["NHOM_NV","Có","Enum","KINH_DOANH/GIAO_HANG/LAI_XE/ADMIN/BOC_XEP_NOI_BO/KHAC"],
     ["MA_VUNG_KD","Không","Text","Vùng KD thấp nhất nhân sự thuộc"],
     ["BO_PHAN","Không","Text","Bộ phận"],
     ["MANG_PHU_TRACH","Không","Text","U/HASU/KHAC"],
     ["NGAY_VAO","Không","Date","Ngày vào làm"],
     ["TRANG_THAI","Có","Enum","DANG_LAM/NGHI"]
   ]
 },
 "dm02a-staff-map": {
   "filename":"DM02A_Ma_nghiep_vu_NV_V1.0.xlsx","sheet":"DM02A_MA_NGHIEP_VU",
   "headers":["MA_NV","LOAI_MA","MA_NGHIEP_VU","HE_THONG","TU_NGAY","DEN_NGAY","TRANG_THAI"],
   "example":['TEST-MA_NV',"MA_MAY_BAN_HANG",'TEST-MA_NGHIEP_VU',"DMS","2026-01-01","","DANG_HIEU_LUC"],
   "rules":[
     ["MA_NV","Có","Text","Mã nhân sự chuẩn VGxxx"],
     ["LOAI_MA","Có","Enum","MA_MAY_BAN_HANG/MA_DMS/MA_MBW/MA_GIAO_HANG/MA_LAI_XE/KHAC"],
     ["MA_NGHIEP_VU","Có","Text","Một MA_NV có thể có nhiều dòng"],
     ["HE_THONG","Có","Text","DMS/MBW/GH/MISA/KHAC"],
     ["TU_NGAY","Có","Date","Bắt đầu hiệu lực"],
     ["DEN_NGAY","Không","Date","Để trống nếu còn hiệu lực"],
     ["TRANG_THAI","Có","Enum","DANG_HIEU_LUC/HET_HIEU_LUC"]
   ]
 },
 "sales-region": {
   "filename":"Co_cau_vung_kinh_doanh_V1.0.xlsx","sheet":"VUNG_KINH_DOANH",
   "headers":["MA_VUNG","TEN_VUNG","MA_VUNG_CHA","CAP_VUNG","MA_QUAN_LY","LOAI_QUAN_LY","TRANG_THAI"],
   "example":["VG_01","Vùng VG01","U",2,"VG110","DOI_TRUONG","HOAT_DONG"],
   "rules":[
     ["MA_VUNG","Có","Text","U, VG_01...VG_04, HASU, KHAC"],
     ["TEN_VUNG","Có","Text","Tên vùng"],
     ["MA_VUNG_CHA","Không","Text","Mã vùng cấp trên"],
     ["CAP_VUNG","Có","Number","1,2,3..."],
     ["MA_QUAN_LY","Không","Text","Mã NV quản lý"],
     ["LOAI_QUAN_LY","Không","Enum","TRUONG_PHONG/DOI_TRUONG/QUAN_LY"],
     ["TRANG_THAI","Có","Enum","HOAT_DONG/NGUNG"]
   ]
 },
 "dm03-customer": {
   "filename":"DM03_Khach_hang_V1.0.xlsx","sheet":"DM03_KHACH_HANG",
   "headers":["MA_KH_CHUAN","MA_KH_MBW","MA_KH_DMS","TEN_KH","SDT_CHINH","DIA_CHI","MA_NVBH_01","MA_NVBH_02","MA_NHOM_CUA_HANG","MA_LOAI_CUA_HANG","MA_DIA_BAN","TRANG_THAI"],
   "example":['TEST-MA_KH_CHUAN','TEST-MA_KH_MBW','TEST-MA_KH_DMS','TEST-TEN_KH','TEST-SDT_CHINH','TEST-DIA_CHI','TEST-MA_NVBH_01',"","SS_FAMILY_GROCERY","SS_FG_URBAN","T001","BINH_THUONG"],
   "rules":[
     ["MA_KH_CHUAN","Có","Text","Mã KH nội bộ Mini App"],
     ["MA_KH_MBW","Không","Text","Mã KH MBW, ví dụ V00015"],
     ["MA_KH_DMS","Không","Text","Mã KH DMS"],
     ["TEN_KH","Có","Text","Tên khách hàng"],
     ["SDT_CHINH","Có","Text","10 số; tín hiệu nhận diện mạnh"],
     ["DIA_CHI","Không","Text","Địa chỉ/đường phố của khách hàng"],
     ["MA_NVBH_01","Không","Text","Mã nghiệp vụ bán hàng thứ nhất; đối chiếu qua DM02A"],
     ["MA_NVBH_02","Không","Text","Mã nghiệp vụ bán hàng thứ hai; để trống nếu không có"],
     ["MA_NHOM_CUA_HANG","Không","Text","Mã nhóm cửa hàng, lấy từ danh mục NHOM_CUA_HANG"],
     ["MA_LOAI_CUA_HANG","Không","Text","Mã loại cửa hàng, phải thuộc MA_NHOM_CUA_HANG đã chọn"],
     ["MA_DIA_BAN","Không","Text","Mã node trong danh mục Địa bàn/Tuyến"],
     ["TRANG_THAI","Có","Enum","BINH_THUONG/NGHI_TRUNG_KH/NGUNG"]
   ]
 },
 "territory": {
   "filename":"Dia_ban_Tuyen_V1.0.xlsx","sheet":"DIA_BAN_TUYEN",
   "headers":["MA_DIA_BAN","TEN_DIA_BAN","MA_CHA","CAP","LOAI","TRANG_THAI"],
   "example":["T001","Tuyến Võ Cường 1","VC",5,"TUYEN","HOAT_DONG"],
   "rules":[
     ["MA_DIA_BAN","Có","Text","Mã node duy nhất"],
     ["TEN_DIA_BAN","Có","Text","Tên địa bàn/tuyến"],
     ["MA_CHA","Không","Text","Node cha"],
     ["CAP","Có","Number","Cấp phân cấp"],
     ["LOAI","Có","Enum","QUOC_GIA/TINH/TP_HUYEN/PHUONG_XA/TUYEN/KHAC"],
     ["TRANG_THAI","Có","Enum","HOAT_DONG/NGUNG"]
   ]
 },
 "basic": {
   "filename":"BU_Nhan_Mang_V1.0.xlsx","sheet":"DM_NEN",
   "headers":["LOAI_DANH_MUC","MA","TEN","MA_CHA","THU_TU","TRANG_THAI"],
   "example":["BU","HC","Home Care","",1,"HOAT_DONG"],
   "rules":[
     ["LOAI_DANH_MUC","Có","Enum","BU/NHAN/MANG/NHOM_NV/NHOM_CUA_HANG/LOAI_CUA_HANG"],
     ["MA","Có","Text","Mã danh mục"],
     ["TEN","Có","Text","Tên hiển thị"],
     ["MA_CHA","Không","Text","Mã danh mục cha; bắt buộc với LOAI_CUA_HANG để liên kết NHOM_CUA_HANG"],
     ["THU_TU","Không","Number","Thứ tự hiển thị"],
     ["TRANG_THAI","Có","Enum","HOAT_DONG/NGUNG"]
   ]
 },
 "users": {
   "filename":"User_Phan_quyen_V1.0.xlsx","sheet":"USER_PHAN_QUYEN",
   "headers":["USERNAME","HO_TEN","MA_NV","VAI_TRO","GH","XB","KQ","BX","DANH_MUC","XUAT_EXCEL","TRANG_THAI"],
   "example":["ketoan01",'TEST-HO_TEN','TEST-MA_NV',"KE_TOAN","EDIT","EDIT","VIEW","EDIT","EDIT",1,"HOAT_DONG"],
   "rules":[
     ["USERNAME","Có","Text","Tên đăng nhập duy nhất"],
     ["HO_TEN","Có","Text","Tên người dùng"],
     ["VAI_TRO","Có","Enum","SUPER_ADMIN/ADMIN/KE_TOAN/THU_QUY/THU_KHO/QUAN_LY/VIEWER/HCNS"],
     ["MA_NV","Không","Text","Liên kết DM02 nếu là nhân sự nội bộ"],
     ["TRANG_THAI","Có","Enum","HOAT_DONG/KHOA"]
   ]
 }
}

def build_master_template(key:str):
    spec = MASTER_TEMPLATES[key]
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    for c,h in enumerate(spec["headers"],1):
        cell=ws.cell(1,c,h)
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="1769E0")
        cell.alignment=Alignment(horizontal="center",vertical="center")
        ws.column_dimensions[get_column_letter(c)].width=max(16,min(28,len(h)+4))
    for c,v in enumerate(spec["example"],1):
        ws.cell(2,c,v)
    ws.freeze_panes="A2"

    guide=wb.create_sheet("HUONG_DAN")
    guide["A1"]=f"VĨNH GIANG – {spec['filename']}"
    guide["A1"].font=Font(bold=True,size=14,color="FFFFFF")
    guide["A1"].fill=PatternFill("solid",fgColor="1769E0")
    guide.merge_cells("A1:D1")
    guide.append([])
    guide.append(["TRƯỜNG","BẮT BUỘC","KIỂU","MÔ TẢ"])
    for c in guide[3]:
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")
    for row in spec["rules"]:
        guide.append(row)
    for i,w in enumerate([24,12,14,55],1):
        guide.column_dimensions[get_column_letter(i)].width=w
    bio=BytesIO()
    wb.save(bio)
    return spec["filename"], bio.getvalue()
