from fastapi import APIRouter, HTTPException, UploadFile, File, Header
from pathlib import Path
from openpyxl import Workbook, load_workbook
from io import BytesIO
import sqlite3, json, hashlib, datetime
from .dms_mbw_template_v2 import (
    DETAIL_HEADERS as CANONICAL_DETAIL_HEADERS,
    BONUS_HEADERS as CANONICAL_REWARD_HEADERS,
    build_dms_mbw_v2_template,
)

router=APIRouter(prefix="/api/qtgh-v11",tags=["QTGH BA V1.1"])

ROOT=Path(__file__).resolve().parent.parent
DB=ROOT/"data"/"miniapp.db"

DATA_HEADERS=list(CANONICAL_DETAIL_HEADERS)
REWARD_HEADERS=list(CANONICAL_REWARD_HEADERS)

def conn():
    c=sqlite3.connect(DB)
    c.row_factory=sqlite3.Row
    c.execute("PRAGMA busy_timeout=5000")
    return c

def parse_date(v):
    if v is None or v=="": return None
    if isinstance(v,datetime.datetime): return v.date()
    if isinstance(v,datetime.date): return v
    s=str(v).strip()
    for f in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y"):
        try:return datetime.datetime.strptime(s[:10],f).date()
        except Exception:pass
    try:return datetime.date.fromisoformat(s[:10])
    except Exception:return None

def master_rows(category):
    c=conn(); rows=[]
    try:
        for r in c.execute("SELECT business_key,payload_json FROM master_data_generic WHERE category=?",(category,)):
            try:p=json.loads(r["payload_json"] or "{}")
            except Exception:p={}
            p["_business_key"]=r["business_key"]
            rows.append(p)
    finally:c.close()
    return rows

def master_map(category,*keys):
    out={}
    for p in master_rows(category):
        for k in keys:
            v=p.get(k)
            if v not in (None,""): out[str(v).strip().upper()]=p
    return out

def active_banks():
    c=conn()
    rows=[dict(r) for r in c.execute("""SELECT id,ma_tknh,bank_name,account_no,account_name,status
      FROM qtgh_receiving_bank_account WHERE status='ACTIVE'""").fetchall()]
    c.close();return rows

def find_bank(code):
    code=str(code or "").strip().upper()
    if not code:
        return None
    for b in active_banks():
        candidates={
            str(b.get("id") or "").strip().upper(),
            str(b.get("ma_tknh") or "").strip().upper(),
            str(b.get("bank_name") or "").strip().upper(),
            str(b.get("account_no") or "").strip().upper(),
        }
        if code in candidates:
            return b
    return None

def dm02_maps():
    dm02=master_rows("dm02-staff"); by_ma={}
    for p in dm02:
        ma=str(p.get("MA_NV") or p.get("_business_key") or "").strip().upper()
        if ma:by_ma[ma]=p
    dm02a=master_rows("dm02a-staff-map"); by_nv={}
    for p in dm02a:
        for k in ("MA_NGHIEP_VU","MA_NV_NGHIEP_VU","MA_MAY","MA_NV_MAY"):
            v=p.get(k)
            if v not in (None,""):by_nv[str(v).strip().upper()]=p
    return by_ma,by_nv

def resolve_staff(code):
    code=str(code or "").strip().upper()
    if not code:return None
    by_ma,by_nv=dm02_maps()
    if code in by_ma:return by_ma[code]
    m=by_nv.get(code)
    if m:
        ma=str(m.get("MA_NV") or "").strip().upper()
        if ma in by_ma:return by_ma[ma]
    return None

def resolve_customer(code):
    code=str(code or "").strip().upper()
    if not code:return None
    for p in master_rows("dm03-customer"):
        vals=[p.get("MA_KH_CHUAN"),p.get("MA_KH_MBW"),p.get("MA_KH_DMS"),p.get("_business_key")]
        if any(str(v or "").strip().upper()==code for v in vals):return p
    return None

def resolve_product(code):
    return master_map("dm01-product","MA_SKU").get(str(code or "").strip().upper())

def build_template():
    # Một nguồn sinh duy nhất cho XB và luồng chuẩn bị/in chứng từ.
    # Không chèn dữ liệu mẫu vào sheet nghiệp vụ để tránh seed nhầm vào Pilot.
    _, data=build_dms_mbw_v2_template()
    return data

@router.get("/template-v11.xlsx")
def template_v11():
    from fastapi.responses import Response
    return Response(
      content=build_template(),
      media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
      headers={"Content-Disposition":'attachment; filename="02_Template_DMS_MBW_GiaoThat_BanBu_InChungTu_V2_1.xlsx"'}
    )

@router.post("/prep/preview")
async def prep_preview(file:UploadFile=File(...),uploaded_by:str|None=Header(default="pilot",alias="X-Uploaded-By")):
    raw=await file.read();checksum=hashlib.sha256(raw).hexdigest()
    c=conn();dup=c.execute("SELECT prep_batch_id,status FROM qtgh_prep_batch WHERE source_checksum=?",(checksum,)).fetchone();c.close()
    if dup:raise HTTPException(409,f"FR04: File đã tồn tại trong PREP BATCH {dup['prep_batch_id']}")

    try:wb=load_workbook(BytesIO(raw),data_only=True)
    except Exception as e:raise HTTPException(400,f"Không đọc được Excel: {e}")

    # V2.1 dùng CHI_TIET_GIAO_BU. Vẫn đọc DATA của bản V1.1 cũ để tương thích.
    detail_sheet=next((s for s in ("CHI_TIET_GIAO_BU","DATA") if s in wb.sheetnames),None)
    if not detail_sheet:
        raise HTTPException(400,"Thiếu sheet bắt buộc: CHI_TIET_GIAO_BU")
    if "THUONG_CAP_DON" not in wb.sheetnames:
        raise HTTPException(400,"Thiếu sheet bắt buộc: THUONG_CAP_DON")

    ws=wb[detail_sheet];headers=[str(c.value or "").strip() for c in ws[1]]
    missing=[h for h in DATA_HEADERS if h not in headers]
    if missing:raise HTTPException(400,f"{detail_sheet} thiếu cột: "+", ".join(missing))
    idx={h:headers.index(h) for h in headers}

    rw=wb["THUONG_CAP_DON"];rheaders=[str(c.value or "").strip() for c in rw[1]]
    rmissing=[h for h in REWARD_HEADERS if h not in rheaders]
    if rmissing:raise HTTPException(400,"THUONG_CAP_DON thiếu cột: "+", ".join(rmissing))
    ridx={h:rheaders.index(h) for h in rheaders}

    rewards={};reward_errors=[]
    for rn,vals in enumerate(rw.iter_rows(min_row=2,values_only=True),start=2):
        if not any(v not in (None,"") for v in vals):continue
        d={h:vals[ridx[h]] for h in REWARD_HEADERS}
        order=str(d["SO_HOA_DON"] or "").strip()
        if not order:
            reward_errors.append({"row_no":rn,"errors":["THUONG_CAP_DON: thiếu SO_HOA_DON"]});continue
        if order in rewards:
            reward_errors.append({"row_no":rn,"errors":[f"THUONG_CAP_DON: trùng SO_HOA_DON {order}"]});continue
        rewards[order]={
          "THUONG_TRUNG_BAY":float(d.get("THUONG_TRUNG_BAY") or 0),
          "THUONG_TICH_LUY":float(d.get("THUONG_TICH_LUY") or 0),
          # Trường giảm trừ cấp đơn đã bỏ theo BA; bỏ qua nếu file cũ còn cột này.
          "GIAM_TRU_CAP_DON_KHAC":0
        }

    rows=[];errors=list(reward_errors)
    for rn,values in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(v not in (None,"") for v in values):continue
        d={h:values[idx[h]] for h in DATA_HEADERS}
        od=parse_date(d["NGAY_DON"]);pd=parse_date(d["NGAY_GIAO"])
        order=str(d["SO_HOA_DON"] or "").strip();sku=str(d["MA_SKU"] or "").strip()
        kh=str(d["MA_KH_NGUON"] or "").strip()
        nvkd=str(d["MA_NV_KD_NGHIEP_VU"] or "").strip()
        nvgh=str(d["MA_NV_GIAO_HANG_NGHIEP_VU"] or "").strip()
        driver=str(d["MA_NV_LAI_XE_NGHIEP_VU"] or "").strip()
        vehicle=str(d["PHUONG_TIEN"] or "").strip();lot=str(d["LO_GIAO_HANG"] or "").strip()
        bank_code=str(d["MA_TKNH"] or "").strip()

        rew=rewards.get(order,{"THUONG_TRUNG_BAY":0,"THUONG_TICH_LUY":0,"GIAM_TRU_CAP_DON_KHAC":0})
        d.update(rew)

        er=[]
        if not od:er.append("VAL01: NGAY_DON không hợp lệ")
        if not pd:er.append("VAL02: NGAY_GIAO không hợp lệ")
        elif od and pd!=od+datetime.timedelta(days=1):er.append("VAL02: NGAY_GIAO phải bằng NGAY_DON + 1")
        if not order:er.append("Thiếu SO_HOA_DON")
        if not sku:er.append("Thiếu MA_SKU")
        if sku and not resolve_product(sku):er.append(f"VAL03: SKU {sku} chưa mapping DM01")
        if kh and not resolve_customer(kh):er.append(f"VAL03: KH {kh} chưa mapping DM03")
        if nvkd and not resolve_staff(nvkd):er.append(f"VAL03: NVKD {nvkd} chưa mapping DM02A/DM02")
        if nvgh and not resolve_staff(nvgh):er.append(f"VAL03: NVGH {nvgh} chưa mapping DM02A/DM02")
        if driver and not resolve_staff(driver):er.append(f"VAL03: Lái xe {driver} chưa mapping DM02A/DM02")
        if not lot:er.append("VAL03: thiếu mã lô dự kiến")
        if not vehicle:er.append("VAL03: thiếu phương tiện dự kiến")
        if not find_bank(bank_code):er.append(f"VAL03: tài khoản nhận {bank_code} không hợp lệ")

        rows.append({"row_no":rn,"valid":not er,"errors":er,"data":d})
        if er:errors.append({"row_no":rn,"errors":er})

    return {
      "filename":file.filename,"checksum":checksum,
      "total_rows":len(rows),"valid_rows":sum(1 for r in rows if r["valid"]),
      "error_rows":sum(1 for r in rows if not r["valid"])+len(reward_errors),
      "reward_orders":len(rewards),"rows":rows,"errors":errors[:500]
    }
