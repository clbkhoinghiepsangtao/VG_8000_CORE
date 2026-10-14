
import sqlite3, json
from pathlib import Path
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def _db():
    return Path(__file__).resolve().parent.parent/"data"/"miniapp.db"

def _conn():
    c=sqlite3.connect(str(_db()))
    c.row_factory=sqlite3.Row
    return c

def _table_exists(c,t):
    return c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(t,)).fetchone() is not None

def _columns(c,t):
    return [r["name"] for r in c.execute(f"PRAGMA table_info({t})").fetchall()]

def _json_from_row(row):
    for k in ("data_json","row_json","payload_json","json_data","data"):
        if k in row.keys() and row[k]:
            try:
                v=json.loads(row[k])
                if isinstance(v,dict):
                    return v
            except:
                pass
    # fallback: physical columns as dict
    return {k:row[k] for k in row.keys()}

def _load_master_dataset(c, wanted_keys):
    """
    Flexible reader for generic master-data storage.
    Supports common dataset-key columns and JSON payload columns.
    """
    candidates=["master_data_generic","master_data","master_rows","dm_master"]
    out=[]
    for t in candidates:
        if not _table_exists(c,t):
            continue
        cols=_columns(c,t)
        keycol=None
        for x in ("dataset_key","master_key","category","data_type","source_type","table_key","dm_key"):
            if x in cols:
                keycol=x;break

        if keycol:
            qs=",".join("?" for _ in wanted_keys)
            try:
                rows=c.execute(f"SELECT * FROM {t} WHERE {keycol} IN ({qs})",tuple(wanted_keys)).fetchall()
            except:
                rows=[]
        else:
            try:
                rows=c.execute(f"SELECT * FROM {t}").fetchall()
            except:
                rows=[]

        for r in rows:
            d=_json_from_row(r)
            # when table has no dataset key, infer from row contents later
            d["_TABLE"]=t
            if keycol:
                d["_DATASET"]=r[keycol]
            out.append(d)
    return out

def _v(d,*keys):
    for k in keys:
        if k in d and d[k] not in (None,""):
            return str(d[k]).strip()
    return ""

def _upper(v):
    return str(v or "").strip().upper()

def _master_indexes(c):
    rows=_load_master_dataset(c,[
        "dm01-product","product",
        "dm02a-staff-map","staffmap","staff-map",
        "dm03-customer","customer"
    ])

    sku=set()
    nvmap={}
    kh_codes={}
    kh_phone={}
    kh_status={}

    for d in rows:
        ds=_upper(d.get("_DATASET",""))

        # Product
        ma_sku=_v(d,"MA_SKU","SKU","ma_sku")
        if ma_sku and ("PRODUCT" in ds or "DM01" in ds or "TEN_SKU" in d or "TEN_SAN_PHAM" in d):
            sku.add(_upper(ma_sku))

        # Staff business-code mapping
        ma_nv=_v(d,"MA_NV","ma_nv")
        ma_nghiep_vu=_v(d,"MA_NGHIEP_VU","MA_NV_NGHIEP_VU","MA_NV_NGUON","ma_nghiep_vu")
        if ma_nv and ma_nghiep_vu:
            nvmap[_upper(ma_nghiep_vu)]=_upper(ma_nv)

        # Customer mapping
        ma_kh_chuan=_v(d,"MA_KH_CHUAN","ma_kh_chuan")
        if ma_kh_chuan:
            for k in ("MA_KH_DMS","MA_DMS","MA_KH_MBW","MA_MBW","MA_KH_NGUON_1","MA_KH_NGUON_2"):
                code=_v(d,k)
                if code:
                    kh_codes[_upper(code)]=_upper(ma_kh_chuan)
            sdt=_v(d,"SDT_CHINH","SDT","SO_DIEN_THOAI")
            if sdt:
                kh_phone.setdefault(sdt,[]).append(_upper(ma_kh_chuan))
            kh_status[_upper(ma_kh_chuan)]=_v(d,"TRANG_THAI","STATUS")

    # Fallback: if generic rows did not expose product set, inspect any table with MA_SKU column.
    if not sku:
        for trow in c.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall():
            t=trow["name"]
            try:
                cols=_columns(c,t)
            except:
                continue
            if "MA_SKU" in cols and t not in ("xb_src_row_v2","xb_dms_mbw_row","xb_misa_bu_row"):
                try:
                    for r in c.execute(f"SELECT DISTINCT MA_SKU FROM {t} WHERE MA_SKU IS NOT NULL").fetchall():
                        sku.add(_upper(r["MA_SKU"]))
                except:
                    pass

    return {
        "sku":sku,
        "nvmap":nvmap,
        "kh_codes":kh_codes,
        "kh_phone":kh_phone,
        "kh_status":kh_status,
    }

def _source_rows(c):
    out=[]

    # DMS Xả
    if _table_exists(c,"xb_src_batch_v2") and _table_exists(c,"xb_src_row_v2"):
        batches=[r["batch_id"] for r in c.execute(
            "SELECT batch_id FROM xb_src_batch_v2 WHERE UPPER(status)='COMMITTED'"
        ).fetchall()]
        for bid in batches:
            for r in c.execute(
                "SELECT sheet_name,row_no,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='VALID'",
                (bid,)
            ).fetchall():
                if _upper(r["sheet_name"])!="CHI_TIET_XA":
                    continue
                try:d=json.loads(r["payload_json"])
                except:continue
                out.append({"source":"DMS-XA","batch":bid,"row_no":r["row_no"],"data":d})

    # DMS/MBW
    if _table_exists(c,"xb_dms_mbw_batch") and _table_exists(c,"xb_dms_mbw_row"):
        batches=[r["batch_id"] for r in c.execute(
            "SELECT batch_id FROM xb_dms_mbw_batch WHERE UPPER(status)='COMMITTED'"
        ).fetchall()]
        for bid in batches:
            for r in c.execute(
                "SELECT row_no,payload_json,status FROM xb_dms_mbw_row WHERE batch_id=? AND status='VALID'",
                (bid,)
            ).fetchall():
                try:d=json.loads(r["payload_json"])
                except:continue
                sheet=_upper(d.get("SHEET",""))
                if sheet=="THUONG_CAP_DON":
                    continue
                out.append({"source":"DMS-MBW","batch":bid,"row_no":r["row_no"],"data":d})

    # MISA
    if _table_exists(c,"xb_misa_bu_batch") and _table_exists(c,"xb_misa_bu_row"):
        batches=[r["batch_id"] for r in c.execute(
            "SELECT batch_id FROM xb_misa_bu_batch WHERE UPPER(status)='COMMITTED'"
        ).fetchall()]
        for bid in batches:
            for r in c.execute(
                "SELECT row_no,payload_json,status FROM xb_misa_bu_row WHERE batch_id=? AND status='VALID'",
                (bid,)
            ).fetchall():
                try:d=json.loads(r["payload_json"])
                except:continue
                out.append({"source":"MISA-BU","batch":bid,"row_no":r["row_no"],"data":d})

    return out

def run_mapping_audit():
    c=_conn()
    idx=_master_indexes(c)
    src=_source_rows(c)
    c.close()

    results=[]
    counts={
        "TOTAL":0,
        "MAP_OK":0,
        "CHUA_MAP_SKU":0,
        "CHUA_MAP_NV":0,
        "CHUA_MAP_KH":0,
        "CAN_XAC_NHAN_KH":0,
    }

    for item in src:
        d=item["data"]
        sku=_upper(_v(d,"MA_SKU"))
        nvsrc=_upper(_v(d,
            "MA_NV_NGHIEP_VU",
            "MA_NV_KD_NGHIEP_VU",
            "MA_NV",
            "MA_NHAN_VIEN"
        ))
        khsrc=_upper(_v(d,
            "MA_KH_NGUON",
            "MA_KH_DMS",
            "MA_KH_MBW",
            "MA_KH"
        ))
        phone=_v(d,"SDT","SDT_CHINH","SO_DIEN_THOAI")

        issues=[]
        mapped_nv=idx["nvmap"].get(nvsrc,"") if nvsrc else ""
        mapped_kh=idx["kh_codes"].get(khsrc,"") if khsrc else ""

        if sku and sku not in idx["sku"]:
            issues.append("CHUA_MAP_SKU")
        elif not sku:
            issues.append("CHUA_MAP_SKU")

        if nvsrc and not mapped_nv:
            issues.append("CHUA_MAP_NV")
        elif not nvsrc:
            issues.append("CHUA_MAP_NV")

        # MISA customer can be legitimately irrelevant for XB.
        if item["source"]!="MISA-BU":
            if khsrc:
                if not mapped_kh:
                    issues.append("CHUA_MAP_KH")
            elif phone:
                matches=idx["kh_phone"].get(phone,[])
                if len(matches)==1:
                    mapped_kh=matches[0]
                elif len(matches)>1:
                    issues.append("CAN_XAC_NHAN_KH")
                else:
                    issues.append("CHUA_MAP_KH")
            else:
                issues.append("CHUA_MAP_KH")

        if not issues:
            status="MAP_OK"
            counts["MAP_OK"]+=1
        else:
            status="|".join(dict.fromkeys(issues))
            for x in dict.fromkeys(issues):
                counts[x]=counts.get(x,0)+1

        counts["TOTAL"]+=1
        results.append({
            "SOURCE":item["source"],
            "BATCH":item["batch"],
            "ROW_NO":item["row_no"],
            "MA_SKU_NGUON":sku,
            "MA_NV_NGUON":nvsrc,
            "MA_NV_CHUAN":mapped_nv,
            "MA_KH_NGUON":khsrc,
            "MA_KH_CHUAN":mapped_kh,
            "STATUS":status,
        })

    return {"summary":counts,"rows":results[:500],"total_rows":len(results)}

def mapping_errors_xlsx():
    audit=run_mapping_audit()
    wb=Workbook()
    ws=wb.active
    ws.title="KET_QUA_MAPPING"

    headers=[
        "SOURCE","BATCH","ROW_NO","MA_SKU_NGUON",
        "MA_NV_NGUON","MA_NV_CHUAN",
        "MA_KH_NGUON","MA_KH_CHUAN","STATUS"
    ]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")
        ws.column_dimensions[get_column_letter(i)].width=22

    for r in audit["rows"]:
        if r["STATUS"]=="MAP_OK":
            continue
        ws.append([r.get(h,"") for h in headers])

    s=wb.create_sheet("TONG_HOP")
    s.append(["CHI_TIEU","SO_DONG"])
    for k,v in audit["summary"].items():
        s.append([k,v])

    bio=BytesIO()
    wb.save(bio)
    return "XB_Kiem_Tra_Mapping.xlsx",bio.getvalue()

from collections import defaultdict

def missing_master_lists():
    c=_conn()
    idx=_master_indexes(c)
    src=_source_rows(c)
    c.close()

    sku_missing={}
    nv_missing={}
    kh_missing={}

    def add(bucket, code, source, row_no, batch):
        code=_upper(code)
        if not code:
            return
        if code not in bucket:
            bucket[code]={
                "CODE":code,
                "SOURCES":set(),
                "COUNT":0,
                "FIRST_ROW":row_no,
                "BATCHES":set()
            }
        bucket[code]["SOURCES"].add(source)
        bucket[code]["BATCHES"].add(str(batch))
        bucket[code]["COUNT"]+=1

    for item in src:
        d=item["data"]
        source=item["source"]

        sku=_upper(_v(d,"MA_SKU"))
        nvsrc=_upper(_v(
            d,
            "MA_NV_NGHIEP_VU",
            "MA_NV_KD_NGHIEP_VU",
            "MA_NV",
            "MA_NHAN_VIEN"
        ))
        khsrc=_upper(_v(
            d,
            "MA_KH_NGUON",
            "MA_KH_DMS",
            "MA_KH_MBW",
            "MA_KH"
        ))

        if not sku or sku not in idx["sku"]:
            add(sku_missing,sku,source,item["row_no"],item["batch"])

        if not nvsrc or nvsrc not in idx["nvmap"]:
            add(nv_missing,nvsrc,source,item["row_no"],item["batch"])

        # Customer mapping is not required for MISA-BU in XB.
        if source!="MISA-BU":
            if khsrc and khsrc not in idx["kh_codes"]:
                add(kh_missing,khsrc,source,item["row_no"],item["batch"])

    def normalize(bucket):
        rows=[]
        for code,d in bucket.items():
            rows.append({
                "CODE":code,
                "SOURCES":", ".join(sorted(d["SOURCES"])),
                "COUNT":d["COUNT"],
                "BATCHES":", ".join(sorted(d["BATCHES"])),
                "FIRST_ROW":d["FIRST_ROW"],
            })
        rows.sort(key=lambda x:(-x["COUNT"],x["CODE"]))
        return rows

    return {
        "SKU_THIEU":normalize(sku_missing),
        "NV_THIEU":normalize(nv_missing),
        "KH_THIEU":normalize(kh_missing),
    }

def missing_master_xlsx():
    data=missing_master_lists()

    wb=Workbook()
    # remove default by using it for first sheet
    ws=wb.active
    ws.title="SKU_THIEU"

    def style_sheet(ws,headers,widths):
        for i,h in enumerate(headers,1):
            c=ws.cell(1,i,h)
            c.font=Font(bold=True,color="FFFFFF")
            c.fill=PatternFill("solid",fgColor="1769E0")
            ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions

    # Sheet 1: SKU missing
    headers=["MA_SKU","NGUON_PHAT_SINH","SO_DONG_ANH_HUONG","BATCH","DONG_DAU_TIEN","HUONG_XU_LY"]
    style_sheet(ws,headers,[18,32,18,15,15,42])
    for r in data["SKU_THIEU"]:
        ws.append([
            r["CODE"],r["SOURCES"],r["COUNT"],r["BATCHES"],r["FIRST_ROW"],
            "Bổ sung/kiểm tra tại Danh mục & Quản trị > Sản phẩm (DM01)"
        ])

    # Sheet 2: NV missing
    ws=wb.create_sheet("NV_THIEU")
    headers=["MA_NV_NGHIEP_VU","NGUON_PHAT_SINH","SO_DONG_ANH_HUONG","BATCH","DONG_DAU_TIEN","MA_NV_CHUAN_CAN_BO_SUNG","HUONG_XU_LY"]
    style_sheet(ws,headers,[22,32,18,15,15,26,48])
    for r in data["NV_THIEU"]:
        ws.append([
            r["CODE"],r["SOURCES"],r["COUNT"],r["BATCHES"],r["FIRST_ROW"],"",
            "Xác định nhân sự rồi bổ sung tại Mã nghiệp vụ NV (DM02A); không đổi MA_NV chuẩn"
        ])

    # Sheet 3: Customer missing
    ws=wb.create_sheet("KH_THIEU")
    headers=["MA_KH_NGUON","NGUON_PHAT_SINH","SO_DONG_ANH_HUONG","BATCH","DONG_DAU_TIEN","MA_KH_CHUAN_CAN_MAP","HUONG_XU_LY"]
    style_sheet(ws,headers,[20,32,18,15,15,24,48])
    for r in data["KH_THIEU"]:
        ws.append([
            r["CODE"],r["SOURCES"],r["COUNT"],r["BATCHES"],r["FIRST_ROW"],"",
            "Map vào khách hàng hiện có hoặc bổ sung tại DM03; MISA-Bù không yêu cầu map KH"
        ])

    bio=BytesIO()
    wb.save(bio)
    return "XB_Danh_sach_can_bo_sung_Danh_muc.xlsx",bio.getvalue()
