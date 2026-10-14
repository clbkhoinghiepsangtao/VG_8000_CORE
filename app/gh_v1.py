
import sqlite3, json, datetime, math
from pathlib import Path
from io import BytesIO
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def _db():
    return Path(__file__).resolve().parent.parent/"data"/"miniapp.db"

def _conn():
    c=sqlite3.connect(str(_db()))
    c.row_factory=sqlite3.Row
    return c

def _exists(c,t):
    return c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(t,)).fetchone() is not None

def _u(v):
    return str(v or "").strip().upper()

def _v(d,*keys):
    for k in keys:
        if k in d and d[k] not in (None,""):
            return d[k]
    return None

def _num(v):
    if v in (None,""): return 0.0
    if isinstance(v,str) and v.strip() in ("-","–","—"): return 0.0
    try:return float(v)
    except:
        try:return float(str(v).replace(",",""))
        except:return 0.0

def _date(v):
    if v in (None,""): return None
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        n=float(v)
        if 20000<=n<=80000:
            return datetime.date(1899,12,30)+datetime.timedelta(days=int(n))
    if isinstance(v,datetime.datetime): return v.date()
    if isinstance(v,datetime.date): return v
    s=str(v).strip()
    try:
        if s.replace(".","",1).isdigit():
            n=float(s)
            if 20000<=n<=80000:
                return datetime.date(1899,12,30)+datetime.timedelta(days=int(n))
    except: pass
    for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%Y/%m/%d","%Y-%m-%d %H:%M:%S"):
        try:return datetime.datetime.strptime(s[:19],fmt).date()
        except: pass
    try:return datetime.date.fromisoformat(s[:10])
    except:return None

def _ensure(c):
    c.executescript("""
    CREATE TABLE IF NOT EXISTS gh_settlement(
      settlement_id INTEGER PRIMARY KEY AUTOINCREMENT,
      order_no TEXT NOT NULL,
      order_date TEXT,
      customer_code TEXT,
      sales_code TEXT,
      delivery_code TEXT,
      driver_code TEXT,
      vehicle TEXT,
      delivery_lot TEXT,
      status TEXT NOT NULL DEFAULT 'DRAFT',
      cash_amount REAL NOT NULL DEFAULT 0,
      bank_amount REAL NOT NULL DEFAULT 0,
      customer_debt_amount REAL NOT NULL DEFAULT 0,
      amount_due REAL NOT NULL DEFAULT 0,
      delivery_debt_amount REAL NOT NULL DEFAULT 0,
      overpay_amount REAL NOT NULL DEFAULT 0,
      overpay_reason TEXT,
      created_at TEXT,
      created_by TEXT,
      committed_at TEXT,
      committed_by TEXT,
      adjustment_of INTEGER,
      adjustment_reason TEXT
    );
    CREATE INDEX IF NOT EXISTS ix_gh_settlement_order ON gh_settlement(order_no);
    CREATE TABLE IF NOT EXISTS gh_settlement_line(
      line_id INTEGER PRIMARY KEY AUTOINCREMENT,
      settlement_id INTEGER NOT NULL,
      sku TEXT NOT NULL,
      sku_name TEXT,
      ordered_qty REAL NOT NULL DEFAULT 0,
      return_qty REAL NOT NULL DEFAULT 0,
      actual_qty REAL NOT NULL DEFAULT 0,
      unit_price_prevat REAL NOT NULL DEFAULT 0,
      vat_rate REAL NOT NULL DEFAULT 0,
      discount_amount REAL NOT NULL DEFAULT 0,
      gss_actual REAL NOT NULL DEFAULT 0,
      sales_actual REAL NOT NULL DEFAULT 0,
      FOREIGN KEY(settlement_id) REFERENCES gh_settlement(settlement_id)
    );
    CREATE TABLE IF NOT EXISTS gh_customer_debt(
      debt_id INTEGER PRIMARY KEY AUTOINCREMENT,
      settlement_id INTEGER NOT NULL,
      order_no TEXT,
      customer_code TEXT,
      original_amount REAL NOT NULL,
      paid_amount REAL NOT NULL DEFAULT 0,
      balance REAL NOT NULL,
      status TEXT NOT NULL DEFAULT 'OPEN',
      created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS gh_customer_debt_payment(
      payment_id INTEGER PRIMARY KEY AUTOINCREMENT,
      debt_id INTEGER NOT NULL,
      cash_amount REAL NOT NULL DEFAULT 0,
      bank_amount REAL NOT NULL DEFAULT 0,
      overpay_amount REAL NOT NULL DEFAULT 0,
      overpay_reason TEXT,
      created_at TEXT,
      created_by TEXT
    );
    CREATE TABLE IF NOT EXISTS gh_delivery_debt(
      debt_id INTEGER PRIMARY KEY AUTOINCREMENT,
      settlement_id INTEGER NOT NULL,
      order_no TEXT,
      delivery_code TEXT,
      original_amount REAL NOT NULL,
      paid_amount REAL NOT NULL DEFAULT 0,
      balance REAL NOT NULL,
      status TEXT NOT NULL DEFAULT 'OPEN',
      created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS gh_delivery_debt_payment(
      payment_id INTEGER PRIMARY KEY AUTOINCREMENT,
      debt_id INTEGER NOT NULL,
      cash_amount REAL NOT NULL DEFAULT 0,
      bank_amount REAL NOT NULL DEFAULT 0,
      created_at TEXT,
      created_by TEXT
    );
    CREATE TABLE IF NOT EXISTS gh_cashier_receipt(
      receipt_id INTEGER PRIMARY KEY AUTOINCREMENT,
      ref_type TEXT NOT NULL,
      ref_id INTEGER NOT NULL,
      cash_amount REAL NOT NULL DEFAULT 0,
      bank_amount REAL NOT NULL DEFAULT 0,
      status TEXT NOT NULL DEFAULT 'PENDING',
      received_at TEXT,
      received_by TEXT
    );
    """)
    c.commit()

def _active_batches(c):
    if not _exists(c,"xb_dms_mbw_batch"): return []
    return [r["batch_id"] for r in c.execute(
        "SELECT batch_id FROM xb_dms_mbw_batch WHERE UPPER(status)='COMMITTED'"
    ).fetchall()]

def _product_ratio():
    try:
        from .xb_mapping_audit import _load_master_dataset
        c=_conn(); rows=_load_master_dataset(c,["dm01-product","product"]); c.close()
        out={}
        for d in rows:
            sku=_u(_v(d,"MA_SKU"))
            if not sku: continue
            ratio=_num(_v(d,"SO_SU_THUNG"))
            out[sku]=ratio if ratio>0 else 1.0
        return out
    except:
        return {}

def _source_order_lines(order_no=None):
    c=_conn()
    ratios=_product_ratio()
    out=[]
    if _exists(c,"xb_dms_mbw_row"):
        for bid in _active_batches(c):
            rows=c.execute(
                "SELECT row_no,payload_json FROM xb_dms_mbw_row WHERE batch_id=? AND status='VALID'",
                (bid,)
            ).fetchall()
            for r in rows:
                try:d=json.loads(r["payload_json"])
                except:continue
                if _u(_v(d,"LOAI_DON"))!="L": continue
                ono=str(_v(d,"SO_HOA_DON","MA_DON_DMS","MA_DON","SO_DON") or "").strip()
                if not ono: ono=f"B{bid}-R{r['row_no']}"
                if order_no and ono!=order_no: continue
                sku=_u(_v(d,"MA_SKU"))
                ratio=ratios.get(sku,1.0)
                sl_su=_num(_v(d,"SL_SU","SO_LUONG_SU"))
                sl_thung=_num(_v(d,"SL_THUNG","SO_LUONG_THUNG"))
                qty=sl_su+sl_thung*ratio
                price=_num(_v(d,"DON_GIA_CHUA_VAT","DON_GIA_CHUA_THUE"))
                vat=_num(_v(d,"VAT_RATE","VAT"))
                if vat>1: vat=vat/100.0
                discount=_num(_v(d,"CK_BANG_TIEN","CK_KM_TIEN","CHIET_KHAU"))
                out.append({
                    "order_no":ono,
                    "order_date":_date(_v(d,"NGAY_GIAO","NGAY_DON")),
                    "customer_code":str(_v(d,"MA_KH_NGUON","MA_KH_DMS","MA_KH") or "").strip(),
                    "sales_code":_u(_v(d,"MA_NV_KD_NGHIEP_VU","MA_NV_NGHIEP_VU")),
                    "delivery_code":_u(_v(d,"MA_NV_GIAO_HANG_NGHIEP_VU","MA_NV_GIAO_HANG")),
                    "driver_code":_u(_v(d,"MA_NV_LAI_XE_NGHIEP_VU","MA_NV_LAI_XE")),
                    "vehicle":_u(_v(d,"PHUONG_TIEN","BIEN_SO")),
                    "delivery_lot":_u(_v(d,"LO_GIAO_HANG","LO_GIAO")),
                    "sku":sku,
                    "sku_name":str(_v(d,"TEN_SKU") or "").strip(),
                    "ordered_qty":qty,
                    "price":price,
                    "vat":vat,
                    "discount":discount,
                })
    c.close()
    return out

def list_orders(date_from=None,date_to=None,status=None):
    lines=_source_order_lines()
    agg={}
    for x in lines:
        day=x["order_date"]
        df=_date(date_from) if date_from else None
        dt=_date(date_to) if date_to else None
        if df and (day is None or day<df): continue
        if dt and (day is None or day>dt): continue
        a=agg.setdefault(x["order_no"],{
            "ORDER_NO":x["order_no"],"NGAY_GIAO":str(day) if day else "",
            "CUSTOMER_CODE":x["customer_code"],"SALES_CODE":x["sales_code"],
            "DELIVERY_CODE":x["delivery_code"],"DRIVER_CODE":x["driver_code"],
            "VEHICLE":x["vehicle"],"DELIVERY_LOT":x["delivery_lot"],
            "SO_SKU":set(),"ORDERED_QTY":0.0
        })
        a["SO_SKU"].add(x["sku"]); a["ORDERED_QTY"]+=x["ordered_qty"]
    c=_conn(); _ensure(c)
    sts={}
    for r in c.execute("SELECT order_no,status,settlement_id FROM gh_settlement ORDER BY settlement_id DESC").fetchall():
        sts.setdefault(r["order_no"],dict(r))
    c.close()
    rows=[]
    for ono,a in agg.items():
        st=sts.get(ono,{})
        s=st.get("status","CHUA_QUYET_TOAN")
        if status and s!=status: continue
        rows.append({
            **{k:v for k,v in a.items() if k!="SO_SKU"},
            "SO_SKU":len(a["SO_SKU"]),
            "ORDERED_QTY":round(a["ORDERED_QTY"],4),
            "STATUS":s,
            "SETTLEMENT_ID":st.get("settlement_id")
        })
    rows.sort(key=lambda r:(r["NGAY_GIAO"],r["ORDER_NO"]),reverse=True)
    return {"rows":rows,"count":len(rows)}

def create_draft(order_no,created_by="pilot"):
    c=_conn(); _ensure(c)
    existing=c.execute(
        "SELECT settlement_id,status FROM gh_settlement WHERE order_no=? ORDER BY settlement_id DESC LIMIT 1",
        (order_no,)
    ).fetchone()
    if existing and existing["status"] in ("DRAFT","COMMITTED"):
        sid=existing["settlement_id"]; c.close(); return get_settlement(sid)

    src=_source_order_lines(order_no)
    if not src:
        c.close(); raise ValueError("Không tìm thấy đơn L trong nguồn COMMITTED")
    h=src[0]; now=datetime.datetime.now().isoformat(timespec="seconds")
    cur=c.execute("""INSERT INTO gh_settlement(
        order_no,order_date,customer_code,sales_code,delivery_code,driver_code,vehicle,delivery_lot,
        status,created_at,created_by
    ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",(
        order_no,str(h["order_date"] or ""),h["customer_code"],h["sales_code"],h["delivery_code"],
        h["driver_code"],h["vehicle"],h["delivery_lot"],"DRAFT",now,created_by
    ))
    sid=cur.lastrowid
    for x in src:
        ordered=x["ordered_qty"]; actual=ordered
        ratio=(actual/ordered) if ordered else 0
        disc=x["discount"]*ratio
        gss=actual*x["price"]
        sales=actual*x["price"]*(1+x["vat"])-disc
        c.execute("""INSERT INTO gh_settlement_line(
          settlement_id,sku,sku_name,ordered_qty,return_qty,actual_qty,unit_price_prevat,vat_rate,
          discount_amount,gss_actual,sales_actual
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",(
          sid,x["sku"],x["sku_name"],ordered,0,actual,x["price"],x["vat"],disc,gss,sales
        ))
    amount=c.execute("SELECT COALESCE(SUM(sales_actual),0) x FROM gh_settlement_line WHERE settlement_id=?",(sid,)).fetchone()["x"]
    c.execute("UPDATE gh_settlement SET amount_due=?,cash_amount=? WHERE settlement_id=?",(amount,amount,sid))
    c.commit(); c.close()
    return get_settlement(sid)

def get_settlement(sid):
    c=_conn(); _ensure(c)
    h=c.execute("SELECT * FROM gh_settlement WHERE settlement_id=?",(sid,)).fetchone()
    if not h: c.close(); raise ValueError("Không tìm thấy quyết toán")
    lines=[dict(r) for r in c.execute("SELECT * FROM gh_settlement_line WHERE settlement_id=? ORDER BY line_id",(sid,)).fetchall()]
    out=dict(h); out["lines"]=lines
    c.close(); return out

def update_draft(sid,lines,cash_amount=None,bank_amount=None,customer_debt_amount=None,overpay_reason=None):
    c=_conn(); _ensure(c)
    h=c.execute("SELECT status FROM gh_settlement WHERE settlement_id=?",(sid,)).fetchone()
    if not h: c.close(); raise ValueError("Không tìm thấy quyết toán")
    if h["status"]!="DRAFT": c.close(); raise ValueError("Chứng từ đã chốt, không được sửa trực tiếp")
    for item in lines or []:
        lid=int(item["line_id"]); ret=max(0,_num(item.get("return_qty")))
        row=c.execute("SELECT * FROM gh_settlement_line WHERE line_id=? AND settlement_id=?",(lid,sid)).fetchone()
        if not row: continue
        ordered=row["ordered_qty"]
        if ret>ordered+1e-9:
            c.close(); raise ValueError("Số lượng trả không được lớn hơn số lượng đặt")
        actual=ordered-ret
        ratio=(actual/ordered) if ordered else 0
        # discount stored on full draft line: derive full discount from current actual if needed
        full_disc=row["discount_amount"]/(row["actual_qty"]/ordered) if ordered and row["actual_qty"] else row["discount_amount"]
        disc=full_disc*ratio
        gss=actual*row["unit_price_prevat"]
        sales=actual*row["unit_price_prevat"]*(1+row["vat_rate"])-disc
        c.execute("""UPDATE gh_settlement_line
            SET return_qty=?,actual_qty=?,discount_amount=?,gss_actual=?,sales_actual=? WHERE line_id=?""",
            (ret,actual,disc,gss,sales,lid))
    amount=c.execute("SELECT COALESCE(SUM(sales_actual),0) x FROM gh_settlement_line WHERE settlement_id=?",(sid,)).fetchone()["x"]
    if cash_amount is None:
        cash_amount=c.execute("SELECT cash_amount FROM gh_settlement WHERE settlement_id=?",(sid,)).fetchone()["cash_amount"]
    if bank_amount is None:
        bank_amount=c.execute("SELECT bank_amount FROM gh_settlement WHERE settlement_id=?",(sid,)).fetchone()["bank_amount"]
    if customer_debt_amount is None:
        customer_debt_amount=c.execute("SELECT customer_debt_amount FROM gh_settlement WHERE settlement_id=?",(sid,)).fetchone()["customer_debt_amount"]
    customer_debt_amount=max(0,_num(customer_debt_amount))
    cash_amount=max(0,_num(cash_amount)); bank_amount=max(0,_num(bank_amount))
    required=max(0,amount-customer_debt_amount)
    received=cash_amount+bank_amount
    delivery_debt=max(0,required-received)
    overpay=max(0,received-required)
    if overpay>0 and not str(overpay_reason or "").strip():
        c.close(); raise ValueError("Khách/NV giao hàng trả thừa: bắt buộc ghi lý do")
    c.execute("""UPDATE gh_settlement SET amount_due=?,cash_amount=?,bank_amount=?,customer_debt_amount=?,
      delivery_debt_amount=?,overpay_amount=?,overpay_reason=? WHERE settlement_id=?""",
      (amount,cash_amount,bank_amount,customer_debt_amount,delivery_debt,overpay,str(overpay_reason or "").strip(),sid))
    c.commit(); c.close()
    return get_settlement(sid)

def commit_settlement(sid,user="pilot"):
    c=_conn(); _ensure(c)
    h=c.execute("SELECT * FROM gh_settlement WHERE settlement_id=?",(sid,)).fetchone()
    if not h: c.close(); raise ValueError("Không tìm thấy quyết toán")
    if h["status"]=="COMMITTED": c.close(); return get_settlement(sid)
    if h["status"]!="DRAFT": c.close(); raise ValueError("Trạng thái không cho phép chốt")
    now=datetime.datetime.now().isoformat(timespec="seconds")
    c.execute("UPDATE gh_settlement SET status='COMMITTED',committed_at=?,committed_by=? WHERE settlement_id=?",(now,user,sid))
    if h["customer_debt_amount"]>0:
        c.execute("""INSERT INTO gh_customer_debt(settlement_id,order_no,customer_code,original_amount,balance,status,created_at)
          VALUES(?,?,?,?,?,'OPEN',?)""",(sid,h["order_no"],h["customer_code"],h["customer_debt_amount"],h["customer_debt_amount"],now))
    if h["delivery_debt_amount"]>0:
        c.execute("""INSERT INTO gh_delivery_debt(settlement_id,order_no,delivery_code,original_amount,balance,status,created_at)
          VALUES(?,?,?,?,?,'OPEN',?)""",(sid,h["order_no"],h["delivery_code"],h["delivery_debt_amount"],h["delivery_debt_amount"],now))
    if h["cash_amount"]>0 or h["bank_amount"]>0:
        c.execute("""INSERT INTO gh_cashier_receipt(ref_type,ref_id,cash_amount,bank_amount,status)
          VALUES('SETTLEMENT',?,?,?,'PENDING')""",(sid,h["cash_amount"],h["bank_amount"]))
    c.commit(); c.close()
    return get_settlement(sid)

def customer_debts():
    c=_conn(); _ensure(c)
    rows=[dict(r) for r in c.execute("SELECT * FROM gh_customer_debt ORDER BY debt_id DESC").fetchall()]
    c.close(); return {"rows":rows,"count":len(rows)}

def collect_customer_debt(debt_id,cash,bank,reason="",user="pilot"):
    c=_conn(); _ensure(c)
    d=c.execute("SELECT * FROM gh_customer_debt WHERE debt_id=?",(debt_id,)).fetchone()
    if not d: c.close(); raise ValueError("Không tìm thấy công nợ KH")
    paid=max(0,_num(cash))+max(0,_num(bank)); over=max(0,paid-d["balance"])
    if over>0 and not str(reason).strip():
        c.close(); raise ValueError("Khách trả thừa: bắt buộc ghi lý do")
    applied=min(paid,d["balance"]); newbal=d["balance"]-applied
    now=datetime.datetime.now().isoformat(timespec="seconds")
    c.execute("""INSERT INTO gh_customer_debt_payment(debt_id,cash_amount,bank_amount,overpay_amount,overpay_reason,created_at,created_by)
      VALUES(?,?,?,?,?,?,?)""",(debt_id,max(0,_num(cash)),max(0,_num(bank)),over,str(reason).strip(),now,user))
    c.execute("UPDATE gh_customer_debt SET paid_amount=paid_amount+?,balance=?,status=? WHERE debt_id=?",
              (applied,newbal,"CLOSED" if newbal<=1e-9 else "OPEN",debt_id))
    c.execute("""INSERT INTO gh_cashier_receipt(ref_type,ref_id,cash_amount,bank_amount,status)
      VALUES('CUSTOMER_DEBT',?,?,?,'PENDING')""",(debt_id,max(0,_num(cash)),max(0,_num(bank))))
    c.commit(); c.close(); return customer_debts()

def delivery_debts():
    c=_conn(); _ensure(c)
    rows=[dict(r) for r in c.execute("SELECT * FROM gh_delivery_debt ORDER BY debt_id DESC").fetchall()]
    c.close(); return {"rows":rows,"count":len(rows)}

def collect_delivery_debt(debt_id,cash,bank,user="pilot"):
    c=_conn(); _ensure(c)
    d=c.execute("SELECT * FROM gh_delivery_debt WHERE debt_id=?",(debt_id,)).fetchone()
    if not d: c.close(); raise ValueError("Không tìm thấy công nợ NVGH")
    paid=max(0,_num(cash))+max(0,_num(bank)); applied=min(paid,d["balance"]); newbal=d["balance"]-applied
    now=datetime.datetime.now().isoformat(timespec="seconds")
    c.execute("""INSERT INTO gh_delivery_debt_payment(debt_id,cash_amount,bank_amount,created_at,created_by)
      VALUES(?,?,?,?,?)""",(debt_id,max(0,_num(cash)),max(0,_num(bank)),now,user))
    c.execute("UPDATE gh_delivery_debt SET paid_amount=paid_amount+?,balance=?,status=? WHERE debt_id=?",
              (applied,newbal,"CLOSED" if newbal<=1e-9 else "OPEN",debt_id))
    c.execute("""INSERT INTO gh_cashier_receipt(ref_type,ref_id,cash_amount,bank_amount,status)
      VALUES('DELIVERY_DEBT',?,?,?,'PENDING')""",(debt_id,max(0,_num(cash)),max(0,_num(bank))))
    c.commit(); c.close(); return delivery_debts()

def cashier_list():
    c=_conn(); _ensure(c)
    rows=[dict(r) for r in c.execute("SELECT * FROM gh_cashier_receipt ORDER BY receipt_id DESC").fetchall()]
    c.close(); return {"rows":rows,"count":len(rows)}

def cashier_receive(receipt_id,user="thuquy"):
    c=_conn(); _ensure(c)
    now=datetime.datetime.now().isoformat(timespec="seconds")
    c.execute("UPDATE gh_cashier_receipt SET status='RECEIVED',received_at=?,received_by=? WHERE receipt_id=?",(now,user,receipt_id))
    c.commit(); c.close(); return cashier_list()

def admin_export():
    wb=Workbook()
    c=_conn(); _ensure(c)
    specs=[
      ("QUYET_TOAN","SELECT * FROM gh_settlement ORDER BY settlement_id"),
      ("DONG_SKU","SELECT * FROM gh_settlement_line ORDER BY settlement_id,line_id"),
      ("NO_KH","SELECT * FROM gh_customer_debt ORDER BY debt_id"),
      ("NO_NVGH","SELECT * FROM gh_delivery_debt ORDER BY debt_id"),
      ("THU_QUY","SELECT * FROM gh_cashier_receipt ORDER BY receipt_id"),
    ]
    first=True
    for title,sql in specs:
        rows=c.execute(sql).fetchall()
        ws=wb.active if first else wb.create_sheet(title)
        first=False; ws.title=title
        if rows:
            headers=rows[0].keys()
            for i,h in enumerate(headers,1):
                cc=ws.cell(1,i,h); cc.font=Font(bold=True,color="FFFFFF"); cc.fill=PatternFill("solid",fgColor="1769E0")
            for r in rows: ws.append([r[h] for h in headers])
    c.close()
    bio=BytesIO(); wb.save(bio)
    return "GH_V1_Quan_tri.xlsx",bio.getvalue()
