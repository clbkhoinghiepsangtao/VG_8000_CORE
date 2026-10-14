from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .template_defs import TEMPLATES

def build_template(template_key: str) -> bytes:
    spec=TEMPLATES[template_key]
    wb=Workbook()
    wb.remove(wb.active)
    meta=wb.create_sheet('_META')
    meta['A1']='TEMPLATE_KEY'; meta['B1']=template_key
    meta['A2']='VERSION'; meta['B2']=spec['version']
    meta['A3']='RULE'; meta['B3']='Không đổi tên sheet/cột. Mỗi dòng dữ liệu = 1 SKU.'
    meta.sheet_state='hidden'
    guide=wb.create_sheet('HUONG_DAN')
    guide.append(['VĨNH GIANG – TEMPLATE DỮ LIỆU NGUỒN'])
    guide.append(['Template', template_key])
    guide.append(['Version', spec['version']])
    guide.append(['Nguyên tắc', 'Không đổi tên sheet/cột; không chèn/xóa cột. Chỉ điền dữ liệu nguồn.'])
    guide.append([])
    guide.append(['Sheet','Trường','Bắt buộc','Kiểu'])
    for sname, fields in spec['sheets'].items():
        for name, req, typ in fields: guide.append([sname,name,'Có' if req else 'Không',typ])
    for c in guide[1]:
        c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1769E0')
    guide.column_dimensions['A'].width=22; guide.column_dimensions['B'].width=34; guide.column_dimensions['C'].width=14; guide.column_dimensions['D'].width=20
    for sname, fields in spec['sheets'].items():
        ws=wb.create_sheet(sname)
        headers=[x[0] for x in fields]
        ws.append(headers)
        for c in ws[1]:
            c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1769E0'); c.alignment=Alignment(horizontal='center')
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for idx,(name,req,typ) in enumerate(fields,1):
            ws.column_dimensions[ws.cell(1,idx).column_letter].width=max(13,min(28,len(name)+3))
            if typ.startswith('enum:'):
                vals=typ.split(':',1)[1].split('|')
                dv=DataValidation(type='list', formula1='"'+','.join(vals)+'"', allow_blank=not req)
                ws.add_data_validation(dv); dv.add(f'{ws.cell(2,idx).column_letter}2:{ws.cell(5000,idx).column_letter}5000')
        ws.row_dimensions[1].height=24
    bio=BytesIO(); wb.save(bio); return bio.getvalue()
