
import sqlite3, json, hashlib, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from openpyxl import load_workbook

REQ={"NGAY_DON","MA_SKU","MA_NV_KD_NGHIEP_VU","SO_HOA_DON","LOAI_DON"}

ALIASES={
 "MA_KH_NGUON":{"MA_KH_NGUON","MA_KH"},
 "SL_THUNG":{"SL_THUNG","SL_THU"},
 "SL_KM_THUNG":{"SL_KM_THUNG","SL_KM_TH"},
 "DOANH_SO_PHAI_THU":{"DOANH_SO_PHAI_THU","DOANH_SO_TRUOC","DOANH_SO_TRUOC_CHIET_KHAU"},
}

def _db():
    return Path(__file__).resolve().parent.parent/"data"/"miniapp.db"

def _conn():
    c=sqlite3.connect(str(_db())); c.row_factory=sqlite3.Row
    c.execute("""CREATE TABLE IF NOT EXISTS xb_dms_mbw_batch(
      batch_id INTEGER PRIMARY KEY AUTOINCREMENT,
      file_name TEXT,file_hash TEXT UNIQUE,uploaded_by TEXT,uploaded_at TEXT,
      total_rows INTEGER DEFAULT 0,valid_rows INTEGER DEFAULT 0,error_rows INTEGER DEFAULT 0,status TEXT
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS xb_dms_mbw_row(
      row_id INTEGER PRIMARY KEY AUTOINCREMENT,batch_id INTEGER,row_no INTEGER,
      payload_json TEXT,status TEXT,error_text TEXT
    )""")
    c.commit(); return c

def _h(v):
    return "" if v is None else str(v).strip().upper().replace(" ","_")

def _headers(ws):
    return [_h(c.value) for c in ws[1]]

def _resolve(headers,name):
    hs=set(headers)
    if name in hs:return name
    for a in ALIASES.get(name,set()):
        if a in hs:return a
    return None

def preview_dms_mbw(file_name,raw_bytes,uploaded_by="pilot"):
    hsh=hashlib.sha256(raw_bytes).hexdigest()
    c=_conn()
    dup=c.execute("SELECT batch_id FROM xb_dms_mbw_batch WHERE file_hash=?",(hsh,)).fetchone()
    if dup:
        c.close()
        return {"duplicate":True,"duplicate_batch_id":dup["batch_id"]}

    wb=load_workbook(BytesIO(raw_bytes),data_only=False)
    names={s.upper():s for s in wb.sheetnames}

    detail_name=None
    for candidate in ("CHI_TIET_GIAO_BU","CHI_TIET_GIAO","CHI_TIET"):
        if candidate in names:
            detail_name=names[candidate]
            break
    if detail_name is None:
        for s in wb.sheetnames:
            if s.upper() in ("HUONG_DAN","HƯỚNG_DẪN","THUONG_CAP_DON"):
                continue
            detail_name=s
            break

    structure_errors=[]
    if detail_name is None:
        structure_errors.append("Thiếu sheet chi tiết giao/bù")
    if "THUONG_CAP_DON" not in names:
        structure_errors.append("Thiếu sheet THUONG_CAP_DON")

    detail_headers=[]
    if detail_name:
        detail_headers=_headers(wb[detail_name])
        missing=[x for x in REQ if _resolve(detail_headers,x) is None]
        if missing:
            structure_errors.append("Sheet chi tiết thiếu cột: "+", ".join(sorted(missing)))

    if "THUONG_CAP_DON" in names:
        bonus_headers=_headers(wb[names["THUONG_CAP_DON"]])
        bonus_required={"SO_HOA_DON","MA_NV_KD_NGHIEP_VU"}
        missing_bonus=sorted(bonus_required-set(bonus_headers))
        if missing_bonus:
            structure_errors.append("THUONG_CAP_DON thiếu cột: "+", ".join(missing_bonus))

    now=datetime.datetime.now().isoformat(timespec="seconds")
    cur=c.execute(
        "INSERT INTO xb_dms_mbw_batch(file_name,file_hash,uploaded_by,uploaded_at,status) VALUES(?,?,?,?,?)",
        (unquote(file_name),hsh,uploaded_by,now,"PREVIEW")
    )
    bid=cur.lastrowid
    total=valid=error=0
    preview=[]

    if structure_errors:
        error=1
        preview=[{"sheet":"","row_no":0,"status":"ERROR","error":"; ".join(structure_errors)}]
        c.execute(
            "INSERT INTO xb_dms_mbw_row(batch_id,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?)",
            (bid,0,"{}","ERROR","; ".join(structure_errors))
        )
    else:
        ws=wb[detail_name]
        headers=_headers(ws)
        detail_orders={}
        for rno,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            vals=list(row)
            if all(v is None or str(v).strip()=="" for v in vals):
                continue
            total+=1
            data={h:(vals[i] if i<len(vals) else None) for i,h in enumerate(headers) if h}
            er=[]
            sku=data.get(_resolve(headers,"MA_SKU"))
            nv=data.get(_resolve(headers,"MA_NV_KD_NGHIEP_VU"))
            invoice=str(data.get(_resolve(headers,"SO_HOA_DON")) or "").strip()
            loai=str(data.get(_resolve(headers,"LOAI_DON")) or "").strip().upper()

            if not sku: er.append("Thiếu MA_SKU")
            if not nv: er.append("Thiếu MA_NV_KD_NGHIEP_VU")
            if not invoice: er.append("Thiếu SO_HOA_DON")
            if loai not in ("L","B"): er.append("LOAI_DON phải là L hoặc B")

            if invoice:
                detail_orders.setdefault(invoice,set()).add(loai)

            st="ERROR" if er else "VALID"
            error += 1 if er else 0
            valid += 0 if er else 1
            et="; ".join(er)
            c.execute(
                "INSERT INTO xb_dms_mbw_row(batch_id,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?)",
                (bid,rno,json.dumps({"SHEET":detail_name,**data},ensure_ascii=False,default=str),st,et)
            )
            if len(preview)<90:
                preview.append({"sheet":detail_name,"row_no":rno,"status":st,"error":et})

        ws=wb[names["THUONG_CAP_DON"]]
        headers=_headers(ws)
        seen=set()
        for rno,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            vals=list(row)
            if all(v is None or str(v).strip()=="" for v in vals):
                continue
            total+=1
            data={h:(vals[i] if i<len(vals) else None) for i,h in enumerate(headers) if h}
            er=[]
            invoice=str(data.get("SO_HOA_DON") or "").strip()
            nv=str(data.get("MA_NV_KD_NGHIEP_VU") or "").strip()

            if not invoice: er.append("Thiếu SO_HOA_DON")
            if not nv: er.append("Thiếu MA_NV_KD_NGHIEP_VU")
            if invoice:
                if invoice in seen:
                    er.append("Trùng SO_HOA_DON trong THUONG_CAP_DON")
                seen.add(invoice)
                if invoice not in detail_orders:
                    er.append("SO_HOA_DON không tồn tại trong sheet chi tiết")
                elif "L" not in detail_orders[invoice]:
                    er.append("THUONG_CAP_DON chỉ áp dụng cho đơn LOAI_DON=L")

            st="ERROR" if er else "VALID"
            error += 1 if er else 0
            valid += 0 if er else 1
            et="; ".join(er)
            c.execute(
                "INSERT INTO xb_dms_mbw_row(batch_id,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?)",
                (bid,rno,json.dumps({"SHEET":"THUONG_CAP_DON",**data},ensure_ascii=False,default=str),st,et)
            )
            if len(preview)<100:
                preview.append({"sheet":"THUONG_CAP_DON","row_no":rno,"status":st,"error":et})

    c.execute(
        "UPDATE xb_dms_mbw_batch SET total_rows=?,valid_rows=?,error_rows=? WHERE batch_id=?",
        (total,valid,error,bid)
    )
    c.commit()
    c.close()
    return {"duplicate":False,"batch_id":bid,"total_rows":total,"valid_rows":valid,
            "error_rows":error,"structure_errors":structure_errors,"preview":preview}

def commit_dms_mbw(batch_id):
    c=_conn();b=c.execute("SELECT * FROM xb_dms_mbw_batch WHERE batch_id=?",(batch_id,)).fetchone()
    if not b:c.close();raise ValueError("Batch không tồn tại")
    if b["error_rows"]>0:c.close();raise ValueError("Batch còn lỗi")
    c.execute("UPDATE xb_dms_mbw_batch SET status='COMMITTED' WHERE batch_id=?",(batch_id,));c.commit();c.close()
    return {"batch_id":batch_id,"valid_rows":b["valid_rows"],"status":"COMMITTED"}

def history_dms_mbw():
    c=_conn();rows=c.execute("SELECT batch_id,file_name,uploaded_by,uploaded_at,total_rows,valid_rows,error_rows,status FROM xb_dms_mbw_batch ORDER BY batch_id DESC LIMIT 100").fetchall();c.close()
    return [dict(r) for r in rows]

# ===== DMS_MBW_LEGACY_FIELD_ALIASES_V3 =====
DMS_MBW_LEGACY_FIELD_ALIASES_V3 = {
    "DOANH_SO_TRUOC_VAT": "DOANH_SO_PHAI_THU",
    "DOANH_SO_TRUOC_CK": "DOANH_SO_PHAI_THU",
}
