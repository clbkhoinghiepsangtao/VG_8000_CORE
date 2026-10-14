"""Đọc XLSX; không tạo hoặc sửa workbook nguồn."""
from io import BytesIO
from zipfile import ZipFile,BadZipFile
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
import unicodedata

IDENTIFIERS={'SO_HOA_DON','MA_KH_NGUON','MA_SKU','MA_KH_CHUAN','MA_KH_DMS','MA_KH_MBW','MA_NV','MA_XE','MA_NGHIEP_VU','MA_KH','MA_NV_GIAO_HANG_NGHIEP_VU','MA_NV_LAI_XE_NGHIEP_VU','PHUONG_TIEN','LO_GIAO_HANG','MA_TKNH'}
SOURCE_IDENTIFIERS={'SỐ HÓA ĐƠN','MÃ CỬA HÀNG','MÃ HÀNG HÓA','MÃ NHÂN VIÊN','MÃ NHÂN VIÊN BÁN HÀNG'}
DATE_HEADERS={'NGAY_DON','NGÀY LẬP HOÁ ĐƠN'}

def read_sheet(content,sheet_name=None,*,allow_empty=False,required_headers=None,prefix_headers=None, ignore_columns=()):
    if len(content)>10*1024*1024:raise ValueError('Tệp vượt 10 MB')
    try:
        with ZipFile(BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist())>50*1024*1024:raise ValueError('Dung lượng giải nén vượt 50 MB')
    except BadZipFile:raise ValueError('Tệp không phải XLSX hợp lệ')
    wb=load_workbook(BytesIO(content),read_only=True,data_only=False)
    try:
        if sheet_name and sheet_name not in wb.sheetnames:raise ValueError('Không tìm thấy sheet '+sheet_name)
        if not sheet_name and len(wb.sheetnames)!=1:raise ValueError('Tệp nhiều sheet: chọn tên sheet cần đọc')
        sheet=wb[sheet_name or wb.sheetnames[0]]
        if (sheet.max_column or 0)>100 or (sheet.max_row or 0)>50001:raise ValueError('Tối đa 100 cột, 50.000 dòng dữ liệu')
        it=sheet.iter_rows();first=next(it,None)
        if not first:raise ValueError('Sheet trống')
        if prefix_headers:
            first=first[:len(prefix_headers)]
        first=tuple(c for i,c in enumerate(first) if i not in ignore_columns)
        heads=[unicodedata.normalize('NFC',str(c.value or '')).strip().upper() for c in first]
        if prefix_headers and tuple(heads)!=tuple(prefix_headers):raise ValueError('Các cột đầu phải là: '+', '.join(prefix_headers))
        while heads and not heads[-1]:heads.pop()
        if not heads or any(not h for h in heads) or len(set(heads))!=len(heads):raise ValueError('Tên cột trống hoặc trùng')
        if required_headers and set(required_headers)-set(heads):raise ValueError('Thiếu cột: '+', '.join(sorted(set(required_headers)-set(heads))))
        out=[]
        for rn,cells in enumerate(it,2):
            if rn>50001 or len(cells)>100:raise ValueError('Vượt giới hạn dòng/cột')
            if prefix_headers:cells=cells[:len(prefix_headers)]
            cells=tuple(c for i,c in enumerate(cells) if i not in ignore_columns)
            if all(c.value in (None,'') for c in cells):continue
            data={}
            for h,c in zip(heads,cells):
                if c.data_type=='f':raise ValueError(f'Dòng {rn}, {h}: cần giá trị, không nhận công thức')
                if c.data_type=='e':raise ValueError(f'Dòng {rn}, {h}: ô Excel lỗi')
                value=c.value
                if h in DATE_HEADERS and isinstance(value,(int,float)) and not isinstance(value,bool):
                    if value!=int(value) or not 1<=value<=2958465:raise ValueError(f'Dòng {rn}: ngày Excel không hợp lệ')
                    try:value=from_excel(value,wb.epoch).date().isoformat()
                    except (ValueError,OverflowError,AttributeError):raise ValueError(f'Dòng {rn}: ngày Excel không hợp lệ')
                if h in IDENTIFIERS|SOURCE_IDENTIFIERS and value is not None:
                    if isinstance(value,(int,float)) and not isinstance(value,bool):
                        if value!=int(value) or abs(value)>=10**15:raise ValueError(f'Dòng {rn}: mã dạng số không an toàn')
                        value=str(int(value))
                        if c.number_format and set(c.number_format)=={'0'}:value=value.zfill(len(c.number_format))
                    elif not isinstance(value,str):raise ValueError(f'Dòng {rn}: mã không hợp lệ')
                data[h]=value
            data['_SOURCE_ROW']=rn
            out.append(data)
        if not out and not allow_empty:raise ValueError('Không có dòng dữ liệu')
        return out
    finally:wb.close()
