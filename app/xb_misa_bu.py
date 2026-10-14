
import sqlite3, json, hashlib, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from openpyxl import load_workbook

ALIASES = {
    "MA_NV_NGHIEP_VU": {"MA_NV_NGHIEP_VU","MA_NV","MA_NHAN_VIEN","MA_NV_BAN_HANG","MA_NVBH"},
    "MA_SKU": {"MA_SKU","MA_HANG","MA_VT","MA_VAT_TU","MA_HANG_HOA"},
    "TONG_SL_BAN": {"TONG_SL_BAN","SL_BAN","SO_LUONG_BAN","SO_LUONG","SL_XUAT"},
    "TONG_SL_TRA_LAI": {"TONG_SL_TRA_LAI","SL_TRA_LAI","SO_LUONG_TRA_LAI","SL_TRA","HANG_TRA_LAI"},
    "CK_KM_TIEN": {"CK_KM_TIEN","CHIET_KHAU","TIEN_CHIET_KHAU","CK_TIEN","GIA_TRI_CHIET_KHAU"},
    "GIAM_TRU_DONG_KHAC": {"GIAM_TRU_DONG_KHAC","GIA_TRI_GIAM_GIA","GIAM_GIA","TIEN_GIAM_GIA"},
    "SO_CHUNG_TU": {"SO_CHUNG_TU","SO_CT","SO_HOA_DON","SO_PHIEU"},
    "NGAY_CHUNG_TU": {"NGAY_CHUNG_TU","NGAY_CT","NGAY_HACH_TOAN","NGAY_HOA_DON","NGAY"},
}

REQUIRED = {"MA_NV_NGHIEP_VU","MA_SKU","TONG_SL_BAN"}

def _db():
    return Path(__file__).resolve().parent.parent / "data" / "miniapp.db"

def _conn():
    c=sqlite3.connect(str(_db()))
    c.row_factory=sqlite3.Row
    c.execute("""CREATE TABLE IF NOT EXISTS xb_misa_bu_batch(
      batch_id INTEGER PRIMARY KEY AUTOINCREMENT,
      file_name TEXT,
      file_hash TEXT UNIQUE,
      uploaded_by TEXT,
      uploaded_at TEXT,
      total_rows INTEGER DEFAULT 0,
      valid_rows INTEGER DEFAULT 0,
      error_rows INTEGER DEFAULT 0,
      status TEXT
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS xb_misa_bu_row(
      row_id INTEGER PRIMARY KEY AUTOINCREMENT,
      batch_id INTEGER,
      sheet_name TEXT,
      row_no INTEGER,
      payload_json TEXT,
      status TEXT,
      error_text TEXT
    )""")
    c.commit()
    return c

def _h(v):
    if v is None: return ""
    return str(v).strip().upper().replace(" ","_")

def _headers(ws):
    return [_h(c.value) for c in ws[1]]

def _resolve(headers, canonical):
    hs=set(headers)
    for a in ALIASES.get(canonical,{canonical}):
        if a in hs:
            return a
    return None

def _num(v):
    if v in (None,""): return 0.0
    if isinstance(v,(int,float)): return float(v)
    s=str(v).strip().replace(",","")
    try: return float(s)
    except: return None

def preview_misa_bu(file_name, raw_bytes, uploaded_by="pilot"):
    hsh=hashlib.sha256(raw_bytes).hexdigest()
    c=_conn()

    # Only COMMITTED file is a true duplicate.
    # Failed/PREVIEW copy may be uploaded again after validator/template fixes.
    existing=c.execute(
        "SELECT batch_id,status FROM xb_misa_bu_batch WHERE file_hash=?",
        (hsh,)
    ).fetchone()

    if existing:
        if str(existing["status"] or "").upper()=="COMMITTED":
            c.close()
            return {"duplicate":True,"duplicate_batch_id":existing["batch_id"]}

        # Remove old failed/preview attempt and its rows so UNIQUE(file_hash)
        # does not prevent a fresh validation run.
        old_id=existing["batch_id"]
        c.execute("DELETE FROM xb_misa_bu_row WHERE batch_id=?",(old_id,))
        c.execute("DELETE FROM xb_misa_bu_batch WHERE batch_id=?",(old_id,))
        c.commit()

    wb=load_workbook(BytesIO(raw_bytes),data_only=False)

    ws=None
    sheet_map={s.upper():s for s in wb.sheetnames}
    if "DU_LIEU" in sheet_map:
        ws=wb[sheet_map["DU_LIEU"]]
    elif "DULIEU" in sheet_map:
        ws=wb[sheet_map["DULIEU"]]
    else:
        for s in wb.sheetnames:
            su=s.upper()
            if su in ("HUONG_DAN","HƯỚNG_DẪN"):
                continue
            if s.startswith("_"):
                continue
            cand=wb[s]
            if any(_headers(cand)):
                ws=cand
                break

    if ws is None:
        c.close()
        raise ValueError("Không tìm thấy sheet dữ liệu MISA")

    headers=_headers(ws)
    missing=[x for x in REQUIRED if _resolve(headers,x) is None]
    structure_errors=[]
    if missing:
        structure_errors.append("Thiếu cột bắt buộc: "+", ".join(sorted(missing)))

    now=datetime.datetime.now().isoformat(timespec="seconds")
    cur=c.execute(
        "INSERT INTO xb_misa_bu_batch(file_name,file_hash,uploaded_by,uploaded_at,status) VALUES(?,?,?,?,?)",
        (unquote(file_name),hsh,uploaded_by,now,"PREVIEW")
    )
    bid=cur.lastrowid

    total=valid=error=0
    preview=[]

    if structure_errors:
        error=1
        msg="; ".join(structure_errors)
        c.execute(
            "INSERT INTO xb_misa_bu_row(batch_id,sheet_name,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?,?)",
            (bid,ws.title,0,"{}","ERROR",msg)
        )
        preview.append({"sheet":ws.title,"row_no":0,"status":"ERROR","error":msg})
    else:
        for rno,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            vals=list(row)
            if all(v is None or str(v).strip()=="" for v in vals):
                continue

            total+=1
            raw={h:(vals[i] if i<len(vals) else None) for i,h in enumerate(headers) if h}

            er=[]
            nv=raw.get(_resolve(headers,"MA_NV_NGHIEP_VU"))
            sku=raw.get(_resolve(headers,"MA_SKU"))

            sl_ban=_num(raw.get(_resolve(headers,"TONG_SL_BAN")))
            tra_key=_resolve(headers,"TONG_SL_TRA_LAI")
            original_tra=raw.get(tra_key) if tra_key else None

            if isinstance(original_tra,str) and original_tra.strip() in ("-","–","—"):
                sl_tra=0.0
            else:
                sl_tra=_num(original_tra) if tra_key else 0.0

            if not nv: er.append("Thiếu MA_NV_NGHIEP_VU")
            if not sku: er.append("Thiếu MA_SKU")
            if sl_ban is None: er.append("TONG_SL_BAN không phải số")
            if sl_tra is None: er.append("TONG_SL_TRA_LAI không phải số")

            if sl_ban is not None and sl_tra is not None:
                sl_bu_thuc=sl_ban-sl_tra
                if sl_bu_thuc < 0:
                    er.append("SL trả lại lớn hơn SL bán")
            else:
                sl_bu_thuc=None

            ck_key=_resolve(headers,"CK_KM_TIEN")
            gt_key=_resolve(headers,"GIAM_TRU_DONG_KHAC")
            ck=_num(raw.get(ck_key)) if ck_key else 0.0
            gt=_num(raw.get(gt_key)) if gt_key else 0.0

            if ck is None: er.append("CHIET_KHAU không phải số")
            if gt is None: er.append("GIA_TRI_GIAM_GIA không phải số")

            normalized=dict(raw)
            normalized["SL_BU_THUC"]=sl_bu_thuc
            normalized["CK_KM_TIEN_CHUAN"]=ck if ck is not None else None
            normalized["GIAM_TRU_DONG_KHAC_CHUAN"]=gt if gt is not None else None

            st="ERROR" if er else "VALID"
            if er: error+=1
            else: valid+=1
            et="; ".join(er)

            c.execute(
                "INSERT INTO xb_misa_bu_row(batch_id,sheet_name,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?,?)",
                (bid,ws.title,rno,json.dumps(normalized,ensure_ascii=False,default=str),st,et)
            )

            if len(preview)<100:
                preview.append({"sheet":ws.title,"row_no":rno,"status":st,"error":et})

    c.execute(
        "UPDATE xb_misa_bu_batch SET total_rows=?,valid_rows=?,error_rows=? WHERE batch_id=?",
        (total,valid,error,bid)
    )
    c.commit()
    c.close()

    return {
        "duplicate":False,
        "batch_id":bid,
        "total_rows":total,
        "valid_rows":valid,
        "error_rows":error,
        "structure_errors":structure_errors,
        "preview":preview
    }

def commit_misa_bu(batch_id):
    c=_conn()
    b=c.execute("SELECT * FROM xb_misa_bu_batch WHERE batch_id=?",(int(batch_id),)).fetchone()
    if not b:
        c.close()
        raise ValueError("Batch không tồn tại")
    if int(b["error_rows"] or 0)>0:
        c.close()
        raise ValueError("Batch còn lỗi")
    c.execute("UPDATE xb_misa_bu_batch SET status='COMMITTED' WHERE batch_id=?",(int(batch_id),))
    c.commit()
    c.close()
    return {"batch_id":batch_id,"valid_rows":b["valid_rows"],"status":"COMMITTED"}

def history_misa_bu():
    c=_conn()
    rows=c.execute(
        "SELECT batch_id,file_name,uploaded_by,uploaded_at,total_rows,valid_rows,error_rows,status "
        "FROM xb_misa_bu_batch ORDER BY batch_id DESC LIMIT 100"
    ).fetchall()
    c.close()
    return [dict(r) for r in rows]
