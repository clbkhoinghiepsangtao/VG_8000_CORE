
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DETAIL_HEADERS = [
    "NGAY_DON","NGAY_GIAO","MA_SKU","TEN_SKU","MA_NV_KD_NGHIEP_VU",
    "SO_HOA_DON","DOANH_SO_PHAI_THU","MA_KH_NGUON","SL_SU","SL_KM_THUNG",
    "SL_THUNG","DON_GIA_CHUA_VAT","SL_KM_SU","PHUONG_TIEN","LOAI_DON",
    "MA_NV_GIAO_HANG_NGHIEP_VU","MA_NV_LAI_XE_NGHIEP_VU","LO_GIAO_HANG",
    "DON_APP","MA_TKNH","THU_TU_GIAO"
]

BONUS_HEADERS = [
    "SO_HOA_DON","NGAY_DON","MA_NV_KD_NGHIEP_VU",
    "THUONG_TRUNG_BAY","THUONG_TICH_LUY"
]

def _style_header(ws):
    for c in ws[1]:
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")
        c.alignment=Alignment(horizontal="center",vertical="center")
    ws.freeze_panes="A2"
    for i,c in enumerate(ws[1],1):
        ws.column_dimensions[get_column_letter(i)].width=max(14,min(30,len(str(c.value))+4))

def build_dms_mbw_v2_template():
    wb=Workbook()
    ws=wb.active
    ws.title="HUONG_DAN"
    for r in [
        ["TEMPLATE DMS/MBW – GIAO THẬT / BÁN BÙ – IN CHỨNG TỪ V2.1"],
        [""],
        ["CHI_TIET_GIAO_BU: 1 dòng = 1 SKU."],
        ["LOAI_DON = L: giao thật DMS."],
        ["LOAI_DON = B: bán bù MBW."],
        ["THUONG_CAP_DON: 1 dòng = 1 đơn; chỉ áp dụng cho LOAI_DON=L."],
        ["Đơn B không khai báo thưởng cấp đơn."],
        ["DON_APP: ghi DMS nếu là đơn khách tự đặt trên DMS App; nếu không thì để trống."],
        ["MA_TKNH: bắt buộc khi in; dùng để tạo QR nhận tiền cho từng đơn."],
        ["THU_TU_GIAO: số thứ tự giao trong lô."],
        ["NGAY_GIAO dự kiến = NGAY_DON + 1; D0 luôn tính từ NGAY_DON."],
        ["Không đổi tên sheet hoặc tên cột."],
    ]:
        ws.append(r)
    ws["A1"].font=Font(bold=True,size=14,color="17365D")
    ws.column_dimensions["A"].width=95

    ws=wb.create_sheet("CHI_TIET_GIAO_BU")
    ws.append(DETAIL_HEADERS)
    _style_header(ws)

    ws=wb.create_sheet("THUONG_CAP_DON")
    ws.append(BONUS_HEADERS)
    _style_header(ws)

    bio=BytesIO()
    wb.save(bio)
    return "02_Template_DMS_MBW_GiaoThat_BanBu_InChungTu_V2_1.xlsx", bio.getvalue()
