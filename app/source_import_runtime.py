
import sqlite3, json, hashlib, datetime
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE_CONFIG = {
    "dms-xa": {
        "label": "DMS – Xả",
        "template": "/api/templates/dms-xa.xlsx",
        "required_sheets": ["CHI_TIET_XA", "THUONG_CAP_DON"],
        "required_any": [["MA_DON_DMS", "MA_NV", "MA_SKU"], ["SO_DON", "MA_NV", "MA_SKU"]],
    },
    "dms-mbw": {
        "label": "DMS/MBW – Giao thật / Bán bù",
        "template": "/api/templates/dms-mbw-v2.xlsx",
        "required_sheets": [],
        "required_any": [["MA_NV", "MA_SKU", "LOAI_DON"], ["MA_NVKD", "MA_SKU", "LOAI_DON"]],
    },
    "misa-ban-bu": {
        "label": "MISA – Bán bù",
        "template": "/api/templates/misa-ban-bu.xlsx",
        "required_sheets": [],
        "required_any": [["MA_NV", "MA_SKU"], ["MA_NHAN_VIEN", "MA_SKU"]],
    },
}

def _db_path():
    return Path(__file__).resolve().parent.parent / "data" / "miniapp.db"

def _conn():
    p=_db_path()
    p.parent.mkdir(exist_ok=True)
    c=sqlite3.connect(str(p))
    c.row_factory=sqlite3.Row
    c.execute("""CREATE TABLE IF NOT EXISTS src_batch(
        batch_id INTEGER PRIMARY KEY AUTOINCREMENT,
        source_type TEXT NOT NULL,
        file_name TEXT NOT NULL,
        file_hash TEXT NOT NULL,
        uploaded_by TEXT,
        uploaded_at TEXT NOT NULL,
        total_rows INTEGER NOT NULL DEFAULT 0,
        valid_rows INTEGER NOT NULL DEFAULT 0,
        error_rows INTEGER NOT NULL DEFAULT 0,
        status TEXT NOT NULL,
        UNIQUE(source_type,file_hash)
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS src_row(
        row_id INTEGER PRIMARY KEY AUTOINCREMENT,
        batch_id INTEGER NOT NULL,
        sheet_name TEXT,
        row_no INTEGER,
        payload_json TEXT NOT NULL,
        status TEXT NOT NULL,
        error_text TEXT,
        FOREIGN KEY(batch_id) REFERENCES src_batch(batch_id)
    )""")
    c.commit()
    return c

def _norm_header(v):
    return "" if v is None else str(v).strip().upper().replace(" ", "_")

def _sheet_headers(ws):
    return [_norm_header(c.value) for c in ws[1] if c.value is not None]

def _structure_errors(source_type, wb):
    cfg=SOURCE_CONFIG[source_type]
    errors=[]
    sheet_upper={s.upper():s for s in wb.sheetnames}
    for req in cfg.get("required_sheets",[]):
        if req.upper() not in sheet_upper:
            errors.append("Thiếu sheet bắt buộc: "+req)
    accepted=False
    for s in wb.sheetnames:
        headers=set(_sheet_headers(wb[s]))
        for fieldset in cfg.get("required_any",[]):
            if set(fieldset).issubset(headers):
                accepted=True
                break
        if accepted:
            break
    if not accepted:
        errors.append("Không nhận diện được cấu trúc cột của nguồn đã chọn. Hãy dùng đúng file lấy từ Template dữ liệu nguồn.")
    return errors

def preview_source(source_type, file_name, raw_bytes, uploaded_by="pilot"):
    if source_type not in SOURCE_CONFIG:
        raise ValueError("Nguồn dữ liệu không hợp lệ")
    if not raw_bytes:
        raise ValueError("File rỗng")

    file_hash=hashlib.sha256(raw_bytes).hexdigest()
    conn=_conn()
    dup=conn.execute(
        "SELECT batch_id,status,uploaded_at,file_name FROM src_batch WHERE source_type=? AND file_hash=?",
        (source_type,file_hash)
    ).fetchone()
    if dup:
        conn.close()
        return {"duplicate":True,"duplicate_batch_id":dup["batch_id"],
                "message":"File này đã được nạp trước đó; hệ thống không cho tạo batch trùng."}

    try:
        wb=load_workbook(BytesIO(raw_bytes), data_only=False)
    except Exception as e:
        conn.close()
        raise ValueError("Không đọc được file Excel: "+str(e))

    structure_errors=_structure_errors(source_type,wb)
    now=datetime.datetime.now().isoformat(timespec="seconds")
    cur=conn.execute(
        "INSERT INTO src_batch(source_type,file_name,file_hash,uploaded_by,uploaded_at,status) VALUES(?,?,?,?,?,?)",
        (source_type,file_name,file_hash,uploaded_by,now,"PREVIEW")
    )
    batch_id=cur.lastrowid
    total=valid=error=0
    preview=[]

    if structure_errors:
        msg="; ".join(structure_errors)
        conn.execute(
            "INSERT INTO src_row(batch_id,sheet_name,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?,?)",
            (batch_id,"",0,"{}","ERROR",msg)
        )
        error=1
        preview.append({"sheet":"","row_no":0,"status":"ERROR","error":msg,"data":{}})
    else:
        for sheet_name in wb.sheetnames:
            ws=wb[sheet_name]
            headers=[_norm_header(c.value) for c in ws[1]]
            if not any(headers):
                continue
            for rno,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
                vals=list(row)
                if all(v is None or str(v).strip()=="" for v in vals):
                    continue
                total += 1
                data={}
                for i,h in enumerate(headers):
                    if h:
                        data[h]=vals[i] if i < len(vals) else None

                row_errors=[]
                if source_type in ("dms-xa","dms-mbw"):
                    if not data.get("MA_SKU"): row_errors.append("Thiếu MA_SKU")
                    if not (data.get("MA_NV") or data.get("MA_NVKD")): row_errors.append("Thiếu mã nhân viên")
                elif source_type=="misa-ban-bu":
                    if not data.get("MA_SKU"): row_errors.append("Thiếu MA_SKU")
                    if not (data.get("MA_NV") or data.get("MA_NHAN_VIEN")): row_errors.append("Thiếu mã nhân viên")

                status="ERROR" if row_errors else "VALID"
                err="; ".join(row_errors)
                if status=="VALID": valid += 1
                else: error += 1
                conn.execute(
                    "INSERT INTO src_row(batch_id,sheet_name,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?,?)",
                    (batch_id,sheet_name,rno,json.dumps(data,ensure_ascii=False,default=str),status,err)
                )
                if len(preview)<100:
                    preview.append({"sheet":sheet_name,"row_no":rno,"status":status,"error":err,"data":data})

    conn.execute("UPDATE src_batch SET total_rows=?,valid_rows=?,error_rows=? WHERE batch_id=?",
                 (total,valid,error,batch_id))
    conn.commit(); conn.close()
    return {"duplicate":False,"batch_id":batch_id,"source_type":source_type,
            "total_rows":total,"valid_rows":valid,"error_rows":error,
            "structure_errors":structure_errors,"preview":preview}

def commit_source(batch_id):
    conn=_conn()
    b=conn.execute("SELECT * FROM src_batch WHERE batch_id=?",(batch_id,)).fetchone()
    if not b:
        conn.close(); raise ValueError("Batch không tồn tại")
    if b["status"]=="COMMITTED":
        conn.close(); return {"batch_id":batch_id,"status":"COMMITTED","already_committed":True}
    if b["error_rows"]>0:
        conn.close(); raise ValueError("Batch còn dòng lỗi; chưa được Commit.")
    conn.execute("UPDATE src_batch SET status='COMMITTED' WHERE batch_id=?",(batch_id,))
    conn.commit(); conn.close()
    return {"batch_id":batch_id,"status":"COMMITTED","valid_rows":b["valid_rows"]}

def list_batches(limit=100):
    conn=_conn()
    rows=conn.execute(
        "SELECT batch_id,source_type,file_name,uploaded_by,uploaded_at,total_rows,valid_rows,error_rows,status "
        "FROM src_batch ORDER BY batch_id DESC LIMIT ?",(int(limit),)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]

def error_xlsx(batch_id):
    conn=_conn()
    rows=conn.execute(
        "SELECT sheet_name,row_no,payload_json,error_text FROM src_row "
        "WHERE batch_id=? AND status='ERROR' ORDER BY row_id",(batch_id,)
    ).fetchall()
    conn.close()
    wb=Workbook(); ws=wb.active; ws.title="LOI_UPLOAD"
    headers=["SHEET","ROW_NO","ERROR","DATA_JSON"]
    widths=[22,12,55,90]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="C62828")
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
    for r in rows:
        ws.append([r["sheet_name"],r["row_no"],r["error_text"],r["payload_json"]])
    bio=BytesIO(); wb.save(bio)
    return f"XB_Loi_Batch_{batch_id}.xlsx",bio.getvalue()
