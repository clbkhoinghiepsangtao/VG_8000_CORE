
import sqlite3, json, datetime
from pathlib import Path
from collections import defaultdict
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def _db():
    return Path(__file__).resolve().parent.parent/"data"/"miniapp.db"

def _conn():
    c=sqlite3.connect(str(_db()))
    c.row_factory=sqlite3.Row
    return c

def _table_exists(c,t):
    return c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(t,)).fetchone() is not None

def _upper(v):
    return str(v or "").strip().upper()

def _num(v):
    if v is None or v=="":
        return 0.0
    if isinstance(v,str) and v.strip() in ("-","–","—"):
        return 0.0
    try:
        return float(v)
    except:
        try:
            return float(str(v).replace(",",""))
        except:
            return 0.0

def _date(v):
    if v is None or v=="":
        return None

    # Excel serial date, e.g. 46253 = 2026-08-19.
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        try:
            n=float(v)
            if 20000 <= n <= 80000:
                return datetime.date(1899,12,30)+datetime.timedelta(days=int(n))
        except:
            pass

    if isinstance(v,(datetime.date,datetime.datetime)):
        return v.date() if isinstance(v,datetime.datetime) else v

    s=str(v).strip()

    # Numeric Excel serial stored as text.
    try:
        if s.replace(".","",1).isdigit():
            n=float(s)
            if 20000 <= n <= 80000:
                return datetime.date(1899,12,30)+datetime.timedelta(days=int(n))
    except:
        pass

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%Y-%m-%d %H:%M:%S"
    ):
        try:
            return datetime.datetime.strptime(s[:19],fmt).date()
        except:
            pass

    try:
        return datetime.date.fromisoformat(s[:10])
    except:
        return None

def _v(d,*keys):
    for k in keys:
        if k in d and d[k] not in (None,""):
            return d[k]
    return None

def _active_batches(c,table):
    if not _table_exists(c,table):
        return []
    return [r["batch_id"] for r in c.execute(
        f"SELECT batch_id FROM {table} WHERE UPPER(status)='COMMITTED'"
    ).fetchall()]

def _load_mapping():
    from .xb_mapping_audit import _master_indexes
    c=_conn()
    idx=_master_indexes(c)
    c.close()
    return idx

def _map_nv(idx,d):
    src=_upper(_v(d,
        "MA_NV_NGHIEP_VU",
        "MA_NV_KD_NGHIEP_VU",
        "MA_NV",
        "MA_NHAN_VIEN"
    ))
    return idx["nvmap"].get(src,src)

def _qty_su(d,prefix=""):
    # Prefer already-normalized SU quantity if present
    if prefix=="XA":
        v=_v(d,"SL_BAN_QUY_SU","SL_XA_QUY_SU")
        if v is not None:
            return _num(v)
    if prefix=="BU":
        v=_v(d,"SL_BU_THUC","SL_BU_QUY_SU")
        if v is not None:
            return _num(v)

    # Fallback raw SU only. Carton conversion may already be normalized upstream.
    if prefix=="XA":
        return _num(_v(d,"SL_SU","SO_LUONG_SU"))
    return _num(_v(d,"SL_SU","TONG_SL_BAN","SO_LUONG_SU"))

def _iter_dms_xa(c):
    if not (_table_exists(c,"xb_src_batch_v2") and _table_exists(c,"xb_src_row_v2")):
        return
    for bid in _active_batches(c,"xb_src_batch_v2"):
        rows=c.execute(
            "SELECT sheet_name,row_no,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='VALID'",
            (bid,)
        ).fetchall()
        for r in rows:
            if _upper(r["sheet_name"])!="CHI_TIET_XA":
                continue
            try:d=json.loads(r["payload_json"])
            except:continue
            yield bid,r["row_no"],d

def _iter_dms_mbw(c):
    if not (_table_exists(c,"xb_dms_mbw_batch") and _table_exists(c,"xb_dms_mbw_row")):
        return
    for bid in _active_batches(c,"xb_dms_mbw_batch"):
        rows=c.execute(
            "SELECT row_no,payload_json FROM xb_dms_mbw_row WHERE batch_id=? AND status='VALID'",
            (bid,)
        ).fetchall()
        for r in rows:
            try:d=json.loads(r["payload_json"])
            except:continue
            yield bid,r["row_no"],d

def _iter_misa(c):
    if not (_table_exists(c,"xb_misa_bu_batch") and _table_exists(c,"xb_misa_bu_row")):
        return
    for bid in _active_batches(c,"xb_misa_bu_batch"):
        rows=c.execute(
            "SELECT row_no,payload_json FROM xb_misa_bu_row WHERE batch_id=? AND status='VALID'",
            (bid,)
        ).fetchall()
        for r in rows:
            try:d=json.loads(r["payload_json"])
            except:continue
            yield bid,r["row_no"],d

def quantity_report(date_from=None,date_to=None,ma_nv=None,ma_sku=None):
    idx=_load_mapping()
    c=_conn()
    agg=defaultdict(lambda:{
        "DU_DAU":0.0,
        "XA":0.0,
        "BU_MBW":0.0,
        "BU_MISA":0.0,
    })

    df=_date(date_from) if date_from else None
    dt=_date(date_to) if date_to else None
    f_nv=_upper(ma_nv)
    f_sku=_upper(ma_sku)

    # DMS Xả
    for bid,row_no,d in _iter_dms_xa(c) or []:
        ngay=_v5_event_date(d,"DMS-XA")
        if df and ngay and ngay<df: continue
        if dt and ngay and ngay>dt: continue
        nv=_upper(_map_nv(idx,d))
        sku=_upper(_v(d,"MA_SKU"))
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        agg[(nv,sku)]["XA"] += _qty_su(d,"XA")

    # DMS/MBW - only B is compensation for XB
    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B":
            continue
        ngay=_v5_event_date(d,"DMS-MBW")
        if df and ngay and ngay<df: continue
        if dt and ngay and ngay>dt: continue
        nv=_upper(_map_nv(idx,d))
        sku=_upper(_v(d,"MA_SKU"))
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        agg[(nv,sku)]["BU_MBW"] += _qty_su(d,"BU")

    # MISA Bù
    for bid,row_no,d in _iter_misa(c) or []:
        ngay=_v5_event_date(d,"MISA-BU")
        if df and ngay and ngay<df: continue
        if dt and ngay and ngay>dt: continue
        nv=_upper(_map_nv(idx,d))
        sku=_upper(_v(d,"MA_SKU"))
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        agg[(nv,sku)]["BU_MISA"] += _num(_v(d,"SL_BU_THUC","TONG_SL_BAN"))

    c.close()

    rows=[]
    for (nv,sku),x in agg.items():
        bu=x["BU_MBW"]+x["BU_MISA"]
        du_cuoi=x["DU_DAU"]-x["XA"]+bu
        rows.append({
            "MA_NV":nv,
            "MA_SKU":sku,
            "DU_DAU":round(x["DU_DAU"],4),
            "XA_TRONG_KY":round(x["XA"],4),
            "BU_MBW":round(x["BU_MBW"],4),
            "BU_MISA":round(x["BU_MISA"],4),
            "BAN_BU_TRONG_KY":round(bu,4),
            "DU_CUOI":round(du_cuoi,4),
        })

    rows.sort(key=lambda r:(r["MA_NV"],r["MA_SKU"]))
    total={
        "DU_DAU":round(sum(r["DU_DAU"] for r in rows),4),
        "XA_TRONG_KY":round(sum(r["XA_TRONG_KY"] for r in rows),4),
        "BU_MBW":round(sum(r["BU_MBW"] for r in rows),4),
        "BU_MISA":round(sum(r["BU_MISA"] for r in rows),4),
        "BAN_BU_TRONG_KY":round(sum(r["BAN_BU_TRONG_KY"] for r in rows),4),
        "DU_CUOI":round(sum(r["DU_CUOI"] for r in rows),4),
    }
    return {"rows":rows,"total":total,"count":len(rows)}

def quantity_report_xlsx(date_from=None,date_to=None,ma_nv=None,ma_sku=None):
    rep=quantity_report(date_from,date_to,ma_nv,ma_sku)
    wb=Workbook()
    ws=wb.active
    ws.title="SO_LUONG_XA_BU"
    headers=["MA_NV","MA_SKU","DU_DAU","XA_TRONG_KY","BU_MBW","BU_MISA","BAN_BU_TRONG_KY","DU_CUOI"]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")
        ws.column_dimensions[get_column_letter(i)].width=20
    for r in rep["rows"]:
        ws.append([r[h] for h in headers])
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    bio=BytesIO()
    wb.save(bio)
    return "XB_UI02_Bao_cao_So_luong.xlsx",bio.getvalue()

def _money(d,*keys):
    return sum(_num(_v(d,k)) for k in keys)

def ckkm_report(date_from=None,date_to=None,ma_nv=None):
    idx=_load_mapping()
    c=_conn()
    agg=defaultdict(lambda:{"THU_CK_KM_XA":0.0,"THU_THUONG_CAP_DON":0.0,"DA_CHI_MBW":0.0,"DA_CHI_MISA":0.0})
    df=_date(date_from) if date_from else None
    dt=_date(date_to) if date_to else None
    f_nv=_upper(ma_nv)

    if _table_exists(c,"xb_src_batch_v2") and _table_exists(c,"xb_src_row_v2"):
        for bid in _active_batches(c,"xb_src_batch_v2"):
            rows=c.execute("SELECT sheet_name,row_no,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='VALID'",(bid,)).fetchall()
            for r in rows:
                try:d=json.loads(r["payload_json"])
                except:continue
                ngay=_v5_event_date(d,"DMS-XA")
                if df and ngay and ngay<df: continue
                if dt and ngay and ngay>dt: continue
                nv=_upper(_map_nv(idx,d))
                if f_nv and nv!=f_nv: continue
                if not nv: continue
                sheet=_upper(r["sheet_name"])
                if sheet=="CHI_TIET_XA":
                    agg[nv]["THU_CK_KM_XA"] += _money(d,"CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI","KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC")
                elif sheet=="THUONG_CAP_DON":
                    agg[nv]["THU_THUONG_CAP_DON"] += _money(d,"THUONG_TRUNG_BAY","THUONG_TICH_LUY","GIAM_TRU_CAP_DON_KHAC")

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B": continue
        ngay=_v5_event_date(d,"DMS-MBW")
        if df and ngay and ngay<df: continue
        if dt and ngay and ngay>dt: continue
        nv=_upper(_map_nv(idx,d))
        if f_nv and nv!=f_nv: continue
        if nv:
            agg[nv]["DA_CHI_MBW"] += _money(d,"CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI","KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC")

    for bid,row_no,d in _iter_misa(c) or []:
        ngay=_v5_event_date(d,"MISA-BU")
        if df and ngay and ngay<df: continue
        if dt and ngay and ngay>dt: continue
        nv=_upper(_map_nv(idx,d))
        if f_nv and nv!=f_nv: continue
        if nv:
            agg[nv]["DA_CHI_MISA"] += _money(d,"CK_KM_TIEN_CHUAN","CK_KM_TIEN","CHIET_KHAU","GIAM_TRU_DONG_KHAC_CHUAN","GIAM_TRU_DONG_KHAC","GIA_TRI_GIAM_GIA")
    c.close()

    rows=[]
    for nv,x in agg.items():
        thu=x["THU_CK_KM_XA"]+x["THU_THUONG_CAP_DON"]
        chi=x["DA_CHI_MBW"]+x["DA_CHI_MISA"]
        rows.append({"MA_NV":nv,"DU_DAU":0.0,"THU_CK_KM_XA":round(x["THU_CK_KM_XA"],2),"THU_THUONG_CAP_DON":round(x["THU_THUONG_CAP_DON"],2),"TONG_THU":round(thu,2),"DA_CHI_MBW":round(x["DA_CHI_MBW"],2),"DA_CHI_MISA":round(x["DA_CHI_MISA"],2),"TONG_DA_CHI":round(chi,2),"DU_CUOI":round(thu-chi,2)})
    rows.sort(key=lambda r:r["MA_NV"])
    total={k:round(sum(r[k] for r in rows),2) for k in ["DU_DAU","THU_CK_KM_XA","THU_THUONG_CAP_DON","TONG_THU","DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","DU_CUOI"]}
    return {"rows":rows,"total":total,"count":len(rows)}

def ckkm_report_xlsx(date_from=None,date_to=None,ma_nv=None):
    rep=ckkm_report(date_from,date_to,ma_nv)
    wb=Workbook(); ws=wb.active; ws.title="CK_KM"
    headers=["MA_NV","DU_DAU","THU_CK_KM_XA","THU_THUONG_CAP_DON","TONG_THU","DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","DU_CUOI"]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1769E0")
    for r in rep["rows"]: ws.append([r[h] for h in headers])
    bio=BytesIO(); wb.save(bio); return "XB_UI02_CK_KM.xlsx",bio.getvalue()

def exception_report():
    from .xb_mapping_audit import run_mapping_audit
    from .xb_replace_batch import replacement_history
    rows=[]
    for r in run_mapping_audit().get("rows",[]):
        if r.get("STATUS")!="MAP_OK":
            rows.append({"NHOM":"MAPPING","NGUON":r.get("SOURCE",""),"BATCH":r.get("BATCH",""),"DONG":r.get("ROW_NO",""),"MA_NV":r.get("MA_NV_NGUON",""),"MA_SKU":r.get("MA_SKU_NGUON",""),"MA_KH":r.get("MA_KH_NGUON",""),"LOI":r.get("STATUS","")})
    for b in replacement_history():
        if _upper(b.get("status"))=="PREVIEW" and int(b.get("error_rows") or 0)>0:
            rows.append({"NHOM":"BATCH","NGUON":b.get("source_label",""),"BATCH":b.get("batch_id",""),"DONG":"","MA_NV":"","MA_SKU":"","MA_KH":"","LOI":f"PREVIEW còn {b.get('error_rows',0)} lỗi / chưa chốt"})
    return {"rows":rows[:1000],"summary":{"TOTAL":len(rows),"MAPPING":sum(1 for r in rows if r["NHOM"]=="MAPPING"),"BATCH":sum(1 for r in rows if r["NHOM"]=="BATCH")}}

def history_report():
    from .xb_replace_batch import replacement_history
    return {"rows":replacement_history()}

def _in_period(value, df=None, dt=None):
    d=_date(value)
    if df is None and dt is None:
        return True
    if d is None:
        return False
    if df is not None and d < df:
        return False
    if dt is not None and d > dt:
        return False
    return True

def _sku_meta_indexes():
    from .xb_mapping_audit import _load_master_dataset
    c=_conn()
    rows=_load_master_dataset(c,["dm01-product","product"])
    c.close()
    out={}
    for d in rows:
        sku=_upper(_v(d,"MA_SKU","SKU","ma_sku"))
        if not sku:
            continue
        out[sku]={
            "BU": _upper(_v(d,"MA_BU","BU","BU_CODE","MA_MANG","MANG_KINH_DOANH")) or "",
            "NHOM_NHAN_HIEU": _upper(_v(d,"NHOM_NHAN_HIEU","NHOM_NHAN","NHAN_HIEU","BRAND_GROUP","NHOM_BRAND")) or "",
        }
    return out

def quantity_group_report(date_from=None,date_to=None,mode="detail",ma_nv=None,ma_sku=None,bu=None,nhom_nhan_hieu=None):
    idx=_load_mapping()
    meta=_sku_meta_indexes()
    c=_conn()
    df=_date(date_from) if date_from else None
    dt=_date(date_to) if date_to else None
    f_nv=_upper(ma_nv); f_sku=_upper(ma_sku); f_bu=_upper(bu); f_brand=_upper(nhom_nhan_hieu)
    agg=defaultdict(lambda:{"DU_DAU":0.0,"XA_TRONG_KY":0.0,"BU_MBW":0.0,"BU_MISA":0.0})

    def allow(sku):
        m=meta.get(_upper(sku),{})
        if f_bu and _upper(m.get("BU"))!=f_bu: return False
        if f_brand and _upper(m.get("NHOM_NHAN_HIEU"))!=f_brand: return False
        return True

    def key_for(day,nv,sku):
        if mode=="month": return day.strftime("%Y-%m")
        if mode=="bu": return meta.get(sku,{}).get("BU") or "(CHƯA KHAI BÁO)"
        if mode=="brand_group": return meta.get(sku,{}).get("NHOM_NHAN_HIEU") or "(CHƯA KHAI BÁO)"
        return (nv,sku)

    for bid,row_no,d in _iter_dms_xa(c) or []:
        raw=_v(d,"NGAY_DON","NGAY_CHUNG_TU")
        if not _in_period(raw,df,dt): continue
        day=_date(raw); nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        if not allow(sku): continue
        agg[key_for(day,nv,sku)]["XA_TRONG_KY"] += _qty_su(d,"XA")

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B": continue
        raw=_v(d,"NGAY_DON","NGAY_GIAO")
        if not _in_period(raw,df,dt): continue
        day=_date(raw); nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        if not allow(sku): continue
        agg[key_for(day,nv,sku)]["BU_MBW"] += _qty_su(d,"BU")

    for bid,row_no,d in _iter_misa(c) or []:
        raw=_v(d,"NGAY_CHUNG_TU","NGAY_DON")
        if not _in_period(raw,df,dt): continue
        day=_date(raw); nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        if not allow(sku): continue
        agg[key_for(day,nv,sku)]["BU_MISA"] += _num(_v(d,"SL_BU_THUC","TONG_SL_BAN"))
    c.close()

    rows=[]
    for key,x in agg.items():
        bu_total=x["BU_MBW"]+x["BU_MISA"]
        row={
            "DU_DAU":0.0,
            "XA_TRONG_KY":round(x["XA_TRONG_KY"],4),
            "BU_MBW":round(x["BU_MBW"],4),
            "BU_MISA":round(x["BU_MISA"],4),
            "BAN_BU_TRONG_KY":round(bu_total,4),
            "DU_CUOI":round(-x["XA_TRONG_KY"]+bu_total,4),
        }
        if mode=="detail":
            row["MA_NV"],row["MA_SKU"]=key
        else:
            row["NHOM"]=key
        rows.append(row)
    rows.sort(key=lambda r:str(r.get("NHOM",r.get("MA_NV",""))))
    total={k:round(sum(r[k] for r in rows),4) for k in ["DU_DAU","XA_TRONG_KY","BU_MBW","BU_MISA","BAN_BU_TRONG_KY","DU_CUOI"]}
    return {"mode":mode,"rows":rows,"total":total,"count":len(rows)}

def _period_bounds_v2(date_from=None,date_to=None,month=None):
    if month:
        y,m=[int(x) for x in str(month).split("-")[:2]]
        start=datetime.date(y,m,1)
        end=(datetime.date(y+1,1,1) if m==12 else datetime.date(y,m+1,1))-datetime.timedelta(days=1)
        return start,end
    return (_date(date_from) if date_from else None,_date(date_to) if date_to else None)

def _product_meta_v2():
    from .xb_mapping_audit import _load_master_dataset
    c=_conn(); rows=_load_master_dataset(c,["dm01-product","product"]); c.close()
    out={}
    for d in rows:
        sku=_upper(_v(d,"MA_SKU","SKU","ma_sku"))
        if not sku: continue
        out[sku]={
            "BU":_upper(_v(d,"MA_BU","BU","BU_CODE","MA_MANG","MANG_KINH_DOANH")) or "",
            "NHOM_SAN_PHAM":_upper(_v(d,"MA_NHOM_SAN_PHAM","NHOM_SAN_PHAM","TEN_NHOM_SAN_PHAM","NHOM_NHAN_HIEU","NHOM_NHAN","NHAN_HIEU","BRAND_GROUP","NHOM_BRAND")) or "",
            "TEN_SKU":str(_v(d,"TEN_SKU","TEN_SAN_PHAM","SKU_NAME") or "").strip()
        }
    return out

def filter_options_v2():
    meta=_product_meta_v2()
    return {
        "BU":sorted({v["BU"] for v in meta.values() if v["BU"]}),
        "NHOM_SAN_PHAM":sorted({v["NHOM_SAN_PHAM"] for v in meta.values() if v["NHOM_SAN_PHAM"]})
    }

def _qty_events_v2():
    idx=_load_mapping(); c=_conn(); out=[]
    for bid,row_no,d in _iter_dms_xa(c) or []:
        day=_v5_event_date(d,"DMS-XA")
        if day is None: continue
        nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if nv and sku: out.append((day,nv,sku,_qty_su(d,"XA"),0.0,0.0))
    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B": continue
        day=_v5_event_date(d,"DMS-MBW")
        if day is None: continue
        nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if nv and sku: out.append((day,nv,sku,0.0,_qty_su(d,"BU"),0.0))
    for bid,row_no,d in _iter_misa(c) or []:
        day=_v5_event_date(d,"MISA-BU")
        if day is None: continue
        nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if nv and sku: out.append((day,nv,sku,0.0,0.0,_num(_v(d,"SL_BU_THUC","TONG_SL_BAN"))))
    c.close(); return out

def quantity_v2(date_from=None,date_to=None,month=None,ma_nv=None,ma_sku=None,bu=None,nhom_san_pham=None,view="detail"):
    start,end=_period_bounds_v2(date_from,date_to,month)
    f_nv=_upper(ma_nv); f_sku=_upper(ma_sku); f_bu=_upper(bu); f_grp=_upper(nhom_san_pham)
    meta=_product_meta_v2()
    opening=defaultdict(float)
    period=defaultdict(lambda:{"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})

    for day,nv,sku,xa,bmbw,bmisa in _qty_events_v2():
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        m=meta.get(sku,{})
        if f_bu and _upper(m.get("BU"))!=f_bu: continue
        if f_grp and _upper(m.get("NHOM_SAN_PHAM"))!=f_grp: continue
        key=(nv,sku)
        if start is not None and day<start:
            opening[key]+=(-xa+bmbw+bmisa)
            continue
        if start is not None and day<start: continue
        if end is not None and day>end: continue
        period[key]["XA"]+=xa; period[key]["BU_MBW"]+=bmbw; period[key]["BU_MISA"]+=bmisa

    detail=[]
    for key in set(opening)|set(period):
        nv,sku=key; p=period[key]; du=opening[key]; total_bu=p["BU_MBW"]+p["BU_MISA"]; m=meta.get(sku,{})
        detail.append({
            "MA_NV":nv,"MA_SKU":sku,"TEN_SKU":m.get("TEN_SKU",""),"BU":m.get("BU",""),"NHOM_SAN_PHAM":m.get("NHOM_SAN_PHAM",""),
            "DU_DAU":round(du,4),"XA_TRONG_KY":round(p["XA"],4),"BU_MBW":round(p["BU_MBW"],4),"BU_MISA":round(p["BU_MISA"],4),
            "BAN_BU_TRONG_KY":round(total_bu,4),"DU_CUOI":round(du-p["XA"]+total_bu,4)
        })

    if view in ("bu","group"):
        agg=defaultdict(lambda:{"DU_DAU":0.0,"XA_TRONG_KY":0.0,"BU_MBW":0.0,"BU_MISA":0.0})
        for r in detail:
            a=agg[r["MA_SKU"]]
            for k in a: a[k]+=r[k]
        rows=[]
        for sku,a in agg.items():
            m=meta.get(sku,{}); total_bu=a["BU_MBW"]+a["BU_MISA"]
            rows.append({
                "MA_SKU":sku,"TEN_SKU":m.get("TEN_SKU",""),"BU":m.get("BU",""),"NHOM_SAN_PHAM":m.get("NHOM_SAN_PHAM",""),
                "DU_DAU":round(a["DU_DAU"],4),"XA_TRONG_KY":round(a["XA_TRONG_KY"],4),"BU_MBW":round(a["BU_MBW"],4),"BU_MISA":round(a["BU_MISA"],4),
                "BAN_BU_TRONG_KY":round(total_bu,4),"DU_CUOI":round(a["DU_DAU"]-a["XA_TRONG_KY"]+total_bu,4)
            })
        rows.sort(key=lambda r:(r["BU"],r["NHOM_SAN_PHAM"],r["MA_SKU"]))
    else:
        rows=sorted(detail,key=lambda r:(r["MA_NV"],r["MA_SKU"]))

    total={k:round(sum(r[k] for r in rows),4) for k in ["DU_DAU","XA_TRONG_KY","BU_MBW","BU_MISA","BAN_BU_TRONG_KY","DU_CUOI"]}
    return {"period":{"date_from":str(start) if start else None,"date_to":str(end) if end else None,"month":month},"view":view,"rows":rows,"total":total,"count":len(rows)}

def ckkm_v2(date_from=None,date_to=None,month=None,ma_nv=None):
    start,end=_period_bounds_v2(date_from,date_to,month)
    idx=_load_mapping(); c=_conn(); f_nv=_upper(ma_nv)
    opening=defaultdict(float)
    period=defaultdict(lambda:{"THU_CK_KM_XA":0.0,"THU_THUONG_CAP_DON":0.0,"DA_CHI_MBW":0.0,"DA_CHI_MISA":0.0})

    def add(nv,day,thu_line=0.0,thu_bonus=0.0,chi_mbw=0.0,chi_misa=0.0):
        if not nv or day is None: return
        if f_nv and nv!=f_nv: return
        delta=thu_line+thu_bonus-chi_mbw-chi_misa
        if start is not None and day<start:
            opening[nv]+=delta; return
        if end is not None and day>end: return
        period[nv]["THU_CK_KM_XA"]+=thu_line
        period[nv]["THU_THUONG_CAP_DON"]+=thu_bonus
        period[nv]["DA_CHI_MBW"]+=chi_mbw
        period[nv]["DA_CHI_MISA"]+=chi_misa

    if _table_exists(c,"xb_src_batch_v2") and _table_exists(c,"xb_src_row_v2"):
        for bid in _active_batches(c,"xb_src_batch_v2"):
            for r in c.execute("SELECT sheet_name,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='VALID'",(bid,)).fetchall():
                try:d=json.loads(r["payload_json"])
                except:continue
                day=_v5_event_date(d,"DMS-XA"); nv=_upper(_map_nv(idx,d)); sheet=_upper(r["sheet_name"])
                if sheet=="CHI_TIET_XA":
                    val=sum(_num(_v(d,k)) for k in ("CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI","KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC"))
                    add(nv,day,thu_line=val)
                elif sheet=="THUONG_CAP_DON":
                    val=sum(_num(_v(d,k)) for k in ("THUONG_TRUNG_BAY","THUONG_TICH_LUY","GIAM_TRU_CAP_DON_KHAC"))
                    add(nv,day,thu_bonus=val)

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B": continue
        day=_v5_event_date(d,"DMS-MBW"); nv=_upper(_map_nv(idx,d))
        val=sum(_num(_v(d,k)) for k in ("CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI","KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC"))
        add(nv,day,chi_mbw=val)

    for bid,row_no,d in _iter_misa(c) or []:
        day=_v5_event_date(d,"MISA-BU"); nv=_upper(_map_nv(idx,d))
        val=sum(_num(_v(d,k)) for k in ("CK_KM_TIEN_CHUAN","CK_KM_TIEN","CHIET_KHAU","GIAM_TRU_DONG_KHAC_CHUAN","GIAM_TRU_DONG_KHAC","GIA_TRI_GIAM_GIA"))
        add(nv,day,chi_misa=val)
    c.close()

    rows=[]
    for nv in set(opening)|set(period):
        p=period[nv]; du=opening[nv]; thu=p["THU_CK_KM_XA"]+p["THU_THUONG_CAP_DON"]; chi=p["DA_CHI_MBW"]+p["DA_CHI_MISA"]
        rows.append({"MA_NV":nv,"DU_DAU":round(du,2),"THU_CK_KM_XA":round(p["THU_CK_KM_XA"],2),"THU_THUONG_CAP_DON":round(p["THU_THUONG_CAP_DON"],2),
                     "TONG_THU":round(thu,2),"DA_CHI_MBW":round(p["DA_CHI_MBW"],2),"DA_CHI_MISA":round(p["DA_CHI_MISA"],2),"TONG_DA_CHI":round(chi,2),
                     "DU_CUOI":round(du+thu-chi,2)})
    rows.sort(key=lambda r:r["MA_NV"])
    total={k:round(sum(r[k] for r in rows),2) for k in ["DU_DAU","THU_CK_KM_XA","THU_THUONG_CAP_DON","TONG_THU","DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","DU_CUOI"]}
    return {"period":{"date_from":str(start) if start else None,"date_to":str(end) if end else None,"month":month},"rows":rows,"total":total,"count":len(rows)}

def _v3_period(date_from=None,date_to=None,month=None):
    if month:
        y,m=[int(x) for x in str(month).split("-")[:2]]
        start=datetime.date(y,m,1)
        end=(datetime.date(y+1,1,1) if m==12 else datetime.date(y,m+1,1))-datetime.timedelta(days=1)
        return start,end
    return (_date(date_from) if date_from else None,_date(date_to) if date_to else None)

def _v3_meta():
    from .xb_mapping_audit import _load_master_dataset
    c=_conn()
    rows=_load_master_dataset(c,["dm01-product","product"])
    c.close()
    out={}
    for d in rows:
        sku=_upper(_v(d,"MA_SKU","SKU","ma_sku"))
        if not sku:
            continue
        out[sku]={
            "TEN_SKU":str(_v(d,"TEN_SKU","TEN_SAN_PHAM","SKU_NAME") or "").strip(),
            "BU":_upper(_v(d,"MA_BU","BU","BU_CODE","MA_MANG","MANG_KINH_DOANH")) or "",
            "NHAN_HIEU":_upper(_v(d,"NHAN_HIEU","BRAND","TEN_NHAN_HIEU","NHOM_NHAN_HIEU","NHOM_NHAN","NHOM_BRAND")) or "",
        }
    return out

def _v3_options():
    meta=_v3_meta()
    return {
        "BU":sorted({x["BU"] for x in meta.values() if x["BU"]}),
        "NHAN_HIEU":sorted({x["NHAN_HIEU"] for x in meta.values() if x["NHAN_HIEU"]})
    }

def _v3_qty_events():
    idx=_load_mapping()
    c=_conn()
    out=[]

    for bid,row_no,d in _iter_dms_xa(c) or []:
        day=_v5_event_date(d,"DMS-XA")
        if day is None:
            continue
        nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if nv and sku:
            out.append((day,nv,sku,_qty_su(d,"XA"),0.0,0.0))

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B":
            continue
        day=_v5_event_date(d,"DMS-MBW")
        if day is None:
            continue
        nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if nv and sku:
            out.append((day,nv,sku,0.0,_qty_su(d,"BU"),0.0))

    for bid,row_no,d in _iter_misa(c) or []:
        day=_v5_event_date(d,"MISA-BU")
        if day is None:
            continue
        nv=_upper(_map_nv(idx,d)); sku=_upper(_v(d,"MA_SKU"))
        if nv and sku:
            out.append((day,nv,sku,0.0,0.0,_num(_v(d,"SL_BU_THUC","TONG_SL_BAN"))))

    c.close()
    return out

def quantity_v3(date_from=None,date_to=None,month=None,report_type="detail",ma_nv=None,ma_sku=None,bu=None,nhan_hieu=None):
    start,end=_v3_period(date_from,date_to,month)
    meta=_v3_meta()
    f_nv=_upper(ma_nv); f_sku=_upper(ma_sku); f_bu=_upper(bu); f_brand=_upper(nhan_hieu)

    opening=defaultdict(float)
    period=defaultdict(lambda:{"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})

    for day,nv,sku,xa,bmbw,bmisa in _v3_qty_events():
        m=meta.get(sku,{})
        if f_nv and nv!=f_nv: continue
        if f_sku and sku!=f_sku: continue
        if f_bu and _upper(m.get("BU"))!=f_bu: continue
        if f_brand and _upper(m.get("NHAN_HIEU"))!=f_brand: continue

        key=(nv,sku)
        delta=-xa+bmbw+bmisa

        if start is not None and day<start:
            opening[key]+=delta
            continue
        if start is not None and day<start:
            continue
        if end is not None and day>end:
            continue

        period[key]["XA"]+=xa
        period[key]["BU_MBW"]+=bmbw
        period[key]["BU_MISA"]+=bmisa

    detail=[]
    for nv,sku in set(opening)|set(period):
        p=period[(nv,sku)]
        du_dau=opening[(nv,sku)]
        total_bu=p["BU_MBW"]+p["BU_MISA"]
        du_cuoi=du_dau-p["XA"]+total_bu

        # Do not push fully-zero lines.
        if abs(du_dau)<1e-9 and abs(p["XA"])<1e-9 and abs(total_bu)<1e-9 and abs(du_cuoi)<1e-9:
            continue

        m=meta.get(sku,{})
        detail.append({
            "MA_NV":nv,"MA_SKU":sku,"TEN_SKU":m.get("TEN_SKU",""),
            "BU":m.get("BU",""),"NHAN_HIEU":m.get("NHAN_HIEU",""),
            "TON_DAU":round(du_dau,4),
            "XA":round(p["XA"],4),
            "BU_MBW":round(p["BU_MBW"],4),
            "BU_MISA":round(p["BU_MISA"],4),
            "TONG_BU":round(total_bu,4),
            "TON_CUOI":round(du_cuoi,4),
        })

    if report_type=="detail":
        rows=sorted(detail,key=lambda r:(r["MA_NV"],r["MA_SKU"]))
    elif report_type=="employee":
        agg=defaultdict(lambda:{"TON_DAU":0.0,"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})
        for r in detail:
            a=agg[r["MA_NV"]]
            a["TON_DAU"]+=r["TON_DAU"]; a["XA"]+=r["XA"]; a["BU_MBW"]+=r["BU_MBW"]; a["BU_MISA"]+=r["BU_MISA"]
        rows=[]
        for nv,a in sorted(agg.items()):
            tb=a["BU_MBW"]+a["BU_MISA"]; tc=a["TON_DAU"]-a["XA"]+tb
            if abs(a["TON_DAU"])<1e-9 and abs(a["XA"])<1e-9 and abs(tb)<1e-9 and abs(tc)<1e-9:
                continue
            rows.append({"MA_NV":nv,"TON_DAU":round(a["TON_DAU"],4),"XA":round(a["XA"],4),"BU_MBW":round(a["BU_MBW"],4),"BU_MISA":round(a["BU_MISA"],4),"TONG_BU":round(tb,4),"TON_CUOI":round(tc,4)})
    elif report_type=="bu":
        agg=defaultdict(lambda:{"TON_DAU":0.0,"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})
        for r in detail:
            key=(r["BU"] or "(CHƯA KHAI BÁO)",r["MA_SKU"],r["TEN_SKU"])
            a=agg[key]
            a["TON_DAU"]+=r["TON_DAU"]; a["XA"]+=r["XA"]; a["BU_MBW"]+=r["BU_MBW"]; a["BU_MISA"]+=r["BU_MISA"]
        rows=[]
        for (bu0,sku,name),a in sorted(agg.items()):
            tb=a["BU_MBW"]+a["BU_MISA"]; tc=a["TON_DAU"]-a["XA"]+tb
            if abs(a["TON_DAU"])<1e-9 and abs(a["XA"])<1e-9 and abs(tb)<1e-9 and abs(tc)<1e-9:
                continue
            rows.append({"BU":bu0,"MA_SKU":sku,"TEN_SKU":name,"TON_DAU":round(a["TON_DAU"],4),"XA":round(a["XA"],4),"BU_MBW":round(a["BU_MBW"],4),"BU_MISA":round(a["BU_MISA"],4),"TONG_BU":round(tb,4),"TON_CUOI":round(tc,4)})
    else:
        agg=defaultdict(lambda:{"TON_DAU":0.0,"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})
        for r in detail:
            key=(r["NHAN_HIEU"] or "(CHƯA KHAI BÁO)",r["MA_SKU"],r["TEN_SKU"])
            a=agg[key]
            a["TON_DAU"]+=r["TON_DAU"]; a["XA"]+=r["XA"]; a["BU_MBW"]+=r["BU_MBW"]; a["BU_MISA"]+=r["BU_MISA"]
        rows=[]
        for (brand,sku,name),a in sorted(agg.items()):
            tb=a["BU_MBW"]+a["BU_MISA"]; tc=a["TON_DAU"]-a["XA"]+tb
            if abs(a["TON_DAU"])<1e-9 and abs(a["XA"])<1e-9 and abs(tb)<1e-9 and abs(tc)<1e-9:
                continue
            rows.append({"NHAN_HIEU":brand,"MA_SKU":sku,"TEN_SKU":name,"TON_DAU":round(a["TON_DAU"],4),"XA":round(a["XA"],4),"BU_MBW":round(a["BU_MBW"],4),"BU_MISA":round(a["BU_MISA"],4),"TONG_BU":round(tb,4),"TON_CUOI":round(tc,4)})

    total={k:round(sum(r.get(k,0) for r in rows),4) for k in ["TON_DAU","XA","BU_MBW","BU_MISA","TONG_BU","TON_CUOI"]}
    return {"rows":rows,"total":total,"count":len(rows),"report_type":report_type,"period":{"from":str(start) if start else None,"to":str(end) if end else None}}

def ckkm_v3(date_from=None,date_to=None,month=None,ma_nv=None):
    start,end=_v3_period(date_from,date_to,month)
    idx=_load_mapping(); c=_conn(); f_nv=_upper(ma_nv)
    opening=defaultdict(float)
    period=defaultdict(lambda:{"THU_XA":0.0,"THU_THUONG":0.0,"CHI_MBW":0.0,"CHI_MISA":0.0})

    def add(nv,day,thu_xa=0,thu_thuong=0,chi_mbw=0,chi_misa=0):
        if not nv or day is None: return
        if f_nv and nv!=f_nv: return
        delta=thu_xa+thu_thuong-chi_mbw-chi_misa
        if start is not None and day<start:
            opening[nv]+=delta; return
        if end is not None and day>end: return
        p=period[nv]
        p["THU_XA"]+=thu_xa; p["THU_THUONG"]+=thu_thuong; p["CHI_MBW"]+=chi_mbw; p["CHI_MISA"]+=chi_misa

    if _table_exists(c,"xb_src_batch_v2") and _table_exists(c,"xb_src_row_v2"):
        for bid in _active_batches(c,"xb_src_batch_v2"):
            for r in c.execute("SELECT sheet_name,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='VALID'",(bid,)).fetchall():
                try:d=json.loads(r["payload_json"])
                except:continue
                day=_v5_event_date(d,"DMS-XA"); nv=_upper(_map_nv(idx,d)); sh=_upper(r["sheet_name"])
                if sh=="CHI_TIET_XA":
                    val=sum(_num(_v(d,k)) for k in ("CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI","KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC"))
                    add(nv,day,thu_xa=val)
                elif sh=="THUONG_CAP_DON":
                    val=sum(_num(_v(d,k)) for k in ("THUONG_TRUNG_BAY","THUONG_TICH_LUY","GIAM_TRU_CAP_DON_KHAC"))
                    add(nv,day,thu_thuong=val)

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B": continue
        day=_v5_event_date(d,"DMS-MBW"); nv=_upper(_map_nv(idx,d))
        val=sum(_num(_v(d,k)) for k in ("CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI","KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC"))
        add(nv,day,chi_mbw=val)

    for bid,row_no,d in _iter_misa(c) or []:
        day=_v5_event_date(d,"MISA-BU"); nv=_upper(_map_nv(idx,d))
        val=sum(_num(_v(d,k)) for k in ("CK_KM_TIEN_CHUAN","CK_KM_TIEN","CHIET_KHAU","GIAM_TRU_DONG_KHAC_CHUAN","GIAM_TRU_DONG_KHAC","GIA_TRI_GIAM_GIA"))
        add(nv,day,chi_misa=val)
    c.close()

    rows=[]
    for nv in sorted(set(opening)|set(period)):
        p=period[nv]; du=opening[nv]; thu=p["THU_XA"]+p["THU_THUONG"]; chi=p["CHI_MBW"]+p["CHI_MISA"]; dc=du+thu-chi
        if abs(du)<0.005 and abs(thu)<0.005 and abs(chi)<0.005 and abs(dc)<0.005:
            continue
        rows.append({"MA_NV":nv,"TON_DAU":round(du,2),"CK_KM_XA":round(p["THU_XA"],2),"THUONG_CAP_DON":round(p["THU_THUONG"],2),"TONG_QUYEN_LOI":round(thu,2),"DA_CHI_MBW":round(p["CHI_MBW"],2),"DA_CHI_MISA":round(p["CHI_MISA"],2),"TONG_DA_CHI":round(chi,2),"TON_CUOI":round(dc,2)})
    total={k:round(sum(r[k] for r in rows),2) for k in ["TON_DAU","CK_KM_XA","THUONG_CAP_DON","TONG_QUYEN_LOI","DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","TON_CUOI"]}
    return {"rows":rows,"total":total,"count":len(rows),"period":{"from":str(start) if start else None,"to":str(end) if end else None}}

def quantity_v3_xlsx(**kwargs):
    rep=quantity_v3(**kwargs)
    wb=Workbook(); ws=wb.active; ws.title="SO_LUONG"
    rt=kwargs.get("report_type","detail")
    if rt=="detail":
        headers=["MA_NV","MA_SKU","TEN_SKU","BU","NHAN_HIEU","TON_DAU","XA","BU_MBW","BU_MISA","TONG_BU","TON_CUOI"]
    elif rt=="employee":
        headers=["MA_NV","TON_DAU","XA","BU_MBW","BU_MISA","TONG_BU","TON_CUOI"]
    elif rt=="bu":
        headers=["BU","MA_SKU","TEN_SKU","TON_DAU","XA","BU_MBW","BU_MISA","TONG_BU","TON_CUOI"]
    else:
        headers=["NHAN_HIEU","MA_SKU","TEN_SKU","TON_DAU","XA","BU_MBW","BU_MISA","TONG_BU","TON_CUOI"]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1769E0")
    for r in rep["rows"]:
        ws.append([r.get(h,"") for h in headers])
    bio=BytesIO(); wb.save(bio); return "XB_UI02_So_luong.xlsx",bio.getvalue()

def ckkm_v3_xlsx(**kwargs):
    rep=ckkm_v3(**kwargs)
    wb=Workbook(); ws=wb.active; ws.title="CK_KM"
    headers=["MA_NV","TON_DAU","CK_KM_XA","THUONG_CAP_DON","TONG_QUYEN_LOI","DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","TON_CUOI"]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1769E0")
    for r in rep["rows"]:
        ws.append([r.get(h,"") for h in headers])
    bio=BytesIO(); wb.save(bio); return "XB_UI02_CK_KM.xlsx",bio.getvalue()

def _v5_period(date_from=None,date_to=None,month=None):
    if month:
        y,m=[int(x) for x in str(month).split("-")[:2]]
        start=datetime.date(y,m,1)
        end=(datetime.date(y+1,1,1) if m==12 else datetime.date(y,m+1,1))-datetime.timedelta(days=1)
        return start,end
    return (_date(date_from) if date_from else None,
            _date(date_to) if date_to else None)

def _v5_meta():
    from .xb_mapping_audit import _load_master_dataset
    c=_conn()
    rows=_load_master_dataset(c,["dm01-product","product"])
    c.close()

    out={}
    for d in rows:
        sku=_upper(_v(d,"MA_SKU"))
        if not sku:
            continue

        ten=str(_v(d,"TEN_SKU") or "").strip()
        bu=_upper(_v(d,"BU"))
        nhan=_upper(_v(d,"MA_NHAN"))

        quy_cach=_num(_v(d,"SO_SU_THUNG"))
        if quy_cach<=0:
            quy_cach=1.0

        out[sku]={
            "TEN_SKU":ten,
            "BU":bu or "",
            "NHAN_HIEU":nhan or "",
            "SU_THUNG":quy_cach,
        }

    return out

def _v5_event_date(d, source=""):
    # Preferred canonical date fields
    preferred=(
        "NGAY_DON","NGAY_XA","NGAY_BAN","NGAY_GIAO",
        "NGAY_CHUNG_TU","NGAY_CT","NGAY_PHAT_SINH",
        "NGAY_HOA_DON","NGAY"
    )
    for k in preferred:
        if k in d and d.get(k) not in (None,""):
            x=_date(d.get(k))
            if x is not None:
                return x

    # Fuzzy fallback for source files whose date header was normalized differently.
    # Ignore upload/meta timestamps; accept business fields containing NGAY.
    for k,v in d.items():
        ku=str(k).strip().upper()
        if "NGAY" not in ku:
            continue
        if any(x in ku for x in ("UPLOAD","TAO_BATCH","CAP_NHAT","SYSTEM","META")):
            continue
        if v in (None,""):
            continue
        x=_date(v)
        if x is not None:
            return x

    return None

def _v5_qty_events():
    idx=_load_mapping()
    meta=_v5_meta()
    c=_conn()
    out=[]

    def qty_xa(d,sku):
        v=_v(d,"SL_BAN_QUY_SU","SL_XA_QUY_SU")
        if v not in (None,""):
            return _num(v)
        sl_su=_num(_v(d,"SL_SU","SO_LUONG_SU"))
        sl_thung=_num(_v(d,"SL_THUNG","SO_LUONG_THUNG"))
        ratio=meta.get(sku,{}).get("SU_THUNG",1.0)
        return sl_su + sl_thung*ratio

    def qty_bu(d,sku):
        v=_v(d,"SL_BU_THUC","SL_BU_QUY_SU")
        if v not in (None,""):
            return _num(v)
        sl_su=_num(_v(d,"SL_SU","SO_LUONG_SU"))
        sl_thung=_num(_v(d,"SL_THUNG","SO_LUONG_THUNG"))
        ratio=meta.get(sku,{}).get("SU_THUNG",1.0)
        if sl_su!=0 or sl_thung!=0:
            return sl_su + sl_thung*ratio
        ban=_num(_v(d,"TONG_SL_BAN"))
        tra_raw=_v(d,"TONG_SL_TRA_LAI")
        tra=0.0 if isinstance(tra_raw,str) and tra_raw.strip() in ("-","–","—") else _num(tra_raw)
        return ban-tra

    for bid,row_no,d in _iter_dms_xa(c) or []:
        day=_v5_event_date(d,"DMS-XA")
        if day is None:
            continue
        nv=_upper(_map_nv(idx,d))
        sku=_upper(_v(d,"MA_SKU"))
        if nv and sku:
            out.append({
                "DATE":day,"MA_NV":nv,"MA_SKU":sku,
                "XA":qty_xa(d,sku),"BU_MBW":0.0,"BU_MISA":0.0
            })

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B":
            continue
        day=_v5_event_date(d,"DMS-MBW")
        if day is None:
            continue
        nv=_upper(_map_nv(idx,d))
        sku=_upper(_v(d,"MA_SKU"))
        if nv and sku:
            out.append({
                "DATE":day,"MA_NV":nv,"MA_SKU":sku,
                "XA":0.0,"BU_MBW":qty_bu(d,sku),"BU_MISA":0.0
            })

    for bid,row_no,d in _iter_misa(c) or []:
        day=_v5_event_date(d,"MISA-BU")
        if day is None:
            continue
        nv=_upper(_map_nv(idx,d))
        sku=_upper(_v(d,"MA_SKU"))
        if nv and sku:
            out.append({
                "DATE":day,"MA_NV":nv,"MA_SKU":sku,
                "XA":0.0,"BU_MBW":0.0,"BU_MISA":qty_bu(d,sku)
            })

    c.close()
    return out

def report_options_v5():
    # Build employee list directly from active committed source events.
    employees=sorted({
        str(e.get("MA_NV") or "").strip().upper()
        for e in _v5_qty_events()
        if str(e.get("MA_NV") or "").strip()
    })

    meta=_v5_meta()
    bus=sorted({str(x.get("BU") or "").strip().upper() for x in meta.values() if str(x.get("BU") or "").strip()})
    brands=sorted({str(x.get("NHAN_HIEU") or "").strip().upper() for x in meta.values() if str(x.get("NHAN_HIEU") or "").strip()})

    return {
        "NHAN_VIEN":employees,
        "SO_NHAN_VIEN":len(employees),
        "BU":bus,
        "NHAN_HIEU":brands,
    }

def quantity_report_v5(date_from=None,date_to=None,month=None,report_type="employee",ma_nv=None,bu=None,nhan_hieu=None):
    start,end=_v5_period(date_from,date_to,month)
    meta=_v5_meta()
    f_nv=_upper(ma_nv)
    f_bu=_upper(bu)
    f_brand=_upper(nhan_hieu)

    opening=defaultdict(float)
    period=defaultdict(lambda:{"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})

    for e in _v5_qty_events():
        nv=e["MA_NV"]; sku=e["MA_SKU"]; m=meta.get(sku,{})
        if report_type=="employee" and f_nv and nv!=f_nv:
            continue
        if report_type=="bu" and f_bu and _upper(m.get("BU"))!=f_bu:
            continue
        if report_type=="brand" and f_brand and _upper(m.get("NHAN_HIEU"))!=f_brand:
            continue

        key=(nv,sku)
        delta=-e["XA"]+e["BU_MBW"]+e["BU_MISA"]

        if start is not None and e["DATE"]<start:
            opening[key]+=delta
            continue
        if end is not None and e["DATE"]>end:
            continue

        period[key]["XA"]+=e["XA"]
        period[key]["BU_MBW"]+=e["BU_MBW"]
        period[key]["BU_MISA"]+=e["BU_MISA"]

    detail=[]
    for key in set(opening)|set(period):
        nv,sku=key
        p=period[key]
        du_dau=opening[key]
        tong_bu=p["BU_MBW"]+p["BU_MISA"]
        du_cuoi=du_dau-p["XA"]+tong_bu

        # Suppress completely zero lines
        if abs(du_dau)<1e-9 and abs(p["XA"])<1e-9 and abs(tong_bu)<1e-9 and abs(du_cuoi)<1e-9:
            continue

        m=meta.get(sku,{})
        detail.append({
            "MA_NV":nv,
            "MA_SKU":sku,
            "TEN_SKU":m.get("TEN_SKU",""),
            "BU":m.get("BU",""),
            "NHAN_HIEU":m.get("NHAN_HIEU",""),
            "DU_DAU":round(du_dau,4),
            "XA":round(p["XA"],4),
            "BU_MBW":round(p["BU_MBW"],4),
            "BU_MISA":round(p["BU_MISA"],4),
            "TONG_BU":round(tong_bu,4),
            "DU_CUOI":round(du_cuoi,4),
        })

    if report_type=="employee":
        rows=sorted(detail,key=lambda r:(r["MA_NV"],r["MA_SKU"]))
    else:
        # BU / Brand report: aggregate across employees, show SKU lines only
        agg=defaultdict(lambda:{"DU_DAU":0.0,"XA":0.0,"BU_MBW":0.0,"BU_MISA":0.0})
        for r in detail:
            a=agg[r["MA_SKU"]]
            a["DU_DAU"]+=r["DU_DAU"]
            a["XA"]+=r["XA"]
            a["BU_MBW"]+=r["BU_MBW"]
            a["BU_MISA"]+=r["BU_MISA"]

        rows=[]
        for sku,a in agg.items():
            tong_bu=a["BU_MBW"]+a["BU_MISA"]
            du_cuoi=a["DU_DAU"]-a["XA"]+tong_bu
            if abs(a["DU_DAU"])<1e-9 and abs(a["XA"])<1e-9 and abs(tong_bu)<1e-9 and abs(du_cuoi)<1e-9:
                continue
            m=meta.get(sku,{})
            rows.append({
                "MA_SKU":sku,
                "TEN_SKU":m.get("TEN_SKU",""),
                "BU":m.get("BU",""),
                "NHAN_HIEU":m.get("NHAN_HIEU",""),
                "DU_DAU":round(a["DU_DAU"],4),
                "XA":round(a["XA"],4),
                "BU_MBW":round(a["BU_MBW"],4),
                "BU_MISA":round(a["BU_MISA"],4),
                "TONG_BU":round(tong_bu,4),
                "DU_CUOI":round(du_cuoi,4),
            })
        rows.sort(key=lambda r:r["MA_SKU"])

    total={k:round(sum(r.get(k,0) for r in rows),4) for k in ["DU_DAU","XA","BU_MBW","BU_MISA","TONG_BU","DU_CUOI"]}
    return {
        "report_type":report_type,
        "period":{"from":str(start) if start else None,"to":str(end) if end else None},
        "rows":rows,
        "total":total,
        "count":len(rows)
    }

def ckkm_report_v5(date_from=None,date_to=None,month=None,ma_nv=None):
    start,end=_v5_period(date_from,date_to,month)
    idx=_load_mapping()
    c=_conn()
    f_nv=_upper(ma_nv)
    opening=defaultdict(float)
    period=defaultdict(lambda:{"CKKM_XA":0.0,"THUONG":0.0,"CHI_MBW":0.0,"CHI_MISA":0.0})

    def add(nv,day,ckkm=0.0,thuong=0.0,chi_mbw=0.0,chi_misa=0.0):
        if not nv or day is None:
            return
        if f_nv and nv!=f_nv:
            return
        delta=ckkm+thuong-chi_mbw-chi_misa

        if start is not None and day<start:
            opening[nv]+=delta
            return
        if end is not None and day>end:
            return

        p=period[nv]
        p["CKKM_XA"]+=ckkm
        p["THUONG"]+=thuong
        p["CHI_MBW"]+=chi_mbw
        p["CHI_MISA"]+=chi_misa

    if _table_exists(c,"xb_src_batch_v2") and _table_exists(c,"xb_src_row_v2"):
        for bid in _active_batches(c,"xb_src_batch_v2"):
            rows=c.execute(
                "SELECT sheet_name,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='VALID'",
                (bid,)
            ).fetchall()
            for r in rows:
                try:
                    d=json.loads(r["payload_json"])
                except:
                    continue
                day=_v5_event_date(d,"DMS-XA")
                nv=_upper(_map_nv(idx,d))
                sh=_upper(r["sheet_name"])

                if sh=="CHI_TIET_XA":
                    val=sum(_num(_v(d,k)) for k in (
                        "CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI",
                        "KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC"
                    ))
                    add(nv,day,ckkm=val)

                elif sh=="THUONG_CAP_DON":
                    val=sum(_num(_v(d,k)) for k in (
                        "THUONG_TRUNG_BAY","THUONG_TICH_LUY","GIAM_TRU_CAP_DON_KHAC"
                    ))
                    add(nv,day,thuong=val)

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        if _upper(_v(d,"LOAI_DON"))!="B":
            continue
        day=_v5_event_date(d,"DMS-MBW")
        nv=_upper(_map_nv(idx,d))
        val=sum(_num(_v(d,k)) for k in (
            "CK_KM_TIEN","CK_THUONG_MAI","KM_HANG_GIA_TRI",
            "KM_BANG_HANG_GIA_TRI","GIAM_TRU_DONG_KHAC"
        ))
        add(nv,day,chi_mbw=val)

    for bid,row_no,d in _iter_misa(c) or []:
        day=_v5_event_date(d,"MISA-BU")
        nv=_upper(_map_nv(idx,d))
        val=sum(_num(_v(d,k)) for k in (
            "CK_KM_TIEN_CHUAN","CK_KM_TIEN","CHIET_KHAU",
            "GIAM_TRU_DONG_KHAC_CHUAN","GIAM_TRU_DONG_KHAC","GIA_TRI_GIAM_GIA"
        ))
        add(nv,day,chi_misa=val)

    c.close()

    rows=[]
    for nv in sorted(set(opening)|set(period)):
        p=period[nv]
        du_dau=opening[nv]
        quyen_loi=p["CKKM_XA"]+p["THUONG"]
        da_chi=p["CHI_MBW"]+p["CHI_MISA"]
        du_cuoi=du_dau+quyen_loi-da_chi

        if abs(du_dau)<0.005 and abs(quyen_loi)<0.005 and abs(da_chi)<0.005 and abs(du_cuoi)<0.005:
            continue

        rows.append({
            "MA_NV":nv,
            "DU_DAU":round(du_dau,2),
            "CKKM_XA":round(p["CKKM_XA"],2),
            "THUONG_CAP_DON":round(p["THUONG"],2),
            "TONG_QUYEN_LOI":round(quyen_loi,2),
            "DA_CHI_MBW":round(p["CHI_MBW"],2),
            "DA_CHI_MISA":round(p["CHI_MISA"],2),
            "TONG_DA_CHI":round(da_chi,2),
            "DU_CUOI":round(du_cuoi,2),
        })

    total={k:round(sum(r[k] for r in rows),2) for k in [
        "DU_DAU","CKKM_XA","THUONG_CAP_DON","TONG_QUYEN_LOI",
        "DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","DU_CUOI"
    ]}
    return {
        "period":{"from":str(start) if start else None,"to":str(end) if end else None},
        "rows":rows,
        "total":total,
        "count":len(rows)
    }

def quantity_v5_xlsx(**kwargs):
    rep=quantity_report_v5(**kwargs)
    wb=Workbook()
    ws=wb.active
    ws.title="SO_LUONG"

    rt=kwargs.get("report_type","employee")
    if rt=="employee":
        headers=["MA_NV","MA_SKU","TEN_SKU","BU","NHAN_HIEU","DU_DAU","XA","BU_MBW","BU_MISA","TONG_BU","DU_CUOI"]
    else:
        headers=["MA_SKU","TEN_SKU","BU","NHAN_HIEU","DU_DAU","XA","BU_MBW","BU_MISA","TONG_BU","DU_CUOI"]

    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")

    for r in rep["rows"]:
        ws.append([r.get(h,"") for h in headers])

    bio=BytesIO()
    wb.save(bio)
    return "XB_UI02_So_luong.xlsx",bio.getvalue()

def ckkm_v5_xlsx(**kwargs):
    rep=ckkm_report_v5(**kwargs)
    wb=Workbook()
    ws=wb.active
    ws.title="CK_KM"
    headers=["MA_NV","DU_DAU","CKKM_XA","THUONG_CAP_DON","TONG_QUYEN_LOI","DA_CHI_MBW","DA_CHI_MISA","TONG_DA_CHI","DU_CUOI"]

    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")

    for r in rep["rows"]:
        ws.append([r.get(h,"") for h in headers])

    bio=BytesIO()
    wb.save(bio)
    return "XB_UI02_CK_KM.xlsx",bio.getvalue()


# ===== V5.1 FINAL OVERRIDES =====
def report_options_v5():
    c=_conn()
    idx=_load_mapping()
    employees=set()

    for bid,row_no,d in _iter_dms_xa(c) or []:
        nv=_upper(_map_nv(idx,d))
        if nv:
            employees.add(nv)

    for bid,row_no,d in _iter_dms_mbw(c) or []:
        nv=_upper(_map_nv(idx,d))
        if nv:
            employees.add(nv)

    for bid,row_no,d in _iter_misa(c) or []:
        nv=_upper(_map_nv(idx,d))
        if nv:
            employees.add(nv)

    c.close()

    meta=_v5_meta()
    bus=sorted({str(x.get("BU") or "").strip().upper() for x in meta.values() if str(x.get("BU") or "").strip()})
    brands=sorted({str(x.get("NHAN_HIEU") or "").strip().upper() for x in meta.values() if str(x.get("NHAN_HIEU") or "").strip()})

    return {
        "NHAN_VIEN":sorted(employees),
        "SO_NHAN_VIEN":len(employees),
        "BU":bus,
        "NHAN_HIEU":brands,
    }
