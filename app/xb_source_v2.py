
import sqlite3, json, hashlib, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DMS_XA_DETAIL_REQUIRED={"MA_DON_DMS","MA_NV_NGHIEP_VU","MA_SKU"}
DMS_XA_BONUS_REQUIRED={"MA_DON_DMS","MA_NV_NGHIEP_VU"}

def _db_path():
    return Path(__file__).resolve().parent.parent/"data"/"miniapp.db"

def _conn():
    c=sqlite3.connect(str(_db_path())); c.row_factory=sqlite3.Row
    c.execute("CREATE TABLE IF NOT EXISTS xb_src_batch_v2(batch_id INTEGER PRIMARY KEY AUTOINCREMENT,source_type TEXT,file_name TEXT,file_hash TEXT,uploaded_by TEXT,uploaded_at TEXT,total_rows INTEGER DEFAULT 0,valid_rows INTEGER DEFAULT 0,error_rows INTEGER DEFAULT 0,status TEXT,UNIQUE(source_type,file_hash))")
    c.execute("CREATE TABLE IF NOT EXISTS xb_src_row_v2(row_id INTEGER PRIMARY KEY AUTOINCREMENT,batch_id INTEGER,sheet_name TEXT,row_no INTEGER,payload_json TEXT,status TEXT,error_text TEXT)")
    c.commit(); return c

def _h(v):
    return "" if v is None else str(v).strip().upper().replace(" ","_")

def _headers(ws):
    return [_h(c.value) for c in ws[1]]

def _validate_structure(wb):
    errors=[]
    names={s.upper():s for s in wb.sheetnames}
    for req in ("CHI_TIET_XA","THUONG_CAP_DON"):
        if req not in names: errors.append("Thiếu sheet "+req)
    if errors: return errors
    detail=set(_headers(wb[names["CHI_TIET_XA"]]))
    bonus=set(_headers(wb[names["THUONG_CAP_DON"]]))
    md=sorted(DMS_XA_DETAIL_REQUIRED-detail)
    mb=sorted(DMS_XA_BONUS_REQUIRED-bonus)
    if md: errors.append("CHI_TIET_XA thiếu cột: "+", ".join(md))
    if mb: errors.append("THUONG_CAP_DON thiếu cột: "+", ".join(mb))
    return errors

def preview_dms_xa(file_name,raw_bytes,uploaded_by="pilot"):
    file_hash=hashlib.sha256(raw_bytes).hexdigest()
    conn=_conn()
    dup=conn.execute("SELECT batch_id FROM xb_src_batch_v2 WHERE source_type='dms-xa' AND file_hash=?",(file_hash,)).fetchone()
    if dup:
        conn.close(); return {"duplicate":True,"duplicate_batch_id":dup["batch_id"]}
    wb=load_workbook(BytesIO(raw_bytes),data_only=False)
    errs=_validate_structure(wb)
    now=datetime.datetime.now().isoformat(timespec="seconds")
    cur=conn.execute("INSERT INTO xb_src_batch_v2(source_type,file_name,file_hash,uploaded_by,uploaded_at,status) VALUES('dms-xa',?,?,?,?,?)",(unquote(file_name),file_hash,uploaded_by,now,"PREVIEW"))
    bid=cur.lastrowid
    total=valid=error=0; preview=[]
    if errs:
        msg="; ".join(errs)
        conn.execute("INSERT INTO xb_src_row_v2(batch_id,sheet_name,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?,?)",(bid,"",0,"{}","ERROR",msg))
        error=1; preview=[{"sheet":"","row_no":0,"status":"ERROR","error":msg}]
    else:
        for sheet,need_sku in (("CHI_TIET_XA",True),("THUONG_CAP_DON",False)):
            ws=wb[sheet]; headers=_headers(ws); seen=set()
            for rno,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
                vals=list(row)
                if all(v is None or str(v).strip()=="" for v in vals): continue
                total+=1
                data={h:(vals[i] if i<len(vals) else None) for i,h in enumerate(headers) if h}
                er=[]
                if not data.get("MA_DON_DMS"): er.append("Thiếu MA_DON_DMS")
                if not data.get("MA_NV_NGHIEP_VU"): er.append("Thiếu MA_NV_NGHIEP_VU")
                if need_sku and not data.get("MA_SKU"): er.append("Thiếu MA_SKU")
                if sheet=="THUONG_CAP_DON":
                    od=str(data.get("MA_DON_DMS") or "").strip()
                    if od:
                        if od in seen: er.append("Trùng MA_DON_DMS trong THUONG_CAP_DON")
                        seen.add(od)
                st="ERROR" if er else "VALID"
                error+=1 if er else 0; valid+=0 if er else 1
                et="; ".join(er)
                conn.execute("INSERT INTO xb_src_row_v2(batch_id,sheet_name,row_no,payload_json,status,error_text) VALUES(?,?,?,?,?,?)",(bid,sheet,rno,json.dumps(data,ensure_ascii=False,default=str),st,et))
                if len(preview)<100: preview.append({"sheet":sheet,"row_no":rno,"status":st,"error":et})
    conn.execute("UPDATE xb_src_batch_v2 SET total_rows=?,valid_rows=?,error_rows=? WHERE batch_id=?",(total,valid,error,bid))
    conn.commit(); conn.close()
    return {"duplicate":False,"batch_id":bid,"total_rows":total,"valid_rows":valid,"error_rows":error,"structure_errors":errs,"preview":preview}

def commit_batch(batch_id):
    c=_conn(); b=c.execute("SELECT * FROM xb_src_batch_v2 WHERE batch_id=?",(batch_id,)).fetchone()
    if not b: c.close(); raise ValueError("Batch không tồn tại")
    if b["error_rows"]>0: c.close(); raise ValueError("Batch còn lỗi")
    c.execute("UPDATE xb_src_batch_v2 SET status='COMMITTED' WHERE batch_id=?",(batch_id,)); c.commit(); c.close()
    return {"batch_id":batch_id,"valid_rows":b["valid_rows"],"status":"COMMITTED"}

def history():
    c=_conn(); rows=c.execute("SELECT batch_id,source_type,file_name,uploaded_by,uploaded_at,total_rows,valid_rows,error_rows,status FROM xb_src_batch_v2 ORDER BY batch_id DESC LIMIT 100").fetchall(); c.close()
    return [dict(r) for r in rows]

def error_file(batch_id):
    c=_conn(); rows=c.execute("SELECT sheet_name,row_no,error_text,payload_json FROM xb_src_row_v2 WHERE batch_id=? AND status='ERROR'",(batch_id,)).fetchall(); c.close()
    wb=Workbook(); ws=wb.active; ws.title="LOI_UPLOAD"
    heads=["SHEET","ROW_NO","ERROR","DATA_JSON"]
    for i,h in enumerate(heads,1):
        cell=ws.cell(1,i,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="C62828")
        ws.column_dimensions[get_column_letter(i)].width=[22,12,60,100][i-1]
    for r in rows: ws.append([r["sheet_name"],r["row_no"],r["error_text"],r["payload_json"]])
    bio=BytesIO(); wb.save(bio)
    return f"XB_DMS_XA_Loi_Batch_{batch_id}.xlsx",bio.getvalue()
