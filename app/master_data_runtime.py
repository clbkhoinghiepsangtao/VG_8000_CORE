import json, sqlite3, datetime
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from .master_templates import MASTER_TEMPLATES

KEY_FIELDS = {
 'dm01-product':'MA_SKU',
 'price':'MA_SKU|TU_NGAY',
 'dm02-staff':'MA_NV',
 'dm02a-staff-map':'MA_NV|LOAI_MA|MA_NGHIEP_VU|HE_THONG|TU_NGAY',
 'sales-region':'MA_VUNG',
 'dm03-customer':'MA_KH_CHUAN',
 'territory':'MA_DIA_BAN',
 'basic':'LOAI_DANH_MUC|MA',
 'users':'USERNAME'
}

def _db_path():
    return Path(__file__).resolve().parent.parent / 'data' / 'miniapp.db'

def _conn():
    p=_db_path()
    p.parent.mkdir(exist_ok=True)
    c=sqlite3.connect(str(p))
    c.row_factory=sqlite3.Row
    c.execute('CREATE TABLE IF NOT EXISTS master_data_generic (category TEXT NOT NULL, business_key TEXT NOT NULL, payload_json TEXT NOT NULL, updated_at TEXT NOT NULL, PRIMARY KEY(category,business_key))')
    c.commit()
    return c

def _business_key(category,row):
    fields=KEY_FIELDS[category].split('|')
    vals=[]
    for f in fields:
        v=row.get(f)
        if v is None or str(v).strip()=='':
            raise ValueError('Thiếu khóa: '+f)
        vals.append(str(v).strip())
    return '||'.join(vals)

def preview_upload(category, raw_bytes):
    if category not in MASTER_TEMPLATES:
        raise ValueError('Danh mục không hợp lệ')
    spec=MASTER_TEMPLATES[category]
    wb=load_workbook(BytesIO(raw_bytes), data_only=True)
    if spec['sheet'] not in wb.sheetnames:
        raise ValueError('Thiếu sheet bắt buộc: '+spec['sheet'])
    ws=wb[spec['sheet']]
    expected=spec['headers']
    headers=[str(c.value).strip() if c.value is not None else '' for c in ws[1]][:len(expected)]
    if headers != expected:
        raise ValueError('Sai cấu trúc cột. Hãy tải lại template chuẩn từ hệ thống.')
    conn=_conn()
    rows=[]
    counts={'add':0,'update':0,'error':0}
    for r in ws.iter_rows(min_row=2, values_only=True):
        vals=list(r[:len(expected)])
        if all(v is None or str(v).strip()=='' for v in vals):
            continue
        row={expected[i]: vals[i] if i < len(vals) else None for i in range(len(expected))}
        try:
            key=_business_key(category,row)
            exists=conn.execute('SELECT 1 FROM master_data_generic WHERE category=? AND business_key=?',(category,key)).fetchone() is not None
            status='UPDATE' if exists else 'ADD'
            counts['update' if exists else 'add'] += 1
            err=''
        except Exception as e:
            key=''
            status='ERROR'
            counts['error'] += 1
            err=str(e)
        rows.append({'status':status,'key':key,'error':err,'data':row})
    conn.close()
    return {'category':category,'counts':counts,'rows':rows}

def commit_upload(category, rows):
    conn=_conn()
    now=datetime.datetime.now().isoformat(timespec='seconds')
    done=0
    for item in rows:
        if item.get('status')=='ERROR':
            continue
        row=item.get('data') or {}
        key=_business_key(category,row)
        conn.execute('INSERT INTO master_data_generic(category,business_key,payload_json,updated_at) VALUES(?,?,?,?) ON CONFLICT(category,business_key) DO UPDATE SET payload_json=excluded.payload_json, updated_at=excluded.updated_at',(category,key,json.dumps(row,ensure_ascii=False,default=str),now))
        done += 1
    conn.commit()
    conn.close()
    return {'committed':done}

def export_master(category):
    if category not in MASTER_TEMPLATES:
        raise ValueError('Danh mục không hợp lệ')
    spec=MASTER_TEMPLATES[category]
    conn=_conn()
    data=conn.execute('SELECT payload_json FROM master_data_generic WHERE category=? ORDER BY business_key',(category,)).fetchall()
    conn.close()
    wb=Workbook()
    ws=wb.active
    ws.title=spec['sheet']
    for i,h in enumerate(spec['headers'],1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor='1769E0')
        ws.column_dimensions[get_column_letter(i)].width=max(16,min(28,len(h)+4))
    for rr in data:
        d=json.loads(rr['payload_json'])
        ws.append([d.get(h,'') for h in spec['headers']])
    bio=BytesIO()
    wb.save(bio)
    filename=spec['filename'].replace('_V1.0.xlsx','_EXPORT.xlsx')
    return filename,bio.getvalue()

def list_master_rows(category):
    if category not in MASTER_TEMPLATES:
        raise ValueError('Danh mục không hợp lệ')
    conn=_conn()
    data=conn.execute('SELECT payload_json FROM master_data_generic WHERE category=? ORDER BY business_key',(category,)).fetchall()
    conn.close()
    return [json.loads(r['payload_json']) for r in data]


def delete_master_row(category, business_key):
    if category not in MASTER_TEMPLATES:
        raise ValueError('Danh mục không hợp lệ')
    conn = _conn()
    cur = conn.execute('DELETE FROM master_data_generic WHERE category=? AND business_key=?', (category, str(business_key)))
    conn.commit()
    deleted = cur.rowcount
    conn.close()
    return {'deleted': deleted}


DEFAULT_PERMISSION_BY_ROLE = {
    "SUPER_ADMIN": {"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1},
    "ADMIN":       {"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1},
    "KE_TOAN":     {"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1},
    "THU_QUY":     {"GH":"EDIT","XB":"NONE","KQ":"NONE","BX":"NONE","DM":"NONE","EXPORT":0},
    "THU_KHO":     {"GH":"EDIT","XB":"NONE","KQ":"NONE","BX":"EDIT","DM":"NONE","EXPORT":0},
    "QUAN_LY":     {"GH":"VIEW","XB":"VIEW","KQ":"VIEW","BX":"VIEW","DM":"VIEW","EXPORT":1},
    "VIEWER":      {"GH":"VIEW","XB":"VIEW","KQ":"VIEW","BX":"VIEW","DM":"VIEW","EXPORT":0},
    "HCNS":        {"GH":"NONE","XB":"NONE","KQ":"VIEW","BX":"NONE","DM":"EDIT","EXPORT":1},
}

def _ensure_user_permission_table(conn):
    conn.execute(
        "CREATE TABLE IF NOT EXISTS user_permission_matrix ("
        "username TEXT PRIMARY KEY,"
        "gh TEXT NOT NULL DEFAULT 'NONE',"
        "xb TEXT NOT NULL DEFAULT 'NONE',"
        "kq TEXT NOT NULL DEFAULT 'NONE',"
        "bx TEXT NOT NULL DEFAULT 'NONE',"
        "dm TEXT NOT NULL DEFAULT 'NONE',"
        "export_excel INTEGER NOT NULL DEFAULT 0,"
        "updated_at TEXT NOT NULL"
        ")"
    )
    conn.commit()

def _default_permission(role):
    role = str(role or "").strip().upper()
    d = DEFAULT_PERMISSION_BY_ROLE.get(role, {"GH":"NONE","XB":"NONE","KQ":"NONE","BX":"NONE","DM":"NONE","EXPORT":0})
    return dict(d)

def list_users_with_permissions():
    conn = _conn()
    _ensure_user_permission_table(conn)
    rows = conn.execute("SELECT payload_json FROM master_data_generic WHERE category='users' ORDER BY business_key").fetchall()
    result = []
    for rr in rows:
        u = json.loads(rr["payload_json"])
        username = str(u.get("USERNAME","")).strip()
        role = str(u.get("VAI_TRO","")).strip().upper()
        p = conn.execute("SELECT gh,xb,kq,bx,dm,export_excel FROM user_permission_matrix WHERE username=?", (username,)).fetchone()
        if p is None:
            d = _default_permission(role)
            perm = {"GH":d["GH"],"XB":d["XB"],"KQ":d["KQ"],"BX":d["BX"],"DM":d["DM"],"EXPORT":int(d["EXPORT"])}
        else:
            perm = {"GH":p["gh"],"XB":p["xb"],"KQ":p["kq"],"BX":p["bx"],"DM":p["dm"],"EXPORT":int(p["export_excel"])}
        result.append({
            "USERNAME":username,
            "HO_TEN":u.get("HO_TEN",""),
            "MA_NV":u.get("MA_NV",""),
            "VAI_TRO":role,
            "TRANG_THAI":u.get("TRANG_THAI",""),
            "PERMISSION":perm,
        })
    conn.close()
    return result

def save_user_permissions(username, permissions):
    username = str(username or "").strip()
    if not username:
        raise ValueError("Thiếu username")
    conn = _conn()
    _ensure_user_permission_table(conn)
    rr = conn.execute("SELECT payload_json FROM master_data_generic WHERE category='users' AND business_key=?", (username,)).fetchone()
    if rr is None:
        conn.close()
        raise ValueError("User không tồn tại")
    u = json.loads(rr["payload_json"])
    role = str(u.get("VAI_TRO","")).strip().upper()
    if role == "SUPER_ADMIN":
        p = {"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1}
    else:
        allowed = {"NONE","VIEW","EDIT"}
        p = {}
        for k in ("GH","XB","KQ","BX","DM"):
            v = str((permissions or {}).get(k,"NONE")).strip().upper()
            if v not in allowed:
                raise ValueError("Giá trị quyền không hợp lệ: "+k)
            p[k] = v
        if p["KQ"] == "EDIT":
            p["KQ"] = "VIEW"
        p["EXPORT"] = 1 if bool((permissions or {}).get("EXPORT",0)) else 0
    now = datetime.datetime.now().isoformat(timespec="seconds")
    conn.execute(
        "INSERT INTO user_permission_matrix(username,gh,xb,kq,bx,dm,export_excel,updated_at) VALUES(?,?,?,?,?,?,?,?) "
        "ON CONFLICT(username) DO UPDATE SET gh=excluded.gh,xb=excluded.xb,kq=excluded.kq,bx=excluded.bx,dm=excluded.dm,export_excel=excluded.export_excel,updated_at=excluded.updated_at",
        (username,p["GH"],p["XB"],p["KQ"],p["BX"],p["DM"],p["EXPORT"],now)
    )
    conn.commit()
    conn.close()
    return {"saved":1,"username":username,"permission":p}


# VG_USER_PERMISSION_TEMPLATE_OVERRIDE
def _vg_ensure_user_permission_table(conn):
    conn.execute(
        "CREATE TABLE IF NOT EXISTS user_permission_matrix ("
        "username TEXT PRIMARY KEY,"
        "gh TEXT NOT NULL DEFAULT 'NONE',"
        "xb TEXT NOT NULL DEFAULT 'NONE',"
        "kq TEXT NOT NULL DEFAULT 'NONE',"
        "bx TEXT NOT NULL DEFAULT 'NONE',"
        "dm TEXT NOT NULL DEFAULT 'NONE',"
        "export_excel INTEGER NOT NULL DEFAULT 0,"
        "updated_at TEXT NOT NULL)"
    )
    conn.commit()

def _vg_default_perm(role):
    defaults = {
        "SUPER_ADMIN":{"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1},
        "ADMIN":{"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1},
        "KE_TOAN":{"GH":"EDIT","XB":"EDIT","KQ":"VIEW","BX":"EDIT","DM":"EDIT","EXPORT":1},
        "THU_QUY":{"GH":"EDIT","XB":"NONE","KQ":"NONE","BX":"NONE","DM":"NONE","EXPORT":0},
        "THU_KHO":{"GH":"EDIT","XB":"NONE","KQ":"NONE","BX":"EDIT","DM":"NONE","EXPORT":0},
        "QUAN_LY":{"GH":"VIEW","XB":"VIEW","KQ":"VIEW","BX":"VIEW","DM":"VIEW","EXPORT":1},
        "VIEWER":{"GH":"VIEW","XB":"VIEW","KQ":"VIEW","BX":"VIEW","DM":"VIEW","EXPORT":0},
        "HCNS":{"GH":"NONE","XB":"NONE","KQ":"VIEW","BX":"NONE","DM":"EDIT","EXPORT":1},
    }
    return dict(defaults.get(str(role or "").strip().upper(),
                             {"GH":"NONE","XB":"NONE","KQ":"NONE","BX":"NONE","DM":"NONE","EXPORT":0}))

def commit_upload(category, rows):
    conn=_conn()
    now=datetime.datetime.now().isoformat(timespec="seconds")
    done=0
    if category == "users":
        _vg_ensure_user_permission_table(conn)

    for item in rows:
        if item.get("status")=="ERROR":
            continue
        row=item.get("data") or {}
        key=_business_key(category,row)

        conn.execute(
            "INSERT INTO master_data_generic(category,business_key,payload_json,updated_at) "
            "VALUES(?,?,?,?) "
            "ON CONFLICT(category,business_key) DO UPDATE SET "
            "payload_json=excluded.payload_json, updated_at=excluded.updated_at",
            (category,key,json.dumps(row,ensure_ascii=False,default=str),now)
        )

        if category == "users":
            role=str(row.get("VAI_TRO","")).strip().upper()

            if role == "SUPER_ADMIN":
                p=_vg_default_perm(role)
            else:
                allowed={"NONE","VIEW","EDIT"}
                p={
                    "GH":str(row.get("GH","NONE")).strip().upper(),
                    "XB":str(row.get("XB","NONE")).strip().upper(),
                    "KQ":str(row.get("KQ","NONE")).strip().upper(),
                    "BX":str(row.get("BX","NONE")).strip().upper(),
                    "DM":str(row.get("DANH_MUC","NONE")).strip().upper(),
                    "EXPORT":1 if str(row.get("XUAT_EXCEL","0")).strip().lower() in ("1","true","yes","co","có") else 0,
                }
                for k in ("GH","XB","KQ","BX","DM"):
                    if p[k] not in allowed:
                        conn.close()
                        raise ValueError("Quyền không hợp lệ "+k+": "+str(p[k]))
                if p["KQ"]=="EDIT":
                    p["KQ"]="VIEW"

            conn.execute(
                "INSERT INTO user_permission_matrix(username,gh,xb,kq,bx,dm,export_excel,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?) "
                "ON CONFLICT(username) DO UPDATE SET "
                "gh=excluded.gh, xb=excluded.xb, kq=excluded.kq, bx=excluded.bx, "
                "dm=excluded.dm, export_excel=excluded.export_excel, updated_at=excluded.updated_at",
                (str(row.get("USERNAME","")).strip(),
                 p["GH"],p["XB"],p["KQ"],p["BX"],p["DM"],p["EXPORT"],now)
            )

        done += 1

    conn.commit()
    conn.close()
    return {"committed":done}

def export_master(category):
    if category not in MASTER_TEMPLATES:
        raise ValueError("Danh mục không hợp lệ")

    spec=MASTER_TEMPLATES[category]
    conn=_conn()

    wb=Workbook()
    ws=wb.active
    ws.title=spec["sheet"]

    for i,h in enumerate(spec["headers"],1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1769E0")
        ws.column_dimensions[get_column_letter(i)].width=max(16,min(28,len(h)+4))

    if category == "users":
        _vg_ensure_user_permission_table(conn)

        data=conn.execute(
            "SELECT payload_json FROM master_data_generic "
            "WHERE category='users' ORDER BY business_key"
        ).fetchall()

        for rr in data:
            d=json.loads(rr["payload_json"])
            username=str(d.get("USERNAME","")).strip()
            role=str(d.get("VAI_TRO","")).strip().upper()

            p=conn.execute(
                "SELECT gh,xb,kq,bx,dm,export_excel "
                "FROM user_permission_matrix WHERE username=?",
                (username,)
            ).fetchone()

            if p is None:
                q=_vg_default_perm(role)
                gh,xb,kq,bx,dm,ex=q["GH"],q["XB"],q["KQ"],q["BX"],q["DM"],q["EXPORT"]
            else:
                gh,xb,kq,bx,dm,ex=p["gh"],p["xb"],p["kq"],p["bx"],p["dm"],int(p["export_excel"])

            out={
                "USERNAME":d.get("USERNAME",""),
                "HO_TEN":d.get("HO_TEN",""),
                "MA_NV":d.get("MA_NV",""),
                "VAI_TRO":d.get("VAI_TRO",""),
                "GH":gh,
                "XB":xb,
                "KQ":kq,
                "BX":bx,
                "DANH_MUC":dm,
                "XUAT_EXCEL":ex,
                "TRANG_THAI":d.get("TRANG_THAI",""),
            }
            ws.append([out.get(h,"") for h in spec["headers"]])
    else:
        data=conn.execute(
            "SELECT payload_json FROM master_data_generic "
            "WHERE category=? ORDER BY business_key",
            (category,)
        ).fetchall()

        for rr in data:
            d=json.loads(rr["payload_json"])
            ws.append([d.get(h,"") for h in spec["headers"]])

    conn.close()

    bio=BytesIO()
    wb.save(bio)

    filename=spec["filename"].replace("_V1.0.xlsx","_EXPORT.xlsx")
    return filename,bio.getvalue()
