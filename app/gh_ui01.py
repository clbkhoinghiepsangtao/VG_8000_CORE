
import sqlite3, json, datetime
from pathlib import Path
from collections import defaultdict
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def _gh_db():
    return Path(__file__).resolve().parent.parent/"data"/"miniapp.db"

def _gh_conn():
    c=sqlite3.connect(str(_gh_db()))
    c.row_factory=sqlite3.Row
    return c

def _gh_table_exists(c,t):
    return c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(t,)).fetchone() is not None

def _gh_num(v):
    if v in (None,""):
        return 0.0
    if isinstance(v,str) and v.strip() in ("-","–","—"):
        return 0.0
    try:return float(v)
    except:
        try:return float(str(v).replace(",",""))
        except:return 0.0

def _gh_upper(v):
    return str(v or "").strip().upper()

def _gh_v(d,*keys):
    for k in keys:
        if k in d and d[k] not in (None,""):
            return d[k]
    return None

def _gh_date(v):
    if v in (None,""):
        return None
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        try:
            n=float(v)
            if 20000<=n<=80000:
                return datetime.date(1899,12,30)+datetime.timedelta(days=int(n))
        except:pass
    if isinstance(v,datetime.datetime):return v.date()
    if isinstance(v,datetime.date):return v
    s=str(v).strip()
    try:
        if s.replace(".","",1).isdigit():
            n=float(s)
            if 20000<=n<=80000:
                return datetime.date(1899,12,30)+datetime.timedelta(days=int(n))
    except:pass
    for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%Y/%m/%d","%Y-%m-%d %H:%M:%S"):
        try:return datetime.datetime.strptime(s[:19],fmt).date()
        except:pass
    try:return datetime.date.fromisoformat(s[:10])
    except:return None

def _gh_active_batches(c):
    if not _gh_table_exists(c,"xb_dms_mbw_batch"):
        return []
    return [r["batch_id"] for r in c.execute(
        "SELECT batch_id FROM xb_dms_mbw_batch WHERE UPPER(status)='COMMITTED'"
    ).fetchall()]

def _gh_ensure_tables(c):
    c.execute("""
    CREATE TABLE IF NOT EXISTS gh_order_status(
        order_no TEXT PRIMARY KEY,
        status TEXT NOT NULL DEFAULT 'CHUA_QUYET_TOAN',
        updated_at TEXT,
        updated_by TEXT,
        note TEXT
    )
    """)
    c.commit()

def gh_ui01_orders(date_from=None,date_to=None,nv_giao=None,lai_xe=None,phuong_tien=None,lo_giao=None):
    c=_gh_conn()
    _gh_ensure_tables(c)

    df=_gh_date(date_from) if date_from else None
    dt=_gh_date(date_to) if date_to else None
    f_nv=_gh_upper(nv_giao)
    f_lx=_gh_upper(lai_xe)
    f_pt=_gh_upper(phuong_tien)
    f_lo=_gh_upper(lo_giao)

    agg={}
    if _gh_table_exists(c,"xb_dms_mbw_row"):
        for bid in _gh_active_batches(c):
            rows=c.execute(
                "SELECT row_no,payload_json FROM xb_dms_mbw_row WHERE batch_id=? AND status='VALID'",
                (bid,)
            ).fetchall()
            for r in rows:
                try:d=json.loads(r["payload_json"])
                except:continue

                if _gh_upper(_gh_v(d,"LOAI_DON"))!="L":
                    continue

                day=_gh_date(_gh_v(d,"NGAY_GIAO","NGAY_DON"))
                if df and (day is None or day<df):continue
                if dt and (day is None or day>dt):continue

                order_no=str(_gh_v(d,"SO_HOA_DON","MA_DON_DMS","MA_DON","SO_DON") or "").strip()
                if not order_no:
                    order_no=f"B{bid}-R{r['row_no']}"

                nv=_gh_upper(_gh_v(d,"MA_NV_GIAO_HANG_NGHIEP_VU","MA_NV_GIAO_HANG","NV_GIAO_HANG"))
                lx=_gh_upper(_gh_v(d,"MA_NV_LAI_XE_NGHIEP_VU","MA_NV_LAI_XE","LAI_XE"))
                pt=_gh_upper(_gh_v(d,"PHUONG_TIEN","BIEN_SO","XE"))
                lo=_gh_upper(_gh_v(d,"LO_GIAO_HANG","LO_GIAO","CHUYEN_GIAO"))

                if f_nv and nv!=f_nv:continue
                if f_lx and lx!=f_lx:continue
                if f_pt and pt!=f_pt:continue
                if f_lo and lo!=f_lo:continue

                sku=_gh_upper(_gh_v(d,"MA_SKU"))
                sl_su=_gh_num(_gh_v(d,"SL_SU","SO_LUONG_SU"))
                sl_thung=_gh_num(_gh_v(d,"SL_THUNG","SO_LUONG_THUNG"))
                gss=_gh_num(_gh_v(d,"DOANH_SO_PHAI_THU","DOANH_SO_TRUOC_THUE","DOANH_SO"))

                if order_no not in agg:
                    agg[order_no]={
                        "ORDER_NO":order_no,
                        "NGAY_GIAO":day.isoformat() if day else "",
                        "MA_NV_KD":_gh_upper(_gh_v(d,"MA_NV_KD_NGHIEP_VU","MA_NV_NGHIEP_VU")),
                        "MA_NV_GIAO_HANG":nv,
                        "MA_NV_LAI_XE":lx,
                        "PHUONG_TIEN":pt,
                        "LO_GIAO_HANG":lo,
                        "SKU_SET":set(),
                        "SL_SU":0.0,
                        "SL_THUNG":0.0,
                        "DOANH_SO_PHAI_THU":0.0,
                    }
                a=agg[order_no]
                if sku:a["SKU_SET"].add(sku)
                a["SL_SU"]+=sl_su
                a["SL_THUNG"]+=sl_thung
                a["DOANH_SO_PHAI_THU"]+=gss

    status_rows={r["order_no"]:dict(r) for r in c.execute("SELECT * FROM gh_order_status").fetchall()}
    c.close()

    out=[]
    for order_no,a in agg.items():
        st=status_rows.get(order_no,{})
        out.append({
            "ORDER_NO":order_no,
            "NGAY_GIAO":a["NGAY_GIAO"],
            "MA_NV_KD":a["MA_NV_KD"],
            "MA_NV_GIAO_HANG":a["MA_NV_GIAO_HANG"],
            "MA_NV_LAI_XE":a["MA_NV_LAI_XE"],
            "PHUONG_TIEN":a["PHUONG_TIEN"],
            "LO_GIAO_HANG":a["LO_GIAO_HANG"],
            "SO_SKU":len(a["SKU_SET"]),
            "SL_SU":round(a["SL_SU"],4),
            "SL_THUNG":round(a["SL_THUNG"],4),
            "DOANH_SO_PHAI_THU":round(a["DOANH_SO_PHAI_THU"],2),
            "TRANG_THAI":st.get("status","CHUA_QUYET_TOAN"),
        })
    out.sort(key=lambda x:(x["NGAY_GIAO"],x["ORDER_NO"]),reverse=True)
    return {"rows":out,"count":len(out)}

def gh_ui01_xlsx(**kwargs):
    rep=gh_ui01_orders(**kwargs)
    wb=Workbook();ws=wb.active;ws.title="GH_UI01"
    headers=["ORDER_NO","NGAY_GIAO","MA_NV_KD","MA_NV_GIAO_HANG","MA_NV_LAI_XE","PHUONG_TIEN","LO_GIAO_HANG","SO_SKU","SL_SU","SL_THUNG","DOANH_SO_PHAI_THU","TRANG_THAI"]
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h);c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="1769E0")
    for r in rep["rows"]:
        ws.append([r[h] for h in headers])
    bio=BytesIO();wb.save(bio)
    return "GH_UI01_Danh_sach_don_can_quyet_toan.xlsx",bio.getvalue()
